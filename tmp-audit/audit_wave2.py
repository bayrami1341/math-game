# -*- coding: utf-8 -*-
# راستی‌آزمایی سه‌لایهٔ تحویل موج ۲ (v0.3) — تسک ۴۸
# لایهٔ ۲: بازرسی سلولی — مقایسهٔ اسکلت موج ۱ با v0.2 + شمارش لایه‌های تازه
import openpyxl, json, re

V02 = '/home/z/my-project/public/math-game/authoring/khshab-G2-MATH_v0.2.xlsx'
V03 = '/tmp/khshab-G2-MATH_v0.3.xlsx'

wb2 = openpyxl.load_workbook(V02, data_only=True)
wb3 = openpyxl.load_workbook(V03, data_only=True)

out = {}
out['sheets_v02'] = wb2.sheetnames
out['sheets_v03'] = wb3.sheetnames
out['sheet_names_identical'] = wb2.sheetnames == wb3.sheetnames

def norm(v):
    if v is None: return ''
    return str(v).strip()

def find_row(ws, first_prefix, col=1, limit=60):
    for r in range(1, limit):
        if norm(ws.cell(r, col).value).startswith(first_prefix):
            return r
    return None

def data_rows_above_yad(ws, head_idx):
    rows = []
    for r in range(head_idx + 1, ws.max_row + 1):
        first = norm(ws.cell(r, 1).value)
        if not first or first == 'یادداشت':
            break
        rows.append(r)
    return rows

name_map = next(n for n in wb2.sheetnames if 'نقشه' in n)
name_atlas = next(n for n in wb2.sheetnames if 'اطلس' in n)
name_bank = next(n for n in wb2.sheetnames if 'بانک' in n)
name_id = next(n for n in wb2.sheetnames if 'شناسنامه' in n)

# ── ۱) نقشهٔ کتاب ──
m2, m3 = wb2[name_map], wb3[name_map]
h2, h3 = find_row(m2, 'شناسه'), find_row(m3, 'شناسه')
hdr2 = [norm(m2.cell(h2, c).value) for c in range(1, 15)]
hdr3 = [norm(m3.cell(h3, c).value) for c in range(1, 15)]
out['map_header_identical'] = hdr2 == hdr3

rows2 = data_rows_above_yad(m2, h2)
rows3 = data_rows_above_yad(m3, h3)
out['map_rows_v02'] = len(rows2)
out['map_rows_v03'] = len(rows3)

diffs = []
for r2, r3 in zip(rows2, rows3):
    vals2 = [norm(m2.cell(r2, c).value) for c in range(1, 15)]
    vals3 = [norm(m3.cell(r3, c).value) for c in range(1, 15)]
    if vals2 != vals3:
        for c in range(14):
            if vals2[c] != vals3[c]:
                diffs.append({'id': vals2[0], 'col': hdr2[c], 'v02': vals2[c][:40], 'v03': vals3[c][:40]})
out['skeleton_untouched'] = len(diffs) == 0
out['skeleton_diffs'] = diffs[:10]

new_rows = rows3[len(rows2):]
new_meta = []
for r in new_rows:
    new_meta.append({
        'id': norm(m3.cell(r, 1).value),
        'page': norm(m3.cell(r, 2).value),
        'title': norm(m3.cell(r, 4).value)[:30],
        'engine': norm(m3.cell(r, 10).value),
        'anchors': norm(m3.cell(r, 11).value),
        'src': norm(m3.cell(r, 12).value),
        'status': norm(m3.cell(r, 13).value),
    })
out['new_map_count'] = len(new_meta)
out['all_new_yellow'] = all(x['status'] == 'نیاز به تأیید' for x in new_meta)
out['all_new_ai_source'] = all('پیشنهاد AI' in x['src'] for x in new_meta)
out['new_pages'] = [x['page'] for x in new_meta]
out['new_ids_first_last'] = (new_meta[0]['id'], new_meta[-1]['id'])
out['new_engine_compat'] = {}
for x in new_meta:
    out['new_engine_compat'][x['engine']] = out['new_engine_compat'].get(x['engine'], 0) + 1
out['new_meta_sample'] = [new_meta[0], new_meta[1], new_meta[-1]]

# ── ۲) اطلس خطا ──
a2, a3 = wb2[name_atlas], wb3[name_atlas]
ha2, ha3 = find_row(a2, 'کد'), find_row(a3, 'کد')
rowsa2, rowsa3 = data_rows_above_yad(a2, ha2), data_rows_above_yad(a3, ha3)
out['atlas_v02_empty'] = len(rowsa2) == 0
atlas_codes = []
for r in rowsa3:
    atlas_codes.append({
        'code': norm(a3.cell(r, 1).value),
        'name': norm(a3.cell(r, 2).value)[:25],
        'status': norm(a3.cell(r, 10).value),
        'proposed': 'پیشنهاد' in norm(a3.cell(r, 1).value),
    })
out['atlas_count'] = len(atlas_codes)
out['atlas_codes'] = [(x['code'], x['status']) for x in atlas_codes]
out['atlas_proposed'] = [x['code'] for x in atlas_codes if x['proposed']]
out['atlas_all_yellow'] = all(x['status'] == 'نیاز به تأیید' for x in atlas_codes)

# ── ۳) بانک سؤال ──
q2, q3 = wb2[name_bank], wb3[name_bank]
hq2, hq3 = find_row(q2, 'شناسه'), find_row(q3, 'شناسه')
hdrq2 = [norm(q2.cell(hq2, c).value) for c in range(1, 18)]
hdrq3 = [norm(q3.cell(hq3, c).value) for c in range(1, 18)]
out['bank_header_identical'] = hdrq2 == hdrq3
rowsq2, rowsq3 = data_rows_above_yad(q2, hq2), data_rows_above_yad(q3, hq3)
out['bank_v02_empty'] = len(rowsq2) == 0

bank = []
for r in rowsq3:
    bank.append({
        'id': norm(q3.cell(r, 1).value),
        'anchor': norm(q3.cell(r, 2).value),
        'domain': norm(q3.cell(r, 3).value),
        'cpa': norm(q3.cell(r, 4).value),
        'opts': [norm(q3.cell(r, c).value) for c in range(7, 11)],
        'ans': norm(q3.cell(r, 11).value),
        'emap': norm(q3.cell(r, 12).value),
        'src': norm(q3.cell(r, 15).value),
        'status': norm(q3.cell(r, 16).value),
    })
out['bank_count'] = len(bank)
out['bank_all_yellow'] = all(x['status'] == 'نیاز به تأیید' for x in bank)
out['bank_all_ai_src'] = all('پیشنهاد AI' in x['src'] for x in bank)
out['bank_ids_unique'] = len(set(x['id'] for x in bank)) == len(bank)
out['bank_ids_first_last'] = (bank[0]['id'], bank[-1]['id'])
out['bank_ans_balance'] = {L: sum(1 for x in bank if x['ans'] == L) for L in ['الف', 'ب', 'ج', 'د']}
dom = {}
for x in bank: dom[x['domain']] = dom.get(x['domain'], 0) + 1
out['bank_domain'] = dom
out['bank_opts4'] = all(sum(1 for o in x['opts'] if o) == 4 for x in bank)
own = {}
for x in bank: own.setdefault(x['anchor'], []).append(x['id'])
t7_bad = []
for x in new_meta:
    refs = sorted(set(re.findall(r'Q-\d+', x['anchors'])))
    if refs != sorted(own.get(x['id'], [])):
        t7_bad.append({'row': x['id'], 'refs': refs, 'owners': sorted(own.get(x['id'], []))})
out['t7_ok'] = len(t7_bad) == 0
out['t7_bad_sample'] = t7_bad[:6]
out['bank_anchor_dist'] = {k: len(v) for k, v in sorted(own.items())}

# ── ۴) شناسنامه ──
i2, i3 = wb2[name_id], wb3[name_id]
hi2, hi3 = find_row(i2, 'فیلد'), find_row(i3, 'فیلد')
id_diffs = []
r2, r3 = hi2 + 1, hi3 + 1
while r2 <= i2.max_row and r3 <= i3.max_row:
    k2, k3 = norm(i2.cell(r2, 1).value), norm(i3.cell(r3, 1).value)
    if not k2 or not k3: break
    v2 = [norm(i2.cell(r2, c).value) for c in range(2, 5)]
    v3 = [norm(i3.cell(r3, c).value) for c in range(2, 5)]
    if v2 != v3:
        id_diffs.append({'field': k2, 'v02': v2, 'v03': v3})
    r2 += 1; r3 += 1
out['identity_diffs'] = id_diffs

print(json.dumps(out, ensure_ascii=False, indent=1))
