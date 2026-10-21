# -*- coding: utf-8 -*-
# ممیزی موج ۲ — نسخهٔ ۲: تطبیق بر اساس شناسه (نه ترتیب ردیف)
import openpyxl, json, re

V02 = '/home/z/my-project/public/math-game/authoring/khshab-G2-MATH_v0.2.xlsx'
V03 = '/tmp/khshab-G2-MATH_v0.3.xlsx'

wb2 = openpyxl.load_workbook(V02, data_only=True)
wb3 = openpyxl.load_workbook(V03, data_only=True)

def norm(v):
    return '' if v is None else str(v).strip()

def find_row(ws, prefix, limit=60):
    for r in range(1, limit):
        if norm(ws.cell(r, 1).value).startswith(prefix):
            return r
    return None

def rows_by_id(ws, head_idx):
    out = {}
    for r in range(head_idx + 1, ws.max_row + 1):
        k = norm(ws.cell(r, 1).value)
        if not k or k == 'یادداشت':
            break
        out[k] = r
    return out

name_map = next(n for n in wb2.sheetnames if 'نقشه' in n)
name_id = next(n for n in wb2.sheetnames if 'شناسنامه' in n)
m2, m3 = wb2[name_map], wb3[name_map]
h2, h3 = find_row(m2, 'شناسه'), find_row(m3, 'شناسه')
hdr3 = [norm(m3.cell(h3, c).value) for c in range(1, 15)]

ids2 = rows_by_id(m2, h2)
ids3 = rows_by_id(m3, h3)

out = {'skeleton_ids_v02': sorted(ids2.keys()), 'skeleton_ids_in_v03': sorted(ids2.keys() & ids3.keys())}

# تطبیق مقداربه‌مقدار ستون‌های محتوایی (۱-۱۱) + ستون‌های ممیزی (۱۲-۱۴) جداگانه
CONTENT_COLS = list(range(1, 12))
AUDIT_COLS = list(range(12, 15))
content_diffs, audit_diffs = [], []
for rid in sorted(ids2.keys()):
    if rid not in ids3:
        content_diffs.append({'id': rid, 'issue': 'یافت نشد در v0.3'}); continue
    for c in CONTENT_COLS:
        v2, v3 = norm(m2.cell(ids2[rid], c).value), norm(m3.cell(ids3[rid], c).value)
        if v2 != v3:
            content_diffs.append({'id': rid, 'col': hdr3[c-1], 'v02': v2[:60], 'v03': v3[:60]})
    for c in AUDIT_COLS:
        v2, v3 = norm(m2.cell(ids2[rid], c).value), norm(m3.cell(ids3[rid], c).value)
        if v2 != v3:
            audit_diffs.append({'id': rid, 'col': hdr3[c-1], 'v02': v2[:80], 'v03': v3[:80]})
out['content_diffs'] = content_diffs
out['audit_diffs'] = audit_diffs
out['content_untouched'] = len(content_diffs) == 0

# ردیف‌های تازه = v0.3 منهای اسکلت
new_ids = sorted(set(ids3.keys()) - set(ids2.keys()))
out['new_count'] = len(new_ids)
out['new_ids'] = new_ids
new_rows = []
for rid in new_ids:
    r = ids3[rid]
    new_rows.append({
        'id': rid, 'page': norm(m3.cell(r, 2).value), 'chapter': norm(m3.cell(r, 3).value),
        'title': norm(m3.cell(r, 4).value)[:35], 'engine': norm(m3.cell(r, 10).value),
        'anchors': norm(m3.cell(r, 11).value), 'src': norm(m3.cell(r, 12).value)[:45],
        'status': norm(m3.cell(r, 13).value),
    })
out['all_new_yellow'] = all(x['status'] == 'نیاز به تأیید' for x in new_rows)
out['all_new_ai_src'] = all('پیشنهاد AI' in x['src'] for x in new_rows)
out['new_pages_seq'] = [x['page'] for x in new_rows]
eng = {}
for x in new_rows: eng[x['engine']] = eng.get(x['engine'], 0) + 1
out['new_engine'] = eng

# وضعیت و تاریخ اسکلت در v0.3 (ادعای مدل: سبز + تاریخ ثبت شده)
skel_status = {}
for rid in sorted(ids2.keys()):
    r = ids3[rid]
    skel_status[rid] = {'status': norm(m3.cell(r, 13).value), 'src': norm(m3.cell(r, 12).value)[:70]}
out['skeleton_status_v03'] = skel_status

print(json.dumps(out, ensure_ascii=False, indent=1))
