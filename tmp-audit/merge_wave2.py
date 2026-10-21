# -*- coding: utf-8 -*-
# ادغام موج ۲ — گام ۳: اعمال مجدد مُهر سبز رسمی (تسک ۴۴) روی v0.3
#  ۱) ۸ ردیف اسکلت: ستون «منبع» ← متن رسمی از v0.2 محلی (وضعیت تأییدشده در v0.3 درست است)
#  ۲) شناسنامه: ۹ فیلد کتاب‌محور سبز ← مقدار/وضعیت/منبع رسمی از v0.2
#     («وضعیت خشاب» عمداً دست نمی‌خورد — متن v0.3 به‌روز و زردِ مجاز است)
import openpyxl, json

V02 = '/home/z/my-project/public/math-game/authoring/khshab-G2-MATH_v0.2.xlsx'
V03_IN = '/tmp/khshab-G2-MATH_v0.3.xlsx'
V03_OUT = '/tmp/khshab-G2-MATH_v0.3-merged.xlsx'

wb2 = openpyxl.load_workbook(V02)
wb3 = openpyxl.load_workbook(V03_IN)

def norm(v): return '' if v is None else str(v).strip()

def sheet(wb, needle):
    return wb[next(n for n in wb.sheetnames if needle in n)]

def find_row(ws, prefix, limit=60):
    for r in range(1, limit):
        if norm(ws.cell(r, 1).value).startswith(prefix):
            return r
    return None

report = {'rows_restamped': [], 'identity_restamped': []}

# ── ۱) نقشهٔ کتاب: مُهر منبع رسمی روی ۸ ردیف اسکلت ──
m2, m3 = sheet(wb2, 'نقشه'), sheet(wb3, 'نقشه')
h2, h3 = find_row(m2, 'شناسه'), find_row(m3, 'شناسه')

def rows_by_id(ws, head):
    out = {}
    for r in range(head + 1, ws.max_row + 1):
        k = norm(ws.cell(r, 1).value)
        if not k or k == 'یادداشت': break
        out[k] = r
    return out

ids2, ids3 = rows_by_id(m2, h2), rows_by_id(m3, h3)
for rid in sorted(ids2.keys()):
    src_official = m2.cell(ids2[rid], 12).value  # ستون ۱۲ = منبع
    m3.cell(ids3[rid], 12).value = src_official
    report['rows_restamped'].append({'id': rid, 'status_v03': norm(m3.cell(ids3[rid], 13).value)})

# ── ۲) شناسنامه: ۹ فیلد سبز رسمی ──
i2, i3 = sheet(wb2, 'شناسنامه'), sheet(wb3, 'شناسنامه')
hi2, hi3 = find_row(i2, 'فیلد'), find_row(i3, 'فیلد')
GREEN_FIELDS = ['شناسهٔ خشاب', 'عنوان فارسی خشاب', 'درس', 'پایه', 'بستهٔ تجاری',
                'کد کتاب درسی', 'عنوان دقیق کتاب', 'معلم هم‌ساز', 'تاریخ آخرین تأیید سبز']
fields2, fields3 = {}, {}
for r in range(hi2 + 1, i2.max_row + 1):
    k = norm(i2.cell(r, 1).value)
    if not k: break
    fields2[k] = r
for r in range(hi3 + 1, i3.max_row + 1):
    k = norm(i3.cell(r, 1).value)
    if not k: break
    fields3[k] = r
for f in GREEN_FIELDS:
    if f not in fields2 or f not in fields3:
        report['identity_restamped'].append({'field': f, 'error': 'یافت نشد'}); continue
    for c in (2, 3, 4):  # مقدار، رنگ اعتبار، منبع
        i3.cell(fields3[f], c).value = i2.cell(fields2[f], c).value
    report['identity_restamped'].append({
        'field': f,
        'status': norm(i3.cell(fields3[f], 3).value),
        'value': norm(i3.cell(fields3[f], 2).value)[:50],
    })

wb3.save(V03_OUT)
report['saved_to'] = V03_OUT
print(json.dumps(report, ensure_ascii=False, indent=1))
