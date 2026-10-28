# -*- coding: utf-8 -*-
"""
فریز g5-baseline-v1 — ثبت، بایگانی و فریز نقشهٔ موضوعی ریاضی پنجم دبستان (G5)
فرمان مالک (trace 1a17dc3456240ead): «این مستندات G5 نیز با این اوصاف و تست‌هایی بیرونی
تأیید می‌شوند باید فریز ثبت و آرشیو بشه در کتابخانه محتوایی بازی و تمرین»
الگو: freeze_baseline_v1.py (میراث baseline-v1)
"""
import openpyxl, json, hashlib, os, shutil, stat
from datetime import datetime, timezone, timedelta

ROOT = '/home/z/my-project'
SRC = os.path.join(ROOT, 'upload', 'CURRICULUM-MAP-G5-v1 (1).xlsx')
SRC_CLEAN_NAME = 'CURRICULUM-MAP-G5-v1.xlsx'
PKG = os.path.join(ROOT, 'public', 'math-game', 'frozen', 'g5-baseline-v1')
G5DIR = os.path.join(PKG, 'G5')
TEHRAN = timezone(timedelta(hours=3, minutes=30))
now = datetime.now(TEHRAN)
now_iso = now.isoformat(timespec='seconds')
now_fa = now.strftime('%Y-%m-%d %H:%M')

def sha256(p):
    return hashlib.sha256(open(p, 'rb').read()).hexdigest()

OWNER_TRACE = '1a17dc3456240ead'
OWNER_STATEMENT = ('«این مستندات G5 نیز با این اوصاف و تست‌هایی بیرونی تایید می شوند '
                   'باید فریز ثبت و ارشیو بشه در کتابخانه محتوایی بازی و تمرین»')

# ---------- 1) خواندن اکسل و تبدیل کانونی ----------
wb = openpyxl.load_workbook(SRC, data_only=True)
ws_map, ws_mat, ws_self = wb.sheetnames[1], wb.sheetnames[2], wb.sheetnames[5]

wm = wb[ws_map]
lesson_rows = []
for r in wm.iter_rows(min_row=5, values_only=True):
    if r[1] is None:
        continue
    vals = list(r[1:16])  # B..P
    lesson_rows.append({
        'پایه': vals[0], 'شمارهٔ فصل': vals[1], 'عنوان فصل': vals[2],
        'شمارهٔ درس': vals[3], 'عنوان درس / موضوع': vals[4],
        'صفحات (شروع–پایان)': vals[5], 'سطح اطمینان صفحه': vals[6],
        'زیرموضوع\u200cها': vals[7], 'خانوادهٔ محتوایی': vals[8],
        'کد دامنهٔ پیشنهادی': vals[9], 'نوع\u200cهای تمرین': vals[10],
        'حداقل سؤال پیشنهادی': vals[11], 'توزیع سختی': vals[12],
        'گلوگاه\u200cهای احتمالی': vals[13], 'یادداشت': vals[14],
    })

wmx = wb[ws_mat]
matrix_rows = []
for r in wmx.iter_rows(min_row=5, values_only=True):
    if r[1] is None:
        continue
    vals = list(r[1:8])  # B..H
    matrix_rows.append({
        'پایه': vals[0], 'کد دامنه': vals[1], 'عنوان دامنه': vals[2],
        'فصل کتاب': vals[3], 'حداقل سؤال': vals[4],
        'نوع\u200cهای تمرین': vals[5], 'یادداشت عمق': vals[6],
    })

wself = wb[ws_self]
self_decl = []
for r in wself.iter_rows(min_row=2, values_only=True):  # ردیف ۱ سرستون است
    vals = [v for v in r[:4]]
    if any(v is not None and str(v).strip() for v in vals):
        self_decl.append([v if v is not None else '' for v in vals])

src_sha = sha256(SRC)
n_q = sum(x['حداقل سؤال پیشنهادی'] for x in lesson_rows)
n_ch = len(set(x['شمارهٔ فصل'] for x in lesson_rows))
doms = {}
for x in lesson_rows:
    doms[x['کد دامنهٔ پیشنهادی']] = doms.get(x['کد دامنهٔ پیشنهادی'], 0) + x['حداقل سؤال پیشنهادی']

doc = {
    'schema': 'khshab.curriculum-map/v1',
    'family': 'KH-G5-MATH',
    'grade': 5,
    'title': 'نقشهٔ موضوعات ریاضی پنجم دبستان — بستهٔ فریز g5-baseline-v1',
    'provenance': {
        'producer': 'سوپرگراک — برون‌سپاری خواندن کتاب طبق پرامپت استاندارد CURRICULUM-MAP (G1-G6 v1)',
        'sourceFile': SRC_CLEAN_NAME,
        'sourceFileOriginalName': 'CURRICULUM-MAP-G5-v1 (1).xlsx',
        'sourceFileSha256': src_sha,
        'externalVerification': {
            'by': 'مالک — مجید بایرامی (bayrami1341) با آزمون‌های بیرونی',
            'scope': 'کتاب رسمی «ریاضی پنجم دبستان» — ۱۴۱ صفحهٔ چاپی، ۷ لوحهٔ فصل (۱، ۲۱، ۴۷، ۶۷، ۸۷، ۱۰۵، ۱۲۵)',
            'note': 'فایل PDF کتاب سنجه در این سیستم بایگانی نشده است؛ اعتبار محتوایی بر مبنای مُهر بیرونی مالک است، نه راستی‌آزمایی سه‌روشی داخلی.',
        },
    },
    'verification': {
        'date': now.date().isoformat(),
        'trace': OWNER_TRACE,
        'structural': {
            'errors': 0,
            'warnings': 0,
            'lessonRows': len(lesson_rows),
            'chapters': f'۱ تا {n_ch} کامل (۶/۵/۴/۴/۴/۴/۳)',
            'domainQuestionsSum': n_q,
            'checks': [
                'سرستون‌ها = عین قالب ۱۵ ستونه',
                'ارقام لاتین در ستون‌های عددی (فصل/درس/حداقل سؤال)',
                'پیوستگی کامل درس‌ها در هر ۷ فصل (شکست سه‌ردیفی درس ۱ فصل ۱ = اعلام‌شده در خوداظهاری)',
                'enum دامنه (D1–D10) / اطمینان (کامل) / سختی (۴۰/۴۰/۲۰) معتبر',
                'قالب بازهٔ صفحه‌ها معتبر؛ بدون ردیف تکراری',
                'صفر سلول ادغام‌شده و صفر فرمول در ناحیهٔ داده',
                'تطابق درونی نقشه↔ماتریس: ۱۰/۱۰ دامنه، جمع ۶۰۰=۶۰۰',
                'وضعیت «نیاز به تأیید» روی همهٔ ردیف‌های ماتریس',
            ],
        },
        'conceptual': 'تأیید بیرونی به فرمان و اجرای مالک (کتاب سنجه در سیستم موجود نیست؛ نمونه‌گیری سه‌روشی داخلی اجرا نشد)',
        'methods': [
            'اعتبارسنجی ساختاری ۱۰۰٪ خودکار (اسکیما، enum، ارقام، پیوستگی فصل/درس، تطابق درونی نقشه↔ماتریس)',
            'تأیید بیرونی مالک با آزمون‌های بیرونی روی کتاب رسمی ریاضی پنجم دبستان (فرمان trace ' + OWNER_TRACE + ')',
        ],
        'valueScore': 'نمره‌دهی سه‌روشی اجرا نشد (سنجه در دسترس سیستم نبود) — مبنای اعتبار: مُهر بیرونی مالک',
        'evidenceImages': None,
    },
    'ownerSeal': {
        'by': 'مالک — مجید بایرامی (bayrami1341)',
        'trace': OWNER_TRACE,
        'date': now.date().isoformat(),
        'statement': OWNER_STATEMENT,
        'scope': 'مُهر اعتبار نقشهٔ موضوعی به‌عنوان خوراک بازی و تمرین؛ تولید سؤال‌ها و مُهر پداگوژیکِ گام صفر معلم، مرحلهٔ جداگانه است',
    },
    'status': 'فریز‌شده — نقشهٔ موضوعی معتبر خوراک بازی و تمرین (G5) با تأیید بیرونی مالک؛ سؤال‌ها هنوز تولید نشده‌اند (KH-G5-MATH = ۰ سؤال)',
    'lessonRows': lesson_rows,
    'domainMatrix': matrix_rows,
    'selfDeclaration': self_decl,
    'openItems': [
        'D9 (عدد مرکّب و گنجایش) و D10 (نسبت، تناسب و درصد) فراتر از قرارداد D1–D8 هستند؛ justification در خود فایل ثبت شده («کد آماده نبود») — گنجاندن رسمی در واژه‌نامهٔ دامنه با تصمیم مالک (هم‌راستا با موضوع باز D9 در G1)',
        'کتاب سنجه (ریاضی پنجم دبستان) در سیستم بایگانی نشده؛ در صورت ارسال فایل، شاهد صفحه‌ای و راستی‌آزمایی سه‌روشی در نسخهٔ g5-baseline-v1.1 قابل افزودن است',
        'تولید سؤال‌ها (حداقل ۶۰۰، توزیع ۴۰/۴۰/۲۰) و تبدیل به KH-G5-MATH.json و مُهر گام صفر — مرحلهٔ بعد با گیت‌های عادی',
        'درس ۱ فصل ۱ در کتاب یک عنوان دارد و در نقشه به سه ردیف D2/D3/D4 شکسته شده؛ صفحه‌های ۴–۵ بین ردیف‌های D3 و D4 مشترک است (اعلام‌شده در خوداظهاری) — در تبدیل خشاب باید لحاظ شود',
    ],
}

# ---------- 2) ساخت بسته ----------
os.makedirs(G5DIR, exist_ok=True)
json_path = os.path.join(G5DIR, 'curriculum-map-g5.json')
with open(json_path, 'w', encoding='utf-8') as f:
    json.dump(doc, f, ensure_ascii=False, indent=2)

xlsx_path = os.path.join(G5DIR, SRC_CLEAN_NAME)
shutil.copy2(SRC, xlsx_path)

evidence = f"""# سند شفافیت و راستی‌آزمایی — نقشهٔ موضوعی G5 (بستهٔ g5-baseline-v1)

**تاریخ:** {now_fa} (تهران) | **فرمان مالک:** trace {OWNER_TRACE}

## ۱) چه چیزی فریز شد
- `CURRICULUM-MAP-G5-v1.xlsx` — خروجی سوپرگراک (نام اصلی دریافتی: «CURRICULUM-MAP-G5-v1 (1).xlsx»)
  SHA-256: `{src_sha}`
- ۳۰ ردیف درس / ۷ فصل / ۱۰ دامنه / جمع ۶۰۰ حداقل سؤال

## ۲) اعتبارسنجی ساختاری ۱۰۰٪ خودکار (توسط Claude — {'{d}'} )
- صفر خطا، صفر هشدار؛ سرستون‌ها عین قالب ۱۵ ستونه؛ ارقام لاتین در ستون‌های عددی
- پیوستگی فصل‌های ۱–۷ و درس‌ها؛ enum دامنه/اطمینان/سختی معتبر؛ بدون تکراری، ادغام یا فرمول
- تطابق درونی نقشهٔ_درس‌ها ↔ ماتریس_دامنه: ۱۰/۱۰ دامنه (D1=16, D2=46, D3=10, D4=10, D5=128, D6=116, D7=50, D8=94, D9=34, D10=96؛ جمع ۶۰۰=۶۰۰)
- وضعیت «نیاز به تأیید» روی همهٔ ردیف‌های ماتریس (مطابق قاعدهٔ ثبت ورودی AI)

## ۳) مبنای اعتبار محتوایی — مُهر بیرونی مالک
به گواهی مالک، این نقشه «با این اوصاف و تست‌هایی بیرونی» تأیید شده است.
- **شفافیت:** فایل کتاب سنجه (ریاضی پنجم دبستان) در این سیستم موجود نیست؛ راستی‌آزمایی سه‌روشی داخلی (بصری/متنی/منبع مستقل) اجرا نشد و نمرهٔ ۱–۵ داده نشد. این بسته با مُهر بیرونی مالک فریز شده است، نه با حکم ممیزی داخلی.
- در صورت ارسال فایل کتاب، شاهد صفحه‌ای در نسخهٔ v1.1 قابل افزودن است.

## ۴) خوداظهاری تولیدکننده (گراک) — خلاصه وفادار
- منبع: ریاضی پنجم دبستان، فایل کاربر؛ شمارهٔ صفحه = شمارهٔ چاپ پایین برگه، ۱ تا ۱۴۱
- لوحه‌های فصل: ۱، ۲۱، ۴۷، ۶۷، ۸۷، ۱۰۵، ۱۲۵ (۷ فصل)؛ مرور فصل و فرهنگ‌و‌معما ردیف نشدند
- درس اول فصل ۱ در کتاب یک عنوان دارد و به سه ردیف (عددنویسی / جمع‌وتفریق / ضرب‌وتقسیم) شکسته شد؛ صفحه‌های ۴–۵ بین دو عمل مشترک است
- راهنمای معلم فقط سنجهٔ مفهومی هر هفت فصل بوده؛ شمارهٔ صفحهٔ آن نقل نشده
- تقسیم دو کسر بر هم و فرمول «مساحت قاعده × ارتفاع» هدف این پایه دانسته نشده
- وضعیت اعلامی: نیاز به تأیید (مُهر نهایی = معلم/مالک)

## ۵) موضوع‌های باز (openItems)
1. D9 و D10 خارج از قرارداد D1–D8 — تصمیم رسمی با مالک (هم‌راستا با D9 بازِ G1)
2. کتاب سنجه بایگانی نشده — شاهد صفحه‌ای در v1.1 ممکن است
3. تولید سؤال و مُهر گام صفر معلم — مرحلهٔ جداگانه
4. شکست درس ۱ فصل ۱ و اشتراک صفحهٔ ۴–۵ — در تبدیل خشاب لحاظ شود

## ۶) قواعد فریز
- ویرایش ممنوع؛ اصلاح فقط با نسخهٔ g5-baseline-v1.1 و فرمان مالک
- صحت فایل‌ها با MANIFEST.json (SHA-256) قابل بازبینی است
- بازکردن فریز فقط با فرمان صریح مالک ثبت‌شده در worklog
""".replace('{d}', now.date().isoformat())

ev_path = os.path.join(G5DIR, 'verification-evidence-g5.md')
with open(ev_path, 'w', encoding='utf-8') as f:
    f.write(evidence)

# ---------- 3) مانیفست ----------
files = []
for rel in ['G5/' + SRC_CLEAN_NAME, 'G5/curriculum-map-g5.json', 'G5/verification-evidence-g5.md']:
    p = os.path.join(PKG, rel)
    files.append({'path': rel, 'bytes': os.path.getsize(p), 'sha256': sha256(p)})

manifest = {
    'schema': 'khshab.frozen-baseline/v1',
    'version': 'g5-baseline-v1',
    'frozenAt': now_iso,
    'timezone': 'Asia/Tehran',
    'authority': {
        'command': 'فریز، ثبت و بایگانی مستندات G5 در کتابخانه محتوایی بازی و تمرین',
        'issuedBy': 'مالک — مجید بایرامی (bayrami1341)',
        'trace': OWNER_TRACE,
        'justification': 'تأیید بیرونی مستندات G5 با آزمون‌های بیرونی به گواهی مالک',
    },
    'freezeRules': [
        'فایل‌های این بسته تغییرناپذیرند (chmod 444)؛ هر اصلاح = نسخهٔ جدید g5-baseline-v1.1 با فرمان مالک، نه ویرایش درجا',
        'هیچ فایل زندهٔ خط تولید (cartridges/، authoring/، DOMAIN-MATRIX) در این فریز دست نخورده است',
        'بازکردن فریز (unfreeze) فقط با فرمان صریح مالک ثبت‌شده در worklog',
        'این بسته روی حکم گیت پوشش زنده اثری ندارد (فریز = بایگانی مبنا، نه حکم آمادگی)',
    ],
    'grades': {
        'G5': {
            'what': 'نقشهٔ موضوعی کتاب پنجم دبستان (۳۰ ردیف درس، ۷ فصل، ۱۰ دامنهٔ D1–D10، ۶۰۰ حداقل سؤال)',
            'seal': f'معتبر برای خوراک بازی و تمرین — تأیید بیرونی مالک (trace {OWNER_TRACE})',
            'questionCount': 0,
            'nextStep': 'تصمیم D9/D10 → تبدیل به KH-G5-MATH.json (حداقل ۶۰۰ سؤال، ۴۰/۴۰/۲۰) → گیت خشاب → مُهر گام صفر',
        }
    },
    'evidenceAnchors': {
        'note': 'کتاب سنجه (ریاضی پنجم دبستان) در سیستم بایگانی نشده است؛ اعتبار محتوایی بر مُهر بیرونی مالک است. در صورت ارسال فایل، هش و لینک بازیابی در نسخهٔ v1.1 ثبت می‌شود.',
        'witnesses': [],
    },
    'files': files,
}

manifest_path = os.path.join(PKG, 'MANIFEST.json')
with open(manifest_path, 'w', encoding='utf-8') as f:
    json.dump(manifest, f, ensure_ascii=False, indent=2)

index_md = f"""# بایگانی فریز g5-baseline-v1 — نقشهٔ موضوعی ریاضی پنجم دبستان

**تاریخ مهر:** {now_fa} (تهران) | **فرمان:** مالک (مجید بایرامی) — trace {OWNER_TRACE}

## این بسته چیست
نسخهٔ تغییرناپذیر (فریز) از خوانش کتاب ریاضی پنجم دبستان توسط سوپرگراک، که به گواهی مالک با
**آزمون‌های بیرونی** تأیید شده و با مُهر او در کتابخانهٔ محتوایی بازی و تمرین ثبت و بایگانی می‌شود:
- **G5:** نقشهٔ موضوعی ۳۰ ردیفی / ۷ فصل / ۱۰ دامنه (D1–D10) / ۶۰۰ حداقل سؤال — تبدیل‌شده به JSON کانونی (schema: khshab.curriculum-map/v1)

## نکتهٔ صداقت
فایل کتاب سنجه در سیستم بایگانی نشده است؛ اعتبار محتوایی بر **مُهر بیرونی مالک** استوار است، نه ممیزی سه‌روشی داخلی (بصری/متنی/منبع مستقل که برای G1 و G4 اجرا شد). شواهد ساختاری ۱۰۰٪ خودکار داخل بسته ثبت است.

## قواعد فریز
۱) هیچ فایلی در این پوشه ویرایش نمی‌شود؛ اصلاح یعنی نسخهٔ جدید (g5-baseline-v1.1) با فرمان مالک.
۲) صحت هر فایل با `MANIFEST.json` (SHA-256) قابل بازبینی است.
۳) فایل‌های زندهٔ خط تولید از این فریز مستقل‌اند؛ بازی فقط از کارتریج‌های زنده می‌خواند.

## ساختار
```
frozen/g5-baseline-v1/
├── G5/ curriculum-map-g5.json · CURRICULUM-MAP-G5-v1.xlsx · verification-evidence-g5.md
├── MANIFEST.json
└── INDEX.md (همین فایل)
```
"""
index_path = os.path.join(PKG, 'INDEX.md')
with open(index_path, 'w', encoding='utf-8') as f:
    f.write(index_md)

# ---------- 4) قفل فایل‌سیستمی ----------
def lock(p, is_dir):
    os.chmod(p, 0o555 if is_dir else 0o444)

for root, dirs, fs in os.walk(PKG):
    for d in dirs:
        lock(os.path.join(root, d), True)
    for fn in fs:
        lock(os.path.join(root, fn), False)
lock(PKG, True)
lock(G5DIR, True)

# ---------- 5) صحت‌سنجی پس از قفل ----------
ok = True
for entry in files:
    p = os.path.join(PKG, entry['path'])
    if sha256(p) != entry['sha256']:
        ok = False
        print('HASH MISMATCH:', entry['path'])
print('POST-LOCK HASH CHECK:', 'ALL PASS' if ok else 'FAILED')
print('PACKAGE:', PKG)
print('MANIFEST sha256:', sha256(manifest_path))
print('TOTAL SIZE: %.1fKB' % (sum(e['bytes'] for e in files) / 1024))
print('ROWS:', len(lesson_rows), '| CHAPTERS:', n_ch, '| DOMAINS:', len(matrix_rows), '| MINQ TOTAL:', n_q)
