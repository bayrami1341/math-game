# -*- coding: utf-8 -*-
"""چک ماشینی ۲۸ سؤال تازهٔ موج ۳ + هم‌گامی لنگر"""
import openpyxl, re, json
from collections import Counter
wb = openpyxl.load_workbook('public/math-game/authoring/khshab-G2-MATH_v0.4.xlsx', data_only=True)

FA = '۰۱۲۳۴۵۶۷۸۹'
def fa2en(s): return str(s).translate(str.maketrans(FA, '0123456789'))

qs, atlas = {}, set()
ws = wb['اطلس خطا']
for r in range(2, ws.max_row+1):
    k = ws.cell(r,1).value
    if not k or str(k).strip()=='یادداشت': break
    code = str(k).strip().replace('(پیشنهاد)','').strip()
    atlas.add(code)
print('atlas codes:', sorted(atlas))

ws = wb['بانک سؤال']
hdr = None; new = []
for r in range(2, ws.max_row+1):
    k = ws.cell(r,1).value
    if k is None: continue
    if str(k).strip()=='یادداشت': break
    row = [ws.cell(r,c).value for c in range(1, 18)]
    qid = str(row[0]).strip()
    if re.match(r'^Q-10(4[1-9]|5\d|6[0-8])$', qid): new.append(row)

print('new questions:', len(new))
errs, warns = [], []
answers, cpas = Counter(), Counter()
ids = set()
for row in new:
    qid, anchor, scope, cpa, text, struct, a, b, c, d, ans, emap, level, crit, src, color, note = [str(x).strip() if x is not None else '' for x in row]
    ids.add(qid)
    # 1) scope/CPA
    if scope not in ('CORE','DEPTH','SIDE'): errs.append(f'{qid}: دامنه={scope}')
    cpas[cpa]+=1
    if cpa not in ('مجسم','تصویری','علامتی'): errs.append(f'{qid}: CPA={cpa}')
    # 2) answer validity + balance
    answers[ans]+=1
    if ans not in ('الف','ب','ج','د'): errs.append(f'{qid}: پاسخ={ans}')
    # 3) options unique + numeric range
    vals = [a,b,c,d]
    if len(set(vals))<4: errs.append(f'{qid}: گزینه تکراری {vals}')
    for v in vals:
        n = fa2en(re.sub(r'[^0-9۰-۹]', '', v))
        if n and not (0 <= int(n) <= 9999): errs.append(f'{qid}: گزینه خارج دامنه {v}')
    # 4) error map: 3 wrong options → codes in atlas
    refs = re.findall(r'[الفبجد]→\s*(E_[A-Z_]+)', emap)
    if len(refs) != 3: errs.append(f'{qid}: نقشه خطا {len(refs)} سرشاخه {emap!r}')
    for code in refs:
        if code not in atlas: errs.append(f'{qid}: کد {code} بیرون از اطلس')
    # 5) color/source
    if color != 'نیاز به تأیید': errs.append(f'{qid}: رنگ={color}')
    if 'موتور تألیف' not in src: warns.append(f'{qid}: منبع={src!r}')
    # 6) anti-J12 note has latin digits/calc
    if not re.search(r'[0-9]', note): warns.append(f'{qid}: یادداشت بی‌حساب {note!r}')
    # 7) Persian digits in question text (no latin digits)
    if re.search(r'[0-9]', text): warns.append(f'{qid}: رقم لاتین در صورت سؤال {text!r}')
    # 8) M_ model mention for visual CPA
    if cpa in ('مجسم','تصویری') and 'M_' not in (struct+text): warns.append(f'{qid}: CPA تصویری بدون مدل M_')

print('answer balance:', dict(answers))
print('CPA dist:', dict(cpas))
print('unique ids:', len(ids)==len(new))
print('ERRORS:', len(errs)); [print('  ✗', e) for e in errs[:20]]
print('WARNS:', len(warns)); [print('  ⚠', w) for w in warns[:20]]

# 9) anchor sync: L-2xx rows' Q-refs == bank ownership
s1 = wb['نقشهٔ کتاب (S1)']
lmap = {}
for r in range(2, s1.max_row+1):
    k = s1.cell(r,1).value
    if not k or str(k).strip()=='یادداشت': break
    kid = str(k).strip()
    if re.match(r'^L-2(0[1-9]|1[0-8])$', kid):
        lmap[kid] = str(s1.cell(r,11).value or '').strip()
anchor_err = []
owner = {}
for row in new:
    qid, anchor = str(row[0]).strip(), str(row[1]).strip()
    owner.setdefault(anchor, []).append(qid)
for L, refs in lmap.items():
    got = set()
    for m in re.finditer(r'Q-\d+', refs): got.add(m.group())
    # expand ranges like Q-1041-1044
    for m in re.finditer(r'Q-(\d+)-(\d+)', refs):
        got.discard(f'Q-{m.group(1)}-{m.group(2)}')
        got |= {f'Q-{i}' for i in range(int(m.group(1)), int(m.group(2))+1)}
    have = set(owner.get(L, []))
    if got != have: anchor_err.append((L, sorted(got-have), sorted(have-got)))
print('anchor mismatches:', len(anchor_err)); [print('  ✗', a) for a in anchor_err]
# pages per L
pages = {L: str(s1.cell(find_row, 2).value or '') for L, find_row in []}
