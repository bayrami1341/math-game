#!/usr/bin/env python3
"""ادغام رسمی v0.3.1: فایل مدل + بازگردانی مُهرهای رسمی معلم + دو اصلاح بازرس دوم
خروجی: public/math-game/authoring/khshab-G2-MATH_v0.3.1.xlsx
"""
import openpyxl, shutil, sys

SRC = "/tmp/v031.xlsx"           # تحویل مدل (SHA: bb756156…)
V03 = "/home/z/my-project/public/math-game/authoring/khshab-G2-MATH_v0.3.xlsx"  # نسخهٔ رسمی ادغامی قبلی
OUT = "/home/z/my-project/public/math-game/authoring/khshab-G2-MATH_v0.3.1.xlsx"

shutil.copy(SRC, OUT)
wb = openpyxl.load_workbook(OUT, data_only=False)
wb_old = openpyxl.load_workbook(V03, data_only=False)

changes = []
def setcell(ws, coord, val, why):
    old = ws[coord].value
    ws[coord] = val
    changes.append((ws.title, coord, old, val, why))

# ---------- ۱) بازگردانی مُهرهای رسمی (پروتکل تسک ۴۴/۴۸) ----------
wsi, wsi_old = wb["شناسنامهٔ خشاب"], wb_old["شناسنامهٔ خشاب"]
for r in range(3, 8):  # ردیف‌های ۳ تا ۷: شناسه/عنوان/درس/پایه/بسته
    setcell(wsi, f"C{r}", wsi_old[f"C{r}"].value, "بازگردانی مُهر رسمی")
    setcell(wsi, f"D{r}", wsi_old[f"D{r}"].value, "بازگردانی منبع رسمی معلم")
setcell(wsi, "B14", wsi_old["B14"].value, "بازگردانی تاریخ رسمی آخرین تأیید سبز")

wsm, wsm_old = wb["نقشهٔ کتاب (S1)"], wb_old["نقشهٔ کتاب (S1)"]
skeleton = {"L-100", "L-200", "L-300", "L-400", "L-500", "L-600", "L-700", "L-800"}
old_map = {}
for row in wsm_old.iter_rows():
    if row[0].value in skeleton:
        for c in row:
            if c.coordinate.startswith("L") and len(c.coordinate) > 1 and c.coordinate[1:].isdigit():
                old_map[row[0].value] = c.coordinate
for row in wsm.iter_rows():
    if row[0].value in skeleton:
        coord = old_map[row[0].value]
        setcell(wsm, coord, wsm_old[coord].value, f"بازگردانی مُهر رسمی ردیف اسکلت {row[0].value}")

# ---------- ۲) اصلاح L-114 (بند ۲ پیام اصلاحی — کتاب ص ۱۴: شمارش ۴تا۴تا) ----------
for row in wsm.iter_rows():
    if row[0].value == "L-114":
        rr = row[0].row
        setcell(wsm, f"D{rr}", "درس ۵: شمارش ۴تا۴تا (فعّالیت)", "بند ۲ اصلاحیه — بنر/موضوع واقعی صفحهٔ ۱۴")
        setcell(wsm, f"E{rr}", "دانش‌آموز بتواند با شمارش ۴تا۴تا، الگوی رنگی مربع‌ها را ادامه دهد و شکل شابلونی بسازد.", "بند ۲ اصلاحیه")
        setcell(wsm, f"F{rr}", "شمارش، مربّع، رنگ، دسته، خانه، شابلون (فقط واژه‌های صفحهٔ ۱۴)", "واژگان واقعی صفحهٔ ۱۴")
        setcell(wsm, f"G{rr}", "فعّالیت: ادامهٔ الگوی شمارش ۴تا۴تا با رنگ + کار در کلاس (دسته‌های ۵تایی) + خلّق باش", "محتوای واقعی صفحهٔ ۱۴")
        setcell(wsm, f"N{rr}", "شاهد: متن PDF (صفحهٔ ۲۰ فایل C205) و عکس معلّم — بنر «فعّالیت»؛ ادامهٔ الگوی ۴، ۸، … با رنگ‌آمیزی مربّع‌ها؛ کار در کلاس: شمارش خانه‌ها با دسته‌های ۵تایی؛ خلّق باش: ساخت شکل با شابلون ۴ مربّع. (اصلاح بازرس دوم — بند ۲ پیام v0.3.1 که مدل اعمال نکرده بود)", "شاهد بازرس دوم")

# ---------- ۳) اصلاح Q-1011 + شاهد L-108 (کتاب ۱۳ مثلث دارد؛ پنجمین عدد بعد = ۱۸) ----------
wsb = wb["بانک سؤال"]
for row in wsb.iter_rows():
    if row[0].value == "Q-1011":
        rr = row[0].row
        setcell(wsb, f"E{rr}", "شکل زیر ۱۳ مثلث دارد. پنجمین عدد بعد از آن چند است؟", "اصلاح بازرس دوم — شمارش ماشینی: ۱۳ مثلث")
        setcell(wsb, f"G{rr}", "۱۴", "مخفّی: فقط یکی بعد (۱۳+۱)")
        setcell(wsb, f"H{rr}", "۱۶", "مخفّی: کم‌شمردن قدم‌ها (۱۳+۳)")
        setcell(wsb, f"I{rr}", "۱۸", "پاسخ درست (۱۳+۵) — جای پاسخ (ج) ثابت می‌ماند برای توازن")
        setcell(wsb, f"J{rr}", "۸", "مخفّی: تفریق به‌جای ادامه (۱۳−۵)")
        setcell(wsb, f"N{rr}", "پاسخ دقیق: ۱۸ (۱۳+۵)؛ ۱۴ فقط یکی بعد، ۱۶ کم‌شمردن قدم‌ها، ۸ تفریق به‌جای ادامه. (اصلاح بازرس دوم: کتاب ص ۸، سؤال ۳ = ۱۳ مثلث — شمارش ماشینی تصویر؛ مدل ۱۶ نوشته بود)", "پاسخ‌نامهٔ اصلاحی")
        setcell(wsb, f"Q{rr}", "ضد J12 — کار دفترچه‌ای (لاتین): 13+5=18؛ پاسخ جدا حساب شد.", "ضد J12 هم‌گام")
for row in wsm.iter_rows():
    if row[0].value == "L-108":
        rr = row[0].row
        setcell(wsm, f"N{rr}", wsm[f"N{rr}"].value.replace("شکل مثلث‌ها ۱۶ تا، پنجمین عدد بعد = ۲۱", "شکل مثلث‌ها ۱۳ تا، پنجمین عدد بعد = ۱۸ (اصلاح بازرس دوم — شمارش ماشینی تصویر)"), "شاهد هم‌گام با اصلاح Q-1011")

wb.save(OUT)
print(f"ذخیره شد: {OUT}")
print(f"\nتغییرات اعمال‌شده ({len(changes)} سلول):")
for sheet, coord, old, new, why in changes:
    so = (str(old)[:60] + "…") if old and len(str(old)) > 60 else old
    sn = (str(new)[:60] + "…") if new and len(str(new)) > 60 else new
    print(f"  [{sheet}] {coord}: {so!r} → {sn!r}  ({why})")
