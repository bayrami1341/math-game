#!/usr/bin/env python3
"""دیف v0.3.1 ↔ v0.3 — فقط تغییرات مجاز؟"""
import openpyxl, sys

V03 = "/home/z/my-project/public/math-game/authoring/khshab-G2-MATH_v0.3.xlsx"
V031 = "/tmp/v031.xlsx"

wb1 = openpyxl.load_workbook(V03, data_only=False)
wb2 = openpyxl.load_workbook(V031, data_only=False)

print("sheets v0.3 :", wb1.sheetnames)
print("sheets v0.3.1:", wb2.sheetnames)

def rows_of(ws):
    out = {}
    for row in ws.iter_rows():
        key = None
        vals = {}
        for c in row:
            v = c.value
            if v is None: continue
            vals[c.coordinate] = v
            if isinstance(v, str) and (v.startswith("L-") or v.startswith("Q-") or v.startswith("E-")):
                if key is None or len(v) < len(key): key = v
        out[key or f"row{row[0].row}"] = vals
    return out

diffs = {}
if wb1.sheetnames == wb2.sheetnames:
    for name in wb1.sheetnames:
        ws1, ws2 = wb1[name], wb2[name]
        r1, r2 = rows_of(ws1), rows_of(ws2)
        keys = set(r1) | set(r2)
        for k in sorted(keys):
            a, b = r1.get(k, {}), r2.get(k, {})
            if a != b:
                cells = {}
                for coord in sorted(set(a) | set(b)):
                    if a.get(coord) != b.get(coord):
                        cells[coord] = (a.get(coord), b.get(coord))
                diffs.setdefault(name, {})[k] = cells
else:
    print("!! شیت‌ها فرق دارند"); sys.exit(1)

print("\n=== خلاصهٔ دیف ===")
total = 0
for name, rows in diffs.items():
    print(f"\n[{name}] {len(rows)} ردیف تغییر کرده:")
    for k, cells in rows.items():
        total += 1
        print(f"  --- {k} ({len(cells)} سلول):")
        for coord, (a, b) in cells.items():
            sa = (str(a)[:90] + "…") if a and len(str(a)) > 90 else a
            sb = (str(b)[:90] + "…") if b and len(str(b)) > 90 else b
            print(f"      {coord}: OLD={sa!r}")
            print(f"            NEW={sb!r}")
print(f"\nمجموع ردیف‌های تغییر کرده: {total}")
