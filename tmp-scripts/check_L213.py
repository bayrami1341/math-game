# -*- coding: utf-8 -*-
import openpyxl
wb = openpyxl.load_workbook('public/math-game/authoring/khshab-G2-MATH_v0.4.xlsx', data_only=True)
ws = wb['نقشهٔ کتاب (S1)']
for r in range(2, ws.max_row+1):
    k = ws.cell(r,1).value
    if k and str(k).strip() in ('L-213','L-211','L-202'):
        print('==', k, '| صفحه:', ws.cell(r,2).value)
        for c,label in [(4,'درس'),(5,'اهداف'),(7,'نوع فعالیت'),(14,'یادداشت')]:
            print(f'  {label}: {str(ws.cell(r,c).value)[:220]}')
        print()
qb = wb['بانک سؤال']
print('== سؤال‌های L-213 ==')
for r in range(2, qb.max_row+1):
    k = qb.cell(r,1).value
    if not k or str(k).strip()=='یادداشت': break
    if str(qb.cell(r,2).value).strip()=='L-213':
        print(str(k).strip(), '|', str(qb.cell(r,5).value)[:120], '| پاسخ:', qb.cell(r,11).value, '| حساب:', str(qb.cell(r,17).value)[:80])
