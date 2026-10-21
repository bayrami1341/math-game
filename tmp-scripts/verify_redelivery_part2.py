#!/usr/bin/env python3
"""Task 54 — چک‌های تکمیلی: شناسنامه↔v0.3، ردیف شبح CPA، ستون پاسخ، ضد J12 دقیق، لنگرها، متن‌های کامل"""
import openpyxl, re

V03   = "/home/z/my-project/public/math-game/authoring/khshab-G2-MATH_v0.3.xlsx"
REDEL = "/tmp/gate2-redelivery/khshab-G2-MATH_v0.3.1.xlsx"

def rows_by_id(ws, prefix):
    out = {}
    for row in ws.iter_rows():
        aid = row[0].value
        if isinstance(aid, str) and aid.startswith(prefix):
            out[aid] = {c.coordinate[0]: ("" if c.value is None else str(c.value)) for c in row}
    return out

w3 = openpyxl.load_workbook(V03); wr = openpyxl.load_workbook(REDEL)

# ---------- شناسنامه ↔ v0.3 ----------
print("== شناسنامهٔ خشاب: تحویل۲ ↔ v0.3 ==")
a = [[("" if c.value is None else str(c.value)) for c in r] for r in w3["شناسنامهٔ خشاب"].iter_rows()]
b = [[("" if c.value is None else str(c.value)) for c in r] for r in wr["شناسنامهٔ خشاب"].iter_rows()]
for i, (ra, rb) in enumerate(zip(a, b), 1):
    for j, (va, vb) in enumerate(zip(ra, rb)):
        if va != vb:
            print(f"  ردیف{i} ستون{chr(65+j)}: v0.3={va[:60]!r} ↔ تحویل۲={vb[:60]!r}")

# ---------- بانک سؤال ----------
q3, qr = rows_by_id(w3["بانک سؤال"], "Q-"), rows_by_id(wr["بانک سؤال"], "Q-")
print("\n== شناسه‌های بانک تحویل۲ ==", sorted(qr))
print("v0.3:", sorted(q3))

# ---------- CPA: توزیع v0.3 و تحویل۲ ----------
def cpa_dist(q):
    d = {}
    for qid, row in q.items():
        d[row.get("D", "")] = d.get(row.get("D", ""), 0) + 1
    return d
print("\nCPA v0.3:", cpa_dist(q3))
print("CPA تحویل۲:", cpa_dist(qr))

# ردیف‌های CPA تغییرکرده در سؤال‌های قدیمی — با مقدار قدیم/جدید
print("\n-- تغییر CPA سؤال‌های قدیمی (خارج Q-1009..1024) --")
for qid in sorted(q3):
    if qid in qr and qid not in {f"Q-{n}" for n in range(1009, 1025)}:
        v3, vr = q3[qid].get("D", ""), qr[qid].get("D", "")
        if v3 != vr:
            other = {k for k in set(q3[qid]) | set(qr[qid]) if q3[qid].get(k) != qr[qid].get(k)} - {"D"}
            print(f"  {qid}: {v3!r} → {vr!r}" + (f"  !! سایر سلول‌ها: {sorted(other)}" if other else ""))

# ---------- لنگر (ستون B) همهٔ ۴۰ سؤال ----------
print("\n-- لنگر ستون B --")
for qid in sorted(qr):
    print(f"  {qid}: v0.3={q3.get(qid,{}).get('B','—')} → تحویل۲={qr[qid].get('B','—')}")

# ---------- ستون پاسخ: کدام ستون؟ ----------
hdr = [c.value for c in wr["بانک سؤال"][3]]
print("\nسرستون‌ها (ردیف ۳):", hdr)
# فرض: پاسخ در یکی از ستون‌ها با حرف
for k in "ABCDEFGHIJKLMNOPQ":
    vals = {qr[q].get(k, "").strip() for q in qr}
    if vals and vals <= {"الف", "ب", "پ", "د"}:
        bal = {}
        for q in qr: bal[qr[q][k]] = bal.get(qr[q][k], 0) + 1
        print(f"ستون پاسخ: {k} — توزیع تحویل۲: {bal}")
        bal3 = {}
        for q in q3: 
            v = q3[q].get(k, "").strip()
            if v: bal3[v] = bal3.get(v, 0) + 1
        print(f"   توزیع v0.3: {bal3}")
        # جای پاسخ ۱۶ سؤال: تحویل۲ ↔ v0.3
        same = all(q3.get(f"Q-{n}", {}).get(k) == qr.get(f"Q-{n}", {}).get(k) for n in range(1009, 1025))
        print(f"   جای پاسخ Q-1009..1024 مثل v0.3؟ {same}")

# ---------- ضد J12 دقیق: رقم لاتین خارج از ستون یادداشت ----------
print("\n-- رقم لاتین در ستون‌های غیرِ یادداشت (تحویل۲، Q-1009..1024) --")
hits = []
for n in range(1009, 1025):
    qid = f"Q-{n}"
    for k, v in qr[qid].items():
        if k in ("A", "Q"): continue
        if re.search(r"[0-9]", v): hits.append((qid, k, v[:60]))
print(hits if hits else "هیچ ✓ (ارقام لاتین فقط در یادداشت)")

# ---------- متن‌های کامل موارد مناقشه ----------
print("\n== متن کامل سلول‌های مناقشه (تحویل۲) ==")
m = rows_by_id(wr["نقشهٔ کتاب (S1)"], "L-")
for lid, col in (("L-108", "N"), ("L-109", "G"), ("L-109", "H"), ("L-112", "G"), ("L-112", "N"), ("L-118", "D")):
    print(f"  {lid}!{col}: {m.get(lid, {}).get(col, '')}")
print("\n== Q-1011 کامل (تحویل۲) ==")
print("  ", qr["Q-1011"])
print("== Q-1020 کامل (تحویل۲) ==")
print("  ", qr["Q-1020"])
print("== Q-1009 کامل (تحویل۲) ==")
print("  ", qr["Q-1009"])
