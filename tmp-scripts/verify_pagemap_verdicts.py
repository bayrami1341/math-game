# -*- coding: utf-8 -*-
"""دیف نقشهٔ مادر حکم‌شدهٔ آرنا ↔ نقشهٔ اصلی ما — فقط ۳ ستون حکم باید تغییر کنند"""
import openpyxl
ORIG = 'authoring/PAGE-MAP-KH-G2-MATH-v1.xlsx'
VER  = 'tmp-arena-delivery/authoring/PAGE-MAP-KH-G2-MATH-v1_verified.xlsx'
wo, wn = openpyxl.load_workbook(ORIG, data_only=True), openpyxl.load_workbook(VER, data_only=True)
print('sheets:', wo.sheetnames == wn.sheetnames, wn.sheetnames)
ws_o, ws_n = wo[wo.sheetnames[1]], wn[wn.sheetnames[1]]  # «نقشهٔ صفحه‌ها»
hdr_o = [c.value for c in ws_o[1]]; hdr_n = [c.value for c in ws_n[1]]
print('header same:', hdr_o == hdr_n, '| cols:', len(hdr_o), len(hdr_n))
# ستون‌های حکم آرنا = N..P (14,15,16)
diff_rows = []
verdicts = {'✓':0,'✗':0,'؟':0,'empty':0}
for r in range(2, ws_n.max_row+1):
    ko = [ws_o.cell(r,c).value for c in range(1, ws_n.max_column+1)]
    kn = [ws_n.cell(r,c).value for c in range(1, ws_n.max_column+1)]
    v = (kn[13] or '').strip() if kn[13] else ''
    if v in verdicts: verdicts[v]+=1
    else: verdicts['empty']+=1
    for c in range(len(ko)):
        if c in (13,14,15): continue  # حکم/اصلاح/ارجاع مجاز
        if (ko[c] if ko[c] is not None else '') != (kn[c] if kn[c] is not None else ''):
            diff_rows.append((r, c+1, ko[c], kn[c]))
print('verdicts:', verdicts)
print('changes OUTSIDE the 3 verdict columns:', len(diff_rows))
for d in diff_rows[:15]: print('  ', d)
# شمارش مغایرت‌های ✗ به تفکیک فصل
xs = []
for r in range(2, ws_n.max_row+1):
    v = (ws_n.cell(r,14).value or '')
    if v == '✗': xs.append((r, ws_n.cell(r,2).value, ws_n.cell(r,3).value))
print('✗ rows:', len(xs))
