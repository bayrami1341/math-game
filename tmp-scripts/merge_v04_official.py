# -*- coding: utf-8 -*-
"""ادغام رسمی v0.4: پایه = رسمی v0.3.1 (دست‌نخورده) + فقط ردیف‌های تازهٔ آرنا (افزودهٔ محض)"""
import openpyxl
from copy import copy

OFF = 'public/math-game/authoring/khshab-G2-MATH_v0.3.1.xlsx'
NEW = 'tmp-arena-delivery/authoring/khshab-G2-MATH_v0.4.xlsx'
OUT = 'public/math-game/authoring/khshab-G2-MATH_v0.4.xlsx'

wo = openpyxl.load_workbook(OFF)  # با استایل
wn = openpyxl.load_workbook(NEW, data_only=True)

def find_row(ws, key, col=1):
    for r in range(1, ws.max_row+1):
        v = ws.cell(r, col).value
        if v is not None and str(v).strip() == key:
            return r
    return None

def note_row(ws):
    for r in range(1, ws.max_row+1):
        v = ws.cell(r,1).value
        if v is not None and str(v).strip() == 'یادداشت':
            return r
    return None

report = []
for sheet, keys in [('نقشهٔ کتاب (S1)', ['L-%d'%i for i in range(201,219)]),
                    ('اطلس خطا', ['E_BORROW_MISS (پیشنهاد)','E_CARRY_MISS (پیشنهاد)','E_ROUND_NEAREST (پیشنهاد)']),
                    ('بانک سؤال', ['Q-%d'%i for i in range(1041,1069)])]:
    so, sn = wo[sheet], wn[sheet]
    # ترتیب واقعی ردیف‌های تازه در فایل آرنا
    order = []
    for r in range(1, sn.max_row+1):
        k = sn.cell(r,1).value
        if k and str(k).strip() in keys: order.append(str(k).strip())
    assert len(order) == len(keys), f'{sheet}: {len(order)}/{len(keys)}'
    # محل درج: بالای سطر یادداشت (برای S1: بعد از آخرین ردیف رسمی)
    ins = note_row(so)
    assert ins is not None
    so.insert_rows(ins, len(order))
    for i, k in enumerate(order):
        src_r = find_row(sn, k)
        for c in range(1, so.max_column+1):
            v = sn.cell(src_r, c).value
            so.cell(ins+i, c).value = v
            sc = sn.cell(src_r, c)
            if sc.has_style:
                so.cell(ins+i, c)._style = copy(sc._style)
        # ارتفاع ردیف
        if sn.row_dimensions[src_r].height:
            so.row_dimensions[ins+i].height = sn.row_dimensions[src_r].height
    report.append((sheet, len(order), 'درج در ردیف', ins, '→', ins+len(order)-1))

wo.save(OUT)
for r in report: print(r)
print('saved:', OUT)
