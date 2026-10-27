# -*- coding: utf-8 -*-
import openpyxl
wb = openpyxl.load_workbook('public/math-game/authoring/khshab-G2-MATH_v0.4.xlsx')
ws = wb['نقشهٔ کتاب (S1)']
for r in range(2, ws.max_row+1):
    if str(ws.cell(r,1).value).strip()=='L-213':
        old = str(ws.cell(r,7).value)
        new = old.replace('۵۷−۱۴=۴۳', '۳۵−۱۷ (مثال کتاب، حاصل ۱۸)').replace('۲۴+۳۷=۶۱', '۲۴+۳۷=۶۱')
        ws.cell(r,7).value = new
        note = str(ws.cell(r,14).value or '')
        ws.cell(r,14).value = note + ' | بازرس دوم: PDF ص۳۷ «تفریق 35-17» را نشان می‌دهد؛ ادعای «۵۷−۱۴» حکم آرنا نادرست بود و اصلاح شد.'
        print('L-213 نوع فعالیت:', new)
        break
wb.save('public/math-game/authoring/khshab-G2-MATH_v0.4.xlsx')
print('saved')
