# -*- coding: utf-8 -*-
"""به‌روزرسانی رجیستری ۱۸ خشاب → v1.1
- بستن AUD-10: لنگر ف۱ (ص۱۰ الگو/پرش) به ردیف KH-G3-MATH
- ارجاع قالب استاندارد → v1.1 (شیت پروتکل ممیزی)
- وضعیت v0.6 + تصمیم مالک (فاز ۱ موازی با پایلوت؛ موج ب پشت گیت پایلوت)
- فایل v1 دست‌نخورده می‌ماند.
"""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from copy import copy

SRC = "/home/z/my-project/download/cartridge-registry-18-v1.xlsx"
OUT = "/home/z/my-project/download/cartridge-registry-18-v1.1.xlsx"

HDR_FILL   = PatternFill("solid", start_color="548235")
HDR_FONT   = Font(name="Tahoma", size=10, bold=True, color="FFFFFF")
BODY_FONT  = Font(name="Tahoma", size=10)
SMALL_FONT = Font(name="Tahoma", size=9, color="444444")
TITLE_FONT = Font(name="Tahoma", size=13, bold=True, color="2F4F2F")
OK_FILL    = PatternFill("solid", start_color="C6EFCE")
WRAP_RIGHT = Alignment(wrap_text=True, horizontal="right", vertical="top", readingOrder=2)
CENTER     = Alignment(wrap_text=True, horizontal="center", vertical="center", readingOrder=2)
THIN       = Border(*[Side(style="thin", color="BFBFBF")]*4)

wb = load_workbook(SRC)
log = []

# ───────── شیت جدول ۱۸ خشاب — ردیف ۳ (KH-G3-MATH) ─────────
ws = wb["جدول ۱۸ خشاب"]

old = ws["J3"].value
ws["J3"] = "در جریان — خشاب v0.6 (۵ قرمز AUD-01 + ۲ یافتهٔ ممیزی مجدد رفع شد؛ ۲۵ آزمون موتور پاس) + موتور v10.2"
log.append(f"J3 وضعیت: «{old}» → v0.6")

old = ws["K3"].value
ws["K3"] = str(old) + " | ف۱: ص۱۰ الگو/پرش (پیش‌درس — L-PRE)"
log.append("K3 لنگرها: ف۱ ص۱۰ افزوده شد  [بستن AUD-10 — تصدیق نهایی با برگهٔ انسانی]")

old = ws["M3"].value
ws["M3"] = (str(old) + " تصمیم مالک ۲۸ مهر ۱۴۰۵: فاز ۱ (ساخت محتوا) موازی با پایلوت باز شد — "
            "v10.3 = کل کتاب، دو لایه (موج الف تعاملی: تقسیم/جمع‌وتفریق/چهاررقمی/الگو + مبدل؛ "
            "موج ب بانک تمرین: کسر/هندسه/اندازه‌گیری/آمار) و موج ب پشت گیت بازخورد پایلوت می‌ماند.")
log.append("M3 ریسک/یادداشت: تصمیم مالک v10.3 ثبت شد")

# ───────── شیت راهنما ─────────
ws = wb["راهنما"]
for r in range(1, ws.max_row + 1):
    key = ws.cell(row=r, column=1).value
    if key == "نام سند":
        ws.cell(row=r, column=2).value = "شناسنامهٔ ۱۸ خشاب v1.1 — سجلِ وضعیت/اولویت/خانوادهٔ موتور همهٔ خشاب‌ها (به‌روز: ممیزی AUD-01 + تصمیم v10.3)."
    if key == "رابطه با قالب استاندارد":
        ws.cell(row=r, column=2).value = "ساختار تألیف هر خشاب = cartridge-template-standard-v1.1.xlsx (با شیت «پروتکل ممیزی»). این فایل فقط «کی، با چه وضعیتی، با کدام موتور» را نگه می‌دارد."

# سجل نسخه در انتهای راهنما
r = ws.max_row + 2
cell = ws.cell(row=r, column=1, value="سجل نسخه‌ها")
cell.font = Font(name="Tahoma", size=10, bold=True)
r += 1
for k, v in [
    ("v1.0", "ساخت اولیهٔ نقشهٔ ۱۸ خشاب با لنگرهای تأییدشدهٔ کد ۳۰۵."),
    ("v1.1 (۲۸ مهر ۱۴۰۵)", "بستن AUD-10 (لنگر ف۱ ص۱۰)؛ ارجاع قالب → v1.1؛ ثبت ممیزی AUD-01 و وضعیت v0.6؛ ثبت تصمیم مالک: فاز ۱ موازی با پایلوت، موج الف تعاملی و موج ب پشت گیت پایلوت."),
]:
    ws.cell(row=r, column=1, value=k).font = Font(name="Tahoma", size=10, bold=True)
    c = ws.cell(row=r, column=2, value=v)
    c.font = BODY_FONT; c.alignment = WRAP_RIGHT
    r += 1

wb.save(OUT)
print("ذخیره شد:", OUT)
print("\n--- سیاههٔ تغییرات ---")
for line in log:
    print("•", line)

# راستی‌آزمایی reload
wb2 = load_workbook(OUT)
w = wb2["جدول ۱۸ خشاب"]
assert "ف۱: ص۱۰" in str(w["K3"].value), "لنگر ف۱ نیامده"
assert "v0.6" in str(w["J3"].value), "وضعیت v0.6 نیامده"
assert "تصمیم مالک" in str(w["M3"].value), "تصمیم مالک نیامده"
assert "v1.1" in str(wb2["راهنما"].cell(row=8, column=2).value or "") or "v1.1" in str(wb2["راهنما"]["B8"].value or ""), "ارجاع قالب"
print("\nراستی‌آزمایی reload: پاس ✓")
