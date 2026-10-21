# -*- coding: utf-8 -*-
"""اعمال موج الف v0.7 روی خشاب v0.6 → khshab-G3-MATH_v0.7.xlsx
═══════════════════════════════════════════════════════════════════════════════════
جراحی محدود (الگوی apply_v06.py):
- شیت «اطلس خطا»: فقط دو ردیف پیشنهادی جدید (زرد FFEB9C) درج می‌شود — پس از آخرین
  کد پیشنهادی موجود و پیش از پانویس «یادداشت»:
    * E_BORROW_MISS (جاافتادن قرض — خانوادهٔ ۲، تفریق ستونی)
    * E_ZERO_PLACE_MISS (حذف صفر جایگاه — خانوادهٔ ۳، چهاررقمی)
  هر دو با وضعیت «نیاز به تأیید» و منبع «پیشنهاد AI — موج الف v10.3 (موتور errors-v103.js)».
  الگوی سؤال: توصیف الگو بدون ارجاع Q- (T5 ارجاع‌ها را چک می‌کند).
- شیت «سجل ممیزی» (آخرین شیت): فقط سه ردیف جدید AUD-W1/W2/W3 با ساختار ستون‌های موجود.
- هیچ چیز دیگری تغییر نمی‌کند (بانک سؤال، نقشهٔ کتاب، شناسنامه، راهنما، واژه‌نامه).
- فایل v0.6 دست‌نخورده می‌ماند (سابقهٔ ممیزی) — خروجی در v0.7 ذخیره می‌شود.
"""
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from copy import copy

SRC = "/home/z/my-project/download/khshab-G3-MATH_v0.6.xlsx"
OUT = "/home/z/my-project/download/khshab-G3-MATH_v0.7.xlsx"

OK_FILL    = PatternFill("solid", start_color="C6EFCE")
PEND_FILL  = PatternFill("solid", start_color="FFEB9C")
WRAP_RIGHT = Alignment(wrap_text=True, horizontal="right", vertical="top", readingOrder=2)

wb = load_workbook(SRC)
log = []

# ───────────────── ۱) اطلس خطا — دو ردیف پیشنهادی جدید ─────────────────
ws = wb["اطلس خطا"]

# یافتن پانویس («یادداشت» در ستون A) و آخرین ردیف داده، به‌صورت پویا + گارد
footer_row = None
for r in range(1, ws.max_row + 1):
    v = ws.cell(row=r, column=1).value
    if isinstance(v, str) and v.strip() == "یادداشت":
        footer_row = r
        break
assert footer_row is not None, "پانویس «یادداشت» اطلس پیدا نشد"
last_data = None
for r in range(footer_row - 1, 3, -1):  # از پایین به بالا نخستین ردیف دارای کد E-
    v = ws.cell(row=r, column=1).value
    if isinstance(v, str) and v.strip().startswith("E_"):
        last_data = r
        break
assert last_data == 21 and footer_row == 23, f"ساختار غیرمنتظرهٔ اطلس: last_data={last_data}, footer={footer_row}"

# هیچ مرج merged زیر ناحیهٔ درج نیست (فقط A1:K1 و A2:K2) — گارد
for m in ws.merged_cells.ranges:
    assert m.min_row < last_data + 1, f"merged غیرمنتظره در اطلس: {m}"

NEW_ATLAS_ROWS = [
    ["E_BORROW_MISS (پیشنهاد)",
     "جاافتادن قرض",
     "محاسباتی",
     "تفریق ستونی را رقم‌به‌رقم و بدون قرض‌گرفتن از ده‌گان انجام می‌دهد (در هر ستون کوچک را از بزرگ تفریق می‌کند)",
     "۵۲−۲۷ را ۳۵ می‌نویسد (رقم‌به‌رقم بدون قرض)",
     "تفریق ستونی دورقمی با نیاز به قرض (نمونه: ۵۲−۲۷)",
     "بازنمایی مجسم با مکعب‌های ده‌تایی — «از ده‌گان قرض بگیر»",
     "متوسط",
     "پیشنهاد AI — موج الف v10.3 (موتور errors-v103.js)",
     "نیاز به تأیید",
     "کاندید خانوادهٔ ۲ — سبز شدن با تأیید معلم هم‌ساز (برگهٔ انسانی الف)"],
    ["E_ZERO_PLACE_MISS (پیشنهاد)",
     "حذف صفر جایگاه",
     "مفهومی",
     "در خواندن/نوشتن اعداد چهاررقمی، صفرِ یک جایگاه حذف می‌شود و عدد با جایگاه‌های جابه‌جا نوشته/خوانده می‌شود",
     "۴۰۷۰ را ۴۷۰ می‌نویسد/می‌خواند",
     "نوشتن/خواندن چهاررقمی با صفرِ میانی یا پایانی (نمونه: چهار هزار و هفتاد)",
     "کارت جایگاه‌ها با خانهٔ خالی صفر — «صفر نگهبان جایگاه است»",
     "متوسط",
     "پیشنهاد AI — موج الف v10.3 (موتور errors-v103.js)",
     "نیاز به تأیید",
     "کاندید خانوادهٔ ۳ — سبز شدن با تأیید معلم هم‌ساز (برگهٔ انسانی الف)"],
]

STYLE_SRC = 19  # ردیف E_OP_REVERSE (پیشنهادی) برای الگوی ظاهری ردیف‌های زرد
ins_at = last_data + 1  # ردیف ۲۲ — بلافاصله پس از آخرین کد
ws.insert_rows(ins_at, len(NEW_ATLAS_ROWS))
for i, vals in enumerate(NEW_ATLAS_ROWS):
    r = ins_at + i
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=c, value=v)
        src = ws.cell(row=STYLE_SRC, column=c)
        cell.font = copy(src.font); cell.alignment = copy(src.alignment)
        cell.border = copy(src.border); cell.fill = copy(src.fill)
log.append(f"اطلس خطا ردیف {ins_at}: E_BORROW_MISS (پیشنهاد) — جاافتادن قرض، خانوادهٔ ۲، زرد «نیاز به تأیید»")
log.append(f"اطلس خطا ردیف {ins_at + 1}: E_ZERO_PLACE_MISS (پیشنهاد) — حذف صفر جایگاه، خانوادهٔ ۳، زرد «نیاز به تأیید»")

# ───────────────── ۲) سجل ممیزی — سه ردیف جدید AUD-W1/W2/W3 ─────────────────
ws = wb["سجل ممیزی"]

# هدر و آخرین ردیف جدول (ردیفی که «شناسهٔ یافته» و «محل» هر دو پر دارند)
header_row = 2
last_entry = None
for r in range(header_row + 1, ws.max_row + 1):
    a = ws.cell(row=r, column=1).value
    b = ws.cell(row=r, column=2).value
    if isinstance(a, str) and a.strip() and isinstance(b, str) and b.strip():
        last_entry = r
assert last_entry == 12, f"آخرین ردیف جدول سجل غیرمنتظره: {last_entry}"

# ردیف پانویس تک‌سلولی (یادداشت سجل) و مرج merged آن — برای جابه‌جایی درست
note_row = None
for r in range(last_entry + 1, ws.max_row + 1):
    a = ws.cell(row=r, column=1).value
    b = ws.cell(row=r, column=2).value
    if isinstance(a, str) and a.strip() and not (isinstance(b, str) and b.strip()):
        note_row = r
        break
assert note_row == 14, f"ردیف یادداشت سجل غیرمنتظره: {note_row}"

# openpyxl مرج merged را با insert_rows جابه‌جا نمی‌کند — دستی: بازکردن، درج، بستن دوباره
stale = [m for m in list(ws.merged_cells.ranges) if m.min_row >= note_row]
for m in stale:
    ws.unmerge_cells(str(m))

NEW_AUDIT_ROWS = [
    ("AUD-W1", "اطلس خطا — دو ردیف جدید",
     "زرد — کدهای پیشنهادی خانوادهٔ ۲ و ۳ موج الف (E_BORROW_MISS، E_ZERO_PLACE_MISS) با موتور errors-v103.js پیاده شد",
     "افزوده شد به‌عنوان پیشنهاد؛ سبز شدن با برگهٔ انسانی الف",
     "افزوده شد"),
    ("AUD-W2", "بانک سؤال Q-201..Q-206",
     "زرد — برچسب‌های موج الف با موتور جدید بازخوانی می‌شوند؛ بدون تغییر متن",
     "بدون تغییر — صادق‌ماند",
     "ثبت شد"),
    ("AUD-W3", "نسخهٔ خشاب",
     "v0.6 → v0.7 (موج الف)",
     "دو کد اطلس + سجل",
     "اعمال شد"),
]

ws.insert_rows(last_entry + 1, len(NEW_AUDIT_ROWS))
for i, vals in enumerate(NEW_AUDIT_ROWS):
    r = last_entry + 1 + i
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=r, column=c, value=v)
        src = ws.cell(row=header_row + 1, column=c)  # سبک از نخستین ردیف جدول
        cell.font = copy(src.font); cell.alignment = copy(src.alignment)
        cell.border = copy(src.border)
        # قرارداد رنگ سجل (مطابق apply_v06.py): فقط «رفع شد» سبز؛ بقیه زرد محتاطانه
        cell.fill = OK_FILL if vals[-1] == "رفع شد" else PEND_FILL
log.append(f"سجل ممیزی ردیف {last_entry + 1}: AUD-W1 — دو کد پیشنهادی خانوادهٔ ۲/۳ (زرد، «افزوده شد»)")
log.append(f"سجل ممیزی ردیف {last_entry + 2}: AUD-W2 — برچسب‌های موج الف با موتور جدید بازخوانی می‌شود («ثبت شد»)")
log.append(f"سجل ممیزی ردیف {last_entry + 3}: AUD-W3 — نسخهٔ v0.6 → v0.7 («اعمال شد»)")

# بستن دوبارهٔ مرج یادداشت سجل در جایگاه جابه‌جاشده
shift = len(NEW_AUDIT_ROWS)
new_note_row = note_row + shift
note_text = ws.cell(row=new_note_row, column=1).value
assert isinstance(note_text, str) and note_text.strip().startswith("قاعدهٔ"), "یادداشت سجل پس از درج سر جای انتظار نیست"
ws.merge_cells(start_row=new_note_row, start_column=1, end_row=new_note_row + 2, end_column=5)
log.append(f"مرج یادداشت سجل: A{note_row}:E{note_row + 2} → A{new_note_row}:E{new_note_row + 2} (جابه‌جایی با درج)")

wb.save(OUT)
print("ذخیره شد:", OUT)
print("\n--- سیاههٔ تغییرات ---")
for line in log:
    print("•", line)
