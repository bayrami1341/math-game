# -*- coding: utf-8 -*-
"""ادغام رسمی v0.4 (نسخهٔ ۲): پایهٔ رسمی + افزوده‌های آرنا + جابه‌جایی درست mergeها + نرمال‌سازی منبع"""
import openpyxl, re
from copy import copy

OFF = 'public/math-game/authoring/khshab-G2-MATH_v0.3.1.xlsx'
NEW = 'tmp-arena-delivery/authoring/khshab-G2-MATH_v0.4.xlsx'
OUT = 'public/math-game/authoring/khshab-G2-MATH_v0.4.xlsx'

wo = openpyxl.load_workbook(OFF)
wn = openpyxl.load_workbook(NEW, data_only=True)

def find_row(ws, key):
    for r in range(1, ws.max_row+1):
        v = ws.cell(r,1).value
        if v is not None and str(v).strip() == key: return r
    return None

def note_row(ws):
    for r in range(1, ws.max_row+1):
        v = ws.cell(r,1).value
        if v is not None and str(v).strip()=='یادداشت': return r
    return None

report = []
for sheet, keys in [('نقشهٔ کتاب (S1)', ['L-%d'%i for i in range(201,219)]),
                    ('اطلس خطا', ['E_BORROW_MISS (پیشنهاد)','E_CARRY_MISS (پیشنهاد)','E_ROUND_NEAREST (پیشنهاد)']),
                    ('بانک سؤال', ['Q-%d'%i for i in range(1041,1069)])]:
    so, sn = wo[sheet], wn[sheet]
    order = []
    for r in range(1, sn.max_row+1):
        k = sn.cell(r,1).value
        if k and str(k).strip() in keys: order.append(str(k).strip())
    assert len(order)==len(keys), sheet
    ins = note_row(so)
    # ۱) mergeهای زیر نقطهٔ درج را بردار و آدرس‌شان را نگه دار
    to_restore = []
    for m in list(so.merged_cells.ranges):
        if m.min_row >= ins:
            to_restore.append((m.min_row, m.min_col, m.max_row, m.max_col))
            so.unmerge_cells(str(m))
    # ۲) درج + کپی مقدار و استایل
    so.insert_rows(ins, len(order))
    for i, k in enumerate(order):
        src_r = find_row(sn, k)
        for c in range(1, max(so.max_column, sn.max_column)+1):
            v = sn.cell(src_r, c).value
            if isinstance(v, str):
                v = v.replace('پیام AI — موتور تألیف', 'پیشنهاد AI — موتور تألیف')  # عین‌سازی واژهٔ استاندارد
            so.cell(ins+i, c).value = v
            sc = sn.cell(src_r, c)
            if sc.has_style: so.cell(ins+i, c)._style = copy(sc._style)
        if sn.row_dimensions[src_r].height:
            so.row_dimensions[ins+i].height = sn.row_dimensions[src_r].height
    # ۳) بازگردانی mergeها با شیفت +N
    for (r1,c1,r2,c2) in to_restore:
        so.merge_cells(start_row=r1+len(order), start_column=c1, end_row=r2+len(order), end_column=c2)
    report.append((sheet, len(order), ins, '→', ins+len(order)-1, 'merges shifted:', len(to_restore)))

wo.save(OUT)
for r in report: print(r)
print('saved:', OUT)
