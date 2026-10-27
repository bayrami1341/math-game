# -*- coding: utf-8 -*-
"""ساخت اکسل نقشهٔ مادر صفحات کتاب دوم (C205) — نسخهٔ v1 بازرس دوم
Task 55 — فاز صفر پروتکل §6.7: لیست صفحات با داده‌های ریز، پیش از هر کار تألیف.
"""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

M = json.load(open('authoring/PAGE-MAP-KH-G2-MATH-v1.json'))
RECON = json.load(open('authoring/page-map-chunks/reconciliation.json'))

wb = Workbook()

# ---------- استایل‌ها ----------
GREEN = PatternFill('solid', fgColor='D9EAD3')   # اطمینان بالا
AMBER = PatternFill('solid', fgColor='FFF2CC')   # اطمینان متوسط
RED   = PatternFill('solid', fgColor='F4CCCC')   # اطمینان پایین / آشتی‌دهی
HEAD  = PatternFill('solid', fgColor='07575B')
HEADF = Font(bold=True, color='FFFFFF', size=11)
BOLD  = Font(bold=True)
WRAP  = Alignment(wrap_text=True, vertical='top', horizontal='right')
CENTER= Alignment(vertical='top', horizontal='center')
THIN  = Border(*[Side(style='thin', color='BBBBBB')]*4)

def style_sheet(ws, widths, header_row=1):
    ws.sheet_view.rightToLeft = True
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[header_row]:
        if cell.value is not None:
            cell.fill = HEAD; cell.font = HEADF
            cell.alignment = Alignment(vertical='center', horizontal='center', wrap_text=True)
    ws.freeze_panes = ws.cell(row=header_row+1, column=3)

def put(ws, row, vals, border=True):
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.alignment = WRAP
        if border: c.border = THIN
    return row + 1

# ---------- شیت ۱: راهنما ----------
ws = wb.active; ws.title = 'راهنما'
ws.sheet_view.rightToLeft = True
ws.column_dimensions['A'].width = 26; ws.column_dimensions['B'].width = 110
rows = [
    ('نقشهٔ مادر صفحات — کتاب دوم دبستان (C205)', 'KH-G2-MATH · نسخهٔ v1 بازرس دوم · طبق §6.7 «فاز صفر»'),
    ('هدف', 'لیست صفحه‌به‌صفحهٔ کتاب با داده‌های ریز؛ پس از تأیید دو مدل + کلمهٔ معلم، مبنای تمام تألیف موج‌های بعدی است.'),
    ('سازوکار (فاز صفر)', '۱) بازرس دوم (دستیار پروژه) مستقیم از PDF نقشه را می‌سازد (خوانش کور، ۱۰ زیرعامل مستقل) ۲) آرنا ردیف‌به‌ردیف حکم می‌دهد (ستون‌های حکم/اصلاح) ۳) مغایرت‌ها فقط نزد معلم ۴) کلمهٔ نهایی: معلم.'),
    ('offset', 'صفحهٔ ۷ PDF = صفحهٔ ۱ کتاب (ستون «ص PDF» = «ص کتاب» + ۶). کتاب: ص ۱ تا ۱۴۴. پیش‌صفحات PDF ۱–۶ و پس‌صفحات ۱۵۱–۱۵۲.'),
    ('ستون «اطمینان»', 'بالا: خوانش قطعی · متوسط: ابهام جزئی (معمولاً شمارش تصویری یا اسکرامبل RTL) · پایین: ناخوانا (در این نسخه: صفر).'),
    ('ستون «حکم آرنا»', 'خالی است — آرنا پر می‌کند: ✓ (تأیید) یا ✗ (مغایر) یا ؟ (ابته). با هر ✗/؟ حتماً ستون «اصلاح آرنا» و «ارجاع آرنا» پر شود.'),
    ('ستون «اصلاح آرنا»', 'متن اصلاحیِ دقیق آرنا؛ اگر حکم ✓ است بنویسید «—».'),
    ('ستون «ارجاع آرنا»', 'محل دقیق در PDF (صفحه + بخش) که اصلاح از آن برداشت شده؛ برای داوری ماشینی بازرس دوم.'),
    ('صفحات چندبنری', '۴ صفحه (کتاب ۱۶، ۳۴، ۳۳ و ۵۲) دو بازرس دو بنر متفاوت برداشتند — فلگ «آشتی‌دهی» در یادداشت؛ این ۴ ردیف اولویت تأیید معلم‌اند.'),
    ('حق نشر', 'این فایل فقط «توصیف داده‌محور» است؛ متن کتاب کپی نمی‌شود. PDF و عکس کتاب هرگز کامیت نمی‌شوند (upload/ در gitignore).'),
    ('پروندهٔ مادر JSON', 'authoring/PAGE-MAP-KH-G2-MATH-v1.json (152 رکورد؛ منبع ماشینی همین اکسل).'),
]
r = 1
for a, b in rows:
    ws.cell(row=r, column=1, value=a).font = BOLD
    c = ws.cell(row=r, column=2, value=b); c.alignment = WRAP
    if r == 1:
        ws.cell(row=r, column=1).font = Font(bold=True, size=13, color='07575B')
    r += 1

# ---------- شیت ۲: نقشهٔ صفحه‌ها ----------
ws = wb.create_sheet('نقشهٔ صفحه‌ها')
HEADERS = ['ص کتاب', 'ص PDF', 'فصل', 'عنوان بنر (عیناً)', 'نوع صفحه', 'درس',
           'عناصر ریز صفحه', 'تعداد تمرین', 'اعداد کلیدی', 'ابزار/مانیپولاتیو',
           'بازی/گروهی', 'اطمینان خوانش', 'یادداشت بازرس', 'حکم آرنا', 'اصلاح آرنا', 'ارجاع آرنا']
ws.append(HEADERS)
style_sheet(ws, [8, 8, 7, 24, 13, 8, 60, 9, 22, 16, 9, 10, 34, 11, 26, 20])

def fmt_elements(els):
    out = []
    for e in els or []:
        t = f"{e.get('kind','')}؛ {e.get('title','')}".strip('؛ ')
        d = e.get('detail') or ''
        out.append(f"• {t}: {d}" if d else f"• {t}")
    return '\n'.join(out)

for p in M['pages']:
    row = [p.get('bookPage'), p.get('pdfPage'), p.get('chapter') or ('پیش' if p['kind']=='پیش‌صفحه' else 'پس'),
           p.get('banner'), p.get('pageType'), p.get('lessonNo'),
           fmt_elements(p.get('elements')), p.get('exerciseCount') if p['kind']=='محتوا' else None,
           '، '.join(p.get('keyNumbers') or []), '، '.join(p.get('manipulatives') or []),
           'بله' if p.get('isGame') else '—', p.get('confidence'), p.get('notes') or '', '', '', '']
    r = ws.max_row + 1
    ws.append(row)
    fill = None
    if p.get('reconciled'): fill = RED
    elif p.get('confidence') == 'پایین': fill = RED
    elif p.get('confidence') == 'متوسط': fill = AMBER
    else: fill = GREEN
    for i in range(1, len(HEADERS)+1):
        c = ws.cell(row=r, column=i)
        c.border = THIN
        if i not in (7, 9, 13, 15):
            c.alignment = CENTER
        else:
            c.alignment = WRAP
    ws.cell(row=r, column=7).fill = PatternFill('solid', fgColor='F8F9FA')  # متن عناصر خوانا
    for i in (12,):
        ws.cell(row=r, column=i).fill = fill

# ---------- شیت ۳: خلاصهٔ فصلی ----------
from collections import Counter
ws = wb.create_sheet('خلاصهٔ فصلی')
ws.append(['فصل', 'عنوان', 'بازهٔ صفحه', 'تعداد صفحه', 'تمرین شماره‌دار', 'صفحات بازی', 'اطمینان متوسط', 'اطمینان پایین', 'ترکیب نوع صفحه'])
style_sheet(ws, [6, 30, 12, 10, 14, 10, 12, 12, 60])
CH_TITLES = {'۱':'عدد و رقم','۲':'جمع و تفریق اعداد دو رقمی','۳':'اشکال هندسی','۴':'عددهای سه رقمی',
             '۵':'اندازه‌گیری','۶':'جمع و تفریق اعداد سه رقمی','۷':'کسر و احتمال','۸':'آمار و نمودار'}
chapters = {}
for p in M['pages']:
    if p['kind'] != 'محتوا': continue
    ch = p.get('chapter') or '?'
    d = chapters.setdefault(ch, {'n':0,'ex':0,'g':0,'med':0,'low':0,'types':Counter(),'lo':None,'hi':None})
    d['n'] += 1
    d['ex'] += p.get('exerciseCount') or 0
    d['g'] += 1 if p.get('isGame') else 0
    d['med'] += 1 if p.get('confidence')=='متوسط' else 0
    d['low'] += 1 if p.get('confidence')=='پایین' else 0
    d['types'][p.get('pageType') or '?'] += 1
    bp = p.get('bookPage')
    d['lo'] = bp if d['lo'] is None else min(d['lo'], bp)
    d['hi'] = bp if d['hi'] is None else max(d['hi'], bp)
tot = [0,0,0,0,0]
for ch in sorted(chapters):
    d = chapters[ch]
    ws.append([f'فصل {ch}', CH_TITLES.get(ch,''), f"{d['lo']}–{d['hi']}", d['n'], d['ex'], d['g'], d['med'], d['low'],
               '، '.join(f"{k}×{v}" for k,v in d['types'].items())])
    tot = [t+v for t,v in zip(tot,[d['n'],d['ex'],d['g'],d['med'],d['low']])]
ws.append(['جمع', '۸ فصل', '۱–۱۴۴', *tot, ''])
for c in ws[ws.max_row]: c.font = BOLD
ws.append([])
ws.append(['پیش‌صفحات', 'جلد، شناسنامه، سخن امام، سخنی با معلّم، ساختار کتاب، فهرست', 'PDF ۱–۶', 6, '', '', '', '', 'عناوین فهرست: ۸ فصل، اعداد ص ۱/۱۹/۳۷/۵۵/۷۳/۹۱/۱۰۹/۱۲۷ دقیقاً منطبق مرزها'])
ws.append(['پس‌صفحات', 'پشت جلد / پشتیبان', 'PDF ۱۵۱–۱۵۲', 2, '', '', '', '', ''])

# ---------- شیت ۴: آشتی‌دهی هم‌پوشانی ----------
ws = wb.create_sheet('آشتی‌دهی هم‌پوشانی')
ws.append(['ص PDF', 'ص کتاب', 'فیلد مغایر', 'خوانش بازرس ۱', 'خوانش بازرس ۲', 'حکم معلم / آرنا'])
style_sheet(ws, [8, 8, 14, 26, 26, 22])
for r in RECON:
    chunks = [k for k in r.keys() if k.startswith('[')]
    a, b = (r[chunks[0]], r[chunks[1]]) if len(chunks)==2 else ({}, {})
    ws.append([r['pdfPage'], r['pdfPage']-6, '، '.join(r['diffs']),
               json.dumps(a, ensure_ascii=False), json.dumps(b, ensure_ascii=False), ''])
ws.append([])
ws.append(['توضیح: هر ۴ مورد صفحهٔ «چندبنری» است (فرهنگ نوشتن + مرور فصل هم‌زمان)؛ دو بازرس از یک صفحه دو بنر برداشتند.'])

out = 'authoring/PAGE-MAP-KH-G2-MATH-v1.xlsx'
wb.save(out)
print('نوشته شد:', out)
