#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""نهایی‌سازی نقشهٔ مادر G3: اعمال داوری‌ها + ساخت JSON و XLSX — تسک ۵۶ (§6.7)"""
import json, hashlib
from collections import defaultdict, Counter
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

M = json.load(open('tmp-scripts/merged-g3.json', encoding='utf-8'))
pages = {int(k): v for k, v in M['pages'].items()}
reconciliation = M['reconciliation']

def norm_conf(v):
    if v is None: return 'متوسط'
    s = str(v).strip().lower()
    if s in ('بالا', 'high', 'زیاد'): return 'بالا'
    if s in ('متوسط', 'medium', 'mid'): return 'متوسط'
    if s in ('پایین', 'low', 'کم'): return 'پایین'
    try:
        f = float(s)
        return 'بالا' if f >= 0.85 else ('متوسط' if f >= 0.7 else 'پایین')
    except Exception:
        return 'متوسط'

for r in pages.values():
    r['_conf'] = norm_conf(r.get('confidence'))
    r['_reconciled'] = r.get('pdfPage') in {x['pdfPage'] for x in reconciliation}

CH_START = {1:7, 2:25, 3:43, 4:61, 5:79, 6:97, 7:115, 8:133}
CH_END   = {1:24, 2:42, 3:60, 4:78, 5:96, 6:114, 7:132, 8:150}
CH_TITLE = {1:'الگوها', 2:'عددهای چهار رقمی', 3:'عددهای کسری', 4:'ضرب و تقسیم',
            5:'محیط و مساحت', 6:'جمع و تفریق', 7:'آمار و احتمال', 8:'ضرب عددها'}

# ---------- ۱) داوری‌های مستقیم PDF (توسط بازرس دوم — تسک ۵۶) ----------
ARBITER = {
    39: {'exerciseCount': 8, 'note_add': 'داوری بازرس: شمارش C1 (۸) ترجیح داده شد؛ خوانش B2 (۱۶) زیرسؤال‌شمار بود — مغایرت در آشتی‌دهی'},
    49: {'banner': 'کاربرد کسر در اندازه گیری', 'pageType': 'درس', 'note_add': 'داوری PDF: سرصفحهٔ درس «کاربرد کسر در اندازه‌گیری» + ریبون «فعّالیت» و «انجام بده» هر دو در صفحه‌اند'},
    60: {'banner': 'سرگرمی و معّما', 'pageType': 'سرگرمی و معمّا', 'note_add': 'چندبنری: «فرهنگ خواندن» نیز در همین صفحه؛ داوری PDF'},
    70: {'banner': 'فعّالیت', 'pageType': 'فعّالیت', 'note_add': 'داوری PDF: ریبون «فعّالیت» + موضوع «خاصیت‌های ضرب»؛ شمارش ۵ از D1 حفظ شد، خوانش ۴ در آشتی‌دهی'},
    79: {'banner': 'فصل ۵: محیط و مساحت'},
    99: {'banner': 'حلّ مسئلهٔ ساده‌تر', 'pageType': 'حلّ‌مسئله'},
    100: {'banner': 'فعّالیت', 'pageType': 'فعّالیت'},
    119: {'banner': 'کار در کلاس', 'pageType': 'کار در کلاس', 'exerciseCount': 4},
    120: {'banner': 'تمرین', 'pageType': 'تمرین'},
    129: {'banner': 'تمرین', 'pageType': 'تمرین'},
    130: {'banner': 'مرور فصل', 'pageType': 'مرور فصل'},
}
for p, ov in ARBITER.items():
    r = pages[p]
    add = ov.pop('note_add', None)
    r.update(ov)
    if add:
        r['notes'] = ((r.get('notes') or '') + ' | ' + add).strip(' |')
# نرمال‌سازی فاصلهٔ بنرهای فصل
for ch, s in CH_START.items():
    b = pages[s].get('banner')
    if b: pages[s]['banner'] = f'فصل {ch}: {CH_TITLE[ch]}'

# ---------- ۲) خلاصهٔ فصلی ----------
def conf_avg(ch_pages):
    sc = {'بالا': 3, 'متوسط': 2, 'پایین': 1}
    vals = [sc[pages[p].get('_conf', 'متوسط')] for p in ch_pages]
    return round(sum(vals) / len(vals), 2)

ch_summary = []
for ch in range(1, 9):
    ch_pages = list(range(CH_START[ch], CH_END[ch] + 1))
    ex = sum(int(pages[p].get('exerciseCount') or 0) for p in ch_pages)
    games = [p for p in ch_pages if pages[p].get('isGame')]
    types = Counter(pages[p].get('pageType') or 'نامشخص' for p in ch_pages)
    low = sum(1 for p in ch_pages if pages[p].get('_conf') == 'پایین')
    mid = sum(1 for p in ch_pages if pages[p].get('_conf') == 'متوسط')
    ch_summary.append({
        'chapter': ch, 'title': CH_TITLE[ch], 'range': [CH_START[ch], CH_END[ch]],
        'pageCount': len(ch_pages), 'exerciseCount': ex, 'gamePages': games,
        'confAvg': conf_avg(ch_pages), 'low': low, 'mid': mid,
        'typeMix': dict(types.most_common()),
    })

# ---------- ۳) JSON ماشین‌خوان ----------
def rec_public(r):
    d = {k: v for k, v in r.items() if not k.startswith('_')}
    d['confidence'] = r.get('_conf', 'متوسط')
    return d

out = {
    'mapId': 'PAGE-MAP-KH-G3-MATH-v1',
    'book': 'ریاضی سوم دبستان (C305) — سازمان پژوهش و برنامه‌ریزی آموزشی',
    'offset': 0,
    'offsetRule': 'bookPage = pdfPage (شمارهٔ چاپی هر صفحه با شمارهٔ PDF برابر است)',
    'totalPdfPages': 152,
    'bookPages': 152,
    'structure': {'prePages': [1, 6], 'chapters': {str(c): [CH_START[c], CH_END[c]] for c in CH_START}, 'postPages': [151, 152]},
    'chapterTitles': {str(c): CH_TITLE[c] for c in CH_START},
    'builtBy': 'بازرس دوم (تسک ۵۶) — ۱۷ زیرعامل کور مستقل، ادغام و آشتی‌دهی ماشینی',
    'protocol': '§6.7 فاز صفر — دستور معلم: «اول لیست صفحات با داده‌های ریز، بعد هر کار دیگر»',
    'readers': '۱۰+۷ زیرعامل کور (چانک‌های ۸–۱۲ صفحه‌ای با هم‌پوشانی مرزی ۲ صفحه‌ای؛ ۲۰ صفحهٔ تکراری)',
    'stats': {
        'totalExercises': sum(int(r.get('exerciseCount') or 0) for r in pages.values()),
        'gamePages': sorted(p for p, r in pages.items() if r.get('isGame')),
        'confidence': dict(Counter(r.get('_conf') for r in pages.values())),
        'needsVisualCheck': sorted(p for p, r in pages.items() if 'نیازمند تأیید تصویری' in str(r.get('notes') or '')),
        'reconciliationPages': sorted({x['pdfPage'] for x in reconciliation}),
    },
    'chapterSummary': ch_summary,
    'pages': [{'pdfPage': p, **rec_public(pages[p])} for p in range(1, 153)],
}
json.dump(out, open('authoring/PAGE-MAP-KH-G3-MATH-v1.json', 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

# ---------- ۴) XLSX چهارشیتی RTL ----------
wb = openpyxl.Workbook()
F_HEAD = Font(bold=True, size=12)
FILL_HEAD = PatternFill('solid', fgColor='DDEBC8')
FILL_YELLOW = PatternFill('solid', fgColor='FFF2CC')
FILL_GREEN = PatternFill('solid', fgColor='D9EAD3')
RTL = Alignment(horizontal='right', vertical='center', wrap_text=True)

def style_header(ws, row=1):
    for c in ws[row]:
        if c.value: c.font, c.fill, c.alignment = F_HEAD, FILL_HEAD, RTL

# شیت ۱: راهنما
ws = wb.active; ws.title = 'راهنما'
guide = [
    ('نقشهٔ مادر صفحات — کتاب ریاضی سوم دبستان (C305)', 'KH-G3-MATH · نسخهٔ v1 بازرس دوم — پروتکل §6.7 «فاز صفر»'),
    ('هدف', 'لیست صفحه‌به‌صفحهٔ هر ۱۵۲ صفحه با داده‌های ریز؛ پیش از هر کار تألیف، طبق دستور معلم'),
    ('منبع', 'upload/C305.pdf (هرگز کامیت نمی‌شود) — ۱۵۲ صفحهٔ PDF'),
    ('قاعدهٔ صفحه', 'offset=0: bookPage = pdfPage — شمارهٔ چاپی هر صفحه با شمارهٔ PDF برابر است (تفاوت با C205 که offset=+۶ داشت)'),
    ('ساختار', 'پیش‌صفحات ۱–۶ (جلد/شناسنامه/امام خمینی/فهرست/سخنی با معلّم) · ۸ فصل × ۱۸ صفحه در ص ۷–۱۵۰ · پس‌صفحات ۱۵۱–۱۵۲'),
    ('روش ساخت', '۱۷ زیرعامل کور مستقل (بدون دیدن هیچ نقشه/خشاب قبلی) با چانک‌های ۸–۱۲ صفحه‌ای و هم‌پوشانی مرزی ۲ صفحه‌ای؛ ۲۰ صفحهٔ تکراری برای خودچک'),
    ('آشتی‌دهی', 'مغایرت‌های هم‌پوشانی در شیت «آشتی‌دهی»؛ داوری بازرس در یادداشت ردیف؛ حکم نهایی با آرنا/معلم'),
    ('شمارش تمرین', 'آیتم‌های شماره‌دار صفحه (۱ــ ۲ــ …)؛ برخی زیرعامل‌ها زیرسؤال‌ها را هم شمردند — صفحات مغایر در آشتی‌دهی'),
    ('محدودیت شفاف', 'خوانش فقط لایهٔ متنی/چیدمانی بود؛ شمارش عناصر صرفاً تصویری با «نیازمند تأیید تصویری» فلگ شده (۷۴ صفحه)'),
    ('ستون‌های آرنا', 'سه ستون آخر (حکم آرنا / اصلاح آرنا / ارجاع آرنا) خالی برای جلسهٔ حکم ردیف‌به‌ردیف — بریف: ARENA-BRIEF-PAGEMAP-KH-G3-v1.fa.md'),
    ('فایل‌های همراه', 'PAGE-MAP-KH-G3-MATH-v1.json (مادهٔ ماشینی) · PAGE-MAP-KH-G3-MATH-v1.fa.md (گزارش)'),
]
for row in guide: ws.append(row)
ws.column_dimensions['A'].width, ws.column_dimensions['B'].width = 24, 110
for r in ws.iter_rows():
    for c in r: c.alignment = RTL
for c in ws[1]: c.font, c.fill = F_HEAD, FILL_HEAD

# شیت ۲: نقشهٔ صفحه‌ها
ws2 = wb.create_sheet('نقشهٔ صفحه‌ها')
heads = ['ص کتاب', 'ص PDF', 'فصل', 'عنوان بنر (عیناً)', 'نوع صفحه', 'درس',
         'عناصر ریز صفحه', 'تعداد تمرین', 'اعداد کلیدی', 'ابزار/مانیپولاتیو',
         'بازی/گروهی', 'اطمینان خوانش', 'یادداشت بازرس', 'حکم آرنا', 'اصلاح آرنا', 'ارجاع آرنا']
ws2.append(heads)
for p in range(1, 153):
    r = pages[p]
    ch = r.get('chapter')
    elems = ' • '.join(f"[{e.get('kind','؟')}] {e.get('title','')}: {e.get('detail','')}" for e in (r.get('elements') or []))
    ws2.append([p, p, ch if ch else ('پیش‌صفحه' if p <= 6 else 'پس‌صفحه'),
                r.get('banner') or '—', r.get('pageType') or '—', r.get('lessonNo') or '—',
                elems, int(r.get('exerciseCount') or 0),
                '، '.join(str(x) for x in (r.get('keyNumbers') or [])) or '—',
                '، '.join(str(x) for x in (r.get('manipulatives') or [])) or '—',
                '✓' if r.get('isGame') else '—', r.get('_conf'),
                r.get('notes') or '—', '', '', ''])
    row = ws2.max_row
    conf = r.get('_conf')
    if conf == 'متوسط': ws2.cell(row=row, column=12).fill = FILL_YELLOW
    elif conf == 'پایین': ws2.cell(row=row, column=12).fill = FILL_YELLOW
    if r.get('_reconciled') or 'آشتی‌دهی' in str(r.get('notes') or ''):
        ws2.cell(row=row, column=13).fill = FILL_YELLOW
widths = [7, 7, 10, 22, 13, 6, 60, 9, 18, 16, 9, 10, 40, 11, 20, 16]
for i, w in enumerate(widths, 1): ws2.column_dimensions[get_column_letter(i)].width = w
style_header(ws2)
ws2.freeze_panes = 'A2'

# شیت ۳: خلاصهٔ فصلی
ws3 = wb.create_sheet('خلاصهٔ فصلی')
ws3.append(['فصل', 'عنوان', 'بازهٔ صفحه', 'تعداد صفحه', 'تمرین شماره‌دار', 'صفحات بازی',
            'اطمینان متوسط (۱–۳)', 'اطمینان پایین', 'ترکیب نوع صفحه'])
for s in ch_summary:
    ws3.append([s['chapter'], s['title'], f"{s['range'][0]}–{s['range'][1]}", s['pageCount'],
                s['exerciseCount'], '، '.join(map(str, s['gamePages'])) or '—',
                s['confAvg'], s['low'], '؛ '.join(f"{k}: {v}" for k, v in s['typeMix'].items())])
for i, w in enumerate([6, 20, 12, 10, 14, 18, 16, 12, 60], 1): ws3.column_dimensions[get_column_letter(i)].width = w
style_header(ws3)

# شیت ۴: آشتی‌دهی هم‌پوشانی
ws4 = wb.create_sheet('آشتی‌دهی هم‌پوشانی')
ws4.append(['ص PDF', 'ص کتاب', 'فیلد مغایر', 'خوانش بازرس ۱', 'خوانش بازرس ۲', 'حکم معلم / آرنا'])
chunk_pair = {'39': ('B2', 'C1'), '49': ('C1', 'C2'), '50': ('C1', 'C2'), '60': ('D1', 'D2'),
              '70': ('D1', 'D2'), '79': ('D2', 'E1'), '99': ('E2', 'F1'), '100': ('E2', 'F1'),
              '119': ('F2S', 'G1'), '120': ('F2S', 'G1'), '129': ('G1', 'G2'), '130': ('G1', 'G2')}
for x in reconciliation:
    a, b = chunk_pair.get(str(x['pdfPage']), ('؟', '؟'))
    ws4.append([x['pdfPage'], x['pdfPage'], x['field'], f'[{a}] {x["main"]}', f'[{b}] {x["other"]}', ''])
for i, w in enumerate([8, 8, 14, 30, 30, 24], 1): ws4.column_dimensions[get_column_letter(i)].width = w
style_header(ws4)

wb.save('authoring/PAGE-MAP-KH-G3-MATH-v1.xlsx')

# ---------- ۵) SHA-256 ----------
def sha(f):
    h = hashlib.sha256()
    h.update(open(f, 'rb').read())
    return h.hexdigest()
print('SHA256 xlsx:', sha('authoring/PAGE-MAP-KH-G3-MATH-v1.xlsx'))
print('SHA256 json:', sha('authoring/PAGE-MAP-KH-G3-MATH-v1.json'))
print('تمرین کل:', out['stats']['totalExercises'], '| بازی:', out['stats']['gamePages'])
print('اطمینان:', out['stats']['confidence'])
print('تصویری:', len(out['stats']['needsVisualCheck']), '| آشتی‌دهی:', len(out['stats']['reconciliationPages']))
for s in ch_summary:
    print(f"  فصل {s['chapter']} {s['title']}: {s['range']} · تمرین {s['exerciseCount']} · بازی {s['gamePages']} · اطمینان {s['confAvg']}")
