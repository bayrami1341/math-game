# -*- coding: utf-8 -*-
import openpyxl
wb = openpyxl.load_workbook('tmp-arena-delivery/authoring/khshab-G2-MATH_v0.4.xlsx', data_only=True)
ws = wb['بانک سؤال']
for r in range(1, ws.max_row+1):
    k = ws.cell(r,1).value
    if k is None: continue
    s = str(k).strip()
    if s in ('Q-1042','Q-1043','Q-1044'):
        vals = [ws.cell(r,c).value for c in range(1,18)]
        print(f'row {r}: {s}')
        for i,v in enumerate(vals,1): print(f'   col{i}: {repr(v)[:110]}')
        print()
# merged cells?
print('merged:', [str(m) for m in ws.merged_cells.ranges][:20])
# نمونهٔ یک سؤال تصویری هشداردار: Q-1041
for r in range(1, ws.max_row+1):
    if str(ws.cell(r,1).value).strip()=='Q-1041':
        print('Q-1041 struct:', repr(ws.cell(r,6).value)[:200])
        print('Q-1041 text:', repr(ws.cell(r,5).value)[:200])
        print('Q-1041 note:', repr(ws.cell(r,17).value)[:250])
