#!/usr/bin/env python3
"""اصلاح گیت مبدل: هم‌گام‌سازی بازهٔ لنگرهای نقشه با مالکیت واقعی بانک (Q-1020 → L-112)"""
import openpyxl

F = "/home/z/my-project/public/math-game/authoring/khshab-G2-MATH_v0.3.1.xlsx"
wb = openpyxl.load_workbook(F, data_only=False)
ws = wb["نقشهٔ کتاب (S1)"]
for row in ws.iter_rows():
    if row[0].value == "L-110":
        rr = row[0].row
        print(f"L-110 K{rr}: {ws[f'K{rr}'].value!r} →", end=" ")
        ws[f"K{rr}"] = "E_ADD_CONFUSION، E_SUB_CONFUSION (Q-1017-1019)"
        print(repr(ws[f"K{rr}"].value), " (Q-1020 به L-112 منتقل شده — شاهد: سؤال «۱۳−۸+۷» از ص ۱۲)")
    if row[0].value == "L-112":
        rr = row[0].row
        print(f"L-112 K{rr}: {ws[f'K{rr}'].value!r} →", end=" ")
        ws[f"K{rr}"] = "E_ADD_CONFUSION، E_SUB_CONFUSION (Q-1020)"
        print(repr(ws[f"K{rr}"].value))
wb.save(F)
print("ذخیره شد.")
