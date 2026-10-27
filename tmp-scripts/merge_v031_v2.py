#!/usr/bin/env python3
"""ادغام رسمی v0.3.1 — نسخهٔ ۲ (تحویل دوم مدل ddee0d9 / SHA a386ba15…)
پایه: تحویل دوم مدل (CPA منطبق AUD-06، بندهای ۲ و ۴ اصلاحیه اعمال‌شده)
روپوش: تصحیحات مستند بازرس دوم + بازگردانی مُهرهای رسمی معلم (پروتکل تسک ۴۸)
خروجی: public/math-game/authoring/khshab-G2-MATH_v0.3.1.xlsx
"""
import openpyxl, shutil

SRC = "/tmp/gate2-redelivery/khshab-G2-MATH_v0.3.1.xlsx"   # تحویل دوم مدل (a386ba15)
OFF = "/home/z/my-project/public/math-game/authoring/khshab-G2-MATH_v0.3.1.xlsx"  # رسمی فعلی (c2d511a8)
OUT = OFF

# ۱) اول روپوش را به حافظه بخوان (قبل از هر بازنویسی!)
wo = openpyxl.load_workbook(OFF, data_only=False)

# ۲) حالا پایه را جای‌گذاری کن
shutil.copy(SRC, OUT)
wb = openpyxl.load_workbook(OUT, data_only=False)

changes = []
def setcell(ws, coord, val, why):
    old = ws[coord].value
    if old == val:
        return
    ws[coord] = val
    changes.append((ws.title, coord, old, val, why))

# ---------- ۱) بازگردانی مُهرهای رسمی معلم (پروتکل تسک ۴۸) ----------
wsi, wsi_o = wb["شناسنامهٔ خشاب"], wo["شناسنامهٔ خشاب"]
for r in range(3, 8):  # شناسه/عنوان/درس/پایه/بستهٔ تجاری
    setcell(wsi, f"C{r}", wsi_o[f"C{r}"].value, "بازگردانی مُهر رسمی معلم")
    setcell(wsi, f"D{r}", wsi_o[f"D{r}"].value, "بازگردانی منبع رسمی معلم")
setcell(wsi, "B13", wsi_o["B13"].value, "بازگردانی وضعیت رسمی v0.3.1")
setcell(wsi, "B14", wsi_o["B14"].value, "بازگردانی تاریخ رسمی آخرین تأیید سبز")

wsm, wsm_o = wb["نقشهٔ کتاب (S1)"], wo["نقشهٔ کتاب (S1)"]
skeleton = ["L-100", "L-200", "L-300", "L-400", "L-500", "L-600", "L-700", "L-800"]
def find_row(ws, lid):
    for row in ws.iter_rows():
        if row[0].value == lid:
            return row[0].row
    return None

for lid in skeleton:
    rr = find_row(wsm, lid)
    rr_o = find_row(wsm_o, lid)
    setcell(wsm, f"L{rr}", wsm_o[f"L{rr_o}"].value, f"بازگردانی مُهر رسمی اسکلت {lid}")

# ---------- ۲) Q-1011: کتاب ص ۸ = ۱۳ مثلث → ۱۸ (شمارش ماشینی تصویر؛ سه‌بار تأییدشده) ----------
wsb, wsb_o = wb["بانک سؤال"], wo["بانک سؤال"]
def qrow(ws, qid):
    for row in ws.iter_rows():
        if row[0].value == qid:
            return row[0].row
    return None

rr, rr_o = qrow(wsb, "Q-1011"), qrow(wsb_o, "Q-1011")
for col in ("E", "G", "H", "I", "J", "K", "N", "Q"):
    setcell(wsb, f"{col}{rr}", wsb_o[f"{col}{rr_o}"].value, f"اصلاح بازرس دوم Q-1011 ({col}) — ۱۳ مثلث→۱۸")

# ---------- ۳) شاهد L-108 هم‌گام با ۱۳/۱۸ ----------
rr, rr_o = find_row(wsm, "L-108"), find_row(wsm_o, "L-108")
setcell(wsm, f"N{rr}", wsm_o[f"N{rr_o}"].value, "شاهد L-108: ۱۳ مثلث → ۱۸ (نسخهٔ رسمی بازرس)")

# ---------- ۴) L-109 «کامل کن»: قرینه طبق PDF («سه قسمت دیگر… با تاکردن چپ/راست روی هم») ----------
rr = find_row(wsm, "L-109")
setcell(wsm, f"G{rr}",
        "رسم شکل: پرش روی محور تا ۱۶ + ثبت چوب‌خطی + «کامل کن»: تکمیل سه قسمتِ دیگرِ شکل به‌صورت قرینه با خط‌کش و شابلون (با تاخوردن، چپ/راست یا بالا/پایین روی هم می‌افتند)",
        "بند تصحیحی بازرس دوم با متن PDF ص ۱۵ فایل — قرینه")

# ---------- ۵) L-112: «خلّق باش» (PDF: «خلق باش» + شابلون ۴ مثلث) ----------
rr, rr_o = find_row(wsm, "L-112"), find_row(wsm_o, "L-112")
setcell(wsm, f"G{rr}",
        "تمرین: رنگ‌آمیزی و نوشتن جمع + محاسبه روی محور + «خلّق باش»: ساخت شکل با شابلون ۴ مثلث کنار هم",
        "واژهٔ بنر طبق PDF («خلق باش») — خلاق نادرست بود")
setcell(wsm, f"N{rr}", wsm_o[f"N{rr_o}"].value, "شاهد رسمی L-112 (خلّق باش) از نسخهٔ ادغام قبلی")

# ---------- ۶) هم‌گام‌سازی لنگرها: Q-1020→L-112 ⟸ K نقشه ----------
rr, rr_o = find_row(wsm, "L-110"), find_row(wsm_o, "L-110")
setcell(wsm, f"K{rr}", wsm_o[f"K{rr_o}"].value, "لنگر K: (Q-1017-1019) — هم‌گام با B سؤال‌ها")
rr, rr_o = find_row(wsm, "L-112"), find_row(wsm_o, "L-112")
setcell(wsm, f"K{rr}", wsm_o[f"K{rr_o}"].value, "لنگر K: (Q-1020) — هم‌گام با Q-1020!B=L-112")

# ---------- ۷) L-118: عنوان صفحه (PDF خود صفحه: «معمّا» + «فرهنگ خواندن») — از رسمی ----------
rr, rr_o = find_row(wsm, "L-118"), find_row(wsm_o, "L-118")
for col in ("D", "N"):
    setcell(wsm, f"{col}{rr}", wsm_o[f"{col}{rr_o}"].value, f"L-118 ({col}): نسخهٔ رسمی چک‌شدهٔ معلم — فرهنگ خواندن")

# ---------- ۸) L-114: عنوان لختِ پیام اصلاحی + محتوای غنی‌تر رسمی ----------
rr, rr_o = find_row(wsm, "L-114"), find_row(wsm_o, "L-114")
setcell(wsm, f"D{rr}", "شمارش ۴تا۴تا", "بند ۲ اصلاحیه — عنوان عین پیام مُهردار")
for col in ("E", "F", "G", "N"):
    setcell(wsm, f"{col}{rr}", wsm_o[f"{col}{rr_o}"].value, f"L-114 ({col}): محتوای رسمی غنی‌تر بازرس (شامل خلّق باش ص۱۴ و دسته‌های ۵تایی)")

wb.save(OUT)
print(f"ذخیره شد: {OUT}")
print(f"\nتغییرات اعمال‌شده ({len(changes)} سلول):")
for sheet, coord, old, new, why in changes:
    so = (str(old)[:55] + "…") if old and len(str(old)) > 55 else old
    sn = (str(new)[:55] + "…") if new and len(str(new)) > 55 else new
    print(f"  [{sheet}] {coord}: {so!r} → {sn!r}  ({why})")
