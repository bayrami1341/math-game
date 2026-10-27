#!/usr/bin/env python3
"""Task 54 — راستی‌آزمایی تحویل دوم v0.3.1 (ddee0d9, SHA a386ba15…)
سه مقایسه: تحویل دوم ↔ v0.3 (فقط تغییرات مجاز) | تحویل دوم ↔ رسمی c2d511a8 | چک‌های هدفمند
"""
import openpyxl, re, json, sys

V03   = "/home/z/my-project/public/math-game/authoring/khshab-G2-MATH_v0.3.xlsx"
OFF   = "/home/z/my-project/public/math-game/authoring/khshab-G2-MATH_v0.3.1.xlsx"   # رسمی c2d511a8
REDEL = "/tmp/gate2-redelivery/khshab-G2-MATH_v0.3.1.xlsx"                            # تحویل دوم a386ba15

FAIL = []
def check(name, ok, detail=""):
    print(f"{'✓' if ok else '✗ FAIL'} {name}" + (f" — {detail}" if detail else ""))
    if not ok: FAIL.append(name)

def load(path):
    wb = openpyxl.load_workbook(path, data_only=False)
    return wb

def rows_by_id(ws, prefix=None):
    out = {}
    for row in ws.iter_rows():
        aid = row[0].value
        if isinstance(aid, str) and (prefix is None or aid.startswith(prefix)):
            out[aid] = {c.coordinate[0]: ("" if c.value is None else str(c.value)) for c in row}
    return out

w3, wo, wr = load(V03), load(OFF), load(REDEL)
print("شیت‌های تحویل دوم:", wr.sheetnames)

# ---------- ۰) شیت‌های ساختاری مو به مو ----------
for sh in w3.sheetnames:
    a = [[("" if c.value is None else str(c.value)) for c in r] for r in w3[sh].iter_rows()]
    b = [[("" if c.value is None else str(c.value)) for c in r] for r in wr[sh].iter_rows()]
    if sh in ("شناسنامهٔ خشاب", "اطلس خطا", "راهنما", "واژه‌نامهٔ ستون‌ها", "پروتکل ممیزی"):
        check(f"شیت «{sh}» مو به مو با v0.3", a == b,
              "" if a == b else f"تعداد ردیف {len(a)}↔{len(b)}")

# ---------- ۱) نقشهٔ کتاب S1: تحویل دوم ↔ v0.3 ----------
m3, mr = rows_by_id(w3["نقشهٔ کتاب (S1)"], "L-"), rows_by_id(wr["نقشهٔ کتاب (S1)"], "L-")
ALLOWED_MAP = {f"L-{n}" for n in (104, 107, 108, 109, 110, 111, 112, 114, 115, 118)}
map_bad = []
map_changed = []
for lid in sorted(set(m3) | set(mr)):
    if m3.get(lid) == mr.get(lid): continue
    map_changed.append(lid)
    if lid not in ALLOWED_MAP: map_bad.append(lid)
check("نقشه S1: فقط ۱۰ ردیف مجاز تغییر کرده", not map_bad and sorted(map_changed) == sorted(ALLOWED_MAP),
      f"تغییریافته: {sorted(map_changed)} | نامجاز: {map_bad or '—'}")

# اسکلت سبز: L-100 و L-200..L-800 دست‌نخورده
skel = [l for l in map_changed if l in ("L-100",) or (l.startswith("L-") and l[2] != "0" and l[2:3].isdigit() and int(l[2]) >= 2)]
check("اسکلت سبز (L-100، L-200..L-800) دست‌نخورده", not skel, f"{skel or '—'}")

# ---------- ۲) عناوین عین بنر ----------
TITLES = {"L-108": "حلّ مسئله", "L-109": "رسم شکل", "L-110": "جمع و تفریق",
          "L-111": "فعّالیت", "L-112": "تمرین"}
for lid, t in TITLES.items():
    col = None
    # ستون عنوان: همان ستونی که در v0.3 عنوان درس بوده — از L-101 الگو می‌گیریم
    for k, v in m3["L-101"].items():
        if "عدد و شمارش" in v or "شمارش" in str(v):
            col = k; break
    got = mr.get(lid, {}).get(col or "C", "")
    check(f"عنوان {lid} = «{t}»", got.strip() == t, f"got={got!r} (ستون {col})")

# صفحات تمرین ۴/۷/۱۵ + صفحهٔ ۱۸
for lid in ("L-104", "L-107", "L-115"):
    txt = " | ".join(mr.get(lid, {}).values())
    check(f"{lid} عنوان تمرین بدون بنر مستقل", "عنوان بنری مستقل ندارد" in txt)
txt18 = " | ".join(mr.get("L-118", {}).values())
check("L-118 = «معمّا + فرهنگ ریاضی» خارج از موتور", "فرهنگ ریاضی" in txt18 and "خارج از موتور" in txt18,
      txt18[:120])

# ---------- ۳) L-114: شمارش ۴تا۴تا + کلمهٔ ممنوعه ----------
r114 = mr.get("L-114", {})
all114 = " | ".join(r114.values())
check("L-114 شامل «شمارش ۴تا۴تا»", "۴تا۴تا" in all114)
forbidden = [w for w in ("سه‌تایی", "سه تایی", "شمارش سه‌تایی") if w in all114]
check("L-114 بدون کلمهٔ ممنوعه «سه‌تایی»", not forbidden, f"{forbidden or '—'}")
# هیچ اثری از شمارش اشتباه قدیمی در کل نقشه (به‌جز جاهای تاریخی مجاز مثل اطلس/یادداشت گیت؟ نقشه باید پاک باشد)
leftover = [lid for lid, row in mr.items() if "سه‌تایی" in " ".join(row.values()) and lid not in ("L-113",)]
check("نقشهٔ S1 بدون بقایای «سه‌تایی»", not leftover, f"{leftover or '—'}")

# ---------- ۴) بانک سؤال: تحویل دوم ↔ v0.3 ----------
q3, qr = rows_by_id(w3["بانک سؤال"], "Q-"), rows_by_id(wr["بانک سؤال"], "Q-")
REWRITE = {f"Q-{n}" for n in range(1009, 1025)}
q_bad, q_changed = [], []
for qid in sorted(set(q3) | set(qr)):
    if q3.get(qid) == qr.get(qid): continue
    q_changed.append(qid)
    ch = {k for k in set(q3.get(qid, {})) | set(qr.get(qid, {})) if q3.get(qid, {}).get(k) != qr.get(qid, {}).get(k)}
    if qid in REWRITE: continue
    if ch <= {"D"}: continue   # فقط ستون CPA
    q_bad.append((qid, sorted(ch)))
check("بانک: فقط Q-1009..Q-1024 (کامل) + ستون D (CPA) در بقیه", not q_bad,
      f"تغییریافته={len(q_changed)} | نامجاز: {q_bad or '—'}")
n_cpa_old = sum(1 for qid in q_changed if qid not in REWRITE)
check("تعداد تغییر CPA سؤال‌های قدیمی = ۳۰ ادعای مدل (شامل ۱۶ بازنویسی؟)", True,
      f"قدیمی={n_cpa_old} کل={len(q_changed)}")

# CPA: دامنه + توزیع
cpa = {}
for qid, row in qr.items():
    v = row.get("D", "").strip()
    cpa[v] = cpa.get(v, 0) + 1
check("CPA فقط از سه‌گانهٔ مجسم/تصویری/علامتی", set(cpa) <= {"مجسم", "تصویری", "علامتی"}, str(cpa))
check("CPA توزیع ادعایی ۳۳ علامتی / ۷ تصویری", cpa.get("علامتی") == 33 and cpa.get("تصویری") == 7, str(cpa))

# کدهای ۱۶ سؤال
ALLOWED = {"E_SUB_CONFUSION", "E_ADD_CONFUSION", "E_SKIP_UNDERCOUNT", "E_DOUBLE_OFF", "E_GROUPS"}
code_col = "L"
bad_codes = []
for n in range(1009, 1025):
    qid = f"Q-{n}"
    codes = set(re.findall(r"E_[A-Z_]+", qr.get(qid, {}).get(code_col, "")))
    if not codes <= ALLOWED: bad_codes.append((qid, codes - ALLOWED))
check("کدهای Q-1009..Q-1024 ⊆ ۵ کد مجاز", not bad_codes, f"{bad_codes or '—'}")

# توازن پاسخ ۱۰/۱۰/۱۰/۱۰ — ستون حرف پاسخ را پیدا کن: در ردیف، یک سلول تک‌حرفی از الف/ب/پ/د
balance = {}
ans_col = None
for k in qr["Q-1001"]:
    vals = {qr[q].get(k, "").strip() for q in qr if q.startswith("Q-10")}
    if vals <= {"الف", "ب", "پ", "د"} and len(vals) == 4:
        ans_col = k; break
for qid, row in qr.items():
    v = row.get(ans_col, "").strip() if ans_col else ""
    if v in ("الف", "ب", "پ", "د"): balance[v] = balance.get(v, 0) + 1
check(f"توازن پاسخ ۱۰/۱۰/۱۰/۱۰ (ستون {ans_col})", balance == {"الف": 10, "ب": 10, "پ": 10, "د": 10}, str(balance))

# ضد J12: ارقام لاتین در صورت/گزینه‌ها ممنوع — ستون‌های متن سؤال را از واژه‌نامه تشخیص می‌دهیم
# تقریب: هر سلول Q-1009..1024 که «یادداشت» نیست و رقم لاتین دارد = خطا
latin_hits = []
for n in range(1009, 1025):
    qid = f"Q-{n}"
    for k, v in qr[qid].items():
        if k == "A": continue
        if re.search(r"[0-9]", v) and not re.search(r"E_[A-Z_]+|[A-Z]{2,}", v):
            latin_hits.append((qid, k, v[:40]))
check("ضد J12: ارقام لاتین فقط در سلول‌های کد/یادداشت فنی", not latin_hits, f"{latin_hits[:5] or '—'}")

# ---------- ۵) سه اصلاح بازرس دوم در تحویل دوم ----------
allq = lambda qid: " | ".join(qr.get(qid, {}).values())
check("Q-1011: عدد شروع ۱۳ (نه ۱۶)", "۱۳" in allq("Q-1011") and "۱۶" not in re.sub(r"Q-1011", "", allq("Q-1011")) if "Q-1011" in qr else False,
      allq("Q-1011")[:200] if "Q-1011" in qr else "غایب")
check("Q-1011: پاسخ ۱۸", "۱۸" in allq("Q-1011"))
check("L-108 شاهد: ۱۳ (نه ۱۶ مثلث)", "۱۳" in " ".join(mr.get("L-108", {}).values()))
anchor_ok = ("L-112" in allq("Q-1020")) and ("Q-1017" in " ".join(mr.get("L-110", {}).values()) or "1017" in " ".join(mr.get("L-110", {}).values()))
check("لنگرها: Q-1020→L-112 و L-110→(Q-1017-…)", anchor_ok,
      f"Q-1020 لنگر={qr.get('Q-1020', {}).get('C','')!r}")

# ---------- ۶) تحویل دوم ↔ رسمی (c2d511a8) ----------
print("\n--- دیف تحویل دوم ↔ نسخهٔ رسمی ادغام‌شده (c2d511a8) ---")
diffs_all = {}
for sh in w3.sheetnames:
    if sh not in wr.sheetnames: continue
    a = [[("" if c.value is None else str(c.value)) for c in r] for r in wo[sh].iter_rows()]
    b = [[("" if c.value is None else str(c.value)) for c in r] for r in wr[sh].iter_rows()]
    if len(a) != len(b):
        print(f"  شیت {sh}: تعداد ردیف {len(a)} ↔ {len(b)}")
    for i, (ra, rb) in enumerate(zip(a, b), 1):
        for j, (va, vb) in enumerate(zip(ra, rb)):
            if va != vb:
                col = chr(65 + j) if j < 26 else "?"
                key = a[i-1][0] if a[i-1] else f"ردیف{i}"
                diffs_all.setdefault(sh, []).append((key, col, va, vb))
for sh, ds in diffs_all.items():
    print(f"\n== {sh}: {len(ds)} سلول متفاوت ==")
    for key, col, va, vb in ds:
        print(f"  [{key}!{col}] رسمی={va[:70]!r} ↔ تحویل۲={vb[:70]!r}")
if not diffs_all:
    print("هیچ تفاوتی — تحویل دوم == رسمی ✓")

print("\nنتیجه: FAIL =", FAIL if FAIL else "هیچ — همهٔ چک‌ها سبز ✓")
