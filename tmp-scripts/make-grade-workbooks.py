#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""موتور تألیف — تولید ورک‌بوک‌های خالی پایه‌های ۱/۲/۴/۵/۶ از قالب استاندارد v1.1
هم‌خوان با مبدل cartridge-converter:
  - سطر «یادداشت» بلافاصله زیر هدر می‌نشیند (مبدل آنجا می‌ایستد و ردیف نمونه را نمی‌خواند)
  - ردیف‌های نمونه زیر سطر یادداشت نگه داشته می‌شوند (الگوی آموزشی)
"""
import shutil
import openpyxl

TEMPLATE = 'download/cartridge-template-standard-v1.1.xlsx'
OUTDIR = 'public/math-game/authoring'

GRADE_WORDS = {1: 'اول', 2: 'دوم', 4: 'چهارم', 5: 'پنجم', 6: 'ششم'}
FAMILY = {
    1: 'شمارش و ساخت عدد — جدید (پیش‌فرض جدول پایه‌های موتور تألیف؛ نهایی پس از S1)',
    2: 'بسته‌های ده‌تایی + آغاز گروه‌های برابر (پیش‌فرض؛ نهایی پس از S1)',
    4: 'تقسیم طولانی + کسر — جدید (پیش‌فرض؛ نهایی پس از S1)',
    5: 'کسر/اعشار — جدید (پیش‌فرض؛ نهایی پس از S1)',
    6: 'نسبت/معادله — جدید (پیش‌فرض؛ نهایی پس از S1)',
}
ENGINE_DEV = {
    1: 'زیاد — موتور فعلی ضرب‌محور است (پیش‌فرض جدول پایه‌ها؛ نهایی پس از S1)',
    2: 'متوسط — بسته‌های ده‌تایی و آغاز ضرب (پیش‌فرض؛ نهایی پس از S1)',
    4: 'زیاد — کسر/زاویه/تقسیم طولانی (پیش‌فرض؛ نهایی پس از S1)',
    5: 'زیاد — کسر/اعشار/حجم (پیش‌فرض؛ نهایی پس از S1)',
    6: 'زیاد — نسبت/عدد منفی/معادله (پیش‌فرض؛ نهایی پس از S1)',
}

# مقدار هر فیلد شناسنامه (بر اساس برچسب فیلد)
def identity_values(n):
    w = GRADE_WORDS[n]
    return {
        'شناسهٔ خشاب': f'KH-G{n}-MATH',
        'عنوان فارسی خشاب': f'کارگاه ریاضی پایهٔ {w}',
        'درس': 'ریاضی',
        'پایه': w,
        'بستهٔ تجاری': f'بستهٔ پایهٔ {w} (واحد فروش = بستهٔ پایه، نه خشاب جدا)',
        'کد کتاب درسی (سازمان پژوهش)': '؟ — فقط از روی کتاب فیزیکی/سامانهٔ سازمان پر شود (هرگز حدس نزن)',
        'عنوان دقیق کتاب + شمارهٔ چاپ + سال': f'ریاضی {w} دبستان — شمارهٔ چاپ: ؟ — سال: ؟ (از روی کتاب فیزیکی پر شود؛ بدون چاپ، تأیید صفحه بی‌اعتبار است)',
        'خانوادهٔ موتور': FAMILY[n],
        'نیاز به توسعهٔ موتور': ENGINE_DEV[n],
        'وضعیت خشاب': f'طرح — v0.1 (ورک‌بوک خالی خط تولید موتور تألیف)',
        'معلم هم‌ساز / اعتبارسنج': 'باز — منتظر معلمِ در دسترس کتاب فیزیکی',
        'تاریخ آخرین تأیید سبز': '—',
        'یادداشت حق نشر': 'لنگرِ صفحه و هدف آزاد؛ کپی متن/تصویر/تمرین کتاب ممنوع — محتوا خودساخته',
        'یادداشت آزاد': 'پُرکردن با «موتور تألیف خشاب v۱.۰» — سند: public/math-game/authoring/MOTOR-TALIF-KHSHAB-v1.md',
    }

NOTE_GUIDE = ('ردیف واقعی را «بالای این سطر» جای‌گذاری کن؛ ردیف‌های نمونهٔ زیرِ این سطر پاک نشوند '
              '(مبدل نمی‌خواندشان) — پس از اولین تأیید معلم می‌توانی نمونه‌ها را پاک کنی.')

# چیدمان تازهٔ شیت‌های محتوا: هدر → یادداشت → نمونه‌ها
SHEETS = ['نقشهٔ کتاب (S1)', 'اطلس خطا', 'بانک سؤال']

def restructure(ws):
    rows = []
    for r in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        rows.append(list(r))
    # پیدا کردن هدر و سطر یادداشت
    head_idx = None; note_idx = None
    for i, r in enumerate(rows):
        a = str(r[0]).strip() if r[0] is not None else ''
        if head_idx is None and (a.startswith('شناسه (L-)') or a.startswith('کد (E-)') or a.startswith('شناسه (Q-)')):
            head_idx = i
        if a == 'یادداشت':
            note_idx = i
    assert head_idx is not None and note_idx is not None and note_idx > head_idx, ws.title
    header = rows[head_idx]
    note_row = rows[note_idx]
    samples = [r for i, r in enumerate(rows) if head_idx < i < note_idx and any(v is not None for v in r)]
    # بازنویسی: هدر | یادداشت(با راهنما) | نمونه‌ها | (تهی)
    inner = [note_row[0], NOTE_GUIDE] + list(note_row[2:])
    new = [header, inner] + samples
    for i, rowvals in enumerate(new):
        for c, v in enumerate(rowvals):
            ws.cell(row=head_idx + 1 + i, column=c + 1, value=v)
    # پاک کردن سطرهای پس از بخش جدید تا پایان
    for r in range(head_idx + 1 + len(new), ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            ws.cell(row=r, column=c, value=None)

def build(n):
    src = TEMPLATE
    dst = f'{OUTDIR}/khshab-G{n}-MATH_v0.1.xlsx'
    shutil.copyfile(src, dst)
    wb = openpyxl.load_workbook(dst)
    # شناسنامه
    ws = wb['شناسنامهٔ خشاب']
    vals = identity_values(n)
    filled = 0
    for r in range(1, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        if a is None:
            continue
        key = str(a).strip()
        if key in vals:
            ws.cell(row=r, column=2, value=vals[key])
            ws.cell(row=r, column=3, value='نیاز به تأیید')
            ws.cell(row=r, column=4, value='پیش‌پُر موتور تألیف v۱.۰')
            filled += 1
    assert filled == len(vals), f'شناسنامهٔ پایهٔ {n}: فقط {filled}/{len(vals)} فیلد'
    # شیت‌های محتوا
    for name in SHEETS:
        restructure(wb[name])
    # راهنما — ردیف اطلاع‌رسانی
    ws = wb['راهنما']
    r = ws.max_row + 2
    ws.cell(row=r, column=1, value='⚙️ این ورک‌بوک با «موتور تألیف خشاب v۱.۰» ساخته شده است')
    ws.cell(row=r, column=2, value='سند کامل: public/math-game/authoring/MOTOR-TALIF-KHSHAB-v1.md — پُرکردن با پرامپت اجرایی بخش ۱۲ همان سند؛ ردیف واقعی فقط بالای سطر «یادداشت».')
    wb.save(dst)
    wb.close()
    return dst

for n in [1, 2, 4, 5, 6]:
    out = build(n)
    print('✓ ساخته شد:', out)

# کپی در download/ برای هم‌خوانی با رسوم تحویل + نسخهٔ سند موتور
for n in [1, 2, 4, 5, 6]:
    shutil.copyfile(f'{OUTDIR}/khshab-G{n}-MATH_v0.1.xlsx', f'download/khshab-G{n}-MATH_v0.1.xlsx')
print('✓ نسخه‌های download/ هم کپی شدند')
