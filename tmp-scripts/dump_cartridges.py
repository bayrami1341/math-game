# -*- coding: utf-8 -*-
"""Dump کامل سه فایل خشاب برای ممیزی (خروجی متن تحلیل‌پذیر)."""
from openpyxl import load_workbook

FILES = [
    "/home/z/my-project/download/khshab-G3-MATH_v0.5.xlsx",
    "/home/z/my-project/download/cartridge-template-standard-v1.xlsx",
    "/home/z/my-project/download/cartridge-registry-18-v1.xlsx",
]

FILL_NAMES = {
    "C6EFCE": "سبز(OK)", "FFEB9C": "زرد(PEND)", "FFC7CE": "قرمز(ERR)",
    "E2EFDA": "سبزفام(NEW)", "548235": "هدر", "2F4F2F": "عنوان",
}


def fa2en(s):
    if not isinstance(s, str):
        return s
    table = str.maketrans("۰۱۲۳۴۵۶۷۸۹", "0123456789")
    return s.translate(table)


def cellfill(c):
    try:
        rgb = c.fill.start_color.rgb
        if isinstance(rgb, str) and rgb[-6:] in FILL_NAMES:
            return FILL_NAMES[rgb[-6:]]
    except Exception:
        pass
    return ""


for path in FILES:
    wb = load_workbook(path)
    print("=" * 100)
    print("FILE:", path.split("/")[-1])
    for ws in wb.worksheets:
        print("-" * 90)
        print(f"SHEET: {ws.title!r}  dims={ws.dimensions}  RTL={ws.sheet_view.rightToLeft}  maxR={ws.max_row} maxC={ws.max_column}")
        merged = [str(m) for m in ws.merged_cells.ranges]
        if merged:
            print("MERGED:", ", ".join(merged))
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
            vals = []
            for c in row:
                v = c.value
                if v is None:
                    continue
                tag = cellfill(c)
                s = fa2en(str(v)).replace("\n", " ⏎ ")
                vals.append(f"[{c.coordinate}{'|' + tag if tag else ''}] {s}")
            if vals:
                print(" || ".join(vals))
print("=" * 100)
print("DUMP DONE")
