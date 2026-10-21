#!/usr/bin/env python3
"""راستی‌آزمایی ساختاری v0.3.1: بانک سؤال بر اساس Q-id، نقشه بر اساس L-id، شناسنامه"""
import openpyxl
from collections import Counter

V031 = "/tmp/v031.xlsx"
wb = openpyxl.load_workbook(V031, data_only=False)

# ---------- بانک سؤال: keyed by Q-id in column A ----------
ws = wb["بانک سؤال"]
qs = {}
order = []
for row in ws.iter_rows():
    aid = row[0].value
    if isinstance(aid, str) and aid.startswith("Q-"):
        order.append(aid)
        qs[aid] = {c.coordinate[0]: c.value for c in row if c.value is not None}

print(f"تعداد سؤال: {len(qs)} | یکتا: {len(set(qs))} | ترتیب: {order[0]}..{order[-1]}")
dups = [k for k, v in Counter(order).items() if v > 1]
print("تکراری:", dups or "—")
missing = [f"Q-{n}" for n in range(1001, 1041) if f"Q-{n}" not in qs]
print("جامانده:", missing or "—")

# کدهای مجاز
ALLOWED = {"E_SUB_CONFUSION", "E_ADD_CONFUSION", "E_SKIP_UNDERCOUNT", "E_DOUBLE_OFF", "E_GROUPS"}
import re
bad_codes = {}
for qid, cells in qs.items():
    L = cells.get("L", "")
    for m in re.findall(r"E_[A-Z_]+", str(L)):
        if m not in ALLOWED:
            bad_codes.setdefault(qid, []).append(m)
print("کد بیرون از سجل:", bad_codes or "—")

# ضد J12 در همه
no_j12 = [qid for qid, c in qs.items() if "J12" not in str(c.get("Q", ""))]
print("بدون ضد J12:", no_j12 or "—")

# توازن پاسخ‌ها (K = پاسخ صحیح الف/ب/ج/د)
ans = Counter(str(c.get("K", "")) for c in qs.values())
print("توازن پاسخ:", dict(ans))

# لایهٔ CPA (ستون D)
cpa = Counter(str(c.get("D", "")) for c in qs.values())
print("لایهٔ CPA:", dict(cpa))
bad_cpa = [qid for qid, c in qs.items() if str(c.get("D", "")) not in ("مجسم", "تصویری", "علامتی")]
print("CPA بیرون سه‌گانه:", bad_cpa or "—")

# سؤال‌های بازنویسی‌شده: لنگر درس (B)
anchors = {}
for qid, c in qs.items():
    anchors.setdefault(str(c.get("B", "")), []).append(qid)
for a in sorted(anchors):
    print(f"  {a}: {len(anchors[a])} سؤال → {anchors[a]}")

# ---------- نقشهٔ کتاب S1: keyed by L-id ----------
wsm = wb["نقشهٔ کتاب (S1)"]
rows = {}
for row in wsm.iter_rows():
    aid = row[0].value
    if isinstance(aid, str) and aid.startswith("L-"):
        rows[aid] = {c.coordinate[0]: c.value for c in row if c.value is not None}
print(f"\nنقشه: {len(rows)} ردیف")
# L-114: عنوان صفحهٔ ۱۴
l114 = rows.get("L-114", {})
print("L-114 عنوان (D):", l114.get("D"))
print("L-114 هدف (E):", str(l114.get("E"))[:100])
l113 = rows.get("L-113", {})
print("L-113 عنوان (D):", l113.get("D"))
l115 = rows.get("L-115", {})
print("L-115 عنوان (D):", l115.get("D"))

# ---------- شناسنامه: شیت کامل ----------
wsi = wb["شناسنامهٔ خشاب"]
print("\nشناسنامه (همهٔ سلول‌های پر):")
for row in wsi.iter_rows(min_row=1, max_row=16):
    vals = [(c.coordinate, str(c.value)[:70]) for c in row if c.value is not None]
    if vals:
        print(" ", vals)
