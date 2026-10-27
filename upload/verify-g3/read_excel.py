# -*- coding: utf-8 -*-
"""G3 pack verification — stage 1: sheet list + راهنما + قاعده (read-only)."""
import openpyxl, json, sys

PATH = "/home/z/my-project/upload/G3-1403-RUNTIME.xlsx"
wb = openpyxl.load_workbook(PATH, data_only=True)

print("=== SHEET LIST (order) ===")
for i, name in enumerate(wb.sheetnames, 1):
    ws = wb[name]
    print(f"{i}. {name!r}  rows={ws.max_row}  cols={ws.max_column}")

def dump_sheet(name):
    ws = wb[name]
    print(f"\n=== SHEET {name!r} — full dump (rows={ws.max_row}, cols={ws.max_column}) ===")
    for r in range(1, ws.max_row + 1):
        vals = []
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v is None:
                vals.append("")
            else:
                vals.append(str(v).replace("\n", "⏎"))
        line = " | ".join(vals)
        if line.strip(" |"):
            print(f"[{r}] {line}")

if len(sys.argv) > 1 and sys.argv[1] == "rules":
    dump_sheet("قاعده")
else:
    dump_sheet("راهنما")
