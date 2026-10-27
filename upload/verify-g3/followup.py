# -*- coding: utf-8 -*-
"""G3 — follow-up: the extra «قطعی نبود» row, ۰۵۷/۳۸ rows, پدر columns consistency."""
import openpyxl, re
from collections import defaultdict

PATH = "/home/z/my-project/upload/G3-1403-RUNTIME.xlsx"
wb = openpyxl.load_workbook(PATH, data_only=True)
chk = wb["چک\u200cلیست"]; card = wb["کارت"]

def rowvals(ws, r):
    return [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]

print("=== ۱. ردیف‌هایی که تشخیص‌شان «قطعی نبود» دارد ولی ریاضی‌نیست=روشن‌نیست نیست ===")
father = {rowvals(card, r)[0]: rowvals(card, r)[3] for r in range(2, card.max_row + 1)}
for r in range(2, chk.max_row + 1):
    v = rowvals(chk, r)
    if isinstance(v[8], str) and "قطعی نبود" in v[8] and v[9] != "روشن نیست":
        print(f"  [{r}] id={v[0]} مهارت={v[4]!r} ریاضی‌نیست={v[9]!r}")
        for i, name in enumerate(["شناسه","بخش","صفحهٔ چاپ","کار","مهارت","کودک چه می‌کند","قبول","رد","تشخیص","ریاضی نیست","پدر","اگر این باز…","سقف","کارت تازه","وضعیت","منبع"]):
            if v[i] not in (None, ""):
                print(f"     {name}: {str(v[i])[:160]}")

print("\n=== ۲. ردیف‌های ۰۵۷ و P38-T1 و P140-T2 کامل ===")
for r in range(2, chk.max_row + 1):
    v = rowvals(chk, r)
    if v[0] in ("G3-1403-P26-T2", "G3-1403-P38-T1", "G3-1403-P140-T2", "G3-1403-P37-T2"):
        print(f"  --- [{r}] {v[0]}")
        for i, name in enumerate(["شناسه","بخش","صفحهٔ چاپ","کار","مهارت","کودک چه می‌کند","قبول","رد","تشخیص","ریاضی نیست","پدر","اگر این باز…","سقف","کارت تازه","وضعیت","منبع"]):
            if v[i] not in (None, ""):
                print(f"     {name}: {v[i]}")

print("\n=== ۳. ستون ۱۱ (پدر) و ستون ۱۲ (اگر این باز است این را نزن): هم‌ارزی با کارت‌پدر ===")
bad11, bad12, both = [], [], 0
n11 = n12 = 0
for r in range(2, chk.max_row + 1):
    v = rowvals(chk, r)
    f_of_skill = father.get(v[4]) if v[4] else None
    if v[10]:
        n11 += 1
        if v[10] != f_of_skill:
            bad11.append((v[0], v[4], v[10], f_of_skill))
    if v[11]:
        n12 += 1
        if v[11] != f_of_skill:
            bad12.append((v[0], v[4], v[11], f_of_skill))
    if v[10] and v[11]:
        both += 1
print(f"  col11 پدر filled={n11}  mismatch with father(مهارت) = {len(bad11)} {bad11[:3]}")
print(f"  col12 «اگر این باز…» filled={n12} mismatch = {len(bad12)} {bad12[:3]}")
print(f"  rows where both col11 and col12 filled = {both}")
diff = [(rowvals(chk, r)[0], rowvals(chk, r)[10], rowvals(chk, r)[11]) for r in range(2, chk.max_row + 1)
        if rowvals(chk, r)[10] != rowvals(chk, r)[11]]
print(f"  rows where col11 != col12 = {len(diff)}")
for d in diff[:10]:
    print(f"     {d}")

print("\n=== ۴. آمار مهارت‌ها: چند ردیف به هر کارت ===")
cnt = defaultdict(int)
for r in range(2, chk.max_row + 1):
    v = rowvals(chk, r)
    if v[4]:
        cnt[v[4]] += 1
print("  " + ", ".join(f"{k}={cnt[k]}" for k in sorted(cnt)))
one_task = [k for k, c in cnt.items() if c == 1]
print(f"  کارت‌های تک‌ردیفی ({len(one_task)}) = {sorted(one_task)}")

print("\n=== ۵. چک نهایی: وضعیت کارت باز/بسته و نشانهٔ باز شدن در هیچ ستونی ===")
import re as _re
hits = []
for name in ["راهنما", "قاعده", "چک\u200cلیست", "کارت", "نشست"]:
    ws = wb[name]
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and _re.search(r"(باز شد|بسته شد|باز کرد|بسته ماند|در صف)", v):
                hits.append((name, r, c, v[:80]))
print(f"  matches = {len(hits)}")
for h in hits[:12]:
    print("   ", h)
