# -*- coding: utf-8 -*-
"""G3 pack — full structural + content validation (read-only)."""
import openpyxl, re, json
from collections import Counter, defaultdict

PATH = "/home/z/my-project/upload/G3-1403-RUNTIME.xlsx"
wb = openpyxl.load_workbook(PATH, data_only=True)
S_CHK = "چک\u200cلیست"; S_CARD = "کارت"; S_SES = "نشست"

chk = wb[S_CHK]; card = wb[S_CARD]; ses = wb[S_SES]

def rows(ws):
    for r in range(2, ws.max_row + 1):
        yield r, [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]

P = []  # print results
def out(*a):
    s = " ".join(str(x) for x in a)
    P.append(s); print(s)

out("========== A. چک‌لیست ==========")
chk_rows = list(rows(chk))
out(f"A1. data rows = {len(chk_rows)}  (claim: 287)")

ids = [r[1][0] for r in chk_rows]
out(f"A2. unique ids = {len(set(ids))}/{len(ids)}")

pat = re.compile(r"^G3-1403-P(\d{2,3})-T(\d+)$")
bad = [i for i in ids if not pat.match(i or "")]
out(f"A3. ids not matching pattern G3-1403-P{{page}}-T{{task}} = {len(bad)} {bad[:5]}")

pages = [int(pat.match(i).group(1)) for i in ids]
tasks = [int(pat.match(i).group(2)) for i in ids]
out(f"A4. page range = {min(pages)}..{max(pages)}  (claim: lessons 7..150, bookends 151/152)")
out(f"A5. tasks on pages 7..150 = {sum(1 for p in pages if p <= 150)} ; bookend rows 151/152 = {sum(1 for p in pages if p >= 151)}")

# per-page contiguity & gaps
by_page = defaultdict(list)
for p, t in zip(pages, tasks):
    by_page[p].append(t)
missing_pages = [p for p in range(7, 151) if p not in by_page]
out(f"A6. missing pages in 7..150 = {missing_pages} (expect none)")
nonseq = []
for p, ts in by_page.items():
    if ts != list(range(1, len(ts) + 1)):
        nonseq.append((p, ts))
out(f"A7. pages whose task numbers are not 1..k = {len(nonseq)} {nonseq[:3]}")

status = [r[1][14] for r in chk_rows]
out(f"A8. وضعیت distribution = {dict(Counter(status))}  (claim: all نوشته‌شد)")

skill = [r[1][4] for r in chk_rows]
sk_rows = [(i, s) for i, s in zip(ids, skill) if s]
empty_sk = [(i, s) for i, s in zip(ids, skill) if not s]
out(f"A9. rows with مهارت = {len(sk_rows)} ; rows with empty مهارت = {len(empty_sk)}")
sk_vals = set(s for _, s in sk_rows)
out(f"A10. distinct مهارت codes = {len(sk_vals)}  (claim: 45 cards)")
out(f"     codes = {sorted(sk_vals)}")
pat_sk = re.compile(r"^SK-[A-Z0-9-]+$")
out(f"A11. مهارت codes violating SK- pattern = {[s for s in sk_vals if not pat_sk.match(s)]}")

rni = [r[1][9] for r in chk_rows]  # ریاضی نیست col
out(f"A12. ریاضی‌نیست col distribution = {dict(Counter(rni))}")
dn = [i for i, r in zip(ids, chk_rows) if r[1][9] == "روشن نیست"]
out(f"     rows marked روشن‌نیست = {len(dn)}  (claim: 54 rows number not definite)")
dz = [r[1][8] for r in chk_rows]  # تشخیص
cnt_noNum = sum(1 for d in dz if isinstance(d, str) and "قطعی نبود" in d)
out(f"A13. تشخیص contains «قطعی نبود» = {cnt_noNum} ; match with روشن‌نیست rows = {cnt_noNum == len(dn)}")
empty_diag = [i for i, d in zip(ids, dz) if not (d or "").strip()]
out(f"A14. rows with empty تشخیص = {len(empty_diag)} {empty_diag[:5]}")

# two no-card mistakes P78 / P80
p78 = [(i, s) for i, s in zip(ids, skill) if pat.match(i) and int(pat.match(i).group(1)) == 78]
p80 = [(i, s) for i, s in zip(ids, skill) if pat.match(i) and int(pat.match(i).group(1)) == 80]
out(f"A15. page 78 rows & their مهارت = {p78}")
out(f"     page 80 rows & their مهارت = {p80}")

r057 = [(i, s, chk_rows[k][1][5], chk_rows[k][1][8]) for k, (i, s) in enumerate(zip(ids, skill)) if "۰۵۷" in str(chk_rows[k][1][6] or "") + str(chk_rows[k][1][5] or "") + str(chk_rows[k][1][7] or "")]
out(f"A16. rows containing ۰۵۷ = {r057}")

r38 = [(i, s) for i, s in zip(ids, skill) if pat.match(i) and int(pat.match(i).group(1)) == 38]
out(f"A17. page 38 rows & their مهارت = {r38}  (راهنما: P38 روی حذف رقم ولی نزدیک‌تر؛ کارت تازه نساز)")

kt = [r[1][13] for r in chk_rows]  # کارت تازه
out(f"A18. کارت‌تازه distribution = {dict(Counter(kt))}")
az = [r[1][11] for r in chk_rows]  # اگر این باز است این را نزن
out(f"A19. col12 «اگر این باز است…» non-empty = {sum(1 for a in az if a)}")
sq = [r[1][12] for r in chk_rows]
out(f"A20. col13 «سقف» non-empty = {sum(1 for a in sq if a)}")
pd_ = [(i, r[1][10]) for i, r in zip(ids, chk_rows) if r[1][10]]
out(f"A21. col11 «پدر» non-empty count = {len(pd_)}  sample={pd_[:3]}")
src = Counter(r[1][15] for r in chk_rows)
out(f"A22. منبع values = {dict(src)}")

out("\n========== B. کارت ==========")
card_rows = list(rows(card))
out(f"B1. card rows = {len(card_rows)}  (claim: 45)")
cids = [r[1][0] for r in card_rows]
out(f"B2. unique card ids = {len(set(cids))}/{len(cids)}")
out(f"B3. card ids == مهارت codes in چک‌لیست ? {set(cids) == sk_vals}  diff={(set(cids) ^ sk_vals) if set(cids) != sk_vals else '—'}")
cstat = Counter(r[1][7] for r in card_rows)
out(f"B4. card وضعیت = {dict(cstat)}  (claim: no card opened/closed → all نوشته‌شد)")
father = {r[1][0]: r[1][3] for r in card_rows}
dangling = [(c, f) for c, f in father.items() if f and f not in set(cids)]
out(f"B5. dangling پدر refs = {dangling} (expect none)")
out(f"B6. پدر map = { {c: f for c, f in father.items() if f} }")
# cycles
def has_cycle():
    for c in father:
        seen = set(); x = c
        while father.get(x):
            if x in seen: return c
            seen.add(x); x = father[x]
    return None
out(f"B7. cycle in پدر chain = {has_cycle()}")
birth = {r[1][0]: r[1][5] for r in card_rows}
first_task_page = {}
for s, p in zip(skill, pages):
    if s:
        first_task_page[s] = min(first_task_page.get(s, 999), p)
mism = {c: (birth[c], first_task_page.get(c)) for c in birth if birth[c] != first_task_page.get(c)}
out(f"B8. card تولد != first-task page → {mism if mism else 'none (all 45 exact)'}")
fviol = []
for c, f in father.items():
    if f and birth[c] and birth[f] and int(birth[c]) < int(birth[f]):
        fviol.append((c, f, birth[c], birth[f]))
out(f"B9. «پدر دیرتر از فرزند به دنیا آمده» violations = {fviol} (expect none — قاعدهٔ ۳)")
qclose = [(r[1][0], r[1][6]) for r in card_rows]
out(f"B10. قاعدهٔ بستن empty = {[c for c, q in qclose if not (q or '').strip()]}")
out(f"     قاعدهٔ بستن duplicates = {[k for k, v in Counter(q for _, q in qclose).items() if v > 1]} (prompt: هر کارت جملهٔ بستن خودش)")
saq = [r[1][4] for r in card_rows]
out(f"B11. card «سقف» non-empty = {sum(1 for a in saq if a)}")
ntype = Counter(r[1][2] for r in card_rows)
out(f"B12. نوع‌کار vocab = {dict(ntype)}")
cnt_map = {}
for c, f in father.items():
    if f: cnt_map.setdefault(f, []).append(c)
out(f"B13. children per father = {cnt_map}")

out("\n========== C. نشست ==========")
ses_rows = list(rows(ses))
out(f"C1. session rows = {len(ses_rows)}  (claim: 287)")
nums = [r[1][0] for r in ses_rows]
out(f"C2. session numbers 1..{len(nums)} contiguous = {nums == list(range(1, len(nums) + 1))}")
sp = [r[1][1] for r in ses_rows]
out(f"C3. session pages sequence == چک‌لیست pages sequence = {sp == pages}")
after = Counter(r[1][5] for r in ses_rows)
out(f"C4. وضعیت بعد distribution = {dict(after)}  (claim: 226 یادداشت / 7 دست‌نخورده / 54 ریاضی‌نیست)")
ses_skill = [r[1][4] for r in ses_rows]
out(f"C5. نشست مهارت column == چک‌لیست مهارت column (same order) = {ses_skill == skill}")
rn = [r[1][3] for r in ses_rows]
out(f"C6. راهنما col distribution = {dict(Counter(rn))}")
res = Counter(r[1][6] for r in ses_rows)
out(f"C7. نتیجه vocab = {dict(res)}")
opened = [ (r[1][0], r[1][2]) for r in ses_rows if isinstance(r[1][6], str) and ("باز شد" in r[1][6] or "بسته شد" in r[1][6] or "باز کرد" in r[1][6] or "بست" in r[1][6])]
out(f"C8. sessions containing open/close verbs = {len(opened)} (claim: هیچ کارتی باز/بسته نشده)")

# detail lists
out("\n========== D. detail lists ==========")
dn_types = [(i, chk_rows[k][1][8][:40]) for k, (i, s) in enumerate(zip(ids, skill)) if not s and chk_rows[k][1][9] != "روشن نیست"]
out(f"D1. empty-مهارت rows NOT marked روشن‌نیست (real tasks w/o card) = {len(dn_types)}")
for x in dn_types: out("   ", x)
hand = [(r[1][0], r[1][1], r[1][2], r[1][6]) for r in ses_rows if r[1][5] == "دست نخورد"]
out(f"D2. the 7 «دست نخورد» sessions:")
for x in hand: out("   ", x)
out(f"D3. کارت‌تازه=بله rows = {[(i) for i, k in zip(ids, kt) if k and k != 'نه']}")
