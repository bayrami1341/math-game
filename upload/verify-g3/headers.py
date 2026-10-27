# -*- coding: utf-8 -*-
"""G3 — headers + samples of چک‌لیست / کارت / نشست."""
import openpyxl

PATH = "/home/z/my-project/upload/G3-1403-RUNTIME.xlsx"
wb = openpyxl.load_workbook(PATH, data_only=True)

for name in ["چک\u200cلیست", "کارت", "نشست"]:
    ws = wb[name]
    print(f"\n=== {name!r} HEADER (cols={ws.max_column}) ===")
    hdr = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    for i, h in enumerate(hdr, 1):
        print(f"  col{i}: {h!r}")
    print(f"--- sample rows 2..6 and last 3 ---")
    rows = list(range(2, 7)) + list(range(ws.max_row - 2, ws.max_row + 1))
    for r in rows:
        vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        vals = ["" if v is None else str(v).replace("\n", "⏎") for v in vals]
        print(f"[{r}] " + " | ".join(vals))
