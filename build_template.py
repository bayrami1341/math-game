# -*- coding: utf-8 -*-
"""قالب استاندارد استخراج موضوعات ریاضی دبستان G1..G6 — سازگار با خط لولهٔ خشاب"""
import sys, os, json

XLSX_SKILL_DIR = "/home/z/my-project/skills/xlsx"
for sub in [XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, "templates")]:
    if sub not in sys.path:
        sys.path.insert(0, sub)

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from base import (use_palette_explicit, setup_sheet, style_header_row, style_data_row,
                  font_caption, align_text)

use_palette_explicit("professional")

OUT = "/home/z/my-project/download/TEMPLATE-CURRICULUM-MAP-G1-G6-v1.xlsx"
MATRIX = "/home/z/my-project/public/math-game/authoring/DOMAIN-MATRIX.json"

wb = openpyxl.Workbook()
wb.properties.creator = "Z.ai"

def rtl(ws):
    ws.sheet_view.rightToLeft = True

# ---------- Sheet 1: راهنما ----------
ws = wb.active
ws.title = "راهنما"
rtl(ws)
setup_sheet(ws, title="قالب استاندارد نقشهٔ موضوعات ریاضی دبستان — پایه‌های ۱ تا ۶ (v1)", last_col=3)
guide = [
    ("این قالب چیست", "فرم استانداردِ «برون‌سپاری خواندن کتاب‌ها»: خروجی دستیار AI (مثل سوپرگراک) باید دقیقاً با ستون‌های دو شیت «نقشهٔ_درس‌ها» و «ماتریس_دامنه» پر شود."),
    ("ورک‌فلو بعد از تحویل", "۱) اعتبارسنجی ساختاری ۱۰۰٪ خودکار  ۲) نمونه‌گیری تصادفی ~۱۰٪ بازبینی مفهومی  ۳) ثبت با وضعیت «نیاز به تأیید»  ۴) تبدیل با مبدل رسمی و گیت پوشش  ۵) مُهر نهایی معلم/مالک. هیچ محتوای خارجی مستقیم در خشاب merge نمی‌شود."),
    ("قانون طلایی صداقت", "شمارهٔ صفحه را اگر دقیق نمی‌دانی «تقریبی» بنویس و ستون «سطح اطمینان صفحه» را پر کن؛ جعل شمارهٔ صفحه دقیق = رد کل بچ."),
    ("قاعدهٔ یک ردیف", "هر درس (یا زیرموضوع معنادار) یک ردیف — نه فصل‌به‌فصل. خانه‌ها را merge نکن؛ فرمول ننویس؛ فقط مقدار."),
    ("اعداد", "محتوای متنی با ارقام فارسی (۱۲۳). ستون‌های عددی (شمارهٔ فصل/درس، حداقل سؤال) با ارقام لاتین ساده تا اعتبارسنجی خودکار کار کند."),
    ("کدهای دامنه", "D1 تا D8 طبق واژه‌نامه؛ دامنهٔ جدید فقط با کد D9+ و justification در ستون یادداشت."),
    ("منبع مرجع", "شیت «نمونهٔ_مرجع» از ماتریس رسمی زندهٔ G2/G3 پر شده است — کیفیت و سبکِ نوع‌های تمرین را با آن تطبیق بده."),
    ("فایل خروجی مورد انتظار", "نام پیشنهادی: CURRICULUM-MAP-G1-G6-v{n}.xlsx — UTF-8، بدون محافظت صفحه، بدون سلول ادغام‌شده در ناحیهٔ داده."),
]
r = 4
ws.cell(row=r, column=2, value="عنوان"); ws.cell(row=r, column=3, value="شرح")
style_header_row(ws, r, 2, 3)
for i, (k, v) in enumerate(guide):
    r += 1
    ws.cell(row=r, column=2, value=k); ws.cell(row=r, column=3, value=v)
    style_data_row(ws, r, 2, 3, i)
    ws.cell(row=r, column=2).font = font_caption()
    ws.cell(row=r, column=2).alignment = align_text()
    ws.cell(row=r, column=3).alignment = align_text()
ws.column_dimensions["B"].width = 24
ws.column_dimensions["C"].width = 110

# ---------- Sheet 2: نقشهٔ درس‌ها (main) ----------
ws = wb.create_sheet("نقشهٔ_درس‌ها")
rtl(ws)
HEAD = ["پایه", "شمارهٔ فصل", "عنوان فصل", "شمارهٔ درس", "عنوان درس / موضوع",
        "صفحات (شروع–پایان)", "سطح اطمینان صفحه", "زیرموضوع‌ها (؛ جدا)",
        "خانوادهٔ محتوایی", "کد دامنهٔ پیشنهادی", "نوع‌های تمرین (؛ جدا)",
        "حداقل سؤال پیشنهادی", "توزیع سختی", "گلوگاه‌های احتمالی", "یادداشت"]
setup_sheet(ws, title="نقشهٔ درس‌ها — یک ردیف به‌ازای هر درس/موضوع (۶ پایه)", last_col=len(HEAD) + 1)
for c, h in enumerate(HEAD, start=2):
    ws.cell(row=4, column=c, value=h)
style_header_row(ws, 4, 2, len(HEAD) + 1)
widths = [8, 11, 22, 10, 34, 15, 14, 30, 16, 14, 34, 12, 12, 28, 24]
for i, w in enumerate(widths):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i + 2)].width = w
ws.freeze_panes = "B5"
MAXROW = 500
dvs = [
    ("B", '"G1,G2,G3,G4,G5,G6"'),
    ("H", '"کامل,تقریبی,نامطمئن"'),
    ("J", '"عدد و محاسبات,کسر,هندسه,اندازه‌گیری,آمار و احتمال,الگو و تقارن"'),
    ("K", '"D1,D2,D3,D4,D5,D6,D7,D8,D9,D10"'),
    ("N", '"آسان,متوسط,دشوار,۴۰/۴۰/۲۰"'),
]
for col, f in dvs:
    dv = DataValidation(type="list", formula1=f, allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv)
    dv.add(f"{col}5:{col}{MAXROW}")

# ---------- Sheet 3: ماتریس دامنه ----------
ws = wb.create_sheet("ماتریس_دامنه")
rtl(ws)
HEAD3 = ["پایه", "کد دامنه", "عنوان دامنه", "فصل کتاب", "حداقل سؤال",
         "نوع‌های تمرین (؛ جدا)", "یادداشت عمق"]
setup_sheet(ws, title="ماتریس دامنه — یک ردیف به‌ازای هر دامنه در هر پایه", last_col=len(HEAD3) + 1)
for c, h in enumerate(HEAD3, start=2):
    ws.cell(row=4, column=c, value=h)
style_header_row(ws, 4, 2, len(HEAD3) + 1)
for i, w in enumerate([8, 10, 26, 10, 12, 60, 36]):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i + 2)].width = w
ws.freeze_panes = "B5"
for col, f in [("B", '"G1,G2,G3,G4,G5,G6"'), ("C", '"D1,D2,D3,D4,D5,D6,D7,D8,D9,D10"')]:
    dv = DataValidation(type="list", formula1=f, allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv)
    dv.add(f"{col}5:{col}200")

# ---------- Sheet 4: نمونهٔ مرجع (زنده از DOMAIN-MATRIX) ----------
ws = wb.create_sheet("نمونهٔ_مرجع")
rtl(ws)
HEAD4 = ["خشاب", "کد دامنه", "عنوان دامنه", "فصل", "حداقل سؤال", "نوع‌های تمرین (؛ جدا)"]
setup_sheet(ws, title="نمونهٔ مرجع — ماتریس رسمی زندهٔ G3 و G2 (کیفیت ملاک)", last_col=len(HEAD4) + 1)
for c, h in enumerate(HEAD4, start=2):
    ws.cell(row=4, column=c, value=h)
style_header_row(ws, 4, 2, len(HEAD4) + 1)
for i, w in enumerate([16, 10, 26, 8, 11, 90]):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i + 2)].width = w
ws.freeze_panes = "B5"
with open(MATRIX) as f:
    m = json.load(f)
r = 4
idx = 0
for fam in ["KH-G3-MATH", "KH-G2-MATH"]:
    fd = m["families"][fam]
    for code in fd["domains"]:
        d = fd["domains"][code]
        r += 1
        vals = [fam, code, d["title"], d.get("chapter", ""), d.get("min", ""),
                "؛ ".join(d.get("qTypes", []))]
        for c, v in enumerate(vals, start=2):
            ws.cell(row=r, column=c, value=v)
        style_data_row(ws, r, 2, len(HEAD4) + 1, idx)
        ws.cell(row=r, column=7).alignment = align_text()
        idx += 1

# ---------- Sheet 5: واژه‌نامه ----------
ws = wb.create_sheet("واژه‌نامه")
rtl(ws)
setup_sheet(ws, title="واژه‌نامه و قواعد اعتبارسنجی", last_col=3)
terms = [
    ("D1..D8 — قاعدهٔ نقشه‌بندی", "الگو/تقارن→D1 · عددنویسی و ارزش مکانی→D2 · جمع و تفریق→D3 · ضرب و تقسیم→D4 · کسر→D5 · هندسه (شکل/محیط/مساحت)→D6 · آمار و احتمال→D7 · دامنهٔ محاسباتی اضافی پایه (مثل ضرب عددهای چندرقمی)→D8. هر پایه کدها را به ترتیب فصل‌های کتاب خودش می‌گیرد."),
    ("۶ قالب پاسخ موتور", "چهارگزینه‌ای · گزینه‌ای · عدد (پاسخ کوتاه) · کسر/رنگ‌آمیزی · محور/زوج‌سازی · شکل/نمودار. نوع تمرین هر ردیف باید به یکی از این قالب‌ها قابل تبدیل باشد."),
    ("توزیع سختی", "پیش‌فرض بچ‌ها: ۴۰٪ آسان / ۴۰٪ متوسط / ۲۰٪ دشوار. در ستون توزیع سختی فقط یکی از: آسان، متوسط، دشوار، ۴۰/۴۰/۲۰."),
    ("سطح اطمینان صفحه", "کامل = مطمئن از صفحهٔ دقیق کتاب · تقریبی = نزدیک ولی نه قطعی · نامطمئن = فقط ساختار فصل. گزارش پوشش «شاهد صفحه‌ای» فقط از ردیف‌های کامل/تقریبی ساخته می‌شود."),
    ("وضعیت تأیید", "هر محتوای ورودی (چه از AI چه از ربات‌ها) با «نیاز به تأیید» ثبت می‌شود؛ «تأییدشده» فقط با مُهر معلم/مالک."),
    ("حداقل سؤال", "مرجع فعلی: G2 = 242 (۶ دامنه) · G3 = 452 (۸ دامنه). برای پایه‌های دیگر پیشنهاد خودت با همین منطق وزن‌دهی؛ نهایی‌اش با مُهر مالک."),
    ("ممنوع", "سلول ادغام‌شده در ناحیهٔ داده · فرمول · ارقام فارسی در ستون‌های عددی · ردیف خالی وسط داده · تغییر نام ستون‌ها یا جابه‌جایی ترتیب‌شان."),
]
r = 4
ws.cell(row=r, column=2, value="اصطلاح / قاعده"); ws.cell(row=r, column=3, value="تعریف")
style_header_row(ws, r, 2, 3)
for i, (k, v) in enumerate(terms):
    r += 1
    ws.cell(row=r, column=2, value=k); ws.cell(row=r, column=3, value=v)
    style_data_row(ws, r, 2, 3, i)
    ws.cell(row=r, column=3).alignment = align_text()
ws.column_dimensions["B"].width = 26
ws.column_dimensions["C"].width = 115

wb.save(OUT)
print("saved:", OUT)

# ---------- QA ----------
import subprocess
res = subprocess.run(["python3", "/home/z/my-project/skills/xlsx/xlsx.py", "validate", OUT],
                     capture_output=True, text=True)
print(res.stdout[-1500:], res.stderr[-300:])
