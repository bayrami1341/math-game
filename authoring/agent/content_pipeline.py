"""Git-versioned, source-bound handoff contracts for the three content stages.

The repository stores source metadata, structured page observations, valuation
records, and draft game-item references. Original PDFs, page images/crops, long
verbatim quotations, and private reader-pass packets are not copied into Git.

A PASS from this module means only that artifact links and stage boundaries are
structurally consistent. This module never creates a teacher verdict, edits an
existing workbook, changes a color, or authorizes release.
"""

from __future__ import annotations

import json
import math
import re
from functools import lru_cache
from pathlib import Path
from typing import Any

from .models import file_sha256

SOURCE_SCHEMA = "khshab.content-source/v1"
EXTRACTION_SCHEMA = "khshab.extraction-ledger/v1"
VALUATION_SCHEMA = "khshab.valuation-bundle/v1"
GAME_PACK_SCHEMA = "khshab.game-content-pack/v1"
REPORT_SCHEMA = "khshab.content-pipeline-report/v1"
_HEX64 = re.compile(r"^[0-9a-f]{64}$")
_TOKEN = re.compile(r"^[A-Za-z][A-Za-z0-9_.:-]{0,127}$")
_ALLOWED_EVIDENCE_KINDS = {"pdf_text_bbox", "pdf_crop_150dpi", "pdf_page_render"}
_ALLOWED_ELEMENT_KINDS = {"TASK", "EXAMPLE", "DEFINITION", "RULE", "VISUAL", "INSTRUCTION", "OTHER"}
_ALLOWED_REVIEW_STATES = {"DRAFT", "REQUIRES_HUMAN", "HUMAN_APPROVED", "REJECTED"}
_STAGE_SCHEMAS = {
    "source": "content-source-v1.schema.json",
    "offset_proof": "offset-proof-v1.schema.json",
    "extraction": "extraction-ledger-v1.schema.json",
    "valuation": "valuation-bundle-v1.schema.json",
    "game_pack": "game-content-pack-v1.schema.json",
    "verdict_ledger": "verdict-ledger-v1.schema.json",
}
_FORBIDDEN_CONTENT_KEYS = {
    "raw_text", "raw_reading", "page_text", "ocr_text", "ocr_output", "transcript",
    "transcription", "quote", "quoted_text", "verbatim_text", "image_bytes", "image_data",
    "image_base64", "image_path", "crop_path", "pdf_bytes", "pdf_base64", "crop_bytes",
}


def _issue(code: str, **details: Any) -> dict[str, Any]:
    return {"code": code, **details}


def _is_within(path: Path, parent: Path) -> bool:
    try:
        path.relative_to(parent)
        return True
    except ValueError:
        return False


def _read_json(path: Path, label: str) -> tuple[dict[str, Any] | None, dict[str, Any] | None]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        return None, _issue(f"{label}_UNREADABLE", exception_type=type(exc).__name__)
    if not isinstance(value, dict):
        return None, _issue(f"{label}_ROOT_NOT_OBJECT")
    return value, None


def _privacy_issues(value: Any, *, label: str, path: str = "$") -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    if isinstance(value, dict):
        for key, child in value.items():
            key_text = str(key).strip().lower()
            child_path = f"{path}.{key_text}"
            if key_text in _FORBIDDEN_CONTENT_KEYS or any(
                marker in key_text for marker in ("base64", "image_path", "crop_path", "raw_text", "raw_read", "ocr")
            ):
                issues.append(_issue("GIT_CONTENT_FIELD_FORBIDDEN", artifact=label, field_path=child_path))
            issues.extend(_privacy_issues(child, label=label, path=child_path))
    elif isinstance(value, list):
        for index, child in enumerate(value):
            issues.extend(_privacy_issues(child, label=label, path=f"{path}[{index}]"))
    elif isinstance(value, str) and len(value.split()) > 10:
        issues.append(_issue("GIT_LONG_TEXT_FIELD_FORBIDDEN", artifact=label, field_path=path))
    return issues


def _token(value: Any) -> bool:
    return isinstance(value, str) and _TOKEN.fullmatch(value.strip()) is not None


def _sha(value: Any) -> bool:
    return isinstance(value, str) and _HEX64.fullmatch(value.lower()) is not None


def _safe_excerpt(value: Any) -> bool:
    return isinstance(value, str) and len(value.split()) <= 10


@lru_cache(maxsize=1)
def _schema_validator_resources() -> tuple[Any, dict[str, dict[str, Any]], Any]:
    """Load and check the versioned Draft 2020-12 contracts and local references."""
    from jsonschema import Draft202012Validator
    from referencing import Registry, Resource

    schema_dir = Path(__file__).resolve().parent / "schemas"
    schemas: dict[str, dict[str, Any]] = {}
    resources: list[tuple[str, Any]] = []
    for path in sorted(schema_dir.glob("*.schema.json")):
        schema = json.loads(path.read_text(encoding="utf-8"))
        # Absolute file URIs make the sibling-schema $refs deterministic in installs/tests.
        schema["$id"] = path.resolve().as_uri()
        Draft202012Validator.check_schema(schema)
        schemas[path.name] = schema
        resources.append((path.resolve().as_uri(), Resource.from_contents(schema)))
    return Draft202012Validator, schemas, Registry().with_resources(resources)


def _schema_issues(document: dict[str, Any], *, artifact: str, schema_filename: str) -> list[dict[str, Any]]:
    try:
        from jsonschema import FormatChecker

        validator_type, schemas, registry = _schema_validator_resources()
        schema = schemas[schema_filename]
        validator = validator_type(schema, registry=registry, format_checker=FormatChecker())
        errors = sorted(
            validator.iter_errors(document),
            key=lambda error: tuple(str(part) for part in error.absolute_path),
        )
    except Exception as exc:
        return [_issue("PIPELINE_SCHEMA_VALIDATION_UNAVAILABLE", artifact=artifact, exception_type=type(exc).__name__)]
    return [
        _issue(
            "PIPELINE_ARTIFACT_SCHEMA_INVALID",
            artifact=artifact,
            schema=schema_filename,
            field_path="/" + "/".join(str(part) for part in error.absolute_path),
            validator=str(error.validator),
        )
        for error in errors[:100]
    ]


def _report(
    *,
    status: str,
    reason_code: str,
    issues: list[dict[str, Any]],
    binding: dict[str, Any] | None = None,
    hashes: dict[str, str] | None = None,
    metrics: dict[str, Any] | None = None,
) -> dict[str, Any]:
    return {
        "schema_version": REPORT_SCHEMA,
        "gate_id": "CONTENT_PIPELINE_TRACE",
        "status": status,
        "reason_code": reason_code,
        "traceability_only": True,
        "source_binding": binding or {},
        "artifact_hashes": hashes or {},
        "metrics": metrics or {},
        "issues": issues[:250],
        "safeguards": {
            "human_verdict_created": False,
            "workbook_modified": False,
            "color_change_authorized": False,
            "section3_release_authorized": False,
        },
    }


def _source_binding(source: dict[str, Any]) -> dict[str, Any]:
    return {
        "source_id": source.get("source_id"),
        "book_code": source.get("book_code"),
        "edition_id": source.get("edition_id"),
        "manifest_id": source.get("manifest_id"),
        "manifest_sha256": source.get("manifest_sha256"),
        "source_pdf_sha256": source.get("source_pdf_sha256"),
    }


def _check_binding(document: dict[str, Any], expected: dict[str, Any], label: str) -> list[dict[str, Any]]:
    observed = document.get("source_binding")
    if not isinstance(observed, dict):
        return [_issue(f"{label}_SOURCE_BINDING_MISSING")]
    return [
        _issue(f"{label}_SOURCE_BINDING_MISMATCH", field=key)
        for key, value in expected.items()
        if observed.get(key) != value
    ]


def _validate_source(source: dict[str, Any]) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    if source.get("schema_version") != SOURCE_SCHEMA:
        issues.append(_issue("SOURCE_SCHEMA_UNSUPPORTED"))
    for key in ("source_id", "book_code", "subject_id", "edition_id", "manifest_id"):
        if not _token(source.get(key)):
            issues.append(_issue("SOURCE_FIELD_INVALID", field=key))
    for key in ("manifest_sha256", "source_pdf_sha256"):
        if not _sha(source.get(key)):
            issues.append(_issue("SOURCE_SHA_INVALID", field=key))
    for key in ("grade", "source_pdf_size_bytes", "pdf_page_count"):
        if type(source.get(key)) is not int or source[key] < 1:
            issues.append(_issue("SOURCE_INTEGER_INVALID", field=key))
    rights = source.get("rights_status")
    if rights not in {"AUTHORIZED", "RESTRICTED", "UNVERIFIED", "SYNTHETIC_TEST_ONLY"}:
        issues.append(_issue("SOURCE_RIGHTS_STATUS_INVALID"))

    asset = source.get("source_asset")
    if not isinstance(asset, dict) or asset.get("storage") != "PRIVATE_EXTERNAL" or not isinstance(asset.get("locator_ref"), str) or not asset["locator_ref"].strip():
        issues.append(_issue("SOURCE_ASSET_MUST_BE_EXTERNAL_REFERENCE"))

    pagination = source.get("pagination")
    if not isinstance(pagination, dict):
        return issues + [_issue("SOURCE_PAGINATION_INVALID")]
    keys = (
        "body_pdf_page_start", "body_pdf_page_end", "book_page_start", "book_page_end",
        "pdf_minus_book_offset",
    )
    if any(type(pagination.get(key)) is not int for key in keys):
        issues.append(_issue("SOURCE_PAGINATION_BOUNDS_INVALID"))
    elif type(source.get("pdf_page_count")) is not int:
        issues.append(_issue("SOURCE_PAGE_COUNT_INVALID"))
    else:
        pdf_start = pagination["body_pdf_page_start"]
        pdf_end = pagination["body_pdf_page_end"]
        book_start = pagination["book_page_start"]
        book_end = pagination["book_page_end"]
        offset = pagination["pdf_minus_book_offset"]
        if (
            pdf_start < 1 or pdf_end < pdf_start or pdf_end > source["pdf_page_count"]
            or book_start < 1 or book_end < book_start
            or pdf_start - offset != book_start or pdf_end - offset != book_end
            or pdf_end - pdf_start != book_end - book_start
        ):
            issues.append(_issue("SOURCE_OFFSET_MAPPING_INVALID"))
    if pagination.get("offset_proof_status") not in {"TWO_READER_MATCHED", "SYNTHETIC_TEST_ONLY"}:
        issues.append(_issue("SOURCE_OFFSET_PROOF_NOT_VERIFIED"))
    if not _sha(pagination.get("offset_proof_sha256")):
        issues.append(_issue("SOURCE_OFFSET_PROOF_SHA_INVALID"))
    manifest_path = source.get("manifest_path")
    if (
        not isinstance(manifest_path, str) or not manifest_path.strip() or "\x00" in manifest_path
        or Path(manifest_path).is_absolute() or ".." in Path(manifest_path).parts
    ):
        issues.append(_issue("SOURCE_MANIFEST_PATH_INVALID"))
    proof_path = pagination.get("offset_proof_path")
    if (
        not isinstance(proof_path, str) or not proof_path.strip() or "\x00" in proof_path
        or Path(proof_path).is_absolute() or ".." in Path(proof_path).parts
    ):
        issues.append(_issue("SOURCE_OFFSET_PROOF_PATH_INVALID"))
    return issues


def _validate_manifest_payload(manifest: Any, source: dict[str, Any]) -> list[dict[str, Any]]:
    issues: list[dict[str, Any]] = []
    if not isinstance(manifest, dict):
        return [_issue("SOURCE_MANIFEST_INVALID")]
    if manifest.get("manifest_id") != source.get("manifest_id"):
        issues.append(_issue("SOURCE_MANIFEST_ID_MISMATCH"))
    if str(manifest.get("book_code", "")).upper() != str(source.get("book_code", "")).upper():
        issues.append(_issue("SOURCE_MANIFEST_BOOK_CODE_MISMATCH"))
    pdf_meta = manifest.get("source_pdf")
    if not isinstance(pdf_meta, dict):
        issues.append(_issue("SOURCE_MANIFEST_PDF_METADATA_MISSING"))
        pdf_meta = {}
    for manifest_key, source_key in (
        ("sha256", "source_pdf_sha256"),
        ("size_bytes", "source_pdf_size_bytes"),
        ("pdf_page_count", "pdf_page_count"),
    ):
        observed = pdf_meta.get(manifest_key)
        expected = source.get(source_key)
        if manifest_key == "sha256":
            observed = str(observed or "").lower()
            expected = str(expected or "").lower()
        if observed != expected:
            issues.append(_issue("SOURCE_MANIFEST_PDF_METADATA_MISMATCH", field=manifest_key))
    manifest_pagination = manifest.get("pagination")
    source_pagination = source.get("pagination")
    if not isinstance(manifest_pagination, dict) or not isinstance(source_pagination, dict):
        issues.append(_issue("SOURCE_MANIFEST_PAGINATION_MISSING"))
    else:
        for key in (
            "body_pdf_page_start", "body_pdf_page_end", "book_page_start", "book_page_end",
            "pdf_minus_book_offset",
        ):
            if manifest_pagination.get(key) != source_pagination.get(key):
                issues.append(_issue("SOURCE_MANIFEST_PAGINATION_MISMATCH", field=key))
    if manifest.get("edition_id") is not None and manifest.get("edition_id") != source.get("edition_id"):
        issues.append(_issue("SOURCE_MANIFEST_EDITION_MISMATCH"))
    return issues


def _verify_source_pdf(
    *,
    pdf_path: str | Path | None,
    repo_root: Path,
    source: dict[str, Any],
    manifest: dict[str, Any],
) -> tuple[list[dict[str, Any]], bool, str | None]:
    if pdf_path is None:
        if source.get("rights_status") == "AUTHORIZED":
            return [_issue("SOURCE_PDF_BYTES_NOT_VERIFIED")], False, None
        return [], False, None
    try:
        path = Path(pdf_path).resolve()
    except (OSError, RuntimeError, ValueError) as exc:
        return [_issue("SOURCE_PDF_PATH_INVALID", exception_type=type(exc).__name__)], False, None
    if _is_within(path, repo_root):
        return [_issue("SOURCE_PDF_MUST_REMAIN_OUTSIDE_REPOSITORY")], False, None
    if not path.is_file():
        return [_issue("SOURCE_PDF_UNAVAILABLE")], False, None
    issues: list[dict[str, Any]] = []
    try:
        sha256 = file_sha256(str(path))
        size_bytes = path.stat().st_size
    except OSError as exc:
        return [_issue("SOURCE_PDF_READ_FAILED", exception_type=type(exc).__name__)], False, None
    if sha256 != source.get("source_pdf_sha256"):
        issues.append(_issue("SOURCE_PDF_SHA_MISMATCH"))
    if size_bytes != source.get("source_pdf_size_bytes"):
        issues.append(_issue("SOURCE_PDF_SIZE_MISMATCH"))
    try:
        import pymupdf

        document = pymupdf.open(str(path))
        page_count = document.page_count
        document.close()
    except Exception as exc:
        return issues + [_issue("SOURCE_PDF_PAGE_COUNT_UNVERIFIABLE", exception_type=type(exc).__name__)], False, sha256
    if page_count != source.get("pdf_page_count"):
        issues.append(_issue("SOURCE_PDF_PAGE_COUNT_MISMATCH"))
    pdf_meta = manifest.get("source_pdf", {}) if isinstance(manifest, dict) else {}
    if sha256 != pdf_meta.get("sha256") or size_bytes != pdf_meta.get("size_bytes") or page_count != pdf_meta.get("pdf_page_count"):
        issues.append(_issue("SOURCE_PDF_MANIFEST_BINDING_MISMATCH"))
    return issues, not issues, sha256


def _validate_extraction(
    extraction: dict[str, Any],
    source: dict[str, Any],
    expected_binding: dict[str, Any],
) -> tuple[list[dict[str, Any]], dict[str, tuple[dict[str, Any], dict[str, Any]]], dict[str, dict[str, Any]]]:
    issues: list[dict[str, Any]] = []
    element_index: dict[str, tuple[dict[str, Any], dict[str, Any]]] = {}
    record_index: dict[str, dict[str, Any]] = {}
    if extraction.get("schema_version") != EXTRACTION_SCHEMA:
        issues.append(_issue("EXTRACTION_SCHEMA_UNSUPPORTED"))
    issues.extend(_check_binding(extraction, expected_binding, "EXTRACTION"))
    pagination = source.get("pagination", {})
    scope = extraction.get("section_scope")
    if not isinstance(scope, dict):
        return issues + [_issue("EXTRACTION_SECTION_SCOPE_INVALID")], element_index, record_index
    scope_keys = ("pdf_page_start", "pdf_page_end", "book_page_start", "book_page_end")
    if not _token(scope.get("section_id")) or any(type(scope.get(key)) is not int for key in scope_keys):
        return issues + [_issue("EXTRACTION_SECTION_SCOPE_INVALID")], element_index, record_index
    pdf_start, pdf_end = scope["pdf_page_start"], scope["pdf_page_end"]
    book_start, book_end = scope["book_page_start"], scope["book_page_end"]
    offset = pagination.get("pdf_minus_book_offset")
    if (
        type(offset) is not int or pdf_start < pagination.get("body_pdf_page_start", 1)
        or pdf_end > pagination.get("body_pdf_page_end", 0) or pdf_end < pdf_start
        or book_start < pagination.get("book_page_start", 1)
        or book_end > pagination.get("book_page_end", 0) or book_end < book_start
        or pdf_start - offset != book_start or pdf_end - offset != book_end
        or pdf_end - pdf_start != book_end - book_start
    ):
        issues.append(_issue("EXTRACTION_SCOPE_OFFSET_MISMATCH"))
    if extraction.get("offset_proof_sha256") != pagination.get("offset_proof_sha256"):
        issues.append(_issue("EXTRACTION_OFFSET_PROOF_BINDING_MISMATCH"))

    expected_pages = set(range(pdf_start, pdf_end + 1))
    seen_pages: set[int] = set()
    global_evidence_ids: set[str] = set()
    page_records = extraction.get("pages")
    if not isinstance(page_records, list):
        return issues + [_issue("EXTRACTION_PAGES_INVALID")], element_index, record_index
    for page in page_records:
        if not isinstance(page, dict):
            issues.append(_issue("EXTRACTION_PAGE_RECORD_INVALID"))
            continue
        record_id = page.get("record_id")
        pdf_page = page.get("pdf_page")
        book_page = page.get("book_page")
        if not _token(record_id) or record_id in record_index:
            issues.append(_issue("EXTRACTION_RECORD_ID_INVALID_OR_DUPLICATE"))
        if type(pdf_page) is not int or pdf_page in seen_pages:
            issues.append(_issue("EXTRACTION_PDF_PAGE_INVALID_OR_DUPLICATE"))
            continue
        seen_pages.add(pdf_page)
        expected_book_page = pdf_page - offset if type(offset) is int else None
        if pdf_page not in expected_pages or type(book_page) is not int or book_page != expected_book_page:
            issues.append(_issue("EXTRACTION_PAGE_MAPPING_MISMATCH", pdf_page=pdf_page))
        if page.get("printed_page_number") != book_page:
            issues.append(_issue("EXTRACTION_PRINTED_PAGE_MISMATCH", pdf_page=pdf_page))
        if not _sha(page.get("page_image_sha256")):
            issues.append(_issue("EXTRACTION_PAGE_IMAGE_SHA_INVALID", pdf_page=pdf_page))
        if page.get("page_status") not in {"CURRICULAR_CONTENT", "NO_CURRICULAR_TASK"}:
            issues.append(_issue("EXTRACTION_PAGE_STATUS_NOT_COMPLETE", pdf_page=pdf_page))
        if page.get("reading_status") != "COMPLETE":
            issues.append(_issue("EXTRACTION_PAGE_TEXT_NOT_COMPLETE", pdf_page=pdf_page))
        if page.get("visual_scan_status") not in {"COMPLETE", "NOT_APPLICABLE"}:
            issues.append(_issue("EXTRACTION_PAGE_VISUAL_NOT_COMPLETE", pdf_page=pdf_page))

        evidence_refs = page.get("evidence_refs")
        evidence_ids: set[str] = set()
        if not isinstance(evidence_refs, list) or not evidence_refs:
            issues.append(_issue("EXTRACTION_PAGE_EVIDENCE_REQUIRED", pdf_page=pdf_page))
            evidence_refs = []
        for evidence in evidence_refs:
            if not isinstance(evidence, dict):
                issues.append(_issue("EXTRACTION_EVIDENCE_RECORD_INVALID", pdf_page=pdf_page))
                continue
            evidence_id = evidence.get("evidence_id")
            bbox = evidence.get("bbox_normalized")
            valid_bbox = (
                isinstance(bbox, list) and len(bbox) == 4
                and all(type(v) in (int, float) and math.isfinite(float(v)) for v in bbox)
                and 0 <= float(bbox[0]) < float(bbox[2]) <= 1
                and 0 <= float(bbox[1]) < float(bbox[3]) <= 1
            )
            if not _token(evidence_id) or evidence_id in evidence_ids or evidence_id in global_evidence_ids:
                issues.append(_issue("EXTRACTION_EVIDENCE_ID_INVALID_OR_DUPLICATE", pdf_page=pdf_page))
            if evidence.get("purpose") not in (None, "PRINTED_PAGE_NUMBER"):
                issues.append(_issue("EXTRACTION_EVIDENCE_PURPOSE_INVALID", pdf_page=pdf_page))
            if evidence.get("kind") not in _ALLOWED_EVIDENCE_KINDS:
                issues.append(_issue("EXTRACTION_EVIDENCE_KIND_INVALID", pdf_page=pdf_page))
            if evidence.get("source_pdf_sha256") != expected_binding.get("source_pdf_sha256"):
                issues.append(_issue("EXTRACTION_EVIDENCE_SOURCE_MISMATCH", pdf_page=pdf_page))
            if evidence.get("pdf_page") != pdf_page or evidence.get("book_page") != book_page:
                issues.append(_issue("EXTRACTION_EVIDENCE_PAGE_MISMATCH", pdf_page=pdf_page))
            if not valid_bbox or not _sha(evidence.get("sha256")):
                issues.append(_issue("EXTRACTION_EVIDENCE_REGION_INVALID", pdf_page=pdf_page))
            if evidence.get("kind") == "pdf_page_render" and evidence.get("sha256") != page.get("page_image_sha256"):
                issues.append(_issue("EXTRACTION_PAGE_RENDER_HASH_MISMATCH", pdf_page=pdf_page))
            if _token(evidence_id):
                evidence_ids.add(evidence_id)
                global_evidence_ids.add(evidence_id)

        elements = page.get("elements")
        if not isinstance(elements, list):
            issues.append(_issue("EXTRACTION_ELEMENTS_INVALID", pdf_page=pdf_page))
            elements = []
        if page.get("page_status") == "CURRICULAR_CONTENT" and not elements:
            issues.append(_issue("CURRICULAR_PAGE_HAS_NO_EXTRACTED_ELEMENTS", pdf_page=pdf_page))
        if page.get("page_status") == "NO_CURRICULAR_TASK" and elements:
            issues.append(_issue("NONCURRICULAR_PAGE_HAS_ELEMENTS", pdf_page=pdf_page))
        for element in elements:
            if not isinstance(element, dict):
                issues.append(_issue("EXTRACTION_ELEMENT_INVALID", pdf_page=pdf_page))
                continue
            element_id = element.get("element_id")
            if not _token(element_id) or element_id in element_index:
                issues.append(_issue("EXTRACTION_ELEMENT_ID_INVALID_OR_DUPLICATE", pdf_page=pdf_page))
            if element.get("element_kind") not in _ALLOWED_ELEMENT_KINDS:
                issues.append(_issue("EXTRACTION_ELEMENT_KIND_INVALID", pdf_page=pdf_page))
            excerpt = element.get("short_excerpt", "")
            if not _safe_excerpt(excerpt):
                issues.append(_issue("EXTRACTION_EXCERPT_OVER_TEN_WORDS", pdf_page=pdf_page))
            refs = element.get("evidence_refs")
            if not isinstance(refs, list) or not refs or any(ref not in evidence_ids for ref in refs):
                issues.append(_issue("EXTRACTION_ELEMENT_EVIDENCE_UNBOUND", pdf_page=pdf_page))
            for list_key in ("concept_candidate_ids", "operation_ids", "number_domain_ids"):
                values = element.get(list_key)
                if not isinstance(values, list) or any(not _token(value) for value in values):
                    issues.append(_issue("EXTRACTION_ELEMENT_TOKEN_LIST_INVALID", field=list_key, pdf_page=pdf_page))
            if _token(element_id) and _token(record_id):
                element_index[element_id] = (page, element)
        if _token(record_id):
            record_index[record_id] = page

    for missing in sorted(expected_pages - seen_pages):
        issues.append(_issue("EXTRACTION_PAGE_MISSING", pdf_page=missing))
    for extra in sorted(seen_pages - expected_pages):
        issues.append(_issue("EXTRACTION_PAGE_OUTSIDE_SECTION", pdf_page=extra))
    return issues, element_index, record_index


def _validate_offset_proof(
    proof: dict[str, Any] | None,
    source: dict[str, Any],
    extraction: dict[str, Any],
    expected_binding: dict[str, Any],
) -> list[dict[str, Any]]:
    if not isinstance(proof, dict):
        return [_issue("OFFSET_PROOF_UNREADABLE")]
    issues: list[dict[str, Any]] = []
    if proof.get("schema_version") != "khshab.offset-proof/v1":
        issues.append(_issue("OFFSET_PROOF_SCHEMA_UNSUPPORTED"))
    if proof.get("status") not in {"TWO_READER_MATCHED", "SYNTHETIC_TEST_ONLY"}:
        issues.append(_issue("OFFSET_PROOF_STATUS_INVALID"))
    issues.extend(_check_binding(proof, expected_binding, "OFFSET_PROOF"))
    pagination = source.get("pagination", {})
    if proof.get("status") != pagination.get("offset_proof_status"):
        issues.append(_issue("OFFSET_PROOF_STATUS_BINDING_MISMATCH"))
    if proof.get("pdf_minus_book_offset") != pagination.get("pdf_minus_book_offset"):
        issues.append(_issue("OFFSET_PROOF_VALUE_MISMATCH"))
    scope = extraction.get("section_scope")
    if not isinstance(scope, dict) or any(
        type(scope.get(key)) is not int for key in ("pdf_page_start", "pdf_page_end")
    ):
        return issues + [_issue("OFFSET_PROOF_SECTION_SCOPE_INVALID")]
    if proof.get("section_id") != scope.get("section_id"):
        issues.append(_issue("OFFSET_PROOF_SECTION_MISMATCH"))
    expected_pages = range(scope["pdf_page_start"], scope["pdf_page_end"] + 1)
    expected_count = len(expected_pages)
    if type(proof.get("expected_page_count")) is not int or proof.get("expected_page_count") != expected_count:
        issues.append(_issue("OFFSET_PROOF_EXPECTED_PAGE_COUNT_MISMATCH"))
    if type(proof.get("matched_page_count")) is not int or proof.get("matched_page_count") != expected_count:
        issues.append(_issue("OFFSET_PROOF_MATCHED_PAGE_COUNT_MISMATCH"))
    readers = proof.get("reader_pass_sha256")
    if (
        not isinstance(readers, list) or len(readers) != 2
        or any(not _sha(value) for value in readers)
        or len(set(readers)) != 2
    ):
        issues.append(_issue("OFFSET_PROOF_READER_PASSES_INVALID"))

    pages = extraction.get("pages", [])
    pages_by_pdf = {p.get("pdf_page"): p for p in pages if isinstance(p, dict) and type(p.get("pdf_page")) is int}
    anchors = proof.get("page_anchors")
    if not isinstance(anchors, list):
        return issues + [_issue("OFFSET_PROOF_ANCHORS_INVALID")]
    anchors_by_pdf: dict[int, dict[str, Any]] = {}
    for anchor in anchors:
        if not isinstance(anchor, dict):
            issues.append(_issue("OFFSET_PROOF_ANCHOR_INVALID"))
            continue
        pdf_page = anchor.get("pdf_page")
        if type(pdf_page) is not int or pdf_page in anchors_by_pdf:
            issues.append(_issue("OFFSET_PROOF_ANCHOR_PAGE_INVALID_OR_DUPLICATE"))
            continue
        anchors_by_pdf[pdf_page] = anchor
        book_page = anchor.get("book_page")
        if (
            book_page != anchor.get("printed_page_number")
            or book_page != pdf_page - pagination.get("pdf_minus_book_offset", 0)
            or not _sha(anchor.get("evidence_sha256"))
            or not _token(anchor.get("evidence_id"))
        ):
            issues.append(_issue("OFFSET_PROOF_ANCHOR_MAPPING_INVALID", pdf_page=pdf_page))
        page = pages_by_pdf.get(pdf_page)
        if page is None or page.get("book_page") != book_page:
            issues.append(_issue("OFFSET_PROOF_ANCHOR_NOT_IN_EXTRACTION", pdf_page=pdf_page))
            continue
        evidence_refs = page.get("evidence_refs", [])
        matched = next((
            evidence for evidence in evidence_refs
            if isinstance(evidence, dict) and evidence.get("evidence_id") == anchor.get("evidence_id")
        ), None)
        if (
            matched is None
            or matched.get("purpose") != "PRINTED_PAGE_NUMBER"
            or matched.get("sha256") != anchor.get("evidence_sha256")
            or matched.get("kind") not in {"pdf_text_bbox", "pdf_crop_150dpi"}
        ):
            issues.append(_issue("OFFSET_PROOF_ANCHOR_EVIDENCE_UNBOUND", pdf_page=pdf_page))
    expected_set = set(expected_pages)
    if set(anchors_by_pdf) != expected_set:
        issues.append(_issue("OFFSET_PROOF_PAGE_COVERAGE_MISMATCH"))
    return issues


def _validate_valuation(
    valuation: dict[str, Any],
    expected_binding: dict[str, Any],
    source: dict[str, Any],
    extraction_sha256: str,
    element_index: dict[str, tuple[dict[str, Any], dict[str, Any]]],
    record_index: dict[str, dict[str, Any]],
) -> tuple[list[dict[str, Any]], dict[str, dict[str, Any]]]:
    issues: list[dict[str, Any]] = []
    valuation_index: dict[str, dict[str, Any]] = {}
    if valuation.get("schema_version") != VALUATION_SCHEMA:
        issues.append(_issue("VALUATION_SCHEMA_UNSUPPORTED"))
    issues.extend(_check_binding(valuation, expected_binding, "VALUATION"))
    if valuation.get("extraction_bundle_sha256") != extraction_sha256:
        issues.append(_issue("VALUATION_EXTRACTION_SHA_MISMATCH"))
    guide = valuation.get("teacher_guide")
    if not isinstance(guide, dict):
        issues.append(_issue("VALUATION_TEACHER_GUIDE_BINDING_MISSING"))
        guide = {}
    if not _token(guide.get("guide_id")) or not _sha(guide.get("guide_sha256")):
        issues.append(_issue("VALUATION_TEACHER_GUIDE_BINDING_INVALID"))
    if guide.get("book_code") != source.get("book_code") or guide.get("grade") != source.get("grade") or guide.get("subject_id") != source.get("subject_id"):
        issues.append(_issue("VALUATION_TEACHER_GUIDE_EDITION_MISMATCH"))

    objectives = valuation.get("objectives")
    objective_index: dict[str, dict[str, Any]] = {}
    if not isinstance(objectives, list) or not objectives:
        issues.append(_issue("VALUATION_OBJECTIVES_REQUIRED"))
        objectives = []
    for objective in objectives:
        if not isinstance(objective, dict):
            issues.append(_issue("VALUATION_OBJECTIVE_INVALID"))
            continue
        objective_id = objective.get("objective_id")
        if not _token(objective_id) or objective_id in objective_index:
            issues.append(_issue("VALUATION_OBJECTIVE_ID_INVALID_OR_DUPLICATE"))
            continue
        level = objective.get("level")
        parent_id = objective.get("parent_objective_id")
        if level not in {"GENERAL", "SPECIFIC", "BEHAVIORAL"}:
            issues.append(_issue("VALUATION_OBJECTIVE_LEVEL_INVALID", objective_id=objective_id))
        if level == "GENERAL" and parent_id is not None:
            issues.append(_issue("GENERAL_OBJECTIVE_MUST_BE_ROOT", objective_id=objective_id))
        if level in {"SPECIFIC", "BEHAVIORAL"} and not _token(parent_id):
            issues.append(_issue("CHILD_OBJECTIVE_PARENT_REQUIRED", objective_id=objective_id))
        objective_index[objective_id] = objective
    for objective_id, objective in objective_index.items():
        level = objective.get("level")
        parent_id = objective.get("parent_objective_id")
        parent = objective_index.get(parent_id) if _token(parent_id) else None
        expected_parent_level = {"SPECIFIC": "GENERAL", "BEHAVIORAL": "SPECIFIC"}.get(level)
        if expected_parent_level and (parent is None or parent.get("level") != expected_parent_level):
            issues.append(_issue("VALUATION_OBJECTIVE_HIERARCHY_INVALID", objective_id=objective_id))

    records = valuation.get("records")
    if not isinstance(records, list) or not records:
        issues.append(_issue("VALUATION_RECORDS_REQUIRED"))
        records = []
    for record in records:
        if not isinstance(record, dict):
            issues.append(_issue("VALUATION_RECORD_INVALID"))
            continue
        valuation_id = record.get("valuation_id")
        if not _token(valuation_id) or valuation_id in valuation_index:
            issues.append(_issue("VALUATION_ID_INVALID_OR_DUPLICATE"))
            continue
        record_id = record.get("extraction_record_id")
        element_id = record.get("element_id")
        if not _token(record_id) or not _token(element_id) or record_id not in record_index or element_id not in element_index:
            issues.append(_issue("VALUATION_EXTRACTION_REFERENCE_MISSING", valuation_id=valuation_id))
        else:
            page, _element = element_index[element_id]
            if page.get("record_id") != record_id:
                issues.append(_issue("VALUATION_ELEMENT_PAGE_REFERENCE_MISMATCH", valuation_id=valuation_id))
        objective_ids = record.get("objective_ids")
        if (
            not isinstance(objective_ids, list) or not objective_ids
            or any(not _token(obj) or obj not in objective_index for obj in objective_ids)
        ):
            issues.append(_issue("VALUATION_OBJECTIVE_REFERENCE_INVALID", valuation_id=valuation_id))
        elif {objective_index[obj].get("level") for obj in objective_ids} != {"GENERAL", "SPECIFIC", "BEHAVIORAL"}:
            issues.append(_issue("VALUATION_OBJECTIVE_HIERARCHY_INCOMPLETE", valuation_id=valuation_id))
        else:
            objective_set = set(objective_ids)
            for objective_id in objective_ids:
                objective = objective_index[objective_id]
                if objective.get("level") == "BEHAVIORAL":
                    specific_id = objective.get("parent_objective_id")
                    specific = objective_index.get(specific_id, {})
                    general_id = specific.get("parent_objective_id")
                    if specific_id not in objective_set or general_id not in objective_set:
                        issues.append(_issue("VALUATION_OBJECTIVE_ANCESTOR_MISSING", valuation_id=valuation_id))
        if record.get("role") not in {"PRIMARY", "SUPPORTING", "INCIDENTAL"}:
            issues.append(_issue("VALUATION_ROLE_INVALID", valuation_id=valuation_id))
        if not _token(record.get("concept_id")):
            issues.append(_issue("VALUATION_CONCEPT_ID_INVALID", valuation_id=valuation_id))
        for field in ("operation_id", "number_domain_id"):
            value = record.get(field)
            if value is not None and not _token(value):
                issues.append(_issue("VALUATION_SCOPE_ID_INVALID", field=field, valuation_id=valuation_id))
        prerequisites = record.get("prerequisite_ids", [])
        if not isinstance(prerequisites, list) or any(not _token(value) for value in prerequisites):
            issues.append(_issue("VALUATION_PREREQUISITE_IDS_INVALID", valuation_id=valuation_id))
        if record.get("assessment_status") not in {"IN_SCOPE", "OUT_OF_SCOPE", "UNKNOWN"}:
            issues.append(_issue("VALUATION_ASSESSMENT_STATUS_INVALID", valuation_id=valuation_id))
        if record.get("confidence") not in {"HIGH", "MEDIUM", "LOW", "UNOBSERVABLE"}:
            issues.append(_issue("VALUATION_CONFIDENCE_INVALID", valuation_id=valuation_id))
        if record.get("review_status") not in _ALLOWED_REVIEW_STATES:
            issues.append(_issue("VALUATION_REVIEW_STATUS_INVALID", valuation_id=valuation_id))
        verdict_id = record.get("human_verdict_id")
        if record.get("review_status") == "HUMAN_APPROVED" and not _token(verdict_id):
            issues.append(_issue("HUMAN_APPROVAL_REFERENCE_REQUIRED", valuation_id=valuation_id))
        if record.get("review_status") != "HUMAN_APPROVED" and verdict_id is not None:
            issues.append(_issue("UNAPPROVED_RECORD_HAS_VERDICT_REFERENCE", valuation_id=valuation_id))
        if record.get("review_status") == "HUMAN_APPROVED" and (
            record.get("assessment_status") != "IN_SCOPE" or record.get("confidence") != "HIGH"
        ):
            issues.append(_issue("HUMAN_APPROVAL_FOR_UNCERTAIN_CONTENT_FORBIDDEN", valuation_id=valuation_id))
        valuation_index[valuation_id] = record
    return issues, valuation_index


def _validate_verdict_ledger(
    ledger: dict[str, Any] | None,
    valuation_sha256: str,
    valuation_index: dict[str, dict[str, Any]],
) -> tuple[list[dict[str, Any]], set[str]]:
    if ledger is None:
        approved = [key for key, value in valuation_index.items() if value.get("review_status") == "HUMAN_APPROVED"]
        if approved:
            return [_issue("HUMAN_VERDICT_LEDGER_REQUIRED", valuation_ids=approved)], set()
        return [], set()
    issues: list[dict[str, Any]] = []
    if ledger.get("schema_version") != "1.0.0":
        return [_issue("VERDICT_LEDGER_SCHEMA_UNSUPPORTED")], set()
    raw_verdicts = ledger.get("verdicts")
    if not isinstance(raw_verdicts, list):
        return [_issue("VERDICT_LEDGER_INVALID")], set()
    if type(ledger.get("verdict_count")) is not int or ledger.get("verdict_count") != len(raw_verdicts):
        issues.append(_issue("VERDICT_LEDGER_COUNT_MISMATCH"))
    verified_ids: set[str] = set()
    by_id: dict[str, dict[str, Any]] = {}
    for verdict in raw_verdicts:
        if not isinstance(verdict, dict) or not _token(verdict.get("verdict_id")):
            issues.append(_issue("VERDICT_LEDGER_RECORD_INVALID"))
            continue
        verdict_id = verdict["verdict_id"]
        if verdict_id in by_id:
            issues.append(_issue("VERDICT_LEDGER_DUPLICATE_ID"))
            continue
        by_id[verdict_id] = verdict
    for valuation_id, record in valuation_index.items():
        if record.get("review_status") != "HUMAN_APPROVED":
            continue
        verdict_id = record.get("human_verdict_id")
        verdict = by_id.get(verdict_id)
        if not isinstance(verdict, dict):
            issues.append(_issue("HUMAN_VERDICT_NOT_FOUND", valuation_id=valuation_id))
            continue
        target_ids = verdict.get("target_ids")
        if (
            verdict.get("scope") != "ROW_APPROVED"
            or verdict.get("actor_role") != "human_teacher"
            or verdict.get("authenticated") is not True
            or verdict.get("auth_method") not in {"KHSHAB_TEACHER_TOKEN_FILE", "X-Teacher-Token"}
            or verdict.get("invalidated") is not False
            or verdict.get("target_sha256") != valuation_sha256
            or not isinstance(target_ids, list)
            or not target_ids
            or any(not _token(target_id) for target_id in target_ids)
            or valuation_id not in target_ids
        ):
            issues.append(_issue("HUMAN_VERDICT_BINDING_INVALID", valuation_id=valuation_id))
            continue
        verified_ids.add(valuation_id)
    return issues, verified_ids


def _validate_game_pack(
    game_pack: dict[str, Any],
    expected_binding: dict[str, Any],
    valuation_sha256: str,
    valuation_index: dict[str, dict[str, Any]],
    element_index: dict[str, tuple[dict[str, Any], dict[str, Any]]],
    verified_approvals: set[str],
) -> tuple[list[dict[str, Any]], int]:
    issues: list[dict[str, Any]] = []
    if game_pack.get("schema_version") != GAME_PACK_SCHEMA:
        issues.append(_issue("GAME_PACK_SCHEMA_UNSUPPORTED"))
    issues.extend(_check_binding(game_pack, expected_binding, "GAME_PACK"))
    if game_pack.get("valuation_bundle_sha256") != valuation_sha256:
        issues.append(_issue("GAME_PACK_VALUATION_SHA_MISMATCH"))
    items = game_pack.get("items")
    if not isinstance(items, list):
        return issues + [_issue("GAME_PACK_ITEMS_INVALID")], 0
    seen_ids: set[str] = set()
    blocked = 0
    for item in items:
        if not isinstance(item, dict):
            issues.append(_issue("GAME_ITEM_INVALID"))
            continue
        game_id = item.get("game_item_id")
        if not _token(game_id) or game_id in seen_ids:
            issues.append(_issue("GAME_ITEM_ID_INVALID_OR_DUPLICATE"))
        seen_ids.add(game_id) if _token(game_id) else None
        valuation_ids = item.get("valuation_ids")
        if (
            not isinstance(valuation_ids, list) or not valuation_ids
            or any(not _token(value) or value not in valuation_index for value in valuation_ids)
        ):
            issues.append(_issue("GAME_ITEM_VALUATION_REFERENCE_INVALID", game_item_id=game_id))
            continue
        expected_record_ids = {valuation_index[value].get("extraction_record_id") for value in valuation_ids}
        extraction_record_id = item.get("extraction_record_id")
        if not _token(extraction_record_id) or extraction_record_id not in expected_record_ids:
            issues.append(_issue("GAME_ITEM_EXTRACTION_REFERENCE_MISMATCH", game_item_id=game_id))
        element_id = item.get("element_id")
        if not _token(element_id) or element_id not in element_index:
            issues.append(_issue("GAME_ITEM_ELEMENT_REFERENCE_MISSING", game_item_id=game_id))
        objective_ids = item.get("objective_ids")
        valuation_objectives = {
            objective for value in valuation_ids for objective in valuation_index[value].get("objective_ids", [])
            if _token(objective)
        }
        if (
            not isinstance(objective_ids, list) or not objective_ids
            or any(not _token(value) for value in objective_ids)
            or not set(objective_ids).issubset(valuation_objectives)
        ):
            issues.append(_issue("GAME_ITEM_OBJECTIVE_REFERENCE_INVALID", game_item_id=game_id))
        if item.get("activity_kind") not in {"PRACTICE", "GAME", "TEACHER_TOOL"}:
            issues.append(_issue("GAME_ITEM_ACTIVITY_KIND_INVALID", game_item_id=game_id))
        for field in ("mechanic_id", "difficulty_id"):
            if not _token(item.get(field)):
                issues.append(_issue("GAME_ITEM_CONFIGURATION_ID_INVALID", field=field, game_item_id=game_id))
        tags = item.get("personalization_tags")
        if not isinstance(tags, list) or any(not _token(tag) for tag in tags):
            issues.append(_issue("GAME_ITEM_PERSONALIZATION_TAGS_INVALID", game_item_id=game_id))
        if item.get("generation_status") != "DRAFT":
            issues.append(_issue("GAME_ITEM_MUST_REMAIN_DRAFT", game_item_id=game_id))
        unapproved = [
            value for value in valuation_ids
            if (
                valuation_index[value].get("review_status") != "HUMAN_APPROVED"
                or value not in verified_approvals
                or valuation_index[value].get("assessment_status") != "IN_SCOPE"
                or valuation_index[value].get("confidence") != "HIGH"
            )
        ]
        if unapproved:
            blocked += 1
            issues.append(_issue("GAME_ITEM_BLOCKED_UNAPPROVED_CONTENT", game_item_id=game_id, valuation_ids=unapproved))
    return issues, blocked


def audit_content_pipeline(
    *,
    repo_root: str | Path,
    source_path: str | Path,
    extraction_path: str | Path,
    valuation_path: str | Path,
    game_pack_path: str | Path,
    verdict_ledger_path: str | Path | None = None,
    source_pdf_path: str | Path | None = None,
) -> dict[str, Any]:
    """Validate Git-versioned content artifacts and cross-stage lineage.

    Source, manifest, extraction, valuation and game artifacts are versioned
    inside the repository. For an AUTHORIZED source, callers must also provide
    the exact PDF from private external storage so its SHA, size and page count
    can be checked; synthetic fixtures may omit that binary. A successful
    result is traceability-only.
    """
    root = Path(repo_root).resolve()
    paths = {
        "source": Path(source_path).resolve(),
        "extraction": Path(extraction_path).resolve(),
        "valuation": Path(valuation_path).resolve(),
        "game_pack": Path(game_pack_path).resolve(),
    }
    if verdict_ledger_path is not None:
        paths["verdict_ledger"] = Path(verdict_ledger_path).resolve()
    outside = [name for name, path in paths.items() if not _is_within(path, root)]
    if outside:
        return _report(
            status="FAIL", reason_code="PIPELINE_ARTIFACT_OUTSIDE_REPOSITORY",
            issues=[_issue("PIPELINE_ARTIFACT_OUTSIDE_REPOSITORY", artifacts=outside)],
        )
    missing = [name for name, path in paths.items() if not path.is_file()]
    if missing:
        return _report(
            status="BLOCKED", reason_code="PIPELINE_INPUT_UNAVAILABLE",
            issues=[_issue("PIPELINE_INPUT_UNAVAILABLE", artifacts=missing)],
        )
    try:
        hashes = {name: file_sha256(str(path)) for name, path in paths.items()}
    except OSError as exc:
        return _report(
            status="BLOCKED", reason_code="PIPELINE_HASH_FAILED",
            issues=[_issue("PIPELINE_HASH_FAILED", exception_type=type(exc).__name__)],
        )
    documents: dict[str, dict[str, Any]] = {}
    read_issues: list[dict[str, Any]] = []
    for key in ("source", "extraction", "valuation", "game_pack"):
        payload, error = _read_json(paths[key], key.upper())
        if error:
            read_issues.append(error)
        elif payload is not None:
            documents[key] = payload
    verdict_ledger: dict[str, Any] | None = None
    if "verdict_ledger" in paths:
        verdict_ledger, error = _read_json(paths["verdict_ledger"], "VERDICT_LEDGER")
        if error:
            read_issues.append(error)
    if read_issues or any(key not in documents for key in ("source", "extraction", "valuation", "game_pack")):
        return _report(
            status="BLOCKED", reason_code="PIPELINE_INPUT_INVALID",
            issues=read_issues or [_issue("PIPELINE_INPUT_INVALID")], hashes=hashes,
        )

    source = documents["source"]
    extraction = documents["extraction"]
    valuation = documents["valuation"]
    game_pack = documents["game_pack"]
    privacy_issues: list[dict[str, Any]] = []
    for name, document_payload in documents.items():
        privacy_issues.extend(_privacy_issues(document_payload, label=name))
    if verdict_ledger is not None:
        privacy_issues.extend(_privacy_issues(verdict_ledger, label="verdict_ledger"))
    if privacy_issues:
        return _report(
            status="BLOCKED", reason_code=privacy_issues[0]["code"], issues=privacy_issues,
            binding=_source_binding(source), hashes=hashes,
        )
    schema_issues: list[dict[str, Any]] = []
    for artifact in ("source", "extraction", "valuation", "game_pack"):
        schema_issues.extend(_schema_issues(
            documents[artifact], artifact=artifact, schema_filename=_STAGE_SCHEMAS[artifact]
        ))
        if schema_issues and schema_issues[-1]["code"] == "PIPELINE_SCHEMA_VALIDATION_UNAVAILABLE":
            break
    if verdict_ledger is not None and not schema_issues:
        schema_issues.extend(_schema_issues(
            verdict_ledger, artifact="verdict_ledger", schema_filename=_STAGE_SCHEMAS["verdict_ledger"]
        ))
    if schema_issues:
        return _report(
            status="BLOCKED", reason_code=schema_issues[0]["code"], issues=schema_issues,
            binding=_source_binding(source), hashes=hashes,
        )
    binding = _source_binding(source)
    source_issues = _validate_source(source)
    if source_issues:
        return _report(
            status="BLOCKED", reason_code=source_issues[0]["code"], issues=source_issues,
            binding=binding, hashes=hashes,
        )
    manifest_path = (root / source["manifest_path"]).resolve()
    if not _is_within(manifest_path, root) or not manifest_path.is_file():
        return _report(
            status="BLOCKED", reason_code="SOURCE_MANIFEST_ARTIFACT_UNAVAILABLE",
            issues=[_issue("SOURCE_MANIFEST_ARTIFACT_UNAVAILABLE")], binding=binding, hashes=hashes,
        )
    try:
        manifest_sha256 = file_sha256(str(manifest_path))
    except OSError as exc:
        return _report(
            status="BLOCKED", reason_code="SOURCE_MANIFEST_HASH_FAILED",
            issues=[_issue("SOURCE_MANIFEST_HASH_FAILED", exception_type=type(exc).__name__)],
            binding=binding, hashes=hashes,
        )
    hashes["manifest"] = manifest_sha256
    manifest, manifest_error = _read_json(manifest_path, "SOURCE_MANIFEST")
    if manifest_error or manifest is None:
        return _report(
            status="BLOCKED", reason_code="SOURCE_MANIFEST_UNREADABLE",
            issues=[manifest_error or _issue("SOURCE_MANIFEST_UNREADABLE")], binding=binding, hashes=hashes,
        )
    manifest_issues = []
    if manifest_sha256 != source.get("manifest_sha256"):
        manifest_issues.append(_issue("SOURCE_MANIFEST_SHA_MISMATCH"))
    manifest_issues.extend(_validate_manifest_payload(manifest, source))
    pdf_issues, source_pdf_verified, source_pdf_sha256 = _verify_source_pdf(
        pdf_path=source_pdf_path, repo_root=root, source=source, manifest=manifest
    )
    manifest_issues.extend(pdf_issues)
    if source_pdf_sha256:
        hashes["source_pdf"] = source_pdf_sha256
    manifest_verified = manifest_sha256 == source.get("manifest_sha256") and not _validate_manifest_payload(manifest, source)
    if manifest_issues:
        return _report(
            status="BLOCKED", reason_code=manifest_issues[0]["code"], issues=manifest_issues,
            binding=binding, hashes=hashes,
            metrics={
                "manifest_file_sha256_verified": manifest_verified,
                "source_pdf_bytes_verified": source_pdf_verified,
                "section3_input_eligible": False,
                "section3_release_authorized": False,
            },
        )

    proof_rel = source["pagination"]["offset_proof_path"]
    proof_path = (root / proof_rel).resolve()
    if not _is_within(proof_path, root) or not proof_path.is_file():
        return _report(
            status="BLOCKED", reason_code="OFFSET_PROOF_ARTIFACT_UNAVAILABLE",
            issues=[_issue("OFFSET_PROOF_ARTIFACT_UNAVAILABLE")], binding=binding, hashes=hashes,
        )
    try:
        proof_sha = file_sha256(str(proof_path))
    except OSError as exc:
        return _report(
            status="BLOCKED", reason_code="OFFSET_PROOF_HASH_FAILED",
            issues=[_issue("OFFSET_PROOF_HASH_FAILED", exception_type=type(exc).__name__)],
            binding=binding, hashes=hashes,
        )
    hashes["offset_proof"] = proof_sha
    if proof_sha != source["pagination"]["offset_proof_sha256"]:
        return _report(
            status="BLOCKED", reason_code="OFFSET_PROOF_SHA_MISMATCH",
            issues=[_issue("OFFSET_PROOF_SHA_MISMATCH")], binding=binding, hashes=hashes,
        )
    offset_proof, proof_error = _read_json(proof_path, "OFFSET_PROOF")
    if proof_error or offset_proof is None:
        return _report(
            status="BLOCKED", reason_code="OFFSET_PROOF_UNREADABLE",
            issues=[proof_error or _issue("OFFSET_PROOF_UNREADABLE")], binding=binding, hashes=hashes,
        )
    proof_privacy = _privacy_issues(offset_proof, label="offset_proof")
    if proof_privacy:
        return _report(
            status="BLOCKED", reason_code=proof_privacy[0]["code"], issues=proof_privacy,
            binding=binding, hashes=hashes,
        )
    proof_schema_issues = _schema_issues(
        offset_proof, artifact="offset_proof", schema_filename=_STAGE_SCHEMAS["offset_proof"]
    )
    if proof_schema_issues:
        return _report(
            status="BLOCKED", reason_code=proof_schema_issues[0]["code"],
            issues=proof_schema_issues, binding=binding, hashes=hashes,
        )

    issues: list[dict[str, Any]] = []
    extraction_issues, element_index, record_index = _validate_extraction(extraction, source, binding)
    issues.extend(extraction_issues)
    issues.extend(_validate_offset_proof(offset_proof, source, extraction, binding))
    valuation_issues, valuation_index = _validate_valuation(
        valuation, binding, source, hashes["extraction"], element_index, record_index
    )
    issues.extend(valuation_issues)
    verdict_issues, verified_approvals = _validate_verdict_ledger(
        verdict_ledger, hashes["valuation"], valuation_index
    )
    issues.extend(verdict_issues)
    game_issues, blocked_game_items = _validate_game_pack(
        game_pack, binding, hashes["valuation"], valuation_index, element_index, verified_approvals
    )
    issues.extend(game_issues)

    if source.get("rights_status") not in {"AUTHORIZED", "SYNTHETIC_TEST_ONLY"}:
        issues.append(_issue("SOURCE_RIGHTS_NOT_CLEARED_FOR_SECTION3"))
    status = "PASS" if not issues else "BLOCKED"
    reason_code = "PIPELINE_TRACE_VALID_DRAFT_ONLY" if status == "PASS" else issues[0]["code"]
    metrics = {
        "extraction_page_records": len(extraction.get("pages", [])) if isinstance(extraction.get("pages"), list) else 0,
        "extracted_elements": len(element_index),
        "valuation_records": len(valuation_index),
        "game_items": len(game_pack.get("items", [])) if isinstance(game_pack.get("items"), list) else 0,
        "game_items_blocked_by_approval": blocked_game_items,
        "human_approval_ledger_matches": len(verified_approvals),
        "section3_input_eligible": (
            bool(game_pack.get("items")) and status == "PASS" and blocked_game_items == 0
            and source.get("rights_status") == "AUTHORIZED"
            and manifest_verified and source_pdf_verified
            and source["pagination"].get("offset_proof_status") == "TWO_READER_MATCHED"
        ),
        "section3_release_authorized": False,
        "manifest_file_sha256_verified": manifest_verified,
        "source_pdf_bytes_verified": source_pdf_verified,
        "offset_reader_independence_cryptographically_verified": False,
        "rights_status": source.get("rights_status"),
    }
    return _report(
        status=status, reason_code=reason_code, issues=issues, binding=binding,
        hashes=hashes, metrics=metrics,
    )


def export_extraction_ledger_xlsx(
    *,
    repo_root: str | Path,
    source_path: str | Path,
    extraction_path: str | Path,
    output_path: str | Path,
    source_pdf_path: str | Path | None = None,
) -> dict[str, Any]:
    """Create a new, Git-local Excel snapshot from a structured extraction ledger.

    Existing files are never overwritten. The JSON ledger remains canonical;
    this workbook is a readable, hash-linked snapshot and includes no PDF/image
    payloads or source text beyond the permitted short excerpt.
    """
    root = Path(repo_root).resolve()
    source_record_path = Path(source_path).resolve()
    extraction_file = Path(extraction_path).resolve()
    output = Path(output_path).resolve()
    sidecar = output.with_suffix(output.suffix + ".sha256")
    if not all(_is_within(path, root) for path in (source_record_path, extraction_file, output)):
        raise ValueError("SOURCE_EXTRACTION_AND_XLSX_MUST_BE_INSIDE_REPOSITORY")
    if output.suffix.lower() != ".xlsx":
        raise ValueError("OUTPUT_MUST_BE_XLSX")
    if output.exists() or sidecar.exists():
        raise FileExistsError("REFUSING_TO_OVERWRITE_EXISTING_ARTIFACT")

    source_record, source_error = _read_json(source_record_path, "SOURCE")
    payload, extraction_error = _read_json(extraction_file, "EXTRACTION")
    if source_error or source_record is None:
        raise ValueError("SOURCE_RECORD_INVALID")
    if extraction_error or payload is None or payload.get("schema_version") != EXTRACTION_SCHEMA:
        raise ValueError("EXTRACTION_LEDGER_INVALID")
    if _privacy_issues(source_record, label="source") or _privacy_issues(payload, label="extraction"):
        raise ValueError("EXTRACTION_PRIVACY_OR_COPYRIGHT_FIELD_FORBIDDEN")
    if _schema_issues(source_record, artifact="source", schema_filename=_STAGE_SCHEMAS["source"]):
        raise ValueError("SOURCE_RECORD_SCHEMA_INVALID")
    if _schema_issues(payload, artifact="extraction", schema_filename=_STAGE_SCHEMAS["extraction"]):
        raise ValueError("EXTRACTION_LEDGER_SCHEMA_INVALID")
    source_issues = _validate_source(source_record)
    if source_issues:
        raise ValueError("SOURCE_RECORD_INVALID")
    if source_record.get("rights_status") not in {"AUTHORIZED", "SYNTHETIC_TEST_ONLY"}:
        raise ValueError("SOURCE_RIGHTS_NOT_CLEARED")
    manifest_file = (root / source_record["manifest_path"]).resolve()
    if not _is_within(manifest_file, root) or not manifest_file.is_file():
        raise ValueError("SOURCE_MANIFEST_ARTIFACT_UNAVAILABLE")
    manifest_sha256 = file_sha256(str(manifest_file))
    if manifest_sha256 != source_record.get("manifest_sha256"):
        raise ValueError("SOURCE_MANIFEST_SHA_MISMATCH")
    manifest, manifest_error = _read_json(manifest_file, "SOURCE_MANIFEST")
    if manifest_error or manifest is None or _validate_manifest_payload(manifest, source_record):
        raise ValueError("SOURCE_MANIFEST_METADATA_INVALID")
    pdf_issues, source_pdf_verified, source_pdf_sha256 = _verify_source_pdf(
        pdf_path=source_pdf_path, repo_root=root, source=source_record, manifest=manifest
    )
    if pdf_issues:
        raise ValueError(pdf_issues[0]["code"])

    proof_rel = source_record["pagination"]["offset_proof_path"]
    proof_file = (root / proof_rel).resolve()
    if not _is_within(proof_file, root) or not proof_file.is_file():
        raise ValueError("OFFSET_PROOF_ARTIFACT_UNAVAILABLE")
    proof_sha256 = file_sha256(str(proof_file))
    if proof_sha256 != source_record["pagination"]["offset_proof_sha256"]:
        raise ValueError("OFFSET_PROOF_SHA_MISMATCH")
    proof, proof_error = _read_json(proof_file, "OFFSET_PROOF")
    if proof_error or proof is None:
        raise ValueError("OFFSET_PROOF_INVALID")
    if _privacy_issues(proof, label="offset_proof"):
        raise ValueError("OFFSET_PROOF_PRIVACY_FIELD_FORBIDDEN")
    if _schema_issues(proof, artifact="offset_proof", schema_filename=_STAGE_SCHEMAS["offset_proof"]):
        raise ValueError("OFFSET_PROOF_SCHEMA_INVALID")

    binding = _source_binding(source_record)
    extraction_issues, element_index, record_index = _validate_extraction(payload, source_record, binding)
    proof_issues = _validate_offset_proof(proof, source_record, payload, binding)
    if extraction_issues or proof_issues:
        raise ValueError("EXTRACTION_SOURCE_OR_OFFSET_BINDING_INVALID")
    pages = payload.get("pages")
    if not isinstance(pages, list):
        raise ValueError("EXTRACTION_PAGES_INVALID")
    source_record_sha256 = file_sha256(str(source_record_path))
    extraction_sha256 = file_sha256(str(extraction_file))
    source_binding = payload["source_binding"]
    rows: list[list[Any]] = []
    for page in pages:
        if not isinstance(page, dict):
            raise ValueError("EXTRACTION_PAGE_RECORD_INVALID")
        elements = page.get("elements", [])
        if not isinstance(elements, list):
            raise ValueError("EXTRACTION_ELEMENTS_INVALID")
        if not elements:
            rows.append([
                payload.get("bundle_id"), page.get("record_id"), page.get("pdf_page"),
                page.get("book_page"), page.get("printed_page_number"), page.get("page_status"),
                "", "", "", "", "", "",
            ])
        for element in elements:
            if not isinstance(element, dict):
                raise ValueError("EXTRACTION_ELEMENT_INVALID")
            excerpt = element.get("short_excerpt", "")
            if not _safe_excerpt(excerpt):
                raise ValueError("EXTRACTION_EXCERPT_OVER_TEN_WORDS")
            rows.append([
                payload.get("bundle_id"), page.get("record_id"), page.get("pdf_page"),
                page.get("book_page"), page.get("printed_page_number"), page.get("page_status"),
                element.get("element_id"), element.get("element_kind"),
                ", ".join(element.get("concept_candidate_ids", [])),
                ", ".join(element.get("operation_ids", [])),
                excerpt,
                ", ".join(element.get("evidence_refs", [])),
            ])
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from .excel_exporter import save_workbook_deterministically
    except Exception as exc:  # pragma: no cover - dependency/import failure path
        raise RuntimeError(f"EXCEL_EXPORT_UNAVAILABLE:{type(exc).__name__}") from exc

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "دفتر استخراج"
    ws.sheet_view.rightToLeft = True
    headers = [
        "bundle_id", "record_id", "pdf_page", "book_page", "printed_page_number", "page_status",
        "element_id", "element_kind", "concept_candidate_ids", "operation_ids", "short_excerpt",
        "evidence_refs",
    ]
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
    widths = [24, 24, 12, 12, 20, 22, 26, 18, 34, 28, 42, 52]
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    meta = wb.create_sheet("ردیابی")
    meta.sheet_view.rightToLeft = True
    metadata_rows = [
        ["field", "value"],
        ["workbook_kind", "DERIVED_EXTRACTION_SNAPSHOT"],
        ["source_record_path", source_record_path.relative_to(root).as_posix()],
        ["source_record_sha256", source_record_sha256],
        ["manifest_path", manifest_file.relative_to(root).as_posix()],
        ["manifest_file_sha256", manifest_sha256],
        ["source_pdf_bytes_verified", source_pdf_verified],
        ["source_pdf_actual_sha256", source_pdf_sha256],
        ["canonical_json_path", extraction_file.relative_to(root).as_posix()],
        ["canonical_json_sha256", extraction_sha256],
        ["ledger_schema", payload.get("schema_version")],
        ["bundle_id", payload.get("bundle_id")],
        ["offset_proof_path", proof_file.relative_to(root).as_posix()],
        ["offset_proof_sha256", proof_sha256],
        ["rights_status", source_record.get("rights_status")],
        ["source_id", source_binding.get("source_id")],
        ["book_code", source_binding.get("book_code")],
        ["edition_id", source_binding.get("edition_id")],
        ["manifest_id", source_binding.get("manifest_id")],
        ["manifest_sha256", source_binding.get("manifest_sha256")],
        ["source_pdf_sha256", source_binding.get("source_pdf_sha256")],
        ["official_workbook_modified", False],
        ["human_verdict_created", False],
        ["section3_release_authorized", False],
    ]
    for row in metadata_rows:
        meta.append(row)
    for cell in meta[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    for row in meta.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
    meta.column_dimensions["A"].width = 34
    meta.column_dimensions["B"].width = 78
    meta.freeze_panes = "A2"
    sha256 = save_workbook_deterministically(wb, output)
    return {
        "path": output.relative_to(root).as_posix(),
        "sha256": sha256,
        "rows": len(rows),
        "source_record": source_record_path.relative_to(root).as_posix(),
        "source_record_sha256": source_record_sha256,
        "manifest": manifest_file.relative_to(root).as_posix(),
        "manifest_sha256": manifest_sha256,
        "source_pdf_bytes_verified": source_pdf_verified,
        "source_pdf_sha256": source_pdf_sha256,
        "canonical_json": extraction_file.relative_to(root).as_posix(),
        "canonical_json_sha256": extraction_sha256,
        "offset_proof_sha256": proof_sha256,
        "overwrote_existing": False,
    }
