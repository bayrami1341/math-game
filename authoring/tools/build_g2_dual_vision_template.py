"""Generate a blank, standard Excel schema for independent visual capture.

No textbook bytes, render, invented printed claim or candidate question is embedded.
Usage: python authoring/tools/build_g2_dual_vision_template.py
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
OUTPUT = ROOT / 'authoring/templates/G2-C205-CH1-VISUAL-CAPTURE-v1.xlsx'
HEADERS = {
    'PAGES': ['page_key', 'source_filename', 'source_sha256_if_computed', 'source_identity_level',
              'pdf_page_index_1based_if_observed', 'printed_page_observed', 'chapter_observed',
              'page_image_sha256_if_computed', 'image_width_px', 'image_height_px',
              'page_complete_Y_N_U', 'pixel_view_status', 'uncertainty_latin_note'],
    'ELEMENTS': ['page_key', 'element_id_local', 'exercise_group_local', 'reading_order',
                 'kind', 'printed_exercise_number_verbatim', 'verbatim_fa', 'visual_description_fa',
                 'analysis_latin_note', 'x1_0_1000', 'y1_0_1000', 'x2_0_1000', 'y2_0_1000',
                 'evidence_mode', 'confidence_0_100', 'uncertainty_status'],
    'LINKS': ['page_key', 'from_element_id_local', 'to_element_id_local', 'link_kind',
              'rationale_fa', 'confidence_0_100', 'uncertainty_status'],
}
RUN_FIELDS = [
    ('schema', 'C205_DUAL_VISION_CAPTURE_v1'),
    ('agent_id', ''), ('run_id', ''), ('book_code', 'C205'), ('chapter_scope', 'FIRST_THREE_ACTUAL_PAGES_OF_CH1'),
    ('source_filename', ''), ('source_sha256_if_computed', ''), ('source_identity_level', 'UNKNOWN'),
    ('pixel_view_method', 'PIXELS_NOT_AVAILABLE'), ('render_or_image_method_latin', ''),
    ('capture_timestamp_utc', ''), ('source_edition_as_observed_fa', ''),
    ('global_limitations_fa', ''),
]


def make_book() -> Workbook:
    wb = Workbook()
    help_sheet = wb.active
    help_sheet.title = 'GUIDE'
    guidance = [
        'C205 chapter 1: independent visual capture — blank standard, not a scanned result.',
        'Give both agents the SAME source image/PDF and separate copies of this template; do not show them each other’s output or candidate game.',
        'Fill RUN; make one PAGES row per actual page; record all visible units in ELEMENTS and directed relations in LINKS.',
        'Shared page_key: C205-CH1-PILOT-01, -02, -03 in source page order. Local element IDs are agent-specific.',
        'Coordinates: 0..1000 relative to FULL page (top-left=0,0); do not guess cropped or invisible regions.',
        'Printed text/digits verbatim in Persian cells; analytical digits and method notes in latin_note.',
        'If image pixels are unavailable, set PIXELS_NOT_AVAILABLE; NEVER claim a visual match from extracted text.',
        'Green/red/orange are produced only AFTER independent capture and source/page preflight, not by either agent.',
        'Keep source PDF, page crops and filled full-text spreadsheets in a secure handoff, OUTSIDE the public Git repo.',
    ]
    for i, text in enumerate(guidance, 1):
        help_sheet.cell(i, 1, text)
    help_sheet.column_dimensions['A'].width = 120
    help_sheet.freeze_panes = 'A2'

    ws = wb.create_sheet('RUN')
    ws.append(['field', 'value', 'help'])
    for name, value in RUN_FIELDS:
        ws.append([name, value, 'Fill only from actual source/tool observation; UNKNOWN is allowed.'])
    ws.column_dimensions['A'].width = 36
    ws.column_dimensions['B'].width = 54
    ws.column_dimensions['C'].width = 75
    ws.freeze_panes = 'B2'
    for name, columns in HEADERS.items():
        sheet = wb.create_sheet(name)
        sheet.append(columns)
        sheet.freeze_panes = 'G2' if name == 'ELEMENTS' else 'B2'
        sheet.auto_filter.ref = f'A1:{get_column_letter(len(columns))}1'
        sheet.row_dimensions[1].height = 40
        for index, column in enumerate(columns, 1):
            sheet.column_dimensions[get_column_letter(index)].width = min(max(len(column) + 2, 17), 48)
        if name == 'ELEMENTS':
            for label in ('x1_0_1000', 'y1_0_1000', 'x2_0_1000', 'y2_0_1000'):
                letter = get_column_letter(columns.index(label) + 1)
                rule = DataValidation(type='whole', operator='between', formula1=0, formula2=1000, allow_blank=True)
                rule.error = 'Use a page-relative integer from 0 to 1000.'
                rule.errorTitle = 'Invalid bounding box'
                rule.showErrorMessage = True
                sheet.add_data_validation(rule)
                rule.add(f'{letter}2:{letter}10001')
            _list(sheet, columns, 'kind', 'TITLE,INSTRUCTION,QUESTION_TEXT,EXERCISE_NUMBER,NUMBER,BLANK,ILLUSTRATION,TABLE,CHART,BANNER,OTHER')
            _list(sheet, columns, 'evidence_mode', 'PIXELS,TEXT_LAYER,OCR,UNKNOWN')
        if name == 'LINKS':
            _list(sheet, columns, 'link_kind', 'REFERS_TO,INSTRUCTS,USES_VALUE_FROM,FILL_THIS_BLANK,ILLUSTRATES,ORDER_AFTER,OTHER')
        if name == 'PAGES':
            _list(sheet, columns, 'page_complete_Y_N_U', 'Y,N,U')
            _list(sheet, columns, 'pixel_view_status', 'RENDERED_PDF,PAGE_IMAGES,PIXELS_NOT_AVAILABLE')
            _list(sheet, columns, 'source_identity_level', 'VERIFIED_SHA,GIVEN_UNVERIFIED,UNKNOWN')
        if 'uncertainty_status' in columns:
            _list(sheet, columns, 'uncertainty_status', 'OBSERVED,UNCERTAIN,UNREADABLE')
        if 'confidence_0_100' in columns:
            letter = get_column_letter(columns.index('confidence_0_100') + 1)
            rule = DataValidation(type='whole', operator='between', formula1=0, formula2=100, allow_blank=True)
            rule.showErrorMessage = True
            sheet.add_data_validation(rule)
            rule.add(f'{letter}2:{letter}10001')
    for sheet in wb:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='17324A')
            cell.alignment = Alignment(wrap_text=True, vertical='center')
    return wb


def _list(sheet, columns: list[str], name: str, values: str) -> None:
    col = get_column_letter(columns.index(name) + 1)
    dv = DataValidation(type='list', formula1=f'"{values}"', allow_blank=True)
    dv.showErrorMessage = True
    sheet.add_data_validation(dv)
    dv.add(f'{col}2:{col}10001')


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    make_book().save(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    main()
