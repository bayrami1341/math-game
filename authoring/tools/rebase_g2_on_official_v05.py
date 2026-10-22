"""Rebase only nonoverlapping G2 draft pages onto official v0.5.
Usage: python authoring/tools/rebase_g2_on_official_v05.py --official /tmp/g2-official-v05.xlsx
Do not publish or certify this workbook; original approvals and all official rows remain unchanged.
"""
import argparse,re
from pathlib import Path
from hashlib import sha256
from collections import Counter
from copy import copy
from datetime import datetime
from zipfile import ZipFile,ZipInfo,ZIP_DEFLATED
import openpyxl
R=Path(__file__).resolve().parents[1]
p=argparse.ArgumentParser();p.add_argument('--official',type=Path,required=True);a=p.parse_args()
assert sha256(a.official.read_bytes()).hexdigest()=='f1f339a09a8cba3363c8c526e088063c2d733c62c5e7e4c236eba871f97d2656'
draft=R/'khshab-G2-MATH_v0.7_teacher-followup-draft.xlsx'
assert sha256(draft.read_bytes()).hexdigest()=='6f68b09bc1174e0d697bc2ad6ac085503c1c70a721bb8fd0acd000dfe4a250de'
mapfile=R/'PAGE-MAP-KH-G2-MATH-v1.2_teacher-followup.xlsx'
assert sha256(mapfile.read_bytes()).hexdigest()=='3bfd31b6f3ebfb4f6d8391c7dedd6c25b164b64140c32977fb2f3b2731db70b2'
w=openpyxl.load_workbook(a.official);d=openpyxl.load_workbook(draft)
assert w.sheetnames==d.sheetnames
fa_to_en=str.maketrans('۰۱۲۳۴۵۶۷۸۹','0123456789')
def page(v):
 t=str(v or '').translate(fa_to_en)
 return int(t) if t.isdigit() else None
def newrow(sh,at,source):
 ranges=[(x.min_row,x.max_row,x.min_col,x.max_col) for x in list(sh.merged_cells.ranges) if x.min_row>=at]
 for x in ranges:sh.unmerge_cells(start_row=x[0],end_row=x[1],start_column=x[2],end_column=x[3])
 sh.insert_rows(at)
 for x in ranges:sh.merge_cells(start_row=x[0]+1,end_row=x[1]+1,start_column=x[2],end_column=x[3])
 for col,c in enumerate(source,1):
  t=sh.cell(at,col);t.value=c.value
  if c.has_style:t._style=copy(c._style)
  if c.hyperlink:t._hyperlink=copy(c.hyperlink)
  if c.comment:t.comment=copy(c.comment)

S=w['نقشهٔ کتاب (S1)'];D=d['نقشهٔ کتاب (S1)']
existing={page(S.cell(r,2).value):S.cell(r,1).value for r in range(3,65) if page(S.cell(r,2).value)}
assert set(existing)==set(range(1,37))|set(range(55,73))
dpages={page(D.cell(r,2).value):r for r in range(3,155) if page(D.cell(r,2).value)}
assert set(dpages)==set(range(1,145))
missing=sorted(set(dpages)-set(existing))
assert len(missing)==90 and missing==list(range(37,55))+list(range(73,145))
# Mapping is confined to the twelve draft-only questions (chapter 3 and page 80).
renames={f'Q-{1069+i}':f'Q-{1093+i}' for i in range(12)}
oldids={w['بانک سؤال'].cell(r,1).value for r in range(4,96)}
assert not (set(renames.values())&oldids)
for i,pn in enumerate(missing):
 newrow(S,65+i,D[dpages[pn]])
 refs=S.cell(65+i,11)
 if refs.value and refs.value!='—':
  for old,new in renames.items():refs.value=re.sub(r'(?<![A-Z0-9-])'+re.escape(old)+r'(?![0-9])',new,refs.value)
  assert not any(re.search(r'(?<![A-Z0-9-])'+re.escape(old)+r'(?![0-9])',refs.value) for old in renames)
 assert S.cell(65+i,13).value=='نیاز به تأیید'
assert S.cell(155,1).value=='یادداشت'
# None of the six new proposed error codes exists in the official atlas.
A=w['اطلس خطا'];DA=d['اطلس خطا']
assert A.cell(17,1).value=='یادداشت'
atlasrows=[DA[r] for r in range(15,21)]
existingcodes={A.cell(r,1).value for r in range(4,17)}
assert not ({row[0].value for row in atlasrows}&existingcodes)
for i,row in enumerate(atlasrows):newrow(A,17+i,row)
assert A.cell(23,1).value=='یادداشت'
Q=w['بانک سؤال'];DQ=d['بانک سؤال']
assert Q.cell(96,1).value=='یادداشت'
for i in range(12):
 row=DQ[72+i];assert row[0].value==f'Q-{1069+i}'
 newrow(Q,96+i,row);Q.cell(96+i,1,renames[row[0].value]);Q.cell(96+i,17,str(row[16].value)+f' شناسهٔ قدیمی {row[0].value} به علت برخورد با خشاب رسمی v0.5 بازشماره‌گذاری شد؛ ردیف زرد است.')
assert Q.cell(108,1).value=='یادداشت'
# Ensure all original official values and styles (including signed/green rows) survive.
o=openpyxl.load_workbook(a.official)
for original in o:
 target=w[original.title]
 at,shift={'نقشهٔ کتاب (S1)':(65,90),'اطلس خطا':(17,6),'بانک سؤال':(96,12)}.get(original.title,(999999,0))
 for row in original:
  for c in row:
   t=target.cell(c.row+(shift if c.row>=at else 0),c.column)
   assert (c.value,c.style_id)==(t.value,t.style_id),(original.title,c.coordinate)
 for merged in original.merged_cells.ranges:
  exp=(merged.min_row+(shift if merged.min_row>=at else 0),merged.max_row+(shift if merged.max_row>=at else 0),merged.min_col,merged.max_col)
  assert any((x.min_row,x.max_row,x.min_col,x.max_col)==exp for x in target.merged_cells.ranges),(original.title,merged)
assert len({S.cell(r,1).value for r in range(3,155)})==152
assert len({Q.cell(r,1).value for r in range(4,108)})==104
assert S['E13'].value==o['نقشهٔ کتاب (S1)']['E13'].value
out=R/'khshab-G2-MATH_v0.8_official-rebased-draft.xlsx'
w.properties.modified=datetime(2026,10,22,0,0,0);w.save(out)
canon=out.with_suffix('.canonical-tmp')
with ZipFile(out,'r') as src,ZipFile(canon,'w') as dst:
 for name in src.namelist():
  info=ZipInfo(name,date_time=(1980,1,1,0,0,0));info.compress_type=ZIP_DEFLATED;info.external_attr=0o600 << 16
  data=src.read(name)
  if name=='docProps/core.xml':data=re.sub(rb'<dcterms:modified[^>]*>[^<]*</dcterms:modified>',b'<dcterms:modified xsi:type="dcterms:W3CDTF">2026-10-22T00:00:00Z</dcterms:modified>',data)
  dst.writestr(info,data)
canon.replace(out)
h=sha256(out.read_bytes()).hexdigest();(R/(out.name+'.sha256')).write_text(f'{h}  {out.name}\n')
print('rebased on official v0.5: +90 S1, +6 E, +12 Q, collisions 12 resolved; SHA',h)
