"""Generate an evidence-first review queue; no auto-approval or source PDF copying."""
import csv
from pathlib import Path
from hashlib import sha256
from collections import Counter
import openpyxl
R=Path(__file__).resolve().parents[1]
x=R/'khshab-G2-MATH_v0.8_official-rebased-draft.xlsx';m=R/'PAGE-MAP-KH-G2-MATH-v1.2_teacher-followup.xlsx'
assert sha256(x.read_bytes()).hexdigest()=='5a21328cad2c58d8a4fcff47effd63e0d5c3a0cd06601dc7aafbded78e4012d7'
assert sha256(m.read_bytes()).hexdigest()=='3bfd31b6f3ebfb4f6d8391c7dedd6c25b164b64140c32977fb2f3b2731db70b2'
w=openpyxl.load_workbook(x,read_only=True,data_only=True);s=w['نقشهٔ کتاب (S1)'];mapsh=openpyxl.load_workbook(m,read_only=True,data_only=True)['نقشهٔ صفحه‌ها'];q=w['بانک سؤال']
fa=str.maketrans('۰۱۲۳۴۵۶۷۸۹','0123456789')
by_page={int(str(s.cell(r,2).value).translate(fa)):r for r in range(3,155) if str(s.cell(r,2).value).translate(fa).isdigit()}
qcounts=Counter(q.cell(r,2).value for r in range(4,108))
confirmed={57:'مورد ۲: ۸۰ سکّهٔ ۱ ریالی (حکم محدود معلم)',80:'۳ تمرین شماره‌دار/۴ تساوی/جدول در انتها (حکم محدود معلم)',89:'شمارهٔ تمرین‌ها ۳ تا ۶ (حکم محدود معلم)',138:'عنوان تمرین (حکم محدود معلم)',139:'عنوان انتخاب نمودار (حکم محدود معلم)'}
out=R/'G2-REVIEW-QUEUE-v0.8.csv'
with out.open('w',newline='',encoding='utf-8-sig') as f:
 wr=csv.writer(f,lineterminator="\n");wr.writerow(['صفحهٔ کتاب','صفحهٔ PDF','فصل','شناسهٔ L','محل در S1','عنوان پیشنهادی','منشأ ردیف','هدف اولیهٔ عمومی؟','شمار سؤال','حکم N نقشهٔ پیشین','حکم محدود معلم','اقدام بازرس دوم','وضعیت گیت'])
 for p in range(37,145):
  r=by_page[p];id=s.cell(r,1).value;generic='بله' if 'فعالیت «' in str(s.cell(r,5).value) else 'خیر'
  orig='رسمی v0.5' if 55<=p<=72 else 'پیش‌نویس ادغامی v0.8'
  action=('بازنویسی هدف + مقابلهٔ تصویر و عددها' if generic=='بله' else 'مقابلهٔ سطر با صفحه و کنترل سازگاری موتور')
  if mapsh.cell(p+8,14).value=='✗':action+='؛ بررسی توضیح اصلاحی O/P نقشه'
  wr.writerow([p,p+6,s.cell(r,3).value,id,r,s.cell(r,4).value,orig,generic,qcounts[id],mapsh.cell(p+8,14).value,confirmed.get(p,'—'),action,'در انتظار بازرس؛ بدون مهر'])
assert len(by_page)==144
print('review queue: 108 pages, 74 generic goals, 5 limited teacher rulings; no green',out)
