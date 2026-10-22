"""Append provisional G2 page anchors 53..144 and a small page-80 question set.
Source: C205.pdf (SHA pinned) and existing corrected page map. Never writes PDF/JSON to repo.
Usage: python authoring/tools/extend_g2_full_draft.py --pdf /path/to/C205.pdf
"""
import argparse, ast, operator, re
from collections import Counter
from copy import copy
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED
from datetime import datetime
import openpyxl
import pymupdf

R=Path(__file__).resolve().parents[1]
p=argparse.ArgumentParser();p.add_argument('--pdf',type=Path,required=True);args=p.parse_args()
assert sha256(args.pdf.read_bytes()).hexdigest()=='a93934b6c418d4c6efe5ebb109bbcfe4f5fb6024fa1f3b1b309a4ff6da2746de'
base=R/'khshab-G2-MATH_v0.5_verified.xlsx'
assert sha256(base.read_bytes()).hexdigest()=='22d05e789d8ed41c8ec106e1e0c5d4d90051d13237200c9a0b65c28d9be3a1da'
mapfile=R/'PAGE-MAP-KH-G2-MATH-v1.1_verified.xlsx'
assert sha256(mapfile.read_bytes()).hexdigest()=='2c54be93a4ec016db61186ed7eed31fc34d51268d3e75a21adfcd6dfa4b6e839'
w=openpyxl.load_workbook(base)
m=openpyxl.load_workbook(mapfile,read_only=True,data_only=True)['نقشهٔ صفحه‌ها']
pdf=pymupdf.open(args.pdf);assert len(pdf)==152
fa=str.maketrans('0123456789','۰۱۲۳۴۵۶۷۸۹')
def f(s):return str(s).translate(fa)
def add(sh,at,vals,style):
 # Merge ranges below the insertion point are moved explicitly (openpyxl does not do this).
 moved=[(x.min_row,x.max_row,x.min_col,x.max_col) for x in list(sh.merged_cells.ranges) if x.min_row>=at]
 for x in moved:sh.unmerge_cells(start_row=x[0],end_row=x[1],start_column=x[2],end_column=x[3])
 sh.insert_rows(at)
 for x in moved:sh.merge_cells(start_row=x[0]+1,end_row=x[1]+1,start_column=x[2],end_column=x[3])
 for j,v in enumerate(vals,1):
  c=sh.cell(at,j,v);t=sh.cell(style,j)
  if t.has_style:c._style=copy(t._style)
  c.alignment=copy(t.alignment)

# Proven against rendered PDF pages; the O-column corrections of the map are
# NOT authoritative where they contradict the printed page number and artwork.
overrides={
57:('فعّالیت: پول و الگویابی','ارزش سکه‌های ۱، ۱۰، ۱۰۰ ریالی را با شمار سکه‌ها مقایسه کند.','سکه، ریال، ارزش پول، الگو','۲۰ سکهٔ ۱ ریالی = ۲ سکهٔ ۱۰ ریالی؛ ۸۰ سکهٔ ۱ ریالی = ۸ سکهٔ ۱۰ ریالی؛ ۱۰۰ سکهٔ ۱ ریالی = ۱۰ سکهٔ ۱۰ ریالی؛ سه الگوی عددی پایین صفحه','PDF چاپی ص ۵۷: مورد ۲ «۸۰ سکهٔ ۱ ریالی»، نه «۸۰ سکهٔ ۱۰ ریالی» در یادداشت O.'),
80:('حلّ مسئله + کامل کن','عدد گمشده را در تساوی‌های جمع و تفریق با حدس و آزمایش پیدا کند.','جای خالی، حدس، آزمایش، جمع، تفریق','۲۰+□=۴۳، □+۱۷=۵۱، ۶۰−□=۲۷، ۲۹+□=۶۳؛ جعبهٔ جدول شگفت‌انگیز پایین صفحه','PDF چاپی ص ۸۰: بنر «حلّ مسئله» بالا، جدول در جعبهٔ «کامل کن».'),
89:('ادامهٔ تمرین اندازه‌گیری (بدون بنر)','ضلع‌ها و خط‌ها را با خط‌کش اندازه بگیرد و نقطهٔ وسط را بیابد.','خط‌کش، ضلع، وسط خط، میلی‌متر','موارد شمارهٔ ۳ تا ۶؛ اندازه‌گیری ضلع شکل‌ها، وسط خط‌ها، حدس میلی‌متر و مقایسهٔ ضلع‌های مستطیل','PDF چاپی ص ۸۹: چهار شمارهٔ ۳،۴،۵،۶؛ یادداشت O شماره‌های ۴ تا ۷ را می‌گوید.'),
99:('زیر مسئله + الگویابی کن','سن محسن و خواهرش را با توجه به فاصلهٔ سنی و گذشت زمان حساب کند.','سن، سال، مسئله، الگوی ساعت','محسن هنگام تولد خواهرش ۸ سال داشت؛ ۱۰ سال بعد سن خواهر ۱۰ و محسن ۱۸ و مجموع ۲۸ است؛ پایین: الگوی ساعت','PDF چاپی ص ۹۹: تاب «زیر مسئله». عقربه‌های الگو مبنای سؤال بانک نیست.'),
135:('روش نمادین + کامل کن','مسیرهای یک مسئله را با نقطه و خط نمایش دهد.','نقطه، خط، پل، مسیر، روش نمادین','پل‌های میان مکان‌های تفریحی دریاچه و کناره؛ جعبهٔ «کامل کن» دو ستون چهار‌تایی با +۲ و −۳','PDF چاپی ص ۱۳۵: هشت جای پاسخ در دو ستون؛ شکل مسیر مبنای سؤال عددی نیست.'),
138:('تمرین + کامل کن','داده‌های درخت و کتاب را در جدول و نمودار تصویری نشان دهد.','نمودار تصویری، جدول، درخت، کتاب','تمرین: سعدی ۵۰، حافظ ۴۰، مولوی ۶۵ درخت؛ برای هر ۱۰ درخت یک نماد؛ نمودار کتابخانه و جعبهٔ +۳ و −۵','PDF چاپی ص ۱۳۸: «تمرین»؛ یادداشت O نقشه به اشتباه با صفحهٔ ۱۳۹ جابه‌جا کرده.'),
139:('انتخاب نمودار — فعّالیت','کارهای روزانهٔ سعید را در شکل زمانی و نمودار ستونی مقایسه کند.','نمودار ستونی، زمان، ساعت','فعالیت کارهای سعید از ۸ صبح تا ۸ بعدازظهر؛ رنگ‌آمیزی یک خانه برای هر ساعت و مقایسهٔ شکل با نمودار ستونی','PDF چاپی ص ۱۳۹: «انتخاب نمودار/فعّالیت»؛ یادداشت O نقشه به اشتباه با صفحهٔ ۱۳۸ جابه‌جا کرده.'),
}
source='پیام AI — موتور تألیف v۱.۰';yellow='نیاز به تأیید'
s=w['نقشهٔ کتاب (S1)'];assert s['A63'].value=='یادداشت'
for idx,page in enumerate(range(53,145)):
 r=page+8;assert m.cell(r,1).value==page and m.cell(r,2).value==page+6
 text=pdf[page+5].get_text();assert text.strip() and len(text)>45,(page,len(text))
 chap=str(m.cell(r,3).value);heading=m.cell(r,4).value;typ=m.cell(r,5).value
 raw=str(m.cell(r,7).value or '');first=raw.split('\n')[0].removeprefix('• ').strip()
 focus=first.split(':',1)[0].strip()[:75]
 title=(str(heading)+' — ' if heading and heading not in ['فعّالیت','تمرین','کار در کلاس','مرور فصل'] else '')+str(typ)
 goal=f'بتواند فعالیت «{focus}» را بر پایهٔ متن و تصویر صفحه توضیح دهد.'
 vocab='، '.join(str(x) for x in [heading if heading else typ, m.cell(r,11).value] if x and x!='None' and x!='—')
 activity=first[:240];visual='M_PAGE_EVIDENCE — '+str(m.cell(r,11).value or 'تصویر/متن کتاب')[:100]+'؛ بازنمایی تصویری، در صورت نیاز، مستقل از متن PDF بازبینی شود'
 issue=''
 if page in overrides:title,goal,vocab,activity,issue=overrides[page];visual='M_PAGE_EVIDENCE — '+('نمودار/تصویر چاپی؛ سؤال وابسته به شمارش تصویر ساخته نشد' if page in (89,135,138,139) else 'متن و شکل چاپی؛ نیازمند طراحی دستی تعامل')
 elif m.cell(r,14).value=='✗':issue='نقشهٔ مادر در ستون N رد شده؛ جزئیات ریز در انتظار بازبینی مستقل است. '+str(m.cell(r,15).value)[:180]
 note=f'ص کتاب {page} / PDF {page+6}، C205.pdf SHA a93934b6…a2746de؛ نقشهٔ مادر ردیف {r}. '+issue+' ردیف تألیف پیشنهادی است؛ نه مهر معلم نه تأیید موتور.'
 vals=[f'L-{10000+page}',f(page),f(chap),title,'دانش‌آموز '+goal,vocab,activity,visual,'متوسط','خارج از موتور','—',source,yellow,note]
 add(s,63+idx,vals,62)
assert s['A155'].value=='یادداشت'
# Four directly visible arithmetic equalities from book p.80; diagnostic labels
# are provisional, not active runtime modules. Wrong options are recomputed.
a=w['اطلس خطا'];assert a['A18'].value=='یادداشت'
errors=[
('E_COPY_KNOWN (پیشنهاد)','کپی‌کردن عدد معلوم','راهبردی','به جای محاسبهٔ مجهول یکی از داده‌های مسئله را می‌نویسد.','در ۲۰+□=۴۳ پاسخ ۲۰ می‌دهد.','تساوی با جای خالی؛ گزینهٔ دادهٔ معلوم','مجهول را جدا از دادهٔ معلوم نام‌گذاری کند.','متوسط'),
('E_COMBINE_GIVENS (پیشنهاد)','جمع‌کردن دو عدد دیده‌شده','مفهومی','به جای حل تساوی دو عددِ چاپ‌شده را با هم جمع می‌کند.','در ۲۰+□=۴۳ پاسخ ۶۳ می‌دهد.','تساوی با جای خالی؛ گزینهٔ جمع دو داده','جای‌گذاری پاسخ در تساوی و آزمایش آن.','متوسط'),
('E_ONE_OFF_INVERSE (پیشنهاد)','اختلاف یک‌واحدی در معکوس‌سازی','محاسباتی','در تفریق معکوس یک واحد کمتر به‌دست می‌آورد.','در ۲۰+□=۴۳ پاسخ ۲۲ به جای ۲۳ می‌دهد.','تساوی با جای خالی؛ گزینهٔ یک کمتر','حاصل را با جای‌گذاری دوباره کنترل کند.','متوسط'),
]
for i,e in enumerate(errors):add(a,18+i,[*e,source,yellow,'پیشنهاد ص ۸۰؛ تشخیص شناختی به بازرس دوم/معلم نیاز دارد؛ در runtime فعال نیست.'],17)
q=w['بانک سؤال'];assert q['A80'].value=='یادداشت'
# (printed form, arithmetic to independently evaluate, correct, copied, added, one-off)
items=[('در تساوی ۲۰ + □ = ۴۳، چه عددی در جای خالی می‌آید؟','43-20',23,20,63,22),
('در تساوی □ + ۱۷ = ۵۱، چه عددی در جای خالی می‌آید؟','51-17',34,17,68,33),
('در تساوی ۶۰ − □ = ۲۷، چه عددی در جای خالی می‌آید؟','60-27',33,60,87,32),
('در تساوی □ + ۲۹ = ۶۳، چه عددی در جای خالی می‌آید؟','63-29',34,29,92,33)]
codes=['E_COPY_KNOWN','E_COMBINE_GIVENS','E_ONE_OFF_INVERSE'];letters='الف ب ج د'.split()
for i,(prompt,work,answer,*wrong) in enumerate(items):
 tree=ast.parse(work,mode='eval').body;assert isinstance(tree,ast.BinOp) and isinstance(tree.op,ast.Sub)
 assert operator.sub(tree.left.value,tree.right.value)==answer
 assert len(set([answer,*wrong]))==4 and all(0<=n<1000 for n in [answer,*wrong])
 pos=i%4;opts=list(zip(wrong,codes));opts.insert(pos,(answer,None))
 # Equations on p.80 independently checked by substitution, not only inverse arithmetic.
 if i==0:assert 20+answer==43
 if i==1:assert answer+17==51
 if i==2:assert 60-answer==27
 if i==3:assert answer+29==63
 mapping='؛ '.join(f'{letters[j]}→{code}' for j,(_,code) in enumerate(opts) if code)
 vals=[f'Q-{1077+i}','L-10080',['CORE','CORE','DEPTH','SIDE'][i],'علامتی',prompt,'چهارگزینه‌ای',*(f(n) for n,_ in opts),letters[pos],mapping,'متوسط',f'پاسخ: {f(answer)}؛ عدد را در تساوی جای‌گذاری کن.',source,yellow,f'ضد J12 (لاتین): {work}={answer}; copied={wrong[0]}; combined={wrong[1]}; one-off={wrong[2]}. ص 80/PDF 86؛ کدها پیشنهادی و محتاج ممیزی دوم.']
 add(q,80+i,vals,79)
s['K90']='، '.join(f'Q-{i}' for i in range(1077,1081)) # page 80: row = page + 10 after starting 53 at 63
assert s['B90'].value=='۸۰' and q['A84'].value=='یادداشت'
# Value+style preservation for every preexisting cell (including seals and notes).
for old in openpyxl.load_workbook(base):
 new=w[old.title];offset={'نقشهٔ کتاب (S1)':(63,92),'اطلس خطا':(18,3),'بانک سؤال':(80,4)}.get(old.title,(10**9,0));at,n=offset
 for row in old:
  for c in row:
   d=new.cell(c.row+(n if c.row>=at else 0),c.column)
   assert c.value==d.value or (old.title=='نقشهٔ کتاب (S1)' and c.coordinate=='K90'),(old.title,c.coordinate)
   assert c.style_id==d.style_id,(old.title,c.coordinate,'style')
assert Counter(s.cell(r,2).value for r in range(47,155))==Counter(f(p) for p in range(37,145))
assert all(s.cell(r,13).value==yellow for r in range(63,155))
assert all(q.cell(r,16).value==yellow for r in range(80,84))
out=R/'khshab-G2-MATH_v0.6_full-draft.xlsx'
w.properties.modified=datetime(2026,10,22,0,0,0)
w.save(out)
# Stable ZIP timestamps/core properties for byte-identical reruns.
canon=out.with_suffix('.canonical-tmp')
with ZipFile(out,'r') as src, ZipFile(canon,'w') as dst:
 for name in src.namelist():
  info=ZipInfo(name,date_time=(1980,1,1,0,0,0));info.compress_type=ZIP_DEFLATED;info.external_attr=0o600 << 16
  data=src.read(name)
  if name=='docProps/core.xml':
   data=re.sub(rb'<dcterms:modified[^>]*>[^<]*</dcterms:modified>',b'<dcterms:modified xsi:type="dcterms:W3CDTF">2026-10-22T00:00:00Z</dcterms:modified>',data)
  dst.writestr(info,data)
canon.replace(out)
check=openpyxl.load_workbook(out,read_only=True,data_only=True)
assert check['نقشهٔ کتاب (S1)']['A155'].value=='یادداشت'
assert check['بانک سؤال']['A84'].value=='یادداشت'
h=sha256(out.read_bytes()).hexdigest();(R/(out.name+'.sha256')).write_text(f'{h}  {out.name}\n')
print('draft 144 page anchors, +92 S1 +3 E +4 Q, sha256',h)
