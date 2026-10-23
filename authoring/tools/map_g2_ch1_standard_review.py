"""Pin chapter-one addresses in the EXISTING official G2 Excel schema.

Does not modify/create an authoring workbook or claim C205 visual verification.
Output is a small review manifest outside the standard workbook (not a new Excel schema).
Usage: python authoring/tools/map_g2_ch1_standard_review.py --out-dir authoring/g2-completion
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import subprocess
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
OFFICIAL_REF = "origin/v10.3-beta:public/math-game/authoring/khshab-G2-MATH_v0.5.xlsx"
OFFICIAL_SHA = "f1f339a09a8cba3363c8c526e088063c2d733c62c5e7e4c236eba871f97d2656"
CANDIDATE = ROOT / "authoring/khshab-G2-MATH_v1.2_page-goals-draft.xlsx"
CANDIDATE_SHA = "18e6b4fa4957e968eca2b3b8066387858a627a99d7d52340064a1e8d8880c618"
MAP_SHEET = "نقشهٔ کتاب (S1)"
Q_SHEET = "بانک سؤال"
FA_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")
EXPECTED_S1 = ["شناسه (L-)", "صفحه/گستره", "فصل", "درس/موضوع", "اهداف رفتاری قابل سنجش",
               "واژگان کلیدی", "نوع فعالیت", "شرح شکل/تصویر", "پتانسیل تمرین",
               "سازگاری موتور", "لنگر Q-/E-", "منبع", "رنگ اعتبار", "یادداشت"]
EXPECTED_Q = ["شناسه (Q-)", "لنگر (L-)", "دامنه", "لایهٔ CPA", "صورت سؤال", "ساختار",
              "گزینهٔ الف", "گزینهٔ ب", "گزینهٔ ج", "گزینهٔ د", "پاسخ",
              "نقشهٔ گزینه→خطای مفهومی", "سطح", "معیار تصحیح (تشریحی)", "منبع", "رنگ اعتبار", "یادداشت"]


def _text(v: object) -> str:
    return str(v or "").translate(FA_DIGITS).strip()


def _headers(ws: openpyxl.worksheet.worksheet.Worksheet, row: int, expected: list[str]) -> bool:
    return [ws.cell(row, i).value for i in range(1, len(expected) + 1)] == expected


def inspect(official: openpyxl.Workbook, candidate: openpyxl.Workbook) -> dict:
    for wb in (official, candidate):
        if not _headers(wb[MAP_SHEET], 2, EXPECTED_S1) or not _headers(wb[Q_SHEET], 3, EXPECTED_Q):
            raise ValueError("STANDARD_WORKBOOK_HEADER_DRIFT")
    base = official[MAP_SHEET]
    draft = candidate[MAP_SHEET]
    base_q = official[Q_SHEET]
    draft_q = candidate[Q_SHEET]
    pages = []
    by_anchor = {}
    for page in range(1, 19):
        row = page + 3
        a, b = base.cell(row, 1).value, draft.cell(row, 1).value
        if a != f"L-{100+page}" or a != b or _text(base.cell(row, 2).value) != str(page) or _text(draft.cell(row, 2).value) != str(page):
            raise ValueError(f"CH1_PAGE_ANCHOR_DRIFT:{page}")
        if _text(base.cell(row, 3).value) != "1" or _text(draft.cell(row, 3).value) != "1":
            raise ValueError("CH1_CHAPTER_DRIFT")
        if any(base.cell(row, col).value != draft.cell(row, col).value for col in range(1, 15)):
            raise ValueError(f"CH1_S1_OFFICIAL_CANDIDATE_DRIFT:{page}")
        if _text(base.cell(row, 13).value) == "تأییدشده" or _text(draft.cell(row, 13).value) == "تأییدشده":
            raise ValueError("UNEXPECTED_GREEN_CH1_PAGE")
        pages.append({"book_page_candidate": page, "pdf_page_expected_unverified": page + 6,
                      "s1_id": a, "s1_row": row, "s1_cells_to_compare": f"D{row}:H{row},K{row}:L{row},N{row}",
                      "s1_stamp_cell_read_only": f"M{row}", "q_ids": [], "q_rows": [],
                      "visual_status": "NOT_SCANNED_IN_ARENA", "scope": "S1_AND_LINKED_Q_NOT_PRINTED_BOOK_VERIFIED"})
        by_anchor[a] = pages[-1]
    q_by_id = {}
    for row in range(4, draft_q.max_row + 1):
        qid = _text(draft_q.cell(row, 1).value)
        anchor = _text(draft_q.cell(row, 2).value)
        if anchor not in by_anchor:
            continue
        if qid in q_by_id or not qid.startswith("Q-"):
            raise ValueError("CH1_Q_ID_DUPLICATE_OR_MISSING")
        q_by_id[qid] = row
        matches = [r for r in range(4, base_q.max_row + 1) if _text(base_q.cell(r, 1).value) == qid]
        if len(matches) != 1 or row != matches[0]:
            raise ValueError(f"CH1_Q_OFFICIAL_ROW_DRIFT:{qid}")
        if any(base_q.cell(row, col).value != draft_q.cell(row, col).value for col in range(1, 18)):
            raise ValueError(f"CH1_Q_OFFICIAL_CANDIDATE_DRIFT:{qid}")
        if _text(base_q.cell(row, 16).value) == "تأییدشده":
            raise ValueError("UNEXPECTED_GREEN_CH1_Q")
        by_anchor[anchor]["q_ids"].append(qid)
        by_anchor[anchor]["q_rows"].append(row)
    if len(q_by_id) != 40 or set(q_by_id) != {f"Q-{i}" for i in range(1001, 1041)}:
        raise ValueError("CH1_Q_COUNT_OR_RANGE_DRIFT")
    # Ensure official sheet has no extra chapter-one questions missing from candidate.
    base_ch1 = [base_q.cell(r, 1).value for r in range(4, base_q.max_row + 1)
                if _text(base_q.cell(r, 2).value) in by_anchor]
    if set(base_ch1) != set(q_by_id):
        raise ValueError("CH1_Q_UNEXPECTED_OFFICIAL_ROWS")
    return {"schema": "G2_CH1_EXISTING_STANDARD_REVIEW_MAP_v1", "official_sha256": OFFICIAL_SHA,
            "candidate_sha256": CANDIDATE_SHA,
            "status": "CH1_OFFICIAL_CANDIDATE_CELL_IDENTICAL_NOT_BOOK_VERIFIED",
            "chapter": 1, "s1_page_rows": 18, "q_rows": len(q_by_id),
            "pages_scanned_by_arena_in_this_run": 0, "new_teacher_approvals": 0,
            "source_offset_verified": False, "workbook_mutated": False,
            "comparison_color_rule": "GREEN/RED/ORANGE is a sidecar review result, NEVER write to S1.M or Q.P",
            "pages": pages}


def make_report(report: dict, out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "G2-CH1-STANDARD-ROW-MAP-v1.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    with (out_dir / "G2-CH1-STANDARD-ROW-MAP-v1.csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, lineterminator="\n")
        writer.writerow(["صفحهٔ چاپی نامزد", "PDF پیشنهادی نامعتبر", "S1_ID", "ردیف S1", "سلول‌های مقابله",
                         "سلول مُهر دست‌نخورده", "شناسه‌های Q", "ردیف‌های Q", "وضعیت تصویر"])
        for p in report["pages"]:
            writer.writerow([p["book_page_candidate"], p["pdf_page_expected_unverified"], p["s1_id"], p["s1_row"],
                             p["s1_cells_to_compare"], p["s1_stamp_cell_read_only"],
                             ";".join(p["q_ids"]), ";".join(map(str,p["q_rows"])), p["visual_status"]])
    lines = ["# فصل اول C205 — تطبیق با **همان اکسل استاندارد** G2 (بدون ساخت قالب رقیب)", "",
             f"پایهٔ رسمی v0.5 `{OFFICIAL_SHA}` و نامزد زرد v1.2 `{CANDIDATE_SHA}`؛ سلول‌های فصل اول S1 و ۴۰ Q در دو نسخه برابرند. **این فقط برابری دو اکسل است، نه مشاهده یا تطابق با صفحهٔ کتاب.**",
             "", "اکسل موجود: `نقشهٔ کتاب (S1)` ردیف ۳ لنگر سبز قدیمی فصل L-100 و ردیف‌های **۴ تا ۲۱** صفحه‌های ۱ تا ۱۸ (L-101..L-118)، همگی زرد؛ `بانک سؤال` Q-1001..Q-1040 ردیف‌های **۴ تا ۴۳**، همه زرد. همان ستون‌های A:N و A:Q و همان شناسه‌ها حفظ می‌شوند.",
             "", "| صفحهٔ نامزد | ردیف S1 | Q موجود | ستون‌های مقابله با تصویر | مُهر |", "|---:|---:|---:|---|---|"]
    for p in report["pages"]:
        lines.append(f"| {p['book_page_candidate']} | {p['s1_row']} (`{p['s1_id']}`) | {len(p['q_ids'])} | D:H، K:L و N؛ سپس Q لنگرشده | M دست‌نخورده |")
    lines += ["", "**نقش ستون‌ها:** D موضوع، E هدف، F واژگان، G نوع فعالیت، H شرح شکل/تصویر، K لنگر Q/E، L منبع، N یادداشت. ستون M اعتبار S1، و ستون P اعتبار Q هرگز با رنگِ تطبیق دو AI جایگزین نمی‌شوند. Qها را با لنگر/متن/گزینه/کلید و دامنهٔ اقتباس بررسی کنید؛ همانند متن چاپی فرض نکنید.",
              "", "## پرامپت جاری برای عامل دوم (کپی کاملِ همان اکسل استاندارد)", "",
              "> فایل رسمی G2 v0.5 با SHA بالا و تصویر کامل صفحه‌های چاپی فصل اول C205 را مستقل بخوان. فقط روی کپیِ شخصی از همان اکسل استاندارد با همان شیت‌ها/ستون‌ها کار کن؛ فایل رسمی را دستکاری نکن و از خروجی عامل دیگر کمک نگیر. ابتدا صفحه‌های ۱ تا ۱۸ را واقعاً به‌صورت تصویر مشاهده کن؛ اگر فقط OCR/متن PDF می‌بینی، عدم دسترسی بصری را ثبت کن و هیچ سلولی را از روی تصویر تأیید نکن. شمارهٔ PDF و چاپی را از تصویر بررسی کن، +۶ را مفروض نگیر. برای هر ردیف S1 از ۴ تا ۲۱، D عنوان/موضوع، E هدف، F واژگان، G فعالیت، H شکل/تصویر، K لینک Q/E، L منشأ، N یادداشت را با *ادعای مشخص و ناحیهٔ مشاهده‌شده* مقابله کن. هرجا متن داخل شکل، رقم، جای‌خالی، جهت، شمار شکل یا رابطهٔ متن به شکل مبهم است، `UNCERTAIN` و دلیل/ناحیه را در گزارش جانبی نام ببر؛ حدس نزن. ۴۰ Q لنگرشدهٔ ردیف‌های ۴ تا ۴۳ بانک سؤال را جداگانه از نظر سازگاری با هدف، محاسبه و اقتباسی‌بودن وارسی کن؛ آن‌ها را تمرین چاپیِ عیناً یکسان فرض نکن. اگر پیشنهادی برای اصلاح داری، آدرس دقیق سلول و متن پیشنهادی را فقط در کپیِ زرد/گزارش جانبی بنویس، نه در سر رسمی؛ شناسه‌ها، کلیدها، `S1.M` و `Q.P`، سبک و شیت‌ها را تغییر نده. خروجی: نسخهٔ بازبینیِ همان workbook + یادداشت محدودِ شاهد و اختلاف با صفحه/آدرس سلول/ناحیه؛ PDF و برش کتاب در Git نباشد. سبز/مهر/انتشار اعلام نکن.",
              "", "## دو خروجی مستقل و داوری بدون توقف", "",
              "- به هر عامل **کپی جداگانه از همان اکسل رسمی v0.5 و تصویرهای همان فصل** بدهید. ابتدا فقط با تصویر، ادعاهای صفحه را مستقل بخواند، سپس هر سلول موجود D:H/K:L/N و Qهای لنگرشده را مقابله کند. عامل‌ها خروجی هم را نبینند و روی فایل رسمی یا نامزد کار نکنند؛ گزارش اختلاف باید به آدرس سلول موجود و شاهد/ناحیهٔ تصویر ارجاع دهد. اگر نیاز به پیشنهاد متن سلول دارد فقط کپیِ زرد را برگرداند؛ M/P/کلید را خودکار تغییر ندهد.",
              "- نتیجهٔ دو عامل در **گزارش جانبی** به ازای `فصل+صفحه+شناسهٔ L/Q+آدرس سلول+ادعای جزئی/ناحیه` همتاسازی می‌شود؛ شناسهٔ پرسش اقتباسی با شمارهٔ تمرین چاپی یکی فرض نمی‌شود. سبز=توافق دو مشاهدهٔ تصویر همان چاپ روی همان ادعا؛ قرمز=بی‌همتای واقعی/نبود جزء در یکی؛ نارنجی=جزء همتا با تفاوت رقم/متن/شکل/رابطه؛ فقدان منبع/پیکسل/صفحهٔ همسان=BLOCKED، نه رنگ. رنگ مقایسه رأی آموزشی، مهر یا انتشار نیست.",
              "- متنِ داخل تصویر، شمار جزء، رقم، جای‌خالی، ترتیب خواندن و پیوند متن↔شکل برای هر ادعای مهم در گزارش جانبی به‌صورت کوتاه با ناحیهٔ تصویر ثبت شود؛ چون اکسل اصلی ستون bounding box ندارد، **ستون تازه‌ای به آن تحمیل نمی‌کنیم**. PDF/crop و اکسل پرشدهٔ کتاب در Git قرار نگیرد.",
              "- بدون پاسخ مالک، موارد نارنجی/قرمز در صف بازرس دوم یا زرد باقی می‌مانند و بقیهٔ ردیف‌ها جلو می‌روند. هیچ صفحه در این اجرا تصویری اسکن نشده؛ offset چاپ/PDF پیشنهادی است.", ""]
    (out_dir / "G2-CH1-STANDARD-ROW-MAP-v1.fa.md").write_text("\n".join(lines), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out-dir", type=Path, required=True)
    args = parser.parse_args()
    try:
        official_bytes = subprocess.run(["git", "show", OFFICIAL_REF], cwd=ROOT, check=True, capture_output=True).stdout
        if hashlib.sha256(official_bytes).hexdigest() != OFFICIAL_SHA or hashlib.sha256(CANDIDATE.read_bytes()).hexdigest() != CANDIDATE_SHA:
            raise ValueError("PINNED_WORKBOOK_SHA_DRIFT")
        official = openpyxl.load_workbook(io.BytesIO(official_bytes), read_only=True, data_only=True)
        candidate = openpyxl.load_workbook(CANDIDATE, read_only=True, data_only=True)
        try:
            report = inspect(official, candidate)
        finally:
            official.close()
            candidate.close()
    except (OSError, ValueError, KeyError, subprocess.CalledProcessError) as exc:
        print("BLOCKED_CH1_STANDARD_MAP: " + str(exc), file=sys.stderr)
        return 2
    make_report(report, args.out_dir if args.out_dir.is_absolute() else ROOT / args.out_dir)
    print(json.dumps({k: report[k] for k in ("status", "s1_page_rows", "q_rows", "pages_scanned_by_arena_in_this_run")}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
