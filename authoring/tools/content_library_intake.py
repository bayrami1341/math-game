"""Private, read-only intake of an Achilles-approved STANDARD workbook.

This archives the exact XLSX bytes outside Git and writes a sidecar index. It does
NOT verify textbook scans, authenticate the approval receipt, change the Excel
schema, stamp anything green, or make a Z.ai API/capability claim.

python authoring/tools/content_library_intake.py --workbook /secure/approved.xlsx \
  --receipt /secure/achilles-receipt.json --out-dir /secure/library
Optional --objectives and --links are separately reviewed curriculum sidecars.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
S1 = "نقشهٔ کتاب (S1)"
Q = "بانک سؤال"
SHEETS = ["راهنما", "شناسنامهٔ خشاب", S1, "اطلس خطا", Q,
          "واژه‌نامهٔ ستون‌ها", "پروتکل ممیزی"]
S1_HEADERS = ["شناسه (L-)", "صفحه/گستره", "فصل", "درس/موضوع", "اهداف رفتاری قابل سنجش",
              "واژگان کلیدی", "نوع فعالیت", "شرح شکل/تصویر", "پتانسیل تمرین",
              "سازگاری موتور", "لنگر Q-/E-", "منبع", "رنگ اعتبار", "یادداشت"]
Q_HEADERS = ["شناسه (Q-)", "لنگر (L-)", "دامنه", "لایهٔ CPA", "صورت سؤال", "ساختار",
             "گزینهٔ الف", "گزینهٔ ب", "گزینهٔ ج", "گزینهٔ د", "پاسخ",
             "نقشهٔ گزینه→خطای مفهومی", "سطح", "معیار تصحیح (تشریحی)", "منبع", "رنگ اعتبار", "یادداشت"]
FA_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")
# Rubric v1 is Arena's PROPOSAL, not an assertion that any guide assigns numbers.
ROLE_POINTS = {"PRIMARY": 100, "SUPPORTING": 65, "ENRICHMENT": 25}
ALIGNMENT_FACTOR = {"DIRECT": 1.0, "TRANSFER": .75, "PREREQUISITE": .4, "INCIDENTAL": .1}
# Repetition is a *priority*, not a number of sessions; actual spacing/mastery is Z.ai's runtime concern.
PRACTICE_FACTOR = {"EMPHASIZED": 1.0, "SUPPORTING": .6, "OPTIONAL": .25}


def digest(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _required(data: dict, key: str) -> str:
    value = data.get(key)
    if not isinstance(value, str) or not value.strip():
        raise ValueError("MISSING_" + key.upper())
    return value.strip()


def _cells(ws, row: int, headers: list[str]) -> dict:
    return {h: ws.cell(row, col).value for col, h in enumerate(headers, start=1)}


def _uid(receipt: dict, kind: str, local_id: str) -> str:
    return f"{receipt['book_id']}:{receipt['edition_id']}:{kind}:{local_id}"


def _read_rows(book, receipt: dict) -> list[dict]:
    if book.sheetnames != SHEETS:
        raise ValueError("STANDARD_SHEET_DRIFT")
    pages, questions, seen = [], [], set()
    for sheet_name, headers, start, prefix, dest in (
        (S1, S1_HEADERS, 3, "L-", pages), (Q, Q_HEADERS, 4, "Q-", questions)
    ):
        ws = book[sheet_name]
        header_row = start - 1
        if [ws.cell(header_row, i).value for i in range(1, len(headers) + 1)] != headers:
            raise ValueError("STANDARD_HEADER_DRIFT:" + sheet_name)
        for row in range(start, ws.max_row + 1):
            local = ws.cell(row, 1).value
            if not isinstance(local, str) or not local.startswith(prefix) or "(نمونه)" in local:
                continue
            if local in seen:
                raise ValueError("DUPLICATE_LOCAL_ID:" + local)
            seen.add(local)
            dest.append({"uid": _uid(receipt, sheet_name, local), "local_id": local,
                         "sheet": sheet_name, "row": row, "fields": _cells(ws, row, headers)})
    if not pages:
        raise ValueError("NO_REAL_S1_ROWS")
    page_by_id = {p["local_id"]: p for p in pages}
    for item in questions:
        anchor = item["fields"][Q_HEADERS[1]]
        if anchor not in page_by_id:
            raise ValueError("QUESTION_ANCHOR_MISSING:" + item["local_id"])
        item["anchor_uid"] = page_by_id[anchor]["uid"]
    return pages + questions


def _validate_objectives(objectives: dict | None) -> dict:
    if objectives is None:
        return {}
    if objectives.get("schema") != "teacher-guide-objectives/v1" or objectives.get("source_kind") != "OFFICIAL_TEACHER_GUIDE":
        raise ValueError("OBJECTIVE_SOURCE_UNVERIFIED")
    sha = _required(objectives, "source_sha256")
    if len(sha) != 64 or any(c not in "0123456789abcdef" for c in sha):
        raise ValueError("INVALID_GUIDE_SHA")
    _required(objectives, "reviewed_by")
    by_id = {}
    for goal in objectives.get("objectives", []):
        key = _required(goal, "id")
        citation = goal.get("citation") or {}
        if (key in by_id or goal.get("role") not in ROLE_POINTS
                or goal.get("level") not in {"GENERAL", "SPECIFIC", "BEHAVIORAL"}
                or not citation.get("page") or not citation.get("section")
                or not goal.get("chapter") or not goal.get("description")):
            raise ValueError("GOAL_WITHOUT_TRACEABLE_GUIDE_CITATION:" + key)
        if "practice_emphasis" in goal:
            repeat_citation = goal.get("practice_citation") or {}
            if (goal["practice_emphasis"] not in PRACTICE_FACTOR or not repeat_citation.get("page")
                    or not repeat_citation.get("section")):
                raise ValueError("REPETITION_WEIGHT_WITHOUT_GUIDE_CITATION:" + key)
        by_id[key] = goal
    if not by_id:
        raise ValueError("EMPTY_VERIFIED_GOALS")
    parent_level = {"SPECIFIC": "GENERAL", "BEHAVIORAL": "SPECIFIC"}
    for key, goal in by_id.items():
        level = goal["level"]
        parent = by_id.get(goal.get("parent_id"))
        if level in parent_level and (not parent or parent["level"] != parent_level[level]
                                      or str(parent["chapter"]) != str(goal["chapter"])):
            raise ValueError("UNTRACEABLE_OBJECTIVE_HIERARCHY:" + key)
    return by_id


def _score(items: list[dict], objectives: dict | None, links: list[dict] | None) -> list[dict]:
    goals = _validate_objectives(objectives)
    by_uid = {item["uid"]: item for item in items}
    matches: dict[str, list[dict]] = defaultdict(list)
    for link in links or []:
        uid = link.get("content_uid")
        gid = link.get("objective_id")
        if uid not in by_uid or gid not in goals or link.get("alignment") not in ALIGNMENT_FACTOR or not link.get("reviewed_by"):
            raise ValueError("UNREVIEWED_OR_UNBOUND_GOAL_LINK")
        item = by_uid[uid]
        if item["sheet"] == Q:
            item = by_uid[item["anchor_uid"]]
        chapter = str(item["fields"][S1_HEADERS[2]] or "").translate(FA_DIGITS)
        if str(goals[gid]["chapter"]).translate(FA_DIGITS) != chapter:
            raise ValueError("GOAL_CHAPTER_MISMATCH")
        matches[uid].append(link)
    scored = []
    for item in items:
        candidate = matches[item["uid"]]
        if not candidate:
            scored.append({"content_uid": item["uid"], "educational_priority_0_100": None,
                           "attention_coefficient_0_1": None, "repetition_priority_coefficient_0_1": None,
                           "assessment": "NEEDS_HUMAN_WEIGHT_OR_GOAL_MAPPING",
                           "engine_readiness": "AWAITING_ZAI_CAPABILITY", "repetition_policy": "NOT_DERIVED",
                           "objective_ids": []})
            continue
        value = max(round(ROLE_POINTS[goals[link["objective_id"]]["role"]] *
                          ALIGNMENT_FACTOR[link["alignment"]]) for link in candidate)
        repeat = [round(ROLE_POINTS[goals[link["objective_id"]]["role"]] / 100 *
                        ALIGNMENT_FACTOR[link["alignment"]] *
                        PRACTICE_FACTOR[goals[link["objective_id"]]["practice_emphasis"]], 3)
                  for link in candidate if "practice_emphasis" in goals[link["objective_id"]]]
        scored.append({"content_uid": item["uid"], "educational_priority_0_100": value,
                       "attention_coefficient_0_1": value / 100,
                       "repetition_priority_coefficient_0_1": max(repeat) if repeat else None,
                       "assessment": "PRIORITY_PROPOSAL_NOT_APPROVAL", "engine_readiness": "AWAITING_ZAI_CAPABILITY",
                       "repetition_policy": "PRIORITY_ONLY_NOT_SESSION_FREQUENCY" if repeat else "NOT_DERIVED",
                       "objective_ids": sorted({m["objective_id"] for m in candidate})})
    return scored


def ingest(workbook: Path, receipt: dict, out_dir: Path, *, objectives: dict | None = None,
           links: list[dict] | None = None) -> dict:
    if out_dir.resolve() == ROOT.resolve() or ROOT.resolve() in out_dir.resolve().parents:
        raise ValueError("PRIVATE_ARCHIVE_MUST_BE_OUTSIDE_REPOSITORY")
    if (receipt.get("schema") != "achilles-approved-workbook/v1" or receipt.get("producer") != "ACHILLES"
            or receipt.get("approval_status") != "APPROVED"):
        raise ValueError("ACHILLES_APPROVAL_RECEIPT_REQUIRED")
    for key in ("book_id", "edition_id", "receipt_id", "approved_by", "workbook_sha256"):
        _required(receipt, key)
    if any(not re.fullmatch(r"[A-Za-z0-9_-]+", receipt[k]) for k in ("book_id", "edition_id")):
        raise ValueError("UNSAFE_ARCHIVE_NAMESPACE")
    if not re.fullmatch(r"[0-9a-f]{64}", receipt["workbook_sha256"]):
        raise ValueError("INVALID_WORKBOOK_SHA")
    if not workbook.is_file() or digest(workbook) != receipt["workbook_sha256"]:
        raise ValueError("APPROVED_WORKBOOK_SHA_MISMATCH")
    book = openpyxl.load_workbook(workbook, read_only=True, data_only=False)
    try:
        items = _read_rows(book, receipt)
    finally:
        book.close()
    if objectives and (objectives.get("book_id") != receipt["book_id"]
                       or objectives.get("edition_id") != receipt["edition_id"]):
        raise ValueError("GUIDE_BOOK_OR_EDITION_MISMATCH")
    assessments = _score(items, objectives, links)
    sha = receipt["workbook_sha256"]
    # Receipt is a trust-boundary declaration, not cryptographic proof of identity.
    manifest = {"schema": "arena-content-library-index/v1", "status": "ARCHIVED_NOT_RUNTIME_APPROVED",
                "approval_trust": "RECEIPT_ATTESTED_NOT_CRYPTOGRAPHICALLY_AUTHENTICATED",
                "book_id": receipt["book_id"], "edition_id": receipt["edition_id"],
                "workbook_sha256": sha, "achilles_receipt_id": receipt["receipt_id"],
                "workbook_sheets": SHEETS, "items": items,
                "expected_chapters": receipt.get("expected_chapters"),
                "observed_chapters": sorted({str(i["fields"][S1_HEADERS[2]]).translate(FA_DIGITS)
                                             for i in items if i["sheet"] == S1}),
                "breadth_depth_status": "NOT_ASSESSED",
                "engine_contract_status": "NOT_NEGOTIATED", "release_authorization": "NOT_GRANTED"}
    unweighted: dict[tuple[str, str], list[str]] = defaultdict(list)
    by_uid = {item["uid"]: item for item in items}
    for assessment in assessments:
        if assessment["educational_priority_0_100"] is not None:
            continue
        item = by_uid[assessment["content_uid"]]
        page = by_uid[item["anchor_uid"]] if item["sheet"] == Q else item
        chapter = str(page["fields"][S1_HEADERS[2]] or "").translate(FA_DIGITS)
        topic = str(page["fields"][S1_HEADERS[3]] or "").strip()
        unweighted[(chapter, topic)].append(item["uid"])
    review_groups = [{"chapter": chap, "topic": topic, "item_count": len(uids),
                      "sample_content_uids": uids[:3], "reason": "NEEDS_GUIDE_GOAL_MAPPING_OR_SCOPED_HUMAN_WEIGHT"}
                     for (chap, topic), uids in sorted(unweighted.items())]
    assessment_doc = {"schema": "arena-curriculum-assessment/v1", "workbook_sha256": sha,
                      "teacher_guide_sha256": objectives.get("source_sha256") if objectives else None,
                      "goal_catalog_sha256": hashlib.sha256(json.dumps(objectives, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest() if objectives else None,
                      "reviewed_links_sha256": hashlib.sha256(json.dumps(links, ensure_ascii=False, sort_keys=True).encode("utf-8")).hexdigest() if links else None,
                      "rubric": "ARENA_PRIORITY_POLICY_v1_NOT_OFFICIAL_GUIDE_NUMBERS",
                      "assessments": assessments, "grouped_review_queue": review_groups,
                      "unmapped_objective_ids": sorted({g["id"] for g in objectives.get("objectives", [])}
                                                      - {link["objective_id"] for link in links or []}) if objectives else [],
                      "release_authorization": "NOT_GRANTED"}
    folder = out_dir / receipt["book_id"] / receipt["edition_id"] / sha
    index_text = json.dumps(manifest, ensure_ascii=False, indent=2, default=str) + "\n"
    assessment_text = json.dumps(assessment_doc, ensure_ascii=False, indent=2, default=str) + "\n"
    assessment_sha = hashlib.sha256(assessment_text.encode("utf-8")).hexdigest()
    index_path = folder / "index.json"
    assessment_path = folder / f"assessment-{assessment_sha}.json"
    if folder.exists():
        if (digest(folder / "approved.xlsx") != sha or not index_path.is_file()
                or index_path.read_text(encoding="utf-8") != index_text):
            raise ValueError("IMMUTABLE_ARCHIVE_COLLISION")
    else:
        folder.mkdir(parents=True)
        shutil.copyfile(workbook, folder / "approved.xlsx")
        index_path.write_text(index_text, encoding="utf-8")
    if assessment_path.exists():
        if assessment_path.read_text(encoding="utf-8") != assessment_text:
            raise ValueError("ASSESSMENT_COLLISION")
    else:
        assessment_path.write_text(assessment_text, encoding="utf-8")
    return {**manifest, "assessment_file": assessment_path.name, "assessments": assessments,
            "grouped_review_queue": review_groups, "unmapped_objective_ids": assessment_doc["unmapped_objective_ids"]}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    for option in ("workbook", "receipt", "out-dir"):
        parser.add_argument("--" + option, type=Path, required=True)
    parser.add_argument("--objectives", type=Path)
    parser.add_argument("--links", type=Path)
    args = parser.parse_args()
    receipt = json.loads(args.receipt.read_text(encoding="utf-8"))
    objectives = json.loads(args.objectives.read_text(encoding="utf-8")) if args.objectives else None
    links = json.loads(args.links.read_text(encoding="utf-8")) if args.links else None
    result = ingest(args.workbook, receipt, args.out_dir, objectives=objectives, links=links)
    print(json.dumps({"status": result["status"], "book_id": result["book_id"],
                      "items": len(result["items"]),
                      "weighted": sum(a["educational_priority_0_100"] is not None for a in result["assessments"]),
                      "engine_contract_status": result["engine_contract_status"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
