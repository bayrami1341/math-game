"""Independent, read-only structural and answer audit for the yellow full-page draft."""
from collections import Counter
from hashlib import sha256
from pathlib import Path
import openpyxl
R=Path(__file__).resolve().parents[1]
old=openpyxl.load_workbook(R/'khshab-G2-MATH_v0.5_verified.xlsx')
new=openpyxl.load_workbook(R/'khshab-G2-MATH_v0.6_full-draft.xlsx')
assert old.sheetnames==new.sheetnames
for name in old.sheetnames:
 a,b=old[name],new[name]
 start,offset={'نقشهٔ کتاب (S1)':(63,92),'اطلس خطا':(18,3),'بانک سؤال':(80,4)}.get(name,(999999,0))
 for row in a:
  for c in row:
   target=b.cell(c.row+(offset if c.row>=start else 0),c.column)
   assert (c.value,c.style_id)==(target.value,target.style_id),(name,c.coordinate)
 for m in a.merged_cells.ranges:
  expected=(m.min_row+(offset if m.min_row>=start else 0),m.max_row+(offset if m.max_row>=start else 0),m.min_col,m.max_col)
  assert any((x.min_row,x.max_row,x.min_col,x.max_col)==expected for x in b.merged_cells.ranges),(name,m)
s=new['نقشهٔ کتاب (S1)'];q=new['بانک سؤال'];e=new['اطلس خطا']
fa=str.maketrans('۰۱۲۳۴۵۶۷۸۹','0123456789')
anchors={int(s.cell(r,2).value.translate(fa)):s.cell(r,1).value for r in range(4,155) if str(s.cell(r,2).value).translate(fa).isdigit()}
assert set(range(1,145))==set(anchors) and len(anchors)==144
assert all(s.cell(r,13).value=='نیاز به تأیید' for r in range(63,155))
assert all(e.cell(r,10).value=='نیاز به تأیید' for r in range(18,21))
assert all(q.cell(r,16).value=='نیاز به تأیید' for r in range(80,84))
assert s['K90'].value=='، '.join(f'Q-{i}' for i in range(1077,1081))
assert [q.cell(r,2).value for r in range(80,84)]==['L-10080']*4
checks=[(20,23,43,lambda a,b:a+b),(34,17,51,lambda a,b:a+b),(60,33,27,lambda a,b:a-b),(34,29,63,lambda a,b:a+b)]
for r,(a,b,expected,fn) in enumerate(checks,80):
 assert fn(a,b)==expected
 opts=[int(str(q.cell(r,c).value).translate(fa)) for c in range(7,11)]
 assert len(set(opts))==4
 key='الف ب ج د'.split().index(q.cell(r,11).value)
 assert opts[key]==[23,34,33,34][r-80]
 assert len(__import__('re').findall('→E_',q.cell(r,12).value))==3
assert s['D148'].value.startswith('تمرین') # book 138 = row 148
assert s['D149'].value.startswith('انتخاب نمودار')
assert '۸۰ سکهٔ ۱ ریالی' in s['G67'].value
assert '۳ تا ۶' in s['G99'].value
assert new['بانک سؤال']['E14'].value.startswith('شکل زیر ۱۳ مثلث دارد')
assert '۳۵−۱۷' in s['G41'].value
h=sha256((R/'khshab-G2-MATH_v0.6_full-draft.xlsx').read_bytes()).hexdigest()
assert (R/'khshab-G2-MATH_v0.6_full-draft.xlsx.sha256').read_text().startswith(h+'  ')
print('PASS: 7 sheets; all old values/styles/merges unchanged; 144 distinct page anchors; 92+3+4 yellow; independent substitution and key checks; SHA',h)
