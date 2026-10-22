"""Evidence-bound depth pass for 16 generic chapter-5 S1 rows; no new approvals.
Usage: python authoring/tools/deepen_g2_ch5.py --pdf /tmp/C205.pdf
"""
import argparse,re
from hashlib import sha256
from copy import copy
from pathlib import Path
from datetime import datetime
from zipfile import ZipFile,ZipInfo,ZIP_DEFLATED
import openpyxl,pymupdf
R=Path(__file__).resolve().parents[1]
p=argparse.ArgumentParser();p.add_argument('--pdf',type=Path,required=True);a=p.parse_args()
base=R/'khshab-G2-MATH_v0.8_official-rebased-draft.xlsx'
assert sha256(base.read_bytes()).hexdigest()=='5a21328cad2c58d8a4fcff47effd63e0d5c3a0cd06601dc7aafbded78e4012d7'
assert sha256(a.pdf.read_bytes()).hexdigest()=='a93934b6c418d4c6efe5ebb109bbcfe4f5fb6024fa1f3b1b309a4ff6da2746de'
w=openpyxl.load_workbook(base);s=w['نقشهٔ کتاب (S1)'];pdf=pymupdf.open(a.pdf)
# page -> (one literal witness in PDF text layer, observable objective, concise activity paraphrase)
items={
73:('نوزادان','بتواند یک کاربرد سنجش قد یا وزن نوزاد را در پایش رشد نام ببرد.','آغاز فصل: گفت‌وگو دربارهٔ اندازه‌گیری قد و وزن برای پیگیری رشد؛ تصویر آغازین مبنای شمارش یا سؤال عددی نیست.'),
74:('گیره','بتواند طول یک شیء را با یک واحد غیراستاندارد میان دو تعداد پیاپی بیان کند.','فعالیت ۱: برآورد طول مداد/شانه/ساعت با گیره، پاک‌کن یا چینه؛ ۲ و ۳: سنجش کتاب با تراش و میز با مداد؛ گفت‌وگو دربارهٔ دشواری واحدهای متفاوت.'),
75:('چینه','بتواند توضیح دهد چرا سنجش یک مداد با گیره، پاک‌کن و چینه عددهای متفاوتی می‌دهد.','کار در کلاس ۱–۳: سنجش یک مداد با سه وسیله و پیشنهاد واحد مشترک؛ کامل کن: حرکت نیم‌ساعتهٔ عقربهٔ بزرگ، سؤال ساعت تصویری نیازمند مشاهدهٔ شکل.'),
76:('چینه','بتواند مقدار طول را به صورت شمار واحد بزرگ و کوچک در یک نوار/چینه بیان کند.','تمرین: چهار مکعبِ کوچک یک چینهٔ بزرگ می‌سازند؛ نمایش طول چند شیء با واحد بزرگ و کوچک؛ کامل کن: حرکت یک‌ربعِ عقربهٔ ساعت.'),
77:('نوار','بتواند با ساخت نوارِ تاخورده طول شیء را بسنجد و علت تفاوت پاسخ دو نفر را با اندازهٔ واحدهایشان توضیح دهد.','فعالیت ۱–۴: نوار کاغذی علی، ساخت نوار شخصی، اندازه‌گیری وسایل و مقایسهٔ اندازهٔ کتاب با نوار دو نفر.'),
78:('نوار','بتواند با تقسیم هر قسمت نوار به چهار قسمت کوچک‌تر، اندازه‌گیری دقیق‌تر را توصیف کند.','فعالیت ۱–۴: تا کردن دوبارهٔ نوار، بیان طول به قسمت بزرگ/کوچک، مقایسهٔ دقت و گفت‌وگو دربارهٔ تفاوت نوارهای دو نفر.'),
79:('خط','بتواند طول خط را با واحد بزرگ و یک‌دهم آن ثبت و دلیل دقیق‌تر شدن سنجش را توضیح دهد.','تمرین ۱–۲: نوار فاطمه با تقسیم هر قسمت به ده قسمت مساوی؛ بیان اندازهٔ خط‌ها با قسمت بزرگ و کوچک و توضیح چرایی دقت.'),
81:('14','بتواند برای ۴ حیوان با ۱۴ پا، حدسِ تعداد گوسفند و شترمرغ را با شمردن پاها بیازماید.','حل مسئله با حدس و آزمایش: ۲ گوسفند و ۲ شترمرغ، ۱۲ پا دارند؛ برای ۱۴ پا، ۳ گوسفند و ۱ شترمرغ به دست می‌آید؛ کامل کنِ ساعت مستقل و تصویری است.'),
82:('سانتی','بتواند با یک واحد مشترکِ سانتی‌متر، دلیل قابل‌مقایسه شدن طول اشیا را توضیح دهد.','فعّالیت ۱–۳: مقایسهٔ دفتر/کتاب و مدادها، سنجیدن مناسب بودن میز برای اتاق؛ معرفی سانتی‌متر و خط‌کش؛ خلّاق باش رسم با خط‌کش.'),
83:('13','بتواند هنگام خواندن خط‌کش، شروع سنجش و محل انتهای خط را کنترل کند.','کار در کلاس: بررسی ادعای ۱۳ سانتی‌متر با خط‌کش ۲۰ سانتی‌متری، اندازه‌گیری خط‌ها و رسم خط‌های ۴ و ۷ سانتی‌متری؛ خلّاق باش ساعت روی گوشهٔ مستطیل.'),
84:('متر','بتواند درازا و پهنای کتاب را با خط‌کش بسنجد و نتیجه را به سانتی‌متر یا بازهٔ دو سانتی‌متر گزارش کند.','تمرین ۱–۳: خواندن بازهٔ طول مداد، اندازه‌گیری درازا/پهنای کتاب و جدول اندازه‌گیری چند شیء با متر خیاطی یا خط‌کش.'),
85:('10','بتواند نشان دهد ۱ سانتی‌متر از ۱۰ میلی‌متر تشکیل می‌شود و طول را با سانتی‌متر و میلی‌متر بیان کند.','فعّالیت فاطمه: تقسیم هر سانتی‌متر خط‌کش به ده قسمت و نام‌گذاری میلی‌متر؛ اندازه‌گیری دو خط با واحد بزرگ و کوچک.'),
86:('10','بتواند ۱۰، ۱۱، ۱۲ و ۲۰ میلی‌متر را با سانتی‌متر و باقیماندهٔ میلی‌متر بازنویسی کند.','کار در کلاس ۱–۴: شمردن خط‌های ریز خط‌کش از صفر، تبدیل میلی‌متر/سانتی‌متر، سنجش خط و کتاب/پاک‌کن با هر دو نمایش.'),
87:('قد','بتواند طول یک خط را حدس بزند، اندازه بگیرد و حدس را با اندازهٔ واقعی مقایسه کند.','تمرین ۱–۳: حدس طول شکل‌ها و خط‌ها، سپس اندازه‌گیری با خط‌کش؛ سنجش قد خود و اعضای خانواده با متر خیاطی.'),
88:('خط','بتواند ترتیب کار با خط‌کش را برای سنجیدن طول یک خط توضیح دهد.','مرور فصل: فرهنگ نوشتن دربارهٔ سنجش خط؛ تمرین سنجش کتاب با پاک‌کن و مقایسهٔ واحدها؛ اندازه‌گیری چند مداد رنگی.'),
90:('12','بتواند ۱۲ قطعهٔ پنج‌مربعی را کنار هم بچیند و سطح مستطیل داده‌شده را بپوشاند.','معمّای دست‌سازهٔ ۱۲ قطعهٔ پنج‌مربعی برای پوشاندن مستطیل؛ فرهنگ خواندن دربارهٔ انجام دادن تمرین‌ها به دست خود دانش‌آموز. فعالیت هندسی در موتور فعلی اجرا نشده.'),
}
assert len(items)==16 and set(items)==set(range(73,91))-{80,89}
fa=str.maketrans('۰۱۲۳۴۵۶۷۸۹','0123456789')
rows={int(str(s.cell(r,2).value).translate(fa)):r for r in range(65,155) if str(s.cell(r,2).value).translate(fa).isdigit()}
for pn,(witness,objective,activity) in items.items():
 r=rows[pn];text=pdf[pn+5].get_text();assert witness in text,(pn,witness)
 assert 'فعالیت «' in s.cell(r,5).value and s.cell(r,13).value=='نیاز به تأیید'
 s.cell(r,5,'دانش‌آموز '+objective)
 s.cell(r,7,activity)
 s.cell(r,14,str(s.cell(r,14).value)+f' | عمق‌کاوی زرد فصل ۵: متن PDF {pn+6} / ص چاپی {pn} و نقشهٔ v1.2؛ شاهد متنی «{witness}». شمارش/اندازهٔ اشکال و کلید تصویری هنوز به بازرس مستقل نیاز دارد؛ مُهر معلم صادر نشده.')
old=openpyxl.load_workbook(base)
for o in old:
 t=w[o.title]
 for line in o:
  for c in line:
   d=t[c.coordinate]
   assert c.style_id==d.style_id
   if o.title=='نقشهٔ کتاب (S1)' and c.row in rows.values() and c.row in [rows[n] for n in items] and c.column in (5,7,14):continue
   assert c.value==d.value,(o.title,c.coordinate)
assert all(s.cell(rows[n],13).value=='نیاز به تأیید' for n in items)
out=R/'khshab-G2-MATH_v0.9_ch5-depth-draft.xlsx';w.properties.modified=datetime(2026,10,22,0,0,0);w.save(out)
canon=out.with_suffix('.canonical-tmp')
with ZipFile(out,'r') as src,ZipFile(canon,'w') as dst:
 for name in src.namelist():
  info=ZipInfo(name,date_time=(1980,1,1,0,0,0));info.compress_type=ZIP_DEFLATED;info.external_attr=0o600 << 16
  data=src.read(name)
  if name=='docProps/core.xml':data=re.sub(rb'<dcterms:modified[^>]*>[^<]*</dcterms:modified>',b'<dcterms:modified xsi:type="dcterms:W3CDTF">2026-10-22T00:00:00Z</dcterms:modified>',data)
  dst.writestr(info,data)
canon.replace(out)
h=sha256(out.read_bytes()).hexdigest();(R/(out.name+'.sha256')).write_text(f'{h}  {out.name}\n')
print('depthened 16 yellow S1 rows (E/G/N), SHA',h)
