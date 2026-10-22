"""Repeatable fail-closed technical G2 gate against the latest official branch.

Run: python authoring/tools/g2_gate.py [--report authoring/G2-GATE-LATEST.fa.md]
Needs: openpyxl, node/npm, git, and an up-to-date origin/v10.3-beta ref.
Only builds a temporary JSON; never publishes or changes workbook approval statuses.
Human evidence, pedagogical review and release are intentionally separate gates.
"""
import argparse,json,re,subprocess,sys,tempfile
from collections import Counter
from hashlib import sha256
from pathlib import Path
import openpyxl

ROOT=Path(__file__).resolve().parents[2]
OFFICIAL_REF='origin/v10.3-beta'
OFFICIAL_PATH='public/math-game/authoring/khshab-G2-MATH_v0.5.xlsx'
PINNED_OFFICIAL='f1f339a09a8cba3363c8c526e088063c2d733c62c5e7e4c236eba871f97d2656'
CANDIDATE=ROOT/'authoring/khshab-G2-MATH_v1.2_page-goals-draft.xlsx'
PINNED_CANDIDATE='18e6b4fa4957e968eca2b3b8066387858a627a99d7d52340064a1e8d8880c618'

p=argparse.ArgumentParser();p.add_argument('--report',type=Path,default=ROOT/'authoring/G2-GATE-LATEST.fa.md');a=p.parse_args()
steps=[];counts={};blocked=[]
def run(cmd,**kwargs):
 r=subprocess.run(cmd,cwd=kwargs.pop('cwd',ROOT),capture_output=True,text=True,**kwargs)
 if r.returncode:raise RuntimeError(f"{' '.join(map(str,cmd))} => {r.returncode}\n{(r.stdout+r.stderr)[-1400:]}")
 return r.stdout
try:
 assert CANDIDATE.is_file() and sha256(CANDIDATE.read_bytes()).hexdigest()==PINNED_CANDIDATE,'Candidate SHA drift: inspect first'
 steps.append('نامزد v1.2 و SHA ثابت: PASS')
 official=subprocess.run(['git','show',f'{OFFICIAL_REF}:{OFFICIAL_PATH}'],cwd=ROOT,capture_output=True,check=True).stdout
 actual=sha256(official).hexdigest()
 assert actual==PINNED_OFFICIAL,f'Official baseline drift: expected {PINNED_OFFICIAL}, got {actual}. Reconcile before publishing.'
 steps.append('مبنای رسمیِ تازه‌خوانده‌شده و SHA ثابت: PASS')
 with tempfile.TemporaryDirectory(prefix='g2-gate-') as folder:
  temp=Path(folder);op=temp/'official.xlsx';op.write_bytes(official)
  run([sys.executable,str(ROOT/'authoring/tools/audit_g2_official_rebase.py'),'--official',str(op)])
  run([sys.executable,str(ROOT/'authoring/tools/audit_g2_page_goals.py')])
  steps.append('ممیزی تفاضل، سبک/مهر و اهداف صفحه‌ای: PASS')
  paths=['tools/cartridge-converter','public/math-game/src/v10/station','public/math-game/cartridges','public/math-game/tests']
  archive=temp/'source.tar'
  with archive.open('wb') as f:
   r=subprocess.run(['git','archive',OFFICIAL_REF,*paths],cwd=ROOT,stdout=f,stderr=subprocess.PIPE)
   if r.returncode:raise RuntimeError(r.stderr.decode())
  runner=temp/'runner';runner.mkdir()
  run(['tar','-xf',str(archive),'-C',str(runner)])
  run(['npm','install','--no-audit','--no-fund','--silent'],cwd=runner/'tools/cartridge-converter',timeout=120)
  tests=run(['node','--test','public/math-game/tests/v10.3-cartridge.test.js','public/math-game/tests/v10.3-families.test.js'],cwd=runner,timeout=120)
  assert re.search(r'# pass 28\b',tests) and re.search(r'# fail 0\b',tests),'Official fixture test count changed; review output'
  steps.append('۲۸/۲۸ آزمون نسخهٔ رسمی: PASS')
  output=run(['node','tools/cartridge-converter/convert.mjs',str(CANDIDATE)],cwd=runner,timeout=120)
  assert '0 قرمز، 2 زرد' in output,output[-800:]
  data=runner/'public/math-game/cartridges/KH-G2-MATH.json';assert data.exists()
  script="""import fs from 'node:fs'; import {validateCartridge} from './public/math-game/src/v10/station/cartridge-schema.js';
const c=JSON.parse(fs.readFileSync(process.argv[1],'utf8')); const v=validateCartridge(c);
if(!v.ok || v.stats.lessons!==152 || v.stats.questions!==104 || v.stats.lessonsConfirmed!==8 || v.stats.questionsConfirmed!==0) {console.error(JSON.stringify(v));process.exit(1)}
console.log(JSON.stringify({stats:v.stats,warnings:v.warnings}));"""
  validated=json.loads(run(['node','--input-type=module','-e',script,str(data)],cwd=runner,timeout=120))
  assert len(validated['warnings'])==2
  run([sys.executable,str(ROOT/'authoring/tools/audit_g2_page_goals.py'),'--json',str(data)])
  counts=validated['stats'];steps.append('مبدل واقعی + اعتبارسنج JSON موقت + ممیزی شناسه‌ها: PASS (۰ قرمز/۲ هشدار قدیمی)')
  # Never present converter's 'JSON published' as a product release: temp dir is destroyed.
  book=openpyxl.load_workbook(CANDIDATE,read_only=True,data_only=True)
  s=book['نقشهٔ کتاب (S1)'];q=book['بانک سؤال'];fa=str.maketrans('۰۱۲۳۴۵۶۷۸۹','0123456789')
  chapter_q=Counter();pages={}
  for row in s.iter_rows(min_row=3,max_row=154,values_only=True):
   pp=str(row[1]).translate(fa)
   if pp.isdigit():pages[row[0]]=str(row[2]).translate(fa)
  for row in q.iter_rows(min_row=4,max_row=107,values_only=True):chapter_q[pages[row[1]]]+=1
  for chap in ['6','7','8']:
   if chapter_q[chap]==0:blocked.append(f'فصل {chap}: صفر سؤال؛ گیت تألیف/موتور باز')
  if counts['questionsConfirmed']==0:blocked.append('۱۰۴ سؤال زرد؛ ممیزی شناختی و حکم معلم برای مُهر باز')
  blocked.append('خوانش مستقلِ تصویر/عددهای صفحه‌ای و حکم انسانی برای ردیف‌های کاربردی باز')
  blocked.append('۲ هشدار اطلس هم‌نام با کد واقعی موتور؛ تصمیم مهندسی/ممیزی باز')
except Exception as e:
 steps.append('گیت فنی: FAIL — '+str(e).replace('\n',' ')[:500]);blocked.insert(0,'گیت فنی شکست خورد؛ از نامزد برای انتشار استفاده نشود')
 failure=True
else:failure=False
report=a.report
if not report.is_absolute():report=ROOT/report
report.parent.mkdir(parents=True,exist_ok=True)
lines=['# گیت خودکار G2 — گزارش آخر','',f'- مبنای رسمی پین‌شده: `{PINNED_OFFICIAL}` (`{OFFICIAL_REF}`)',f'- نامزد زرد: `{CANDIDATE.name}` (`{PINNED_CANDIDATE}`)','- حکم نهایی: **تکنیکی مردود**' if failure else '- حکم نهایی: **تکنیکی پاس؛ پروندهٔ آموزشی/انتشار هنوز سبز نیست**','','## گیت‌های اجراشده']
lines += [f'- {s}' for s in steps]
if counts:lines+=['',f"آمار: {counts['lessons']} درس (سبز {counts['lessonsConfirmed']})، {counts['questions']} سؤال (سبز {counts['questionsConfirmed']})، اطلس {counts['atlasReal']} واقعی و {counts['atlasProposed']} پیشنهادی."]
lines+=['','## بدهی‌های مانع سبزکردن']+[f'- {x}' for x in blocked]+['','این ابزار فقط آزمون مکانیکی است؛ PDF/تصویر یا JSON را در مخزن منتشر نمی‌کند، مُهر آموزشی نمی‌زند، و پس از پایان پاسخِ عامل به‌صورت مستقل تصمیم‌گیر/عامل تألیف نیست.','']
report.write_text('\n'.join(lines),encoding='utf-8')
print('\n'.join(lines))
sys.exit(1 if failure else 0)
