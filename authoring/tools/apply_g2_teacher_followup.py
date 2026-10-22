"""Scoped follow-up of explicit teacher rulings; preserve all old values except whitelist.
Usage: python authoring/tools/apply_g2_teacher_followup.py
"""
from pathlib import Path
from hashlib import sha256
from datetime import datetime
from zipfile import ZipFile,ZipInfo,ZIP_DEFLATED
import re
import openpyxl
R=Path(__file__).resolve().parents[1]
base=R/'khshab-G2-MATH_v0.6_full-draft.xlsx';m=R/'PAGE-MAP-KH-G2-MATH-v1.2_teacher-followup.xlsx'
assert sha256(base.read_bytes()).hexdigest()=='c96cc158db60648d8114bb6c6ff90a130d369f1a71e8c0259eaae2275d759807'
assert sha256(m.read_bytes()).hexdigest()=='3bfd31b6f3ebfb4f6d8391c7dedd6c25b164b64140c32977fb2f3b2731db70b2'
w=openpyxl.load_workbook(base);s=w['نقشهٔ کتاب (S1)'];q=w['بانک سؤال']
updates={
 ('نقشهٔ کتاب (S1)','N67'):'ص کتاب ۵۷ / PDF ۶۳؛ مورد ۲ «۸۰ سکهٔ ۱ ریالی» و ۸ سکهٔ ۱۰ ریالی است؛ معلم در گفت‌وگو صریحاً تأیید کرد. نقشهٔ اصلاحی v1.2 ردیف ۶۵ O/P؛ فقط این حکم تأیید شده، ردیف هنوز زرد است.',
 ('نقشهٔ کتاب (S1)','G90'):'۳ مورد شماره‌دار برای حدس و آزمایش؛ ۱) ۲۰+□=۴۳؛ ۲) □+۱۷=۵۱؛ ۳) دو زیرمورد ۲۹+△=۶۳ و ۶۰−○=۲۷. پایین: جعبهٔ «کامل کن» با جدول شگفت‌انگیز. چهار تساوی ≠ چهار تمرین شماره‌دار.',
 ('نقشهٔ کتاب (S1)','N90'):'ص کتاب ۸۰ / PDF ۸۶؛ تصویر فرستاده‌شدهٔ معلم + تأیید کلامی: ۳ مورد جای‌خالی شماره‌دار (سومی دو زیرمورد)، سپس جدول شگفت‌انگیز. سؤال‌های Q-1077..1080 اقتباس چهارگزینه‌ای از چهار تساوی‌اند، نه چهار تمرین عین کتاب. ردیف و برچسب‌های تشخیصی زردند.',
 ('نقشهٔ کتاب (S1)','N99'):'ص کتاب ۸۹ / PDF ۹۵؛ بدون تاب مستقل؛ موارد ۳ تا ۶. معلم این شماره‌گذاری را در گفت‌وگو تأیید کرد؛ نقشهٔ v1.2 ردیف ۹۷ O/P. بقیهٔ محتوای ردیف هنوز زرد است.',
 ('نقشهٔ کتاب (S1)','N148'):'ص کتاب ۱۳۸ / PDF ۱۴۴؛ عنوان «تمرین» و تخصیص صفحه به تأیید شخصی معلم رسید؛ نقشهٔ v1.2 ردیف ۱۴۶ O/P. تأیید هدف‌ها/نمودارها یا مهر کامل ردیف نیست؛ زرد می‌ماند.',
 ('نقشهٔ کتاب (S1)','N149'):'ص کتاب ۱۳۹ / PDF ۱۴۵؛ عنوان «انتخاب نمودار» و تخصیص صفحه به تأیید شخصی معلم رسید؛ نقشهٔ v1.2 ردیف ۱۴۷ O/P. تأیید هدف‌ها/نمودارها یا مهر کامل ردیف نیست؛ زرد می‌ماند.',
 ('بانک سؤال','E83'):'در تساوی ۲۹ + □ = ۶۳، چه عددی در جای خالی می‌آید؟',
}
assert s['B67'].value=='۵۷' and s['B90'].value=='۸۰' and s['B99'].value=='۸۹'
assert s['B148'].value=='۱۳۸' and s['B149'].value=='۱۳۹' and q['A83'].value=='Q-1080'
for row,sub in [(80,'مورد ۱'),(81,'مورد ۲'),(82,'مورد ۳-ب؛ در کتاب جای خالی با دایره نمایش داده شده'),(83,'مورد ۳-الف؛ در کتاب جای خالی با مثلث نمایش داده شده')]:
 updates[('بانک سؤال',f'Q{row}')]=q[f'Q{row}'].value+f' اقتباس چهارگزینه‌ای از {sub} صفحهٔ ۸۰؛ ۳ سؤال شماره‌دار/۴ تساوی؛ نشانهٔ کتاب برای پاسخ عددی به مربع جای‌خالی ساده‌سازی شده؛ هنوز نیازمند تأیید تناسب آموزشی.'
for (sheet,coord),value in updates.items():w[sheet][coord]=value
old=openpyxl.load_workbook(base)
for a in old:
 b=w[a.title]
 assert a.max_row==b.max_row and a.max_column==b.max_column
 assert {str(x) for x in a.merged_cells.ranges}=={str(x) for x in b.merged_cells.ranges}
 for row in a:
  for c in row:
   t=b[c.coordinate]
   if (a.title,c.coordinate) not in updates:assert c.value==t.value,(a.title,c.coordinate)
   assert c.style_id==t.style_id,(a.title,c.coordinate,'style')
assert len(updates)==11
for r in range(63,155):assert s.cell(r,13).value=='نیاز به تأیید'
for r in range(80,84):assert q.cell(r,16).value=='نیاز به تأیید'
out=R/'khshab-G2-MATH_v0.7_teacher-followup-draft.xlsx';w.properties.modified=datetime(2026,10,22,0,0,0);w.save(out)
canon=out.with_suffix('.canonical-tmp')
with ZipFile(out,'r') as src,ZipFile(canon,'w') as dst:
 for name in src.namelist():
  info=ZipInfo(name,date_time=(1980,1,1,0,0,0));info.compress_type=ZIP_DEFLATED;info.external_attr=0o600 << 16
  data=src.read(name)
  if name=='docProps/core.xml':data=re.sub(rb'<dcterms:modified[^>]*>[^<]*</dcterms:modified>',b'<dcterms:modified xsi:type="dcterms:W3CDTF">2026-10-22T00:00:00Z</dcterms:modified>',data)
  dst.writestr(info,data)
canon.replace(out)
h=sha256(out.read_bytes()).hexdigest();(R/(out.name+'.sha256')).write_text(f'{h}  {out.name}\n')
print('v0.7 11 scoped edits, old seals unchanged, SHA',h)
