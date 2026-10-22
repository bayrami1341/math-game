"""Apply scoped teacher rulings; never alter the 13 raw map columns or old deliverable.
Run: python authoring/tools/amend_g2_page_map.py
"""
from pathlib import Path
import re
from hashlib import sha256
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED
from datetime import datetime
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
OLD=ROOT/'PAGE-MAP-KH-G2-MATH-v1_verified.xlsx'
OUT=ROOT/'PAGE-MAP-KH-G2-MATH-v1.1_verified.xlsx'
assert sha256(OLD.read_bytes()).hexdigest()=='7310680cfe62f4cad8e4222de6f3d32edc2c54bb24a9f586a610ef7e2ac192b8'
w=openpyxl.load_workbook(OLD)
s=w['نقشهٔ صفحه‌ها']
assert [(s.cell(r,1).value,s.cell(r,2).value,s.cell(r,14).value) for r in (39,46,54)]==[(31,37,'✗'),(38,44,'✗'),(46,52,'✗')]
patch={
39:('بنر بالای صفحه «جمع و تفریق دو عدد (ادامه)» + تاب «فعّالیت» است، نه بازنویسیِ فیلد بنر. در مورد ۲ عبارت کتاب «۳۵−۱۷» و حاصل ۱۸ است؛ بخش خامِ ردیف در این مورد درست بود و حکم قبلی آرنا («۵۷−۱۴») پس گرفته می‌شود. مورد ۱: ۳۷=۳۰+۷، ۲۴+۳۰=۵۴، ۵۴+۷=۶۱.','PDF ۳۷ / ص کتاب ۳۱: بنر بالا و موارد ۱–۲؛ حکم کاغذی معلم و بازرس دوم دربارهٔ ۳۵−۱۷.'),
46:('۴ مورد شماره‌دار دارد، نه ۵: ۱ رسم ۴ شکل با شابلون؛ ۲ ترکیب شکل‌ها؛ ۳ نوشتن ضلع/گوشهٔ ۵ شکل شامل نمونهٔ مثلث؛ ۴ مقایسهٔ تعداد ضلع/گوشه و دستور «یک ۶ضلعی رسم کنید». در تصویرِ آمادهٔ این صفحه، ۶ضلعی کشیده نشده است؛ ولی عبارتِ دستور رسم شش‌ضلعی واقعاً وجود دارد. عنوان بالای صفحه «یادآوری شکل‌های هندسی» و تاب «فعّالیت» است.','PDF ۴۴ / ص کتاب ۳۸: بنر بالا؛ مورد ۴ و تصویر پنج شکلِ مورد ۳؛ حکم کتاب کاغذی معلم.'),
54:('بنر بالای صفحه دقیقاً «تقارن» است، نه «تقارن (ادامه)». ۳ مورد شماره‌دار: ۱ تا کردن به دو نیمهٔ مثل هم؛ ۲ رسم خط تقارن برای پوشاندن دو نیمه؛ ۳ تکمیل نیمهٔ دیگر. جعبهٔ «کامل کن» (ساعت ۱۰ و نیم) بی‌شماره است؛ تعداد ۴ در ردیف خام، جعبه را با موارد شماره‌دار یکی گرفته است. عبارت «۳ خط/۳ قسمت» در حکم قبلی آرنا نادرست بود.','PDF ۵۲ / ص کتاب ۴۶: بنر بالا + موارد ۱–۳ + جعبهٔ ساعت؛ حکم کتاب کاغذی معلم.')
}
for row,(fix,ref) in patch.items():
 s.cell(row,15,fix);s.cell(row,16,ref)
# Only O/P of the 3 specified rows may change. N verdict (✗) and all original columns are frozen.
prior=openpyxl.load_workbook(OLD,read_only=True,data_only=True)
for sheet in w:
 old=prior[sheet.title]
 for row in sheet:
  for cell in row:
   if sheet.title=='نقشهٔ صفحه‌ها' and cell.row in patch and cell.column in (15,16):continue
   assert cell.value==old.cell(cell.row,cell.column).value,(sheet.title,cell.coordinate)
w.properties.modified=datetime(2026,10,22,0,0,0)
w.save(OUT)
# openpyxl sets ZIP entry timestamps on each save. Canonicalise package metadata
# so the same cell data and style produce the same SHA-256 on repeated runs.
canon=OUT.with_suffix('.canonical-tmp')
with ZipFile(OUT,'r') as src, ZipFile(canon,'w') as dst:
 for name in src.namelist():
  info=ZipInfo(name,date_time=(1980,1,1,0,0,0))
  info.compress_type=ZIP_DEFLATED
  info.external_attr=0o600 << 16
  data=src.read(name)
  if name=='docProps/core.xml':
   data=re.sub(rb'<dcterms:modified[^>]*>[^<]*</dcterms:modified>',
     b'<dcterms:modified xsi:type="dcterms:W3CDTF">2026-10-22T00:00:00Z</dcterms:modified>',data)
  dst.writestr(info,data)
canon.replace(OUT)
digest=sha256(OUT.read_bytes()).hexdigest()
(ROOT/'PAGE-MAP-KH-G2-MATH-v1.1_verified.sha256').write_text(f'{digest}  {OUT.name}\n',encoding='utf-8')
print('amended map',digest)
