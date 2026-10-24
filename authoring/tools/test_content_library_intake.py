import json
import tempfile
import unittest
from pathlib import Path

import openpyxl

from content_library_intake import Q, S1, SHEETS, Q_HEADERS, S1_HEADERS, digest, ingest


class ContentLibraryIntakeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        book = openpyxl.Workbook()
        book.active.title = SHEETS[0]
        for name in SHEETS[1:]:
            book.create_sheet(name)
        a = book[S1]
        b = book[Q]
        for col, header in enumerate(S1_HEADERS, 1):
            a.cell(2, col, header)
        for col, header in enumerate(Q_HEADERS, 1):
            b.cell(3, col, header)
        for row, content in ((3, ('یادداشت', '', '')), (4, ('L-101', '۱', '۱')),
                             (5, ('L-41 (نمونه)', '', ''))):
            a.cell(row, 1, content[0])
            a.cell(row, 2, content[1])
            a.cell(row, 3, content[2])
        a.cell(4, 4, 'موضوع نمونهٔ مصنوعی')
        a.cell(4, 5, 'هدف مصنوعی')
        b.cell(4, 1, 'یادداشت')
        b.cell(5, 1, 'Q-1001')
        b.cell(5, 2, 'L-101')
        b.cell(5, 5, 'سؤال مصنوعی')
        b.cell(6, 1, 'Q-101 (نمونه)')
        self.workbook = self.root / 'approved.xlsx'
        book.save(self.workbook)
        book.close()
        self.receipt = {"schema": "achilles-approved-workbook/v1", "producer": "ACHILLES",
                        "approval_status": "APPROVED", "book_id": "C205", "edition_id": "test-edition",
                        "receipt_id": "synthetic-receipt", "approved_by": "synthetic-fixture",
                        "workbook_sha256": digest(self.workbook), "expected_chapters": ["1"]}
        self.output = self.root / 'library'

    def test_archive_exact_bytes_and_unknown_without_guide(self):
        result = ingest(self.workbook, self.receipt, self.output)
        folder = self.output / 'C205' / 'test-edition' / self.receipt['workbook_sha256']
        self.assertEqual(digest(folder / 'approved.xlsx'), digest(self.workbook))
        self.assertEqual([x['local_id'] for x in result['items']], ['L-101', 'Q-1001'])
        self.assertEqual(result['items'][1]['anchor_uid'], result['items'][0]['uid'])
        self.assertTrue(all(x['educational_priority_0_100'] is None for x in result['assessments']))
        self.assertTrue(all(x['engine_readiness'] == 'AWAITING_ZAI_CAPABILITY' for x in result['assessments']))
        self.assertEqual(result['breadth_depth_status'], 'NOT_ASSESSED')
        self.assertEqual(len(result['grouped_review_queue']), 1)
        self.assertEqual(result['grouped_review_queue'][0]['item_count'], 2)
        self.assertEqual(json.loads((folder / 'index.json').read_text())['workbook_sha256'], digest(self.workbook))
        first_index = (folder / 'index.json').read_bytes()
        self.assertEqual(ingest(self.workbook, self.receipt, self.output)['assessment_file'], result['assessment_file'])
        self.assertEqual((folder / 'index.json').read_bytes(), first_index)

    def test_weight_only_from_cited_guide_and_reviewed_links(self):
        base = self.receipt.copy()
        content = f"C205:test-edition:{Q}:Q-1001"
        goals = {"schema": "teacher-guide-objectives/v1", "source_kind": "OFFICIAL_TEACHER_GUIDE",
                 "book_id": "C205", "edition_id": "test-edition", "source_sha256": "a" * 64,
                 "reviewed_by": "fixture-only", "objectives": [
                     {"id": "g0", "chapter": "1", "role": "PRIMARY", "level": "GENERAL",
                      "description": "synthetic general goal", "citation": {"page": "1", "section": "A"}},
                     {"id": "gs", "chapter": "1", "role": "PRIMARY", "level": "SPECIFIC", "parent_id": "g0",
                      "description": "synthetic specific goal", "citation": {"page": "2", "section": "A"}},
                     {"id": "g1", "chapter": "1", "role": "PRIMARY", "level": "BEHAVIORAL", "parent_id": "gs",
                      "description": "hypothetical multiplication objective", "citation": {"page": "2", "section": "A"}}]}
        direct = [{"content_uid": content, "objective_id": "g1", "alignment": "DIRECT", "reviewed_by": "fixture"}]
        result = ingest(self.workbook, base, self.output, objectives=goals, links=direct)
        self.assertEqual([x['educational_priority_0_100'] for x in result['assessments']], [None, 100])
        prereq = [{**direct[0], "alignment": "PREREQUISITE"}]
        revised = ingest(self.workbook, base, self.output, objectives=goals, links=prereq)
        self.assertEqual(revised['assessments'][1]['educational_priority_0_100'], 40)
        self.assertEqual(revised['assessments'][1]['attention_coefficient_0_1'], .4)
        self.assertIsNone(revised['assessments'][1]['repetition_priority_coefficient_0_1'])
        self.assertNotEqual(revised['assessment_file'], result['assessment_file'])
        self.assertEqual(revised['assessments'][1]['repetition_policy'], 'NOT_DERIVED')
        goals['objectives'][-1]['practice_emphasis'] = 'EMPHASIZED'
        goals['objectives'][-1]['practice_citation'] = {'page': '3', 'section': 'practice'}
        repetition = ingest(self.workbook, base, self.output, objectives=goals, links=prereq)
        self.assertEqual(repetition['assessments'][1]['repetition_priority_coefficient_0_1'], .4)
        self.assertEqual(revised['engine_contract_status'], 'NOT_NEGOTIATED')

    def test_fail_closed_on_approval_sha_schema_anchor_and_goal(self):
        bad = {**self.receipt, "approval_status": "DRAFT"}
        with self.assertRaisesRegex(ValueError, 'ACHILLES_APPROVAL_RECEIPT_REQUIRED'):
            ingest(self.workbook, bad, self.output)
        bad = {**self.receipt, "workbook_sha256": "b" * 64}
        with self.assertRaisesRegex(ValueError, 'APPROVED_WORKBOOK_SHA_MISMATCH'):
            ingest(self.workbook, bad, self.output)
        bad = {**self.receipt, "book_id": "../../leak"}
        with self.assertRaisesRegex(ValueError, 'UNSAFE_ARCHIVE_NAMESPACE'):
            ingest(self.workbook, bad, self.output)
        with self.assertRaisesRegex(ValueError, 'PRIVATE_ARCHIVE_MUST_BE_OUTSIDE_REPOSITORY'):
            ingest(self.workbook, self.receipt, Path(__file__).resolve().parents[2] / 'archive')
        with self.assertRaisesRegex(ValueError, 'UNREVIEWED_OR_UNBOUND_GOAL_LINK'):
            ingest(self.workbook, self.receipt, self.output, links=[{'objective_id': 'invented'}])
        invalid = {"schema": "teacher-guide-objectives/v1", "source_kind": "OFFICIAL_TEACHER_GUIDE",
                   "book_id": "C205", "edition_id": "test-edition", "source_sha256": "a" * 64,
                   "reviewed_by": "fixture", "objectives": [{"id": "g1", "chapter": "1", "role": "PRIMARY",
                   "level": "BEHAVIORAL", "description": "missing citation"}]}
        with self.assertRaisesRegex(ValueError, 'GOAL_WITHOUT_TRACEABLE_GUIDE_CITATION'):
            ingest(self.workbook, self.receipt, self.output, objectives=invalid)
        invalid['objectives'][0]['citation'] = {'page': '2', 'section': 'A'}
        invalid['objectives'][0]['practice_emphasis'] = 'EMPHASIZED'
        with self.assertRaisesRegex(ValueError, 'REPETITION_WEIGHT_WITHOUT_GUIDE_CITATION'):
            ingest(self.workbook, self.receipt, self.output, objectives=invalid)
        invalid['objectives'][0].pop('practice_emphasis')
        with self.assertRaisesRegex(ValueError, 'UNTRACEABLE_OBJECTIVE_HIERARCHY'):
            ingest(self.workbook, self.receipt, self.output, objectives=invalid)

    def test_duplicate_and_missing_anchor_rejected(self):
        wb = openpyxl.load_workbook(self.workbook)
        wb[S1].cell(6, 1, 'L-101')
        wb.save(self.workbook)
        wb.close()
        self.receipt['workbook_sha256'] = digest(self.workbook)
        with self.assertRaisesRegex(ValueError, 'DUPLICATE_LOCAL_ID'):
            ingest(self.workbook, self.receipt, self.output)
        wb = openpyxl.load_workbook(self.workbook)
        wb[S1].cell(6, 1).value = None
        wb[Q].cell(5, 2, 'L-999')
        wb.save(self.workbook)
        wb.close()
        self.receipt['workbook_sha256'] = digest(self.workbook)
        with self.assertRaisesRegex(ValueError, 'QUESTION_ANCHOR_MISSING'):
            ingest(self.workbook, self.receipt, self.output)


if __name__ == '__main__':
    unittest.main()
