"""Lock the bytes/layout of the user's linked blank G2 cartridge copy."""
import hashlib
import unittest

import openpyxl

from map_g2_ch1_standard_review import ROOT, EXPECTED_Q, EXPECTED_S1

FILE = ROOT / 'authoring/khshab-G2-MATH_v0.1.xlsx'
DRIVE_MD5 = '85955347dc974b2496cd708f645820de'


class LinkedBlankStandardTests(unittest.TestCase):
    def test_exact_linked_bytes_and_blank_chapter_state(self):
        raw = FILE.read_bytes()
        self.assertEqual(len(raw), 24113)
        self.assertEqual(hashlib.md5(raw).hexdigest(), DRIVE_MD5)
        wb = openpyxl.load_workbook(FILE, read_only=True, data_only=True)
        try:
            self.assertEqual(wb.sheetnames, ['راهنما', 'شناسنامهٔ خشاب', 'نقشهٔ کتاب (S1)',
                                             'اطلس خطا', 'بانک سؤال', 'واژه‌نامهٔ ستون‌ها', 'پروتکل ممیزی'])
            s, q = wb['نقشهٔ کتاب (S1)'], wb['بانک سؤال']
            self.assertEqual((s.max_row, s.max_column, q.max_row, q.max_column), (6, 14, 6, 17))
            self.assertEqual([s.cell(2, i).value for i in range(1, 15)], EXPECTED_S1)
            self.assertEqual([q.cell(3, i).value for i in range(1, 18)], EXPECTED_Q)
            self.assertEqual(s.cell(3, 1).value, 'یادداشت')
            self.assertIn('نمونه', s.cell(4, 1).value)
            self.assertEqual(q.cell(4, 1).value, 'یادداشت')
            self.assertIn('نمونه', q.cell(5, 1).value)
        finally:
            wb.close()


if __name__ == '__main__':
    unittest.main()
