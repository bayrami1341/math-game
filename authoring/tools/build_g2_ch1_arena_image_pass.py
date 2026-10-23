"""Record a bounded first visual read from previously uploaded page PNGs.

This is a PARTIAL image-only pass (printed pages 1..12, not C205 PDF bytes).
Only paraphrased S1 observations; NO printed exercise transcription or images in Git.
A v0.1 workbook copy is produced, keeping its seven original sheets/columns,
all sample rows and all yellow status. It is not a replacement for official v0.5.
"""
from __future__ import annotations

import hashlib
import json
import struct
from copy import copy
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
BASE = ROOT / 'authoring/khshab-G2-MATH_v0.1.xlsx'
BASE_SHA = '77b4c48c157af481d4bec3ce7a0255f8975a50b455ff615a49e12e9ac443e41f'
OUT = ROOT / 'authoring/g2-completion'
UPLOADS = Path('/tmp/arena-coding-uploads')
# Page assignment from actual visual reading (not from filename order).
OBSERVATIONS = {
  1: ('010','c3474a5e6f98f2c0d9fae010efd73526821988e35315b0662ff371314a61ea9c','آغاز فصل: عدد و رقم','گفت‌وگو دربارهٔ شمارش و نیاز به عدد','عدد، شمارش','آیا می‌دانید و گفت‌وگو','عکس چوپان و گوسفندان و یک پلاک عددی؛ پرسش پایین صفحه به شمارش در تصویر مربوط است.','بالا: عنوان فصل و پلاک؛ میانه: عکس گله؛ پایین: متن گفت‌وگو. شمارهٔ چاپی در پاصفحه آشکار نیست؛ صفحهٔ ۱ از عنوان/ترتیب استنباط شده.'),
  2: ('003','7c71178eb960d5d3931e0cdb9db9f5f9bce3203154ca33cd62d00a2511ca62aa','عدد و شمارش؛ فعالیت','شمارش منظمِ اشیا و نمایش تعداد به چند روش','شمارش، نمایش عدد','فعالیت و کار در کلاس','بالا-چپ توپ‌های فوتبال برای شمارش؛ پایین چند نمایش تصویریِ تعداد با چوب‌خط، دست، شکل‌های هندسی و اشیای دسته‌شده کنار جای پاسخ دیده می‌شود.','بالا: سؤال شمارش به توپ‌ها اشاره دارد؛ پایین: دستور «هر کدام چه عددی» به شکل‌های جدا و جای‌خالی زیر هر شکل مربوط است. پاصفحه ۲ دیده شد.'),
  3: ('007','83670103837edcb88722c3adbc7cbc44bf0cac97c35c66cf4d834c260180ef7a','شمارش گل‌ها و ثبت در جدول','شمارش نوع‌های متفاوت شکل و ثبت دسته‌ای','شمارش، چوب‌خط، جدول','فعالیت شمارش و جدول شگفت‌انگیز','بالا باغچه‌ای از چند نوع گل؛ میانه جدول نوع گل و جای چوب‌خط/عدد؛ پایین جدول عددی داخل تصویر فیل.','دستور بالای صفحه به گونه‌های گل در باغچه و ردیف‌های جدول میانی وصل است؛ جدول پایین فعالیت دیگری است. شمار جزءِ هر نوع هنوز دوباره‌شماری مستقل نشده؛ پاصفحه ۳ دیده شد.'),
  4: ('009','b65ca655953458ea34d165cd455960f7880c7f89a6edb89105934dfcf87eaa83','تمرین شماره، ساعت و مقایسه','خواندن ساعت، وصل‌کردن عددهای هم‌ارزش، کامل‌کردن الگو و مقایسهٔ مجموعه','ساعت، عدد، مقایسه','تمرین و حدس بزن','بالا سه ساعت با عقربه و جای پاسخ؛ وسط دو ردیف کارت عدد و جای اتصال؛ پایین جای‌خالی‌های عدد و دو مجموعهٔ گل برای مقایسه.','شماره‌های تمرین ۱ تا ۳ به سه ناحیهٔ جدا مربوط‌اند؛ مقایسهٔ گل‌ها در قاب پایین «حدس بزن» است. پاصفحه ۴ دیده شد.'),
  5: ('011','fc15d7527b11c6357f4e4b116014dcba479a2ee8208ae2ffcabb57df09cf2d9a','دسته‌های ده‌تایی و یکی','بازنمایی عدد دو رقمی با دستهٔ ده و یکی','دهگان، یکان، عدد دو رقمی','فعالیت و کار در کلاس','بالا چهار آرایش از میله‌های ده‌خانه‌ای و خانه‌های تکی کنار جدول ارزش مکانی؛ میانه الگوی شمارش ده‌تایی؛ پایین کارت‌های رقم برای ساخت عدد.','هر جدول دهگان/یکی کنار همان آرایش میله‌هاست؛ کارت‌های پایین تمرین ساخت عدد است. پاصفحه ۵ دیده شد.'),
  6: ('016','77e476760837baacebe39a5c9c32e11a37241373f64c5e16dbe669c4051c1f45','شمارش و نوشتن عدد دو رقمی','شمارش اشیا و نوشتن دهگان/یکان به عدد و حروف','دهگان، یکان','فعالیت','بالا مجموعهٔ گل‌ها با جدول ارزش مکانی؛ میانه مثال‌های نوشتن عدد و واژه؛ پایین حباب‌های رقم دهگان/یکان کنار جدول.','سؤال بالا به گل‌های همان ناحیه پیوند دارد؛ پایین برچسب دهگان/یکان به جدول مجاور مربوط است. شمار گل هنوز دوباره‌شماری مستقل نشده؛ پاصفحه ۶ دیده شد.'),
  7: ('013','13539ebcbab9d2150e9a21b72bdec4d40d0d28aeb353459e06bb31bee2c8332a','تمرین اعداد دو رقمی','تکمیل توالی عدد، جدول ارزش مکانی و الگوی رنگی','عدد، رقم، الگو','تمرین و الگویابی کن','بالا چند توالی عددی با جای‌خالی؛ میانه جدول‌های یکان/دهگان و سؤال نوشتن عدد؛ پایین مثلث‌های رنگی و ردیف مربع‌ها برای ادامهٔ الگو.','هر جای‌خالی فقط به توالی یا جدول همان بخش تعلق دارد؛ قاب الگویابی پایین فعالیت دیگری است. پاصفحه ۷ دیده شد.'),
  8: ('014','05426c010401d7fe6cde4e8ad6cc3a8ce771632af68bb2447c6fcd21071c54aa','حل مسئله با شکل','مدل‌سازی مسئلهٔ کم‌کردن و تشخیص تعداد از شکل','مسئله، شکل، شمارش','حل مسئله و حدس بزن','دو مسئلهٔ متنی دربارهٔ کتاب/عکس؛ میله‌های ده‌تایی و یکان برای یکی از آن‌ها، ردیف مثلث‌ها، و قاب حدس با دو مجموعهٔ جانور/حشره.','میله‌های سبز به مسئلهٔ کتاب‌ها وصل‌اند؛ مثلث‌ها برای خواندن عدد و «پنجمین بعد» هستند؛ دو قاب جانور برای مقایسهٔ بیش از ۲۰ است. شمارش تک‌به‌تک شکل‌ها نیازمند بازبین دوم؛ پاصفحه ۸ دیده شد.'),
  9: ('015','ca6346fe3570d488e64784783a82a73e6d5e8335d0bf3cb254db81bbf0a60a4b','رسم شکل و حل مسئله','نمایش مسئله روی محور/چوب‌خط و تکمیل الگوی هندسی','محور اعداد، چوب‌خط، رسم','رسم شکل و کامل کن','بالا مسئلهٔ گذشت روزها با محور اعداد؛ میانه مسئلهٔ کتاب‌ها و چوب‌خط؛ پایین شبکهٔ هندسی با یک ناحیهٔ آبی و دستور تکمیل بخش‌های دیگر.','نقشهٔ پایین یک شکل نیمه‌کامل است، نه شکل آمادهٔ پاسخ؛ جهت خطوط و جای قطعهٔ رنگی برای دستور تکمیل مهم‌اند. پاصفحه ۹ دیده شد.'),
 10: ('012','d8651762e5192956f605a92e446bb2e37ab480f5fcffd15185433fc0fec8dc4c','جمع و تفریق با نمایش‌های گوناگون','محاسبهٔ جمع و تفریق به روش‌های تصویری و عددی','جمع، تفریق، شکل','فعالیت و کار در کلاس','بالا عبارت‌های جمع/تفریق؛ میانه محل‌های رسم چوب‌خط، شکل و رنگ‌کردن خانه‌ها؛ پایین دو محور و عبارت‌های محاسبه.','هر محور به عبارت بلافاصله کنار آن مربوط است؛ خانه‌های رنگی به دستور جمع همان ردیف مربوط‌اند. پاصفحه ۱۰ دیده شد.'),
 11: ('017','77e777a99d9c336d8ad2aec0532859cdf45625f4d731c1bc4b951b4e7de34bbb','بازی جزء مجهول و الگوی شمارش','یافتن جزء مجهول و ادامه‌دادن شمارش با گام ثابت','مهره، شمارش، الگو','فعالیت دو نفره، کار در کلاس و حدس بزن','بالا چند مهره و دستور پنهان‌کردن بخشی؛ میانه الگوهای شمارش و جمع‌های چندجمله‌ای؛ پایین چند ساعت در قاب حدس بزن.','عکس مهره‌ها به فعالیت دو نفره مربوط است؛ ساعت‌های پایین بخش جداگانه‌اند و از عبارت‌های جمع میانه پاسخ‌گیری نمی‌کنند. پاصفحه ۱۱ دیده شد.'),
 12: ('008','fa215bd290dece8463e5de5e93dc7138fc5c72b4c03cd0a60a518682b83bc0fe','تمرین جمع، محور و ساخت شکل','رنگ‌کردن خانه‌ها و نمایش جمع؛ حل با محور؛ ساخت شکل با شابلون','جمع، محور، شابلون','تمرین و خلّق باش','بالا نوارهای خانه‌های رنگی و جای نوشتن جمع؛ میانه دو محور اعداد؛ پایین طرح مثلثی و دستور ساخت شکل با شابلون.','محورها با عبارت‌های جمع/تفریقِ کنار همان ردیف پیوند دارند؛ طرح مثلثیِ پایین نمونهٔ کاردستی است، نه پاسخ عددی عبارت‌ها. پاصفحه ۱۲ دیده شد.'),
}


def image_manifest(uploads: Path = UPLOADS) -> list[dict]:
    images = []
    for page, (number, expected, *_) in OBSERVATIONS.items():
        path = uploads / f'Math-book2-{number}.png'
        raw = path.read_bytes()
        if raw[:8] != b'\x89PNG\r\n\x1a\n' or hashlib.sha256(raw).hexdigest() != expected:
            raise ValueError(f'PAGE_IMAGE_CHANGED_OR_NOT_PNG:{page}')
        width, height = struct.unpack('>II', raw[16:24])
        if width < 600 or height < 800:
            raise ValueError(f'PAGE_IMAGE_TOO_SMALL:{page}')
        images.append({'book_page_candidate': page, 'source_filename_private': path.name,
                       'image_sha256': expected, 'width_px': width, 'height_px': height,
                       'printed_footer_status': 'NOT_VISIBLE_CHAPTER_OPENER' if page == 1 else 'PAGE_NUMBER_VISUALLY_OBSERVED',
                       'source_pdf_sha256_verified': False,
                       'image_to_pdf_edition_binding': 'NOT_VERIFIED'})
    return images


def make_book(images: list[dict]) -> openpyxl.Workbook:
    if len(images) != 12 or [x['book_page_candidate'] for x in images] != list(range(1, 13)):
        raise ValueError('PARTIAL_IMAGE_MANIFEST_DRIFT')
    if hashlib.sha256(BASE.read_bytes()).hexdigest() != BASE_SHA:
        raise ValueError('STANDARD_BASE_SHA_DRIFT')
    book = openpyxl.load_workbook(BASE)
    sheet = book['نقشهٔ کتاب (S1)']
    if sheet['A3'].value != 'یادداشت' or not str(sheet['A4'].value).startswith('L-41 (نمونه)'):
        raise ValueError('STANDARD_BASE_LAYOUT_DRIFT')
    sheet.unmerge_cells('A6:N6')
    sheet.insert_rows(3, 18)
    sheet.merge_cells('A24:N24')
    # The pre-existing example row moved to 22; copy only its presentation style.
    for page in range(1, 19):
        row = page + 2
        for col in range(1, 15):
            dest, sample = sheet.cell(row, col), sheet.cell(22, col)
            dest._style = copy(sample._style)
            dest.alignment = copy(sample.alignment)
        sheet.row_dimensions[row].height = sheet.row_dimensions[22].height
        sheet.cell(row, 1, f'L-{100 + page}')
        sheet.cell(row, 2, str(page).translate(str.maketrans('0123456789', '۰۱۲۳۴۵۶۷۸۹')))
        sheet.cell(row, 3, '۱')
        sheet.cell(row, 13, 'نیاز به تأیید')
        sheet.cell(row, 11, '—')
        if page in OBSERVATIONS:
            number, sha, title, goal, keywords, activity, visual, note = OBSERVATIONS[page]
            for col, value in ((4,title),(5,'پیشنهاد پس از مشاهده: '+goal),(6,keywords),(7,activity),(8,visual),
                               (9,'؟ — گونه‌های واقعی تمرین هنوز ممیزی نشده'),
                               (10,'؟ — قابلیت موتور برای این صفحه آزموده نشده'),
                               (12,f'تصویر خصوصی {number}؛ SHA تصویر {sha[:12]}؛ PDF C205 تطبیق بایتی نشده'),
                               (14,note+' خوانش A محدود؛ نه مهر/نه OCR؛ منشأ بایت PDF نامعلوم.')):
                sheet.cell(row,col,value)
        else:
            for col in (4,5,6,7,8,9,10):sheet.cell(row,col,'؟ — تصویر صفحه در این مجموعه موجود نیست')
            sheet.cell(row,12,'تصویر صفحه موجود نیست؛ PDF C205 در vault آرنا موجود نیست')
            sheet.cell(row,14,'BLOCKED_IMAGE_MISSING؛ هیچ متن/شکل/هدف چاپی برای این صفحه ادعا نمی‌شود.')
    if sheet['A21'].value != 'یادداشت' or sheet['A22'].value != 'L-41 (نمونه)' or sheet['A24'].value != 'یادداشت':
        raise ValueError('STANDARD_SAMPLES_OR_NOTES_LOST')
    return book


def write(uploads: Path = UPLOADS, out: Path = OUT) -> dict:
    images = image_manifest(uploads)
    book = make_book(images)
    out.mkdir(parents=True, exist_ok=True)
    workbook = out / 'G2-CH1-ARENA-IMAGE-PASS-A-PARTIAL-v0.1.xlsx'
    book.save(workbook)
    book.close()
    report = {'schema':'G2_CH1_ARENA_IMAGE_ONLY_PASS_A_v1','status':'PARTIAL_IMAGE_ONLY_NOT_PDF_VERIFIED',
              'standard_base_sha256':BASE_SHA,'standard_workbook_path':workbook.name,
              'official_book_pdf_sha256_verified':False,'agent_A_has_seen_filled_v05_prior_to_this_pass':True,
              'strictly_blind_to_official_baseline':False,'other_agent_output_seen':False,
              'book_pages_expected':18,'visually_read_images_count':12,'pages_missing_images':list(range(13,19)),
              'printed_page_1_inferred_not_footer_verified':True,
              'existing_bank_q_modified':False,'new_questions':0,'new_green_stamps':0,
              'note':'Image-to-C205-PDF binding is not proved. Images were visually inspected separately; no complete text transcription or full-page second audit claimed. Green comparison not available.',
              'images':images}
    (out/'G2-CH1-ARENA-IMAGE-PASS-A-PARTIAL-v1.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    lines = ['# خوانش تصویری مستقل آرنا A — فصل اول G2، محدود به تصویرهای موجود', '',
             '**یافتهٔ تازه:** از ۱۸ PNG در پوشهٔ بارگذاری‌های قبلی، ۶ مورد صفحات مقدماتی/فهرست‌اند و ۱۲ مورد تصویر صفحه‌های چاپی ۱–۱۲ را نشان می‌دهند؛ ترتیب نام فایل برابر ترتیب صفحه نیست. صفحه‌های ۱۳–۱۸ در این مجموعه نیستند. همهٔ تصاویر ۱۲ صفحه با read_file واقعاً دیده شدند؛ خود PDF با SHA رسمی C205 به این تصویرها bind نشده است. این اجرای یک‌نوبتی نه خوانش کل ۱۸ صفحه است و نه حکم چاپ/صفحهٔ PDF.',
             '', 'کپی زردِ همان قالب v0.1 با ۱۸ ردیف S1 ساخته شد؛ برای ۱۲ صفحهٔ دارای تصویر فقط *خلاصهٔ بازگویی‌شده* و پیوندهای بصری محدود و برای شش صفحهٔ فاقد تصویر «نامعلوم» ثبت است. نمونه‌ها، شیت‌ها، کلیدها و مهرها دست‌نخورده‌اند؛ ۰ Q ساخته/تأیید شد. تصویر و متن کامل چاپی وارد Git نشده است.',
             '', '| صفحهٔ چاپیِ پیشنهادی | شناسه | تصویر خصوصی دیده‌شده | خلاصهٔ تصویری/مرز |', '|---:|---|---|---|']
    for page in range(1,19):
        if page in OBSERVATIONS:
            number, _, title, _, _, _, visual, note = OBSERVATIONS[page]
            scope = 'شمارهٔ صفحه از عنوان/ترتیب استنباط شده؛ پاصفحه نیست.' if page == 1 else 'پاصفحه روی تصویر دیده شد.'
            lines.append(f'| {page} | L-{100+page} | Math-book2-{number}.png | {title}؛ {visual} {scope} |')
        else:
            lines.append(f'| {page} | L-{100+page} | **نیست** | `BLOCKED_IMAGE_MISSING`؛ متن و شکل نامعلوم |')
    lines += ['', '**شش PNG خارج از فصل:** `001` صفحهٔ نشر، `002` عنوان، `004`/`005` سخن با معلم، `006` فهرست، `018` صفحهٔ نقل‌قول/تصویر مقدماتی. به‌عنوان صفحه‌های ۱۳–۱۸ جا زده نشده‌اند.',
              '', '**بی‌طرفی و گیت‌ها:** خروجی عامل دوم هنوز دیده نشده، اما آرنا پیش از این دور به اکسل رسمی v0.5 دسترسی داشته است؛ بنابراین این گذر «مستقل از عامل دوم» است، نه **blind نسبت به مبنای رسمی**. توافق آیندهٔ دو عامل بدون بررسی منبع PDF/بازرس دوم، سبزِ تأیید آموزشی نیست. هیچ صفحه‌ای از روی متن OCR/نام فایل جایگزین مشاهدهٔ پیکسل نشده است.',
              '', '**گام بعدی:** PDF همان چاپ را در vault امن آرنا قرار دهید تا SHA/۱۵۲ صفحه و offset چاپی تثبیت، صفحه‌های ۱۳–۱۸ واقعاً رندر، و صفحات ۱–۱۲ با PDF مقابله شوند. خروجی عامل دوم را پس از قفل خوانش A تحویل بگیرید. تا آن زمان وضعیت `PARTIAL_IMAGE_ONLY_NOT_PDF_VERIFIED` باقی می‌ماند.', '']
    (out/'G2-CH1-ARENA-IMAGE-PASS-A-PARTIAL-v1.fa.md').write_text('\n'.join(lines),encoding='utf-8')
    return report


if __name__=='__main__':
    result=write()
    print(json.dumps({k:result[k] for k in ('status','visually_read_images_count','pages_missing_images','new_green_stamps')},ensure_ascii=False))
