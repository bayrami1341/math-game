"""All Git remotes/workbooks in this suite are synthetic and local."""
import hashlib
import json
import subprocess
import tempfile
import unittest
from pathlib import Path

import openpyxl

from content_library_intake import Q, Q_HEADERS, S1, S1_HEADERS, SHEETS
from github_content_exchange_v1 import fetch_pinned, redacted_package


def git(repo, *args):
    return subprocess.check_output(['git', '-C', str(repo), *args], stderr=subprocess.DEVNULL).decode().strip()


class GithubExchangeTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.home = Path(tmp.name)
        self.repo = self.home / 'synthetic-source'
        self.repo.mkdir()
        git(self.repo, 'init', '-q', '-b', 'test-producer')
        git(self.repo, 'config', 'user.name', 'Synthetic Test')
        git(self.repo, 'config', 'user.email', 'fixture@example.test')
        wb = openpyxl.Workbook()
        wb.active.title = SHEETS[0]
        for sheet in SHEETS[1:]:
            wb.create_sheet(sheet)
        for col, header in enumerate(S1_HEADERS, 1):
            wb[S1].cell(2, col, header)
        for col, header in enumerate(Q_HEADERS, 1):
            wb[Q].cell(3, col, header)
        wb[S1].cell(3, 1, 'L-1')
        wb[S1].cell(3, 3, '۱')
        wb[S1].cell(3, 4, 'Only synthetic item')
        wb[Q].cell(4, 1, 'Q-1')
        wb[Q].cell(4, 2, 'L-1')
        wb[Q].cell(4, 5, 'synthetic confidential question')
        source = self.home / 'synthetic.xlsx'
        wb.save(source)
        wb.close()
        self.sha = hashlib.sha256(source.read_bytes()).hexdigest()
        prefix = f'authoring/content-exchange/incoming/G2/C205/fixture-edition/{self.sha}'
        self.workbook_path = prefix + '/extraction-snapshot.xlsx'
        self.receipt_path = prefix + '/extraction-receipt.json'
        self.wb = self.repo / self.workbook_path
        self.rec = self.repo / self.receipt_path
        self.wb.parent.mkdir(parents=True)
        self.wb.write_bytes(source.read_bytes())
        self.receipt = {'schema': 'achilles-extraction-bundle/v1', 'producer': 'ACHILLES',
                        'approval_status': 'EXTRACTION_VERIFIED', 'book_id': 'C205', 'edition_id': 'fixture-edition',
                        'workbook_sha256': self.sha, 'repo_path': self.workbook_path,
                        'storage_rights_status': 'REVIEWED_FOR_PRIVATE_REPO',
                        'rights_decision_ref': 'synthetic-rights-review', 'book_source_sha256': 'a'*64,
                        'evidence_manifest_sha': 'b'*64, 'extraction_scope': 'synthetic-only',
                        'decision_ref': 'synthetic-decision', 'receipt_id': 'test-only',
                        'verified_by': 'test-only'}
        self.save()
        self.archive = self.home / 'private-library'
        self.args = {'producer_ref': 'refs/heads/test-producer', 'producer_commit': self.commit,
                     'workbook_path': self.workbook_path, 'receipt_path': self.receipt_path,
                     'expected_sha256': self.sha, 'out_dir': self.archive,
                     'rights_decision_ref': 'synthetic-rights-review', 'remote': str(self.repo)}

    def save(self):
        self.rec.write_text(json.dumps(self.receipt), encoding='utf8')
        git(self.repo, 'add', '.')
        git(self.repo, 'commit', '-qm', 'synthetic')
        self.commit = git(self.repo, 'rev-parse', 'HEAD')

    def test_fetch_exact_commit_and_redacted_manifest(self):
        result = fetch_pinned(**self.args)
        self.assertEqual(result['producer_commit'], self.commit)
        self.assertEqual(result['workbook_sha256'], self.sha)
        self.assertEqual(result['release_authorization'], 'NOT_GRANTED')
        self.assertEqual(result['content_acceptance'], 'PENDING_ARENA_ASSESSMENT_AND_HUMAN_DECISION')
        manifest, summary = redacted_package(result, self.archive)
        self.assertEqual(manifest['content_acceptance'], 'PENDING_ARENA_ASSESSMENT_AND_HUMAN_DECISION')
        self.assertEqual(manifest['assessment_summary_sha256'], hashlib.sha256(
            (json.dumps(summary, sort_keys=True, ensure_ascii=False, indent=2)+'\n').encode()).hexdigest())
        self.assertEqual(summary['six_factor_status'], 'NOT_COMPUTED')
        self.assertNotIn('confidential question', json.dumps([manifest, summary]))
        self.assertEqual(summary['items_count'], 2)
        # A moving branch does not silently switch the pinned commit.
        (self.repo / 'unrelated.txt').write_text('moved')
        git(self.repo, 'add', 'unrelated.txt')
        git(self.repo, 'commit', '-qm', 'branch moved')
        self.assertEqual(fetch_pinned(**self.args)['producer_commit'], self.commit)

    def test_ref_sha_rights_and_path_fail_closed(self):
        for changes, error in [({'rights_decision_ref': ''}, 'EXTERNAL_RIGHTS_DECISION_REQUIRED_BEFORE_FETCH'),
                               ({'rights_decision_ref': 'other'}, 'ACHILLES_RECEIPT_OR_RIGHTS_MISMATCH'),
                               ({'producer_commit': 'deadbeef'}, 'PINNED_FULL_COMMIT_REQUIRED'),
                               ({'expected_sha256': '0'*64}, 'ACHILLES_RECEIPT_OR_RIGHTS_MISMATCH'),
                               ({'workbook_path': '../secret.xlsx'}, 'UNSAFE_SOURCE_PATH'),
                               ({'producer_ref': 'refs/heads/../other'}, 'INVALID_PRODUCER_REF'),
                               ({'producer_commit': 'f'*40}, 'PINNED_COMMIT_NOT_REACHABLE')]:
            with self.subTest(changes=changes), self.assertRaisesRegex(ValueError, error):
                fetch_pinned(**(self.args | changes))
        self.assertFalse(self.archive.exists())
        self.wb.write_bytes(b'altered after receipt')
        self.save()
        with self.assertRaisesRegex(ValueError, 'PINNED_WORKBOOK_SHA_MISMATCH'):
            fetch_pinned(**(self.args | {'producer_commit': self.commit}))
        self.assertFalse(self.archive.exists())
        self.receipt['storage_rights_status'] = 'UNKNOWN'
        self.save()
        with self.assertRaisesRegex(ValueError, 'ACHILLES_RECEIPT_OR_RIGHTS_MISMATCH'):
            fetch_pinned(**(self.args | {'producer_commit': self.commit}))
        self.assertFalse(self.archive.exists())

    def test_receipt_drift_and_self_reference_fail_before_archive(self):
        self.receipt['workbook_sha256'] = '0'*64
        self.save()
        with self.assertRaisesRegex(ValueError, 'ACHILLES_RECEIPT_OR_RIGHTS_MISMATCH'):
            fetch_pinned(**(self.args | {'producer_commit': self.commit}))
        self.receipt['workbook_sha256'] = self.sha
        self.receipt['git_commit'] = 'f'*40
        self.save()
        with self.assertRaisesRegex(ValueError, 'SELF_REFERENTIAL_RECEIPT_COMMIT'):
            fetch_pinned(**(self.args | {'producer_commit': self.commit}))
        self.assertFalse(self.archive.exists())

    def test_source_not_from_pinned_ref_and_archive_tamper(self):
        result = fetch_pinned(**self.args)
        approved = self.archive / 'C205' / 'fixture-edition' / self.sha / 'approved.xlsx'
        approved.write_bytes(b'tampered')
        with self.assertRaisesRegex(ValueError, 'PRIVATE_ARCHIVE_MISSING_OR_DRIFTED'):
            redacted_package(result, self.archive)
        # The wrong advertised branch must not authorize this commit.
        git(self.repo, 'checkout', '-q', '--orphan', 'unrelated')
        git(self.repo, 'rm', '-q', '-rf', '.')
        (self.repo / 'unrelated.txt').write_text('separate branch')
        git(self.repo, 'add', 'unrelated.txt')
        git(self.repo, 'commit', '-qm', 'other')
        with self.assertRaisesRegex(ValueError, 'PINNED_COMMIT_NOT_REACHABLE'):
            fetch_pinned(**(self.args | {'producer_ref': 'refs/heads/unrelated'}))


if __name__ == '__main__':
    unittest.main()
