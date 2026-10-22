"""Read-only audit of official-v0.5 reconciliation; optional converted JSON.
Usage: python authoring/tools/audit_g2_official_rebase.py --official /tmp/g2-official-v05.xlsx [--json /tmp/generated.json]
"""
import argparse,json,re
from collections import Counter
from hashlib import sha256
from pathlib import Path
import openpyxl
P=argparse.ArgumentParser();P.add_argument('--official',type=Path,required=True);P.add_argument('--json',type=Path);args=P.parse_args()
R=Path(__file__).resolve().parents[1]
out=R/'khshab-G2-MATH_v0.8_official-rebased-draft.xlsx'
assert sha256(args.official.read_bytes()).hexdigest()=='f1f339a09a8cba3363c8c526e088063c2d733c62c5e7e4c236eba871f97d2656'
h=sha256(out.read_bytes()).hexdigest();assert (R/(out.name+'.sha256')).read_text().startswith(h+'  ')
o=openpyxl.load_workbook(args.official);w=openpyxl.load_workbook(out)
assert o.sheetnames==w.sheetnames
for sh in o:
 after=w[sh.title]
 at,shift={'نقشهٔ کتاب (S1)':(65,90),'اطلس خطا':(17,6),'بانک سؤال':(96,12)}.get(sh.title,(999999,0))
 for row in sh:
  for c in row:
   d=after.cell(c.row+(shift if c.row>=at else 0),c.column)
   assert c.value==d.value and c.style_id==d.style_id,(sh.title,c.coordinate)
 for merged in sh.merged_cells.ranges:
  t=(merged.min_row+(shift if merged.min_row>=at else 0),merged.max_row+(shift if merged.max_row>=at else 0),merged.min_col,merged.max_col)
  assert t in [(x.min_row,x.max_row,x.min_col,x.max_col) for x in after.merged_cells.ranges]
fa=str.maketrans('۰۱۲۳۴۵۶۷۸۹','0123456789')
s=w['نقشهٔ کتاب (S1)'];q=w['بانک سؤال'];a=w['اطلس خطا']
rows=[r for r in range(3,155) if str(s.cell(r,2).value).translate(fa).isdigit()]
by_page={int(s.cell(r,2).value.translate(fa)):s.cell(r,1).value for r in rows}
assert len(by_page)==144 and set(by_page)==set(range(1,145))
ids={s.cell(r,1).value for r in range(3,155)}
assert len(ids)==152
assert len({q.cell(r,1).value for r in range(4,108)})==104
assert {q.cell(r,1).value for r in range(96,108)}=={f'Q-{n}' for n in range(1093,1105)}
for r in range(4,108):
 anchor=q.cell(r,2).value;assert anchor in ids,(r,anchor)
for r in rows:
 text=str(s.cell(r,11).value or '')
 refs=re.findall(r'Q-\d+',text)
 for ref in refs:assert q.cell(4+int(ref[2:])-1001,1).value==ref if int(ref[2:])<1093 else ref in {q.cell(x,1).value for x in range(96,108)}
for r in range(65,155):assert s.cell(r,13).value=='نیاز به تأیید'
for r in range(96,108):assert q.cell(r,16).value=='نیاز به تأیید'
assert '۳ مورد شماره‌دار' in s.cell(next(r for r in rows if s.cell(r,1).value=='L-10080'),7).value
assert s.cell(next(r for r in rows if s.cell(r,1).value=='L-10138'),4).value.startswith('تمرین')
assert s.cell(next(r for r in rows if s.cell(r,1).value=='L-10139'),4).value.startswith('انتخاب نمودار')
correct=[11,30,14,10,18,24,15,16,23,34,33,34] # first 8: wave4 book patterns; last 4: page 80
for i,r in enumerate(range(96,108)):
 options=[int(str(q.cell(r,c).value).translate(fa)) for c in range(7,11)]
 assert len(set(options))==4
 idx='الف ب ج د'.split().index(q.cell(r,11).value)
 assert options[idx]==correct[i],(r,options,idx,correct[i])
assert len({a.cell(r,1).value for r in range(4,23)})==19
if args.json:
 j=json.loads(args.json.read_text());assert j['cartridgeId']=='KH-G2-MATH'
 assert len(j['bookMap'])==152 and len(j['questionBank'])==104 and len(j['errorAtlas'])==19
 assert {x['id'] for x in j['questionBank']}=={q.cell(r,1).value for r in range(4,108)}
 assert {x['id'] for x in j['bookMap']}==ids
 assert all(x['status']=='نیاز به تأیید' for x in j['questionBank'])
print('PASS: official 7 sheets values/styles/merges preserved; 144 page anchors; no ID collisions; 104 unique questions; 12 new numeric keys; yellow; SHA',h)
