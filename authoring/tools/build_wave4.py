"""Reproducible DRAFT for G2 chapter 3 pages 37..52, from the published base.
Usage: python authoring/tools/build_wave4.py --base /path/to/official-v0.4.xlsx
The base must be the v10.3-beta public/math-game/authoring workbook (SHA pinned).
Does not mark teacher approvals or publish a cartridge. Raw PDF/photos are not stored.
"""
import argparse,ast,operator,re
from pathlib import Path
from collections import Counter
from copy import copy
from hashlib import sha256
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED
from datetime import datetime
import openpyxl

P=argparse.ArgumentParser();P.add_argument('--base',type=Path,required=True);args=P.parse_args()
BASE_SHA='a2f610585b4f915f3fe01201b2361d5aa607dd015d6c4d57947556245af52d34'
MAP_SHA='2c54be93a4ec016db61186ed7eed31fc34d51268d3e75a21adfcd6dfa4b6e839'
ROOT=Path(__file__).resolve().parents[1]
MAP=ROOT/'PAGE-MAP-KH-G2-MATH-v1.1_verified.xlsx'
OUT=ROOT/'khshab-G2-MATH_v0.5_verified.xlsx'
assert sha256(args.base.read_bytes()).hexdigest()==BASE_SHA,'Wrong base: NEVER start from local old v0.3/v0.4'
assert sha256(MAP.read_bytes()).hexdigest()==MAP_SHA,'Master map teacher amendments changed'
w=openpyxl.load_workbook(args.base)
master=openpyxl.load_workbook(MAP,read_only=True,data_only=True)['نقشهٔ صفحه‌ها']
SOURCE='پیام AI — موتور تألیف v۱.۰';Y='نیاز به تأیید'
fa=str.maketrans('0123456789','۰۱۲۳۴۵۶۷۸۹')
def f(n):return str(n).translate(fa)
def insert(sh,at,values,style_row):
 # openpyxl does NOT shift merged ranges with inserted rows.
 moved=[(x.min_row,x.max_row,x.min_col,x.max_col) for x in list(sh.merged_cells.ranges) if x.min_row>=at]
 for r1,r2,c1,c2 in moved:sh.unmerge_cells(start_row=r1,end_row=r2,start_column=c1,end_column=c2)
 sh.insert_rows(at)
 for r1,r2,c1,c2 in moved:sh.merge_cells(start_row=r1+1,end_row=r2+1,start_column=c1,end_column=c2)
 for j,v in enumerate(values,1):
  c=sh.cell(at,j,v); sample=sh.cell(style_row,j)
  if sample.has_style:c._style=copy(sample._style)
  c.alignment=copy(sample.alignment)

# page, exact page heading (or faithful page-type label), measurable objective,
# page-bound vocabulary, evidence-based activity, abstract M_ description, compat.
# The 16 records are deliberately curated from master-map rows 45-60 + direct
# C205 text/render; this is NOT an automatic PDF-reading agent.
PAGES=[
(37,'فصل ۳: اشکال هندسی','بتواند شکل‌های هندسی را در نقش کاشی‌کاری آغاز فصل شناسایی کند.','شکل، کاشی‌کاری','آغاز فصل: گفت‌وگو دربارهٔ کاشی‌کاری و شکل‌ها','M_GEO_SHAPES — کاشی‌کاریِ آغاز فصل؛ شمار اجزای تصویر مبنای سؤال نیست','خارج از موتور'),
(38,'یادآوری شکل‌های هندسی — فعّالیت','بتواند ضلع‌ها و گوشه‌های چندضلعی‌ها را بشمارد و با شابلون شکل بسازد.','شابلون، مثلث، مربع، مستطیل، دایره، ضلع، گوشه','۴ مورد شماره‌دار؛ مورد ۴ دستور رسم یک ۶ضلعی با شابلون دارد، اما شکل ۶ضلعی آماده در تصویر نیست','M_GEO_SHAPES — پنج شکلِ شمارش ضلع/گوشه، شامل مثلث نمونه؛ هیچ شش‌ضلعی چاپ‌شده‌ای در این بلوک نیست','بانک تمرین'),
(39,'کار در کلاس + کامل کن','بتواند شباهت و تفاوت مربع/مستطیل و دو مثلث را توصیف کند.','مربع، مستطیل، مثلث، شابلون','کار در کلاس: مقایسهٔ شکل‌ها؛ کامل کن: تاخوردگی','M_GEO_SHAPES — دو جفت شکل برای مقایسه؛ خط تا در فعالیت پایانی','خارج از موتور'),
(40,'تمرین + کامل کن','بتواند دو شش‌ضلعی را مقایسه و شکل‌ها را با کلید رنگ طبقه‌بندی کند.','تمرین، شش‌ضلعی، مثلث، مربع، دایره','تمرین: دو شش‌ضلعی و جدول سه رنگ؛ کامل کن: چهار ساعت','M_GEO_SHAPES — شش‌ضلعی منتظم/نامنظم و جدول زرد/قرمز/آبی','خارج از موتور'),
(41,'روابط بین شکل‌ها — فعّالیت','بتواند با قطعه‌های یک مربع بریده‌شده، شکل‌های گوناگون بسازد.','مربع، شکل، قیچی، الگو','فعّالیت: هفت قطعه و بازسازی مربع؛ خلّاق باش: شابلون و الگو','M_GEO_SHAPES — ۷ قطعهٔ بریدنی؛ نمونه‌ها برای کار دست‌آزاد','خارج از موتور'),
(42,'کار در کلاس + خلّاق باش','بتواند با کنار هم گذاشتن شکل‌ها، نمونهٔ پوشاندن سطح را بسازد.','مثلث، مربع، شش‌ضلعی، کاشی‌کاری','کار در کلاس: ۱۰ شکل برای برش و مفهوم کاشی‌کاری؛ خلّاق باش: جعبهٔ معمای کلمات','M_GEO_SHAPES — سه نوع شکلِ قابل استفاده در کاشی‌کاری','خارج از موتور'),
(43,'تمرین','بتواند از توصیف تعداد ضلع/گوشه، شکل را روی نقطه‌ها بسازد.','ضلع، گوشه، نقطه، خط‌کش','تمرین: وصل‌کردن نقطه‌ها + تکمیل نمودار شکل','M_GEO_SHAPES — چهار توصیف برای رسم شکل در شبکهٔ نقطه‌ها','خارج از موتور'),
(44,'حلّ مسئله + کامل کن','بتواند قاعدهٔ الگوی عددی و الگوی دایره‌ها را یافته و ادامه دهد.','الگو، شکل، دایره، عدد','حلّ مسئله: دو مورد شماره‌دار (دایره‌ها و الگوها)؛ کامل کن: ساعت جواد','M_PATTERN — شکل‌های ۱، ۲، ۳ به‌ترتیب ۱، ۳، ۶ دایره؛ الگوهای +۲ و +۵','بانک تمرین'),
(45,'الگویابی + کامل کن','بتواند جملهٔ بعدیِ الگوهای با گام ثابت و گام فزاینده را بیابد.','الگو، عدد، ساعت','الگویابی: مورد ۳ ادامهٔ صفحهٔ قبل؛ سه دنباله به ترتیب صفحه؛ کامل کن: ساعت‌ها','M_PATTERN — دنباله‌های گام ۴، گام ۵، سپس گام فزایندهٔ ۱/۲/۳/۴','بانک تمرین'),
(46,'تقارن — فعّالیت','بتواند خط تقارن را با پوشاندن دو نیمه پیدا و نیمهٔ دیگر شکل را رسم کند.','تقارن، خط، نیمه، شکل','سه مورد شماره‌دار با دو نیمه؛ کامل کن: ساعت ۱۰ و نیمِ بی‌شماره','M_GEO_SHAPES — پروانهٔ دو نیمه و خط تا؛ عنوان عین بنر «تقارن» است','خارج از موتور'),
(47,'کار در کلاس','بتواند با تا کردن نیم‌مربع‌ها شکل بسازد و شکل بدون خط تقارن را تشخیص دهد.','مربع، نیمه، تقارن، ساعت','دو مورد شکل/تا؛ مسئلهٔ ساعت چهار و نیم + سه ساعت','M_GEO_SHAPES — سه مربع بریدنی و صفحهٔ شطرنجی','خارج از موتور'),
(48,'تمرین','بتواند نیمهٔ دیگر شکل را با خط تقارن تکمیل و وسط ضلع‌ها را نشان دهد.','تقارن، شابلون، ضلع، خط‌کش','سه مورد: قرینه، فرش و کاشی‌کاری، وسط ضلع','M_GEO_SHAPES — دو طرح فرش و دیوار برای تکمیل قرینه','خارج از موتور'),
(49,'تقارن (ادامه) — فعّالیت','بتواند با دو خط تقارن، شکل را به چهار قسمت مثل هم تکمیل کند.','تقارن، نیمه، قسمت، ساعت','دو مورد تقارن چهارقسمتی؛ کامل کن: ساعت پنج و ربع','M_GEO_SHAPES — دو خط تقارن راست/چپ و بالا/پایین','خارج از موتور'),
(50,'کار در کلاس','بتواند شکل را به چهار قسمت برابر ببرد و عقربهٔ ساعت را برای ربع نشان دهد.','مربع، قسمت، ساعت، ربع','چهار مورد: ساخت با ۹ تکه، تقسیم به ۴، ساعت، رسم عقربه','M_GEO_SHAPES — سه مربع: دو مربع هر یک چهار قسمت و یکی کامل','خارج از موتور'),
(51,'تمرین','بتواند چهار بخش قرینهٔ فرش یا کاشی‌کاری را کامل کند.','تقارن، فرش، کاشی‌کاری','یک مورد تمرین دست‌آزاد: تکمیل چهار قسمت','M_GEO_SHAPES — دو نمونهٔ نقش برای تکمیل چهار بخش','خارج از موتور'),
(52,'مرور فصل + فرهنگ نوشتن + تمرین','بتواند نیمه یا سه قسمت دیگر شکل را نسبت به خط‌های تقارن کامل کند.','تقارن، شکل، رنگ، تمرین','مرور فصل: فرهنگ نوشتن و دو تمرین تصویری','M_GEO_SHAPES — تکمیل قرینهٔ دو نیمه/چهار قسمت؛ سؤال تصویری نیازمند موتور هندسه است','خارج از موتور'),
]
assert [p[0] for p in PAGES]==list(range(37,53))
s=w['نقشهٔ کتاب (S1)'];assert s['A47'].value=='یادداشت' and s['A22'].value=='L-200' and s['A23'].value=='L-300'
for idx,(page,title,goal,vocab,activity,visual,compat) in enumerate(PAGES):
 row=page+8;assert master.cell(row,1).value==page and master.cell(row,2).value==page+6
 note=f'نقشهٔ مادرِ اصلاح‌شده ردیف {row}، ص کتاب {page} / PDF {page+6}؛ C205.pdf هم‌چاپ. '
 if master.cell(row,14).value=='✗':note+='اصلاح آرنا در ستون O بر دادهٔ خام مقدم است. '
 if page in (38,46):note+='حکم کتاب کاغذی معلم دربارهٔ شش‌ضلعی/عنوان تقارن اعمال شد. '
 note+='موتور فعلی ضرب‌محور است؛ خروجی این ردیف تا تأیید معلم زرد است.'
 v=[f'L-{page+264}',f(page),'۳',title,'دانش‌آموز '+goal,vocab,activity,visual,
    'بالا' if compat=='بانک تمرین' else 'متوسط',compat,'—',SOURCE,Y,note]
 insert(s,47+idx,v,46)

# Diagnostic pattern taxonomy, proposals only; not active in multiplication runtime.
a=w['اطلس خطا'];assert a['A15'].value=='یادداشت'
errors=[
('E_PATTERN_REPEAT (پیشنهاد)','تکرار آخرین جملهٔ الگو','راهبردی','به جای ادامه‌دادن الگو همان جملهٔ آخر را پاسخ می‌دهد.','در ۳،۵،۷،۹ پاسخ ۹ می‌دهد.','الگوی دارای گام مشخص؛ گزینهٔ تکرار آخر','پرش بعدی را روی محور یا شکل نشان دهد.','متوسط'),
('E_PATTERN_STEP (پیشنهاد)','گام نامتناسب الگو','مفهومی','گام واقعی یا رشد گام را از الگوی دیده‌شده استخراج نمی‌کند.','در ۳،۵،۷،۹ پاسخ ۱۰ می‌دهد (+۱ نه +۲).','الگوی گام ثابت یا افزاینده؛ گزینهٔ گام غلط','اختلاف دو جملهٔ پیاپی را جدا ثبت کند.','بالا'),
('E_PATTERN_SKIP (پیشنهاد)','پرش اضافی در الگو','راهبردی','یک جملهٔ اضافه جلوتر از جملهٔ خواسته‌شده می‌رود.','در ۳،۵،۷،۹ پاسخ ۱۳ می‌دهد (پس از ۱۱).','الگوی چندجمله‌ای؛ گزینهٔ یک جمله جلوتر','روی شمارهٔ جملهٔ هدف مکث و پرش‌ها را بشمارد.','متوسط'),
]
for idx,e in enumerate(errors):
 insert(a,15+idx,[*e,SOURCE,Y,'کد پیشنهادی فصل ۳؛ در موتور ضربِ فعلی پیاده نشده؛ نسبت گزینه و سازوکار نیازمند ممیزی دوم است.'],14)

# 2 blocks of 2 CORE + 1 DEPTH + 1 SIDE. (prompt, work, correct, three
# distinct wrong values in fixed order: repeat, wrong-step, skip-next).
BLOCKS={
44:[
('الگوی ۳، ۵، ۷، ۹ را ادامه بده. جملهٔ بعد چند است؟','9+2=11',11,(9,10,13)),
('در الگوی ۱۰، ۱۵، ۲۰، ۲۵، جملهٔ بعد چند است؟','25+5=30',30,(25,26,35)),
('از عدد ۱۰ در الگویی با گام ۲، دو پرش جلو می‌رویم. به چه عددی می‌رسیم؟','10+2+2=14',14,(10,12,16)),
('در الگوی شکل‌هایی با ۱، ۳، ۶ دایره، اگر در هر گام یک دایره بیشتر از گام پیش اضافه شود، شکل بعد چند دایره دارد؟','6+4=10',10,(6,9,15)),
],
45:[
('الگوی ۲، ۶، ۱۰، ۱۴ را ادامه بده. جملهٔ بعد چند است؟','14+4=18',18,(14,15,22)),
('در الگوی ۴، ۹، ۱۴، ۱۹، عدد بعد چند است؟','19+5=24',24,(19,20,29)),
('در الگوی ۱، ۳، ۶، ۱۰، با گام‌هایی که هر بار یک واحد بیشتر می‌شوند، جملهٔ بعد چند است؟','10+5=15',15,(10,14,21)),
('در الگوی ۱، ۲، ۴، ۷، ۱۱، گام‌ها به‌ترتیب ۱، ۲، ۳، ۴ هستند؛ جملهٔ بعد چند است؟','11+5=16',16,(11,15,22)),
]}
q=w['بانک سؤال'];assert q['A72'].value=='یادداشت'
letters=['الف','ب','ج','د'];codes=['E_PATTERN_REPEAT','E_PATTERN_STEP','E_PATTERN_SKIP']
for page,items in BLOCKS.items():
 for i,(prompt,work,correct,wrong) in enumerate(items):
  n=1069+(page-44)*4+i
  assert len(set([correct,*wrong]))==4 and all(0<=x<=9999 for x in [correct,*wrong])
  # Independent restricted arithmetic evaluator of displayed calculation.
  def calc(node):
   if isinstance(node,ast.Expression):return calc(node.body)
   if isinstance(node,ast.Constant) and type(node.value)==int:return node.value
   if isinstance(node,ast.BinOp) and type(node.op) in (ast.Add,ast.Sub,ast.Mult):return {ast.Add:operator.add,ast.Sub:operator.sub,ast.Mult:operator.mul}[type(node.op)](calc(node.left),calc(node.right))
   raise ValueError(work)
  assert calc(ast.parse(work.split('=')[0],mode='eval'))==correct
  pos=(n-1069)%4
  opts=list(zip(wrong,codes));opts.insert(pos,(correct,None))
  mapping='؛ '.join(f'{letters[j]}→{code}' for j,(_,code) in enumerate(opts) if code)
  detail=f'repeat={wrong[0]} (آخرین جمله); wrong-step={wrong[1]} (گام نامتناسب); skip={wrong[2]} (جملهٔ جلوتر)'
  values=[f'Q-{n}',f'L-{page+264}',['CORE','CORE','DEPTH','SIDE'][i],'علامتی',prompt,'چهارگزینه‌ای',
          *(f(v) for v,c in opts),letters[pos],mapping,'آسان' if i<2 else 'متوسط',
          f'پاسخ دقیق: {f(correct)}؛ گام بعدی/تعداد پرش خواسته‌شده را حساب کنید.',SOURCE,Y,
          f'ضد J12 — کار دفترچه‌ای (لاتین): {work}; distractors: {detail}. سؤال خودساخته با لنگر ص {page}/PDF {page+6}؛ برچسب‌گذاری معنایی نیازمند بازرس دوم.']
  insert(q,72+(n-1069),values,71)
for page in BLOCKS:
 ids=[f'Q-{1069+(page-44)*4+i}' for i in range(4)]
 s.cell(47+page-37,11,'، '.join(ids))

# Re-open after save; verify all old values INCLUDING approvals and all 3 anchor rows.
w.properties.modified=datetime(2026,10,22,0,0,0)
w.save(OUT)
# openpyxl sets ZIP entry timestamps on each save. Canonicalise package metadata
# so the same cell data and style produce the same SHA-256 on repeated runs.
canon=OUT.with_suffix('.canonical-tmp')
with ZipFile(OUT,'r') as src, ZipFile(canon,'w') as dst:
 for name in src.namelist():
  info=ZipInfo(name,date_time=(1980,1,1,0,0,0))
  info.compress_type=ZIP_DEFLATED
  info.external_attr=0o600 << 16
  data=src.read(name)
  if name=='docProps/core.xml':
   data=re.sub(rb'<dcterms:modified[^>]*>[^<]*</dcterms:modified>',
     b'<dcterms:modified xsi:type="dcterms:W3CDTF">2026-10-22T00:00:00Z</dcterms:modified>',data)
  dst.writestr(info,data)
canon.replace(OUT)
actual=openpyxl.load_workbook(OUT,read_only=True,data_only=True)
base=openpyxl.load_workbook(args.base,read_only=True,data_only=True)
counts={'نقشهٔ کتاب (S1)':(47,16),'اطلس خطا':(15,3),'بانک سؤال':(72,8)}
for b in base:
 a=actual[b.title];at,offset=counts.get(b.title,(10**9,0))
 for row in b:
  for cell in row:
   if not hasattr(cell,'coordinate'):continue
   target_row=cell.row+(offset if cell.row>=at else 0)
   assert cell.value==a.cell(target_row,cell.column).value,(b.title,cell.coordinate)
 if offset:assert a.cell(at+offset,1).value=='یادداشت'
assert all(actual['نقشهٔ کتاب (S1)'].cell(r,13).value==Y for r in range(47,63))
assert all(actual['اطلس خطا'].cell(r,10).value==Y for r in range(15,18))
assert all(actual['بانک سؤال'].cell(r,16).value==Y for r in range(72,80))
assert all('۵۷−۱۴' not in str(actual[s].cell(r,c).value) for s,r,c in [('نقشهٔ کتاب (S1)',41,7)]) # guard old wrong example
assert all(re.fullmatch(r'[۰-۹]+',str(actual['بانک سؤال'].cell(r,c).value)) for r in range(72,80) for c in range(7,11))
assert Counter(actual['بانک سؤال'].cell(r,11).value for r in range(72,80))==dict.fromkeys(letters,2)
for r in range(72,80):
 x=actual['بانک سؤال'];opts=[x.cell(r,c).value for c in range(7,11)];assert len(set(opts))==4
 correct=letters.index(x.cell(r,11).value)
 refs=re.findall(r'(الف|ب|ج|د)→(E_[A-Z_]+)',x.cell(r,12).value)
 assert {k for k,v in refs}==set(letters)-{letters[correct]}
assert actual['بانک سؤال']['E14'].value.startswith('شکل زیر ۱۳ مثلث دارد')
assert actual['نقشهٔ کتاب (S1)']['G41'].value.find('۳۵−۱۷')!=-1
assert actual['نقشهٔ کتاب (S1)']['K13'].value.find('Q-1020')==-1
assert actual['نقشهٔ کتاب (S1)']['K15'].value.find('Q-1020')!=-1
hash=sha256(OUT.read_bytes()).hexdigest()
(ROOT/'khshab-G2-MATH_v0.5_verified.sha256').write_text(f'{hash}  {OUT.name}\n')
print('DRAFT wave4 from official base: S1=16, atlas=3, Q=8; sha256',hash)
