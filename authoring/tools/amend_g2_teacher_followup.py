"""Apply teacher rulings to map v1.1, changing ONLY O/P in five rows.
Do not promote N verdicts or rewrite the historical v1.1.
"""
from pathlib import Path
from hashlib import sha256
from datetime import datetime
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED
import re
import openpyxl
R=Path(__file__).resolve().parents[1]
base=R/'PAGE-MAP-KH-G2-MATH-v1.1_verified.xlsx'
assert sha256(base.read_bytes()).hexdigest()=='2c54be93a4ec016db61186ed7eed31fc34d51268d3e75a21adfcd6dfa4b6e839'
out=R/'PAGE-MAP-KH-G2-MATH-v1.2_teacher-followup.xlsx'
w=openpyxl.load_workbook(base);s=w['نقشهٔ صفحه‌ها']
patch={
65:('بازپس‌گیری کاملِ ادعای قبلی ستون O دربارهٔ سکه‌های ۱۰/۱۰۰/۱۰۰۰ ریالی در مورد ۲: در مورد ۲ عبارت «ارزش ۸۰ سکهٔ ۱ ریالی با ارزش چند سکهٔ ۱۰ ریالی برابر است؟» و جواب ۸ سکهٔ ۱۰ ریالی است. مورد ۱: ۲۰ سکهٔ ۱ ریالی = ۲ سکهٔ ۱۰ ریالی؛ مورد ۳: ۱۰۰ سکهٔ ۱ ریالی = ۱۰ سکهٔ ۱۰ ریالی. پایین: جعبهٔ «الگویابی کن» با سه دنباله. شمار سکه‌های شکل را جای متن سؤال نگذارید.','ص کتاب ۵۷ / PDF ۶۳، رندر صفحه + تأیید صریح معلم در گفت‌وگو؛ فقط این موردها، نه تأیید کل فصل.'),
88:('بالای صفحه تاب «حلّ مسئله» است؛ دقیقاً ۳ مورد شماره‌دار برای یافتن جای خالی و حدس/آزمایش: ۱) ۲۰+□=۴۳؛ ۲) □+۱۷=۵۱؛ ۳) دو زیرمورد: ۲۹+△=۶۳ و ۶۰−○=۲۷. پایین، جعبهٔ «کامل کن» با جدول شگفت‌انگیز است. چهار تساوی به معنی چهار سؤال شماره‌دار نیست؛ ابعاد جدول از این حکم نتیجه‌گیری نمی‌شود.','ص کتاب ۸۰ / PDF ۸۶، تصویر صفحهٔ فرستاده‌شدهٔ معلم در گفت‌وگو + تأیید کلامی سه سؤال و جدول پایانی؛ تصویر در مخزن منتشر نشده.'),
97:('صفحه بنر/تاب مستقل ندارد (ادامهٔ تمرین صفحهٔ قبل). چهار مورد شماره‌دار دقیقاً ۳ تا ۶ هستند: اندازه‌گیری ضلع شکل‌ها؛ اندازه‌گیری خط و یافتن وسط؛ حدس/اندازهٔ میلی‌متری؛ مقایسهٔ ضلع‌های مستطیل. حکم قبلی ستون O دربارهٔ «۴ تا ۷» باطل شد.','ص کتاب ۸۹ / PDF ۹۵، رندر شماره‌دار صفحه + تأیید صریح معلم در گفت‌وگو.'),
146:('بازپس‌گیری ادعای جابه‌جایی با ص ۱۳۹: ص چاپی ۱۳۸ «تمرین» است؛ جدول درختان خیابان‌های سعدی/حافظ/مولوی، نمودار کتاب‌ها و جعبهٔ «کامل کن». جزئیات تصویریِ سؤال‌ها هنوز به ممیزی مستقل نیاز دارد؛ بنر انتخاب نمودار به این صفحه تعلق ندارد.','ص کتاب ۱۳۸ / PDF ۱۴۴، رندر شماره‌دار + تأیید شخصی معلم در گفت‌وگو دربارهٔ عنوان/تخصیص صفحه.'),
147:('بازپس‌گیری ادعای جابه‌جایی با ص ۱۳۸: ص چاپی ۱۳۹ «انتخاب نمودار» با تاب «فعّالیت» است؛ فعالیت نمایش کارهای سعید با خانه‌های زمانی و نمودار ستونی. «تمرین» جدول درختان متعلق به ص ۱۳۸ است. جزئیات نمودار نیازمند ممیزی مستقل است.','ص کتاب ۱۳۹ / PDF ۱۴۵، رندر شماره‌دار + تأیید شخصی معلم در گفت‌وگو دربارهٔ عنوان/تخصیص صفحه.'),
}
for r,(o,p) in patch.items():
 assert s.cell(r,1).value==r-8 and s.cell(r,14).value=='✗'
 s.cell(r,15,o);s.cell(r,16,p)
old=openpyxl.load_workbook(base)
for a in old:
 b=w[a.title]
 for row in a:
  for cell in row:
   t=b[cell.coordinate]
   if a.title=='نقشهٔ صفحه‌ها' and cell.row in patch and cell.column in (15,16):continue
   assert (cell.value,cell.style_id)==(t.value,t.style_id),(a.title,cell.coordinate)
w.properties.modified=datetime(2026,10,22,0,0,0);w.save(out)
canon=out.with_suffix('.canonical-tmp')
with ZipFile(out,'r') as src,ZipFile(canon,'w') as dst:
 for name in src.namelist():
  info=ZipInfo(name,date_time=(1980,1,1,0,0,0));info.compress_type=ZIP_DEFLATED;info.external_attr=0o600 << 16
  data=src.read(name)
  if name=='docProps/core.xml':data=re.sub(rb'<dcterms:modified[^>]*>[^<]*</dcterms:modified>',b'<dcterms:modified xsi:type="dcterms:W3CDTF">2026-10-22T00:00:00Z</dcterms:modified>',data)
  dst.writestr(info,data)
canon.replace(out)
h=sha256(out.read_bytes()).hexdigest();(R/(out.name+'.sha256')).write_text(f'{h}  {out.name}\n')
print('map v1.2 O/P only',sorted(patch),'sha256',h)
