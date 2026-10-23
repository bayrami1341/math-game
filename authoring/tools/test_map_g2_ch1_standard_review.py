import unittest

from openpyxl import Workbook

from map_g2_ch1_standard_review import EXPECTED_Q, EXPECTED_S1, MAP_SHEET, Q_SHEET, inspect


def fixture():
    wb = Workbook()
    s = wb.active
    s.title = MAP_SHEET
    q = wb.create_sheet(Q_SHEET)
    for i, header in enumerate(EXPECTED_S1, 1):
        s.cell(2, i, header)
    for i, header in enumerate(EXPECTED_Q, 1):
        q.cell(3, i, header)
    for page in range(1, 19):
        r = page + 3
        s.cell(r, 1, f'L-{100+page}')
        s.cell(r, 2, str(page))
        s.cell(r, 3, '۱')
        s.cell(r, 13, 'نیاز به تأیید')
    for i in range(40):
        r = 4 + i
        q.cell(r, 1, f'Q-{1001+i}')
        q.cell(r, 2, 'L-105')
        q.cell(r, 16, 'نیاز به تأیید')
    return wb


class ExistingStandardTests(unittest.TestCase):
    def test_same_standard_without_visual_claims(self):
        result = inspect(fixture(), fixture())
        self.assertEqual(result['s1_page_rows'], 18)
        self.assertEqual(result['q_rows'], 40)
        self.assertEqual(result['pages_scanned_by_arena_in_this_run'], 0)
        self.assertEqual(result['pages'][4]['s1_id'], 'L-105')
        self.assertEqual(len(result['pages'][4]['q_ids']), 40)
        self.assertEqual(result['pages'][4]['s1_stamp_cell_read_only'], 'M8')

    def test_cell_difference_is_not_passed_as_a_visual_verdict(self):
        official, candidate = fixture(), fixture()
        candidate[MAP_SHEET]['H4'] = 'made-up illustration'
        with self.assertRaisesRegex(ValueError, 'CH1_S1_OFFICIAL_CANDIDATE_DRIFT'):
            inspect(official, candidate)
        official, candidate = fixture(), fixture()
        candidate[Q_SHEET]['E4'] = 'different'
        with self.assertRaisesRegex(ValueError, 'CH1_Q_OFFICIAL_CANDIDATE_DRIFT'):
            inspect(official, candidate)

    def test_green_stamp_fails_closed(self):
        official, candidate = fixture(), fixture()
        official[MAP_SHEET]['M4'] = 'تأییدشده'
        candidate[MAP_SHEET]['M4'] = 'تأییدشده'
        with self.assertRaisesRegex(ValueError, 'UNEXPECTED_GREEN_CH1_PAGE'):
            inspect(official, candidate)

    def test_header_and_page_identity_fail_closed(self):
        official, candidate = fixture(), fixture()
        candidate[MAP_SHEET]['H2'] = 'new parallel scheme'
        with self.assertRaisesRegex(ValueError, 'STANDARD_WORKBOOK_HEADER_DRIFT'):
            inspect(official, candidate)
        official, candidate = fixture(), fixture()
        candidate[MAP_SHEET]['B5'] = '۱۳'
        with self.assertRaisesRegex(ValueError, 'CH1_PAGE_ANCHOR_DRIFT'):
            inspect(official, candidate)


if __name__ == '__main__':
    unittest.main()
