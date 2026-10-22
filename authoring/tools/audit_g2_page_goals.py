"""Independent non-mutating check of chapter depth edits and converter artifact.
Usage: python authoring/tools/audit_g2_page_goals.py [--json /tmp/KH-G2-MATH.json]
"""
import argparse,json
from pathlib import Path
from hashlib import sha256
import openpyxl
R=Path(__file__).resolve().parents[1]
a=argparse.ArgumentParser();a.add_argument('--json',type=Path);args=a.parse_args()
p=R/'khshab-G2-MATH_v1.2_page-goals-draft.xlsx'
assert sha256(p.read_bytes()).hexdigest()=='18e6b4fa4957e968eca2b3b8066387858a627a99d7d52340064a1e8d8880c618'
base=openpyxl.load_workbook(R/'khshab-G2-MATH_v0.8_official-rebased-draft.xlsx');final=openpyxl.load_workbook(p)
assert base.sheetnames==final.sheetnames
s=final['نقشهٔ کتاب (S1)'];fa=str.maketrans('۰۱۲۳۴۵۶۷۸۹','0123456789')
rows={int(str(s.cell(r,2).value).translate(fa)):r for r in range(3,155) if str(s.cell(r,2).value).translate(fa).isdigit()}
assert set(rows)==set(range(1,145))
changed_pages=(set(range(73,91))-{80,89})|(set(range(91,109))-{99})|set(range(109,127))|{53,54}|(set(range(127,145))-{135,138,139})
assert len(changed_pages)==68
allowed={(rows[p],c) for p in changed_pages for c in (5,7,14)}
differences=[]
for old in base:
 new=final[old.title]
 assert old.max_row==new.max_row and old.max_column==new.max_column
 assert {str(x) for x in old.merged_cells.ranges}=={str(x) for x in new.merged_cells.ranges}
 for line in old:
  for c in line:
   t=new[c.coordinate]
   assert c.style_id==t.style_id,(old.title,c.coordinate,'style')
   if c.value!=t.value:
    assert old.title=='نقشهٔ کتاب (S1)' and (c.row,c.column) in allowed,(old.title,c.coordinate)
    differences.append((c.row,c.column))
assert set(differences)==allowed and len(differences)==204
assert all(s.cell(r,13).value=='نیاز به تأیید' for r in range(65,155))
assert sum('فعالیت «' in str(s.cell(r,5).value) for r in range(65,155))==0
assert final['بانک سؤال']['E14'].value.startswith('شکل زیر ۱۳ مثلث دارد')
assert final['بانک سؤال']['A108'].value=='یادداشت'
if args.json:
 j=json.loads(args.json.read_text());assert j['cartridgeId']=='KH-G2-MATH'
 assert len(j['bookMap'])==152 and len(j['questionBank'])==104 and len(j['errorAtlas'])==19
 assert {x['id'] for x in j['bookMap']}=={s.cell(r,1).value for r in range(3,155)}
 assert all(x['status']=='نیاز به تأیید' for x in j['questionBank'])
print('PASS: 204 scoped cells on 68 pages; old values/styles/merges and all 104 Q untouched; 0 generic; no new green')
