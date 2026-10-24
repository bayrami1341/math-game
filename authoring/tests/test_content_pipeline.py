from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

import openpyxl
import pymupdf
import pytest

from authoring.agent.cli import main as cli_main
from authoring.agent.content_pipeline import (
    audit_content_pipeline,
    export_extraction_ledger_xlsx,
)
from authoring.agent.models import file_sha256


def _sha(label: str) -> str:
    return hashlib.sha256(label.encode("utf-8")).hexdigest()


def _write_json(path: Path, value: dict[str, Any]) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2) + "\n", encoding="utf-8")
    return path


def _bundle(tmp_path: Path, *, include_synthetic_pdf: bool = False) -> dict[str, Any]:
    repo = tmp_path / "repo"
    content_dir = repo / "authoring" / "content" / "SYNTH-PRINT-1"
    content_dir.mkdir(parents=True)

    source_pdf_path = None
    source_pdf_sha256 = _sha("synthetic external pdf")
    source_pdf_size_bytes = 8123
    pdf_page_count = 12
    if include_synthetic_pdf:
        private_vault = tmp_path / "private-vault"
        private_vault.mkdir(mode=0o700)
        source_pdf_path = private_vault / "SYNTH.pdf"
        document = pymupdf.open()
        for page_number in range(1, pdf_page_count + 1):
            page = document.new_page()
            page.insert_text((40, 50), f"Synthetic test page {page_number}")
        document.save(source_pdf_path)
        document.close()
        source_pdf_sha256 = file_sha256(str(source_pdf_path))
        source_pdf_size_bytes = source_pdf_path.stat().st_size

    manifest = {
        "schema_version": "1.0.0",
        "manifest_id": "SYNTH-MANIFEST-1",
        "book_code": "SYNTH",
        "source_pdf": {
            "filename": "SYNTH.pdf",
            "sha256": source_pdf_sha256,
            "size_bytes": source_pdf_size_bytes,
            "pdf_page_count": pdf_page_count,
        },
        "pagination": {
            "body_pdf_page_start": 10,
            "body_pdf_page_end": 11,
            "book_page_start": 1,
            "book_page_end": 2,
            "pdf_minus_book_offset": 9,
        },
    }
    manifest_path = _write_json(content_dir / "manifest.json", manifest)
    binding = {
        "source_id": "SYNTH-PRINT-1",
        "book_code": "SYNTH",
        "edition_id": "SYNTH-EDITION-1",
        "manifest_id": "SYNTH-MANIFEST-1",
        "manifest_sha256": file_sha256(str(manifest_path)),
        "source_pdf_sha256": source_pdf_sha256,
    }
    pages: list[dict[str, Any]] = []
    anchors: list[dict[str, Any]] = []
    for pdf_page, book_page, suffix in ((10, 1, "A"), (11, 2, "B")):
        page_image_sha = _sha(f"synthetic-page-{pdf_page}")
        number_evidence_id = f"EV-PAGE-NUM-{pdf_page}"
        number_evidence_sha = _sha(f"synthetic-number-region-{pdf_page}")
        content_evidence_id = f"EV-CONTENT-{pdf_page}"
        content_evidence_sha = _sha(f"synthetic-content-region-{pdf_page}")
        record_id = f"PG-{pdf_page}"
        element_id = f"EL-{pdf_page}"
        pages.append({
            "record_id": record_id,
            "pdf_page": pdf_page,
            "book_page": book_page,
            "printed_page_number": book_page,
            "page_image_sha256": page_image_sha,
            "page_status": "CURRICULAR_CONTENT",
            "reading_status": "COMPLETE",
            "visual_scan_status": "NOT_APPLICABLE",
            "evidence_refs": [
                {
                    "evidence_id": f"EV-FULL-{pdf_page}",
                    "kind": "pdf_page_render",
                    "source_pdf_sha256": binding["source_pdf_sha256"],
                    "pdf_page": pdf_page,
                    "book_page": book_page,
                    "bbox_normalized": [0.0, 0.0, 1.0, 1.0],
                    "sha256": page_image_sha,
                },
                {
                    "evidence_id": number_evidence_id,
                    "kind": "pdf_text_bbox",
                    "purpose": "PRINTED_PAGE_NUMBER",
                    "source_pdf_sha256": binding["source_pdf_sha256"],
                    "pdf_page": pdf_page,
                    "book_page": book_page,
                    "bbox_normalized": [0.05, 0.80, 0.25, 0.97],
                    "sha256": number_evidence_sha,
                },
                {
                    "evidence_id": content_evidence_id,
                    "kind": "pdf_text_bbox",
                    "source_pdf_sha256": binding["source_pdf_sha256"],
                    "pdf_page": pdf_page,
                    "book_page": book_page,
                    "bbox_normalized": [0.12, 0.12, 0.88, 0.42],
                    "sha256": content_evidence_sha,
                },
            ],
            "elements": [{
                "element_id": element_id,
                "element_kind": "EXAMPLE" if pdf_page == 10 else "TASK",
                "concept_candidate_ids": ["whole_number_addition"],
                "operation_ids": ["add"],
                "number_domain_ids": ["whole_numbers"],
                "short_excerpt": "Add two whole numbers" if pdf_page == 10 else "Compare whole-number sums",
                "evidence_refs": [content_evidence_id],
            }],
        })
        anchors.append({
            "pdf_page": pdf_page,
            "book_page": book_page,
            "printed_page_number": book_page,
            "evidence_id": number_evidence_id,
            "evidence_sha256": number_evidence_sha,
        })

    extraction = {
        "schema_version": "khshab.extraction-ledger/v1",
        "bundle_id": "SYNTH-EXTRACT-1",
        "source_binding": binding,
        "offset_proof_sha256": "0" * 64,
        "section_scope": {
            "section_id": "SYNTH-SECTION-1",
            "pdf_page_start": 10,
            "pdf_page_end": 11,
            "book_page_start": 1,
            "book_page_end": 2,
        },
        "pages": pages,
    }
    proof = {
        "schema_version": "khshab.offset-proof/v1",
        "proof_id": "SYNTH-OFFSET-PROOF-1",
        "status": "SYNTHETIC_TEST_ONLY",
        "source_binding": binding,
        "section_id": "SYNTH-SECTION-1",
        "pdf_minus_book_offset": 9,
        "expected_page_count": 2,
        "matched_page_count": 2,
        "reader_pass_sha256": [_sha("reader pass A"), _sha("reader pass B")],
        "page_anchors": anchors,
    }
    proof_path = _write_json(content_dir / "offset-proof.json", proof)
    proof_sha = file_sha256(str(proof_path))
    extraction["offset_proof_sha256"] = proof_sha
    extraction_path = _write_json(content_dir / "extraction.json", extraction)

    source = {
        "schema_version": "khshab.content-source/v1",
        "source_id": binding["source_id"],
        "book_code": binding["book_code"],
        "grade": 2,
        "subject_id": "math",
        "edition_id": binding["edition_id"],
        "manifest_id": binding["manifest_id"],
        "manifest_path": manifest_path.relative_to(repo).as_posix(),
        "manifest_sha256": binding["manifest_sha256"],
        "source_pdf_sha256": binding["source_pdf_sha256"],
        "source_pdf_size_bytes": source_pdf_size_bytes,
        "pdf_page_count": pdf_page_count,
        "pagination": {
            "body_pdf_page_start": 10,
            "body_pdf_page_end": 11,
            "book_page_start": 1,
            "book_page_end": 2,
            "pdf_minus_book_offset": 9,
            "offset_proof_status": "SYNTHETIC_TEST_ONLY",
            "offset_proof_sha256": proof_sha,
            "offset_proof_path": proof_path.relative_to(repo).as_posix(),
        },
        "source_asset": {
            "storage": "PRIVATE_EXTERNAL",
            "locator_ref": "private://synthetic-fixture/SYNTH-PRINT-1",
        },
        "rights_status": "SYNTHETIC_TEST_ONLY",
    }
    source_path = _write_json(content_dir / "source.json", source)

    objective_tree = [
        {"objective_id": "OBJ-GENERAL", "level": "GENERAL", "parent_objective_id": None},
        {"objective_id": "OBJ-SPECIFIC", "level": "SPECIFIC", "parent_objective_id": "OBJ-GENERAL"},
        {"objective_id": "OBJ-BEHAVIORAL", "level": "BEHAVIORAL", "parent_objective_id": "OBJ-SPECIFIC"},
    ]
    valuation = {
        "schema_version": "khshab.valuation-bundle/v1",
        "valuation_bundle_id": "SYNTH-VALUATION-1",
        "source_binding": binding,
        "extraction_bundle_sha256": file_sha256(str(extraction_path)),
        "teacher_guide": {
            "guide_id": "SYNTH-GUIDE-1",
            "guide_sha256": _sha("synthetic teacher guide"),
            "book_code": binding["book_code"],
            "grade": 2,
            "subject_id": "math",
        },
        "objectives": objective_tree,
        "records": [
            {
                "valuation_id": f"VAL-{pdf_page}",
                "extraction_record_id": f"PG-{pdf_page}",
                "element_id": f"EL-{pdf_page}",
                "objective_ids": ["OBJ-GENERAL", "OBJ-SPECIFIC", "OBJ-BEHAVIORAL"],
                "concept_id": "whole_number_addition",
                "role": "PRIMARY",
                "operation_id": "add",
                "number_domain_id": "whole_numbers",
                "prerequisite_ids": [],
                "assessment_status": "IN_SCOPE",
                "confidence": "HIGH",
                "review_status": "DRAFT",
                "human_verdict_id": None,
            }
            for pdf_page in (10, 11)
        ],
    }
    valuation_path = _write_json(content_dir / "valuation.json", valuation)
    game_pack = {
        "schema_version": "khshab.game-content-pack/v1",
        "game_pack_id": "SYNTH-GAME-PACK-1",
        "source_binding": binding,
        "valuation_bundle_sha256": file_sha256(str(valuation_path)),
        "items": [],
    }
    game_pack_path = _write_json(content_dir / "game-pack.json", game_pack)
    return {
        "repo": repo,
        "content_dir": content_dir,
        "source_path": source_path,
        "manifest_path": manifest_path,
        "source_pdf_path": source_pdf_path,
        "proof_path": proof_path,
        "extraction_path": extraction_path,
        "valuation_path": valuation_path,
        "game_pack_path": game_pack_path,
        "binding": binding,
    }


def _audit(bundle: dict[str, Any], *, verdict_ledger_path: Path | None = None) -> dict[str, Any]:
    return audit_content_pipeline(
        repo_root=bundle["repo"],
        source_path=bundle["source_path"],
        extraction_path=bundle["extraction_path"],
        valuation_path=bundle["valuation_path"],
        game_pack_path=bundle["game_pack_path"],
        verdict_ledger_path=verdict_ledger_path,
        source_pdf_path=bundle.get("source_pdf_path"),
    )


def _codes(report: dict[str, Any]) -> set[str]:
    return {issue["code"] for issue in report["issues"]}


def test_committed_synthetic_fixture_is_self_consistent_but_ineligible() -> None:
    repo = Path(__file__).resolve().parents[2]
    fixture_dir = repo / "authoring/agent/fixtures/content_pipeline_synthetic"
    bundle = {
        "repo": repo,
        "source_path": fixture_dir / "source.json",
        "extraction_path": fixture_dir / "extraction.json",
        "valuation_path": fixture_dir / "valuation.json",
        "game_pack_path": fixture_dir / "game-pack.json",
    }

    report = _audit(bundle)

    assert report["status"] == "PASS"
    assert len(report["artifact_hashes"]) == 6
    assert report["metrics"]["rights_status"] == "SYNTHETIC_TEST_ONLY"
    assert report["metrics"]["section3_input_eligible"] is False
    assert report["metrics"]["section3_release_authorized"] is False


def test_synthetic_pipeline_trace_pass_is_not_section3_authorization(tmp_path: Path) -> None:
    bundle = _bundle(tmp_path)

    report = _audit(bundle)

    assert report["status"] == "PASS"
    assert report["reason_code"] == "PIPELINE_TRACE_VALID_DRAFT_ONLY"
    assert report["metrics"]["extraction_page_records"] == 2
    assert report["metrics"]["extracted_elements"] == 2
    assert report["metrics"]["section3_input_eligible"] is False
    assert report["metrics"]["section3_release_authorized"] is False
    assert report["safeguards"]["human_verdict_created"] is False
    assert report["safeguards"]["workbook_modified"] is False


def test_authorized_source_without_external_pdf_is_blocked(tmp_path: Path) -> None:
    bundle = _bundle(tmp_path)
    source = json.loads(bundle["source_path"].read_text(encoding="utf-8"))
    source["rights_status"] = "AUTHORIZED"
    _write_json(bundle["source_path"], source)

    report = _audit(bundle)

    assert report["status"] == "BLOCKED"
    assert report["reason_code"] == "SOURCE_PDF_BYTES_NOT_VERIFIED"
    assert report["metrics"]["source_pdf_bytes_verified"] is False
    assert report["metrics"]["section3_input_eligible"] is False


def test_synthetic_external_pdf_hash_size_and_page_count_are_checked(tmp_path: Path) -> None:
    bundle = _bundle(tmp_path, include_synthetic_pdf=True)

    report = _audit(bundle)

    assert report["status"] == "PASS"
    assert report["metrics"]["manifest_file_sha256_verified"] is True
    assert report["metrics"]["source_pdf_bytes_verified"] is True
    assert report["artifact_hashes"]["source_pdf"] == file_sha256(str(bundle["source_pdf_path"]))
    assert report["metrics"]["section3_input_eligible"] is False


def test_tampered_external_pdf_hash_is_fail_closed(tmp_path: Path) -> None:
    bundle = _bundle(tmp_path, include_synthetic_pdf=True)
    with bundle["source_pdf_path"].open("ab") as stream:
        stream.write(b"synthetic tamper")

    report = _audit(bundle)

    assert report["status"] == "BLOCKED"
    assert "SOURCE_PDF_SHA_MISMATCH" in _codes(report)
    assert "SOURCE_PDF_SIZE_MISMATCH" in _codes(report)
    assert report["metrics"]["source_pdf_bytes_verified"] is False


def test_missing_extraction_page_and_offset_anchor_block(tmp_path: Path) -> None:
    bundle = _bundle(tmp_path)
    extraction = json.loads(bundle["extraction_path"].read_text(encoding="utf-8"))
    extraction["pages"].pop()
    _write_json(bundle["extraction_path"], extraction)

    report = _audit(bundle)

    assert report["status"] == "BLOCKED"
    assert "EXTRACTION_PAGE_MISSING" in _codes(report)
    assert "OFFSET_PROOF_ANCHOR_NOT_IN_EXTRACTION" in _codes(report)


def test_stale_manifest_file_hash_blocks_before_content_handoff(tmp_path: Path) -> None:
    bundle = _bundle(tmp_path)
    manifest = json.loads(bundle["manifest_path"].read_text(encoding="utf-8"))
    manifest["book_code"] = "OTHER-SYNTH"
    _write_json(bundle["manifest_path"], manifest)

    report = _audit(bundle)

    assert report["status"] == "BLOCKED"
    assert report["reason_code"] == "SOURCE_MANIFEST_SHA_MISMATCH"
    assert report["metrics"]["section3_input_eligible"] is False


def test_stale_offset_proof_hash_blocks_before_handoff(tmp_path: Path) -> None:
    bundle = _bundle(tmp_path)
    proof = json.loads(bundle["proof_path"].read_text(encoding="utf-8"))
    proof["matched_page_count"] = 1
    _write_json(bundle["proof_path"], proof)

    report = _audit(bundle)

    assert report["status"] == "BLOCKED"
    assert report["reason_code"] == "OFFSET_PROOF_SHA_MISMATCH"


def test_offset_anchor_page_number_mismatch_is_blocked(tmp_path: Path) -> None:
    bundle = _bundle(tmp_path)
    proof = json.loads(bundle["proof_path"].read_text(encoding="utf-8"))
    proof["page_anchors"][0]["printed_page_number"] = 2
    _write_json(bundle["proof_path"], proof)
    source = json.loads(bundle["source_path"].read_text(encoding="utf-8"))
    source["pagination"]["offset_proof_sha256"] = file_sha256(str(bundle["proof_path"]))
    _write_json(bundle["source_path"], source)

    report = _audit(bundle)

    assert report["status"] == "BLOCKED"
    assert "OFFSET_PROOF_ANCHOR_MAPPING_INVALID" in _codes(report)


def test_offset_proof_status_must_match_source_record(tmp_path: Path) -> None:
    bundle = _bundle(tmp_path)
    proof = json.loads(bundle["proof_path"].read_text(encoding="utf-8"))
    proof["status"] = "TWO_READER_MATCHED"
    _write_json(bundle["proof_path"], proof)
    source = json.loads(bundle["source_path"].read_text(encoding="utf-8"))
    source["pagination"]["offset_proof_sha256"] = file_sha256(str(bundle["proof_path"]))
    _write_json(bundle["source_path"], source)

    report = _audit(bundle)

    assert report["status"] == "BLOCKED"
    assert "OFFSET_PROOF_STATUS_BINDING_MISMATCH" in _codes(report)


def test_unknown_fields_are_rejected_by_versioned_schema(tmp_path: Path) -> None:
    bundle = _bundle(tmp_path)
    source = json.loads(bundle["source_path"].read_text(encoding="utf-8"))
    source["unversioned_override"] = "do not accept"
    _write_json(bundle["source_path"], source)

    report = _audit(bundle)

    assert report["status"] == "BLOCKED"
    assert "PIPELINE_ARTIFACT_SCHEMA_INVALID" in _codes(report)
    assert any(issue.get("field_path") == "/" for issue in report["issues"])


def test_valuation_must_hash_exact_extraction_bundle(tmp_path: Path) -> None:
    bundle = _bundle(tmp_path)
    valuation = json.loads(bundle["valuation_path"].read_text(encoding="utf-8"))
    valuation["extraction_bundle_sha256"] = "0" * 64
    _write_json(bundle["valuation_path"], valuation)

    report = _audit(bundle)

    assert report["status"] == "BLOCKED"
    assert "VALUATION_EXTRACTION_SHA_MISMATCH" in _codes(report)


def test_game_item_from_unapproved_valuation_is_blocked(tmp_path: Path) -> None:
    bundle = _bundle(tmp_path)
    game_pack = json.loads(bundle["game_pack_path"].read_text(encoding="utf-8"))
    game_pack["items"] = [{
        "game_item_id": "GAME-DRAFT-1",
        "extraction_record_id": "PG-10",
        "element_id": "EL-10",
        "valuation_ids": ["VAL-10"],
        "objective_ids": ["OBJ-GENERAL", "OBJ-SPECIFIC", "OBJ-BEHAVIORAL"],
        "activity_kind": "PRACTICE",
        "mechanic_id": "guided_choice",
        "difficulty_id": "level_1",
        "personalization_tags": ["needs_scaffold"],
        "generation_status": "DRAFT",
    }]
    _write_json(bundle["game_pack_path"], game_pack)

    report = _audit(bundle)

    assert report["status"] == "BLOCKED"
    assert "GAME_ITEM_BLOCKED_UNAPPROVED_CONTENT" in _codes(report)
    assert report["safeguards"]["section3_release_authorized"] is False


def test_game_ready_status_is_rejected_by_trace_gate(tmp_path: Path) -> None:
    bundle = _bundle(tmp_path)
    game_pack = json.loads(bundle["game_pack_path"].read_text(encoding="utf-8"))
    game_pack["items"] = [{
        "game_item_id": "GAME-UNSAFE-1",
        "extraction_record_id": "PG-10",
        "element_id": "EL-10",
        "valuation_ids": ["VAL-10"],
        "objective_ids": ["OBJ-GENERAL", "OBJ-SPECIFIC", "OBJ-BEHAVIORAL"],
        "activity_kind": "GAME",
        "mechanic_id": "matching",
        "difficulty_id": "level_1",
        "personalization_tags": [],
        "generation_status": "READY",
    }]
    _write_json(bundle["game_pack_path"], game_pack)

    report = _audit(bundle)

    assert report["status"] == "BLOCKED"
    assert "PIPELINE_ARTIFACT_SCHEMA_INVALID" in _codes(report)
    assert any(issue.get("field_path") == "/items/0/generation_status" for issue in report["issues"])


def test_unauthenticated_verdict_record_cannot_approve_game_content(tmp_path: Path) -> None:
    bundle = _bundle(tmp_path)
    valuation = json.loads(bundle["valuation_path"].read_text(encoding="utf-8"))
    valuation["records"][0]["review_status"] = "HUMAN_APPROVED"
    valuation["records"][0]["human_verdict_id"] = "TV-SYNTH-ONLY"
    _write_json(bundle["valuation_path"], valuation)
    game_pack = json.loads(bundle["game_pack_path"].read_text(encoding="utf-8"))
    game_pack["valuation_bundle_sha256"] = file_sha256(str(bundle["valuation_path"]))
    game_pack["items"] = [{
        "game_item_id": "GAME-DRAFT-1",
        "extraction_record_id": "PG-10",
        "element_id": "EL-10",
        "valuation_ids": ["VAL-10"],
        "objective_ids": ["OBJ-GENERAL", "OBJ-SPECIFIC", "OBJ-BEHAVIORAL"],
        "activity_kind": "PRACTICE",
        "mechanic_id": "guided_choice",
        "difficulty_id": "level_1",
        "personalization_tags": [],
        "generation_status": "DRAFT",
    }]
    _write_json(bundle["game_pack_path"], game_pack)

    without_ledger = _audit(bundle)
    assert without_ledger["status"] == "BLOCKED"
    assert "HUMAN_VERDICT_LEDGER_REQUIRED" in _codes(without_ledger)

    verdict_ledger = _write_json(bundle["content_dir"] / "verdict-ledger.json", {
        "schema_version": "1.0.0",
        "verdict_count": 1,
        "verdicts": [{
            "verdict_id": "TV-SYNTH-ONLY",
            "scope": "ROW_APPROVED",
            "actor": "Synthetic fixture only",
            "actor_role": "human_teacher",
            "target_ids": ["VAL-10"],
            "target_sha256": file_sha256(str(bundle["valuation_path"])),
            "timestamp_utc": "2026-01-01T00:00:00Z",
            "notes": "Synthetic test record, not a real human verdict",
            "authenticated": False,
            "auth_method": "KHSHAB_TEACHER_TOKEN_FILE",
            "invalidated": False,
            "invalidation_reason": None,
        }],
    })
    with_ledger = _audit(bundle, verdict_ledger_path=verdict_ledger)
    assert with_ledger["status"] == "BLOCKED"
    assert "HUMAN_VERDICT_BINDING_INVALID" in _codes(with_ledger)
    assert with_ledger["metrics"]["human_approval_ledger_matches"] == 0
    assert with_ledger["metrics"]["section3_input_eligible"] is False
    assert with_ledger["metrics"]["section3_release_authorized"] is False


def test_git_artifacts_reject_long_text_and_raw_payload_keys(tmp_path: Path) -> None:
    bundle = _bundle(tmp_path)
    extraction = json.loads(bundle["extraction_path"].read_text(encoding="utf-8"))
    extraction["pages"][0]["elements"][0]["raw_text"] = "synthetic protected raw content"
    _write_json(bundle["extraction_path"], extraction)

    report = _audit(bundle)

    assert report["status"] == "BLOCKED"
    assert "GIT_CONTENT_FIELD_FORBIDDEN" in _codes(report)
    assert "synthetic protected raw content" not in json.dumps(report, ensure_ascii=False)


def test_artifacts_must_be_versioned_inside_repository(tmp_path: Path) -> None:
    bundle = _bundle(tmp_path)
    outside = tmp_path / "outside-source.json"
    outside.write_bytes(bundle["source_path"].read_bytes())

    report = audit_content_pipeline(
        repo_root=bundle["repo"],
        source_path=outside,
        extraction_path=bundle["extraction_path"],
        valuation_path=bundle["valuation_path"],
        game_pack_path=bundle["game_pack_path"],
    )

    assert report["status"] == "FAIL"
    assert report["reason_code"] == "PIPELINE_ARTIFACT_OUTSIDE_REPOSITORY"


def test_excel_snapshot_is_new_hash_linked_and_never_overwrites(tmp_path: Path) -> None:
    bundle = _bundle(tmp_path, include_synthetic_pdf=True)
    output = bundle["content_dir"] / "extraction-snapshot.xlsx"

    result = export_extraction_ledger_xlsx(
        repo_root=bundle["repo"],
        source_path=bundle["source_path"],
        extraction_path=bundle["extraction_path"],
        source_pdf_path=bundle["source_pdf_path"],
        output_path=output,
    )

    assert result["rows"] == 2
    assert result["overwrote_existing"] is False
    assert result["canonical_json_sha256"] == file_sha256(str(bundle["extraction_path"]))
    assert result["source_pdf_bytes_verified"] is True
    assert result["source_pdf_sha256"] == file_sha256(str(bundle["source_pdf_path"]))
    assert output.is_file()
    assert output.with_suffix(".xlsx.sha256").is_file()
    wb = openpyxl.load_workbook(output, data_only=True)
    ws = wb["دفتر استخراج"]
    assert ws.max_row == 3
    assert ws["A2"].value == "SYNTH-EXTRACT-1"
    assert ws["K2"].value == "Add two whole numbers"
    meta = wb["ردیابی"]
    assert meta["A1"].value == "field"
    assert meta["B7"].value is True
    assert meta["B8"].value == file_sha256(str(bundle["source_pdf_path"]))
    assert meta["B10"].value == file_sha256(str(bundle["extraction_path"]))
    assert meta["B14"].value == json.loads(bundle["extraction_path"].read_text(encoding="utf-8"))["offset_proof_sha256"]
    assert meta["B22"].value is False
    assert meta["B24"].value is False
    wb.close()
    with pytest.raises(FileExistsError):
        export_extraction_ledger_xlsx(
            repo_root=bundle["repo"],
            source_path=bundle["source_path"],
            extraction_path=bundle["extraction_path"],
            output_path=output,
            source_pdf_path=bundle["source_pdf_path"],
        )


def test_excel_export_refuses_stale_offset_proof_before_writing(tmp_path: Path) -> None:
    bundle = _bundle(tmp_path)
    proof = json.loads(bundle["proof_path"].read_text(encoding="utf-8"))
    proof["matched_page_count"] = 1
    _write_json(bundle["proof_path"], proof)
    output = bundle["content_dir"] / "must-not-exist.xlsx"

    with pytest.raises(ValueError, match="OFFSET_PROOF_SHA_MISMATCH"):
        export_extraction_ledger_xlsx(
            repo_root=bundle["repo"],
            source_path=bundle["source_path"],
            extraction_path=bundle["extraction_path"],
            output_path=output,
        )

    assert not output.exists()
    assert not output.with_suffix(".xlsx.sha256").exists()


def test_excel_export_requires_external_pdf_for_authorized_source(tmp_path: Path) -> None:
    bundle = _bundle(tmp_path)
    source = json.loads(bundle["source_path"].read_text(encoding="utf-8"))
    source["rights_status"] = "AUTHORIZED"
    _write_json(bundle["source_path"], source)
    output = bundle["content_dir"] / "must-not-export.xlsx"

    with pytest.raises(ValueError, match="SOURCE_PDF_BYTES_NOT_VERIFIED"):
        export_extraction_ledger_xlsx(
            repo_root=bundle["repo"],
            source_path=bundle["source_path"],
            extraction_path=bundle["extraction_path"],
            output_path=output,
        )

    assert not output.exists()


def test_extraction_export_cli_dispatch_creates_only_new_snapshot(tmp_path: Path, capsys) -> None:
    bundle = _bundle(tmp_path)
    output = bundle["content_dir"] / "cli-snapshot.xlsx"
    rc = cli_main([
        "export-extraction-xlsx",
        "--repo-root", str(bundle["repo"]),
        "--source-record", str(bundle["source_path"]),
        "--extraction-ledger", str(bundle["extraction_path"]),
        "--output", str(output),
    ])
    result = json.loads(capsys.readouterr().out)

    assert rc == 0
    assert result["status"] == "PASS"
    assert result["canonical_json_sha256"] == file_sha256(str(bundle["extraction_path"]))
    assert output.is_file()
    assert output.with_suffix(".xlsx.sha256").is_file()


def test_pipeline_cli_dispatch_reads_synthetic_github_bundle(tmp_path: Path, capsys) -> None:
    bundle = _bundle(tmp_path)
    rc = cli_main([
        "content-pipeline-audit",
        "--repo-root", str(bundle["repo"]),
        "--source-record", str(bundle["source_path"]),
        "--extraction-ledger", str(bundle["extraction_path"]),
        "--valuation-bundle", str(bundle["valuation_path"]),
        "--game-pack", str(bundle["game_pack_path"]),
    ])
    report = json.loads(capsys.readouterr().out)

    assert rc == 0
    assert report["status"] == "PASS"
    assert report["traceability_only"] is True
    assert report["safeguards"]["section3_release_authorized"] is False
