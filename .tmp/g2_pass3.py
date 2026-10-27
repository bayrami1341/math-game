#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""پاس ۳ — ریشه‌یابی نکات باقی‌ماندهٔ G2"""
import re, collections
from openpyxl import load_workbook

XL = '/home/z/my-project/upload/G2-1404-RUNTIME.xlsx'
wb = load_workbook(XL, data_only=True)
Z = '\u200c'

def s(v):
    return '' if v is None else ' '.join(str(v).split())

def sheet_rows(name):
    return [[s(v) for v in r] for r in wb[name].iter_rows(values_only=True)]

CHK = sheet_rows('چک‌لیست')
KRT = sheet_rows('کارت')
NSH = sheet_rows('نشست')
idx = {h: i for i, h in enumerate(CHK[0])}
kidx = {h: i for i, h in enumerate(KRT[0])}
nidx = {h: i for i, h in enumerate(NSH[0])}
data = [r for r in CHK[1:] if any(c for c in r)]
kdata = [r for r in KRT[1:] if any(c for c in r)]
ndata = [r for r in NSH[1:] if any(c for c in r)]
ID_RE = re.compile(r'^G2-1404-P(\d+)-T(\d+)$')

print('=== ۱) الگوی تشخیص با قالب درست (دنبالهٔ ثابت) ===')
ok = bad = 0
bad_ex = []
tail1 = '. بار اول است. کارت باز نمی' + Z + 'شود.'
patA = 'نشانهٔ رد دیده نشد. کارت باز نمی' + Z + 'شود.'
patR = 'عدد این ردیف قطعی نبود. تشخیص نساز و وقت نساز.'
for r in data:
    d = r[idx['تشخیص']]
    if r[idx['ریاضی نیست']]:
        if d == patR:
            ok += 1
        else:
            bad += 1
            bad_ex.append((r[idx['شناسه']], d))
        continue
    if r[idx['رد']] == 'اگر جواب غیر از قبول بدهد':
        if d == patA:
            ok += 1
        else:
            bad += 1
            bad_ex.append((r[idx['شناسه']], d))
        continue
    exp = 'چه دیدی: ' + r[idx['رد']] + '. قدم بعد: ' + r[idx['قبول']] + tail1
    if d == exp:
        ok += 1
    else:
        bad += 1
        bad_ex.append((r[idx['شناسه']], 'انتظار: ' + exp + ' || واقعی: ' + d))
print(f'ok={ok} bad={bad}')
for e in bad_ex[:6]:
    print('  ', e)

print()
print('=== ۲) ردیف‌های بدون‌مهارتِ غیرِروشن‌نیست ===')
weird = [r for r in data if not r[idx['مهارت']] and not r[idx['ریاضی نیست']]]
print('count:', len(weird))
for r in weird:
    i = data.index(r)
    n = ndata[i]
    print(f"  {r[idx['شناسه']]} | کودک: {r[idx['کودک چه می' + Z + 'کند']][:45]} | قبول: {r[idx['قبول']][:40]} | رد: {r[idx['رد']][:40]}")
    print(f"      تشخیص: {r[idx['تشخیص']][:90]} | نشست: چه‌شد={n[nidx['چه شد']][:30]} وضعیت={n[nidx['وضعیت بعد']]} نتیجه={n[nidx['نتیجه']][:40]}")

print()
print('=== ۳) جزئیات SK-PACK-LEFT / SK-TWO-STEP ===')
cards = {r[kidx['شناسهٔ مهارت']]: (i, r) for i, r in enumerate(kdata)}
for cid in ['SK-TWO-STEP', 'SK-PACK-LEFT']:
    i, r = cards[cid]
    print(f"  {cid}: ردیف کارت={i+2} | نام={r[kidx['نام']]} | پدر={r[kidx['پدر']]} | تولد={r[kidx['صفحهٔ تولد']]} | وضعیت={r[kidx['وضعیت']]}")
# اولین کار هر دو در چک‌لیست
for cid in ['SK-TWO-STEP', 'SK-PACK-LEFT']:
    firsts = [r[idx['شناسه']] for r in data if r[idx['مهارت']] == cid]
    print(f'  {cid}: کارها در چک‌لیست = {firsts}')

print()
print('=== ۴) ستون پدر در برابر «اگر این باز است این را نزن» ===')
same = diff = 0
diff_ex = []
for r in data:
    a, b = r[idx['پدر']], r[idx['اگر این باز است این را نزن']]
    if a == b:
        same += 1
    else:
        diff += 1
        if len(diff_ex) < 6:
            diff_ex.append((r[idx['شناسه']], 'پدر=' + a, 'نزن=' + b))
print(f'برابر: {same} | متفاوت: {diff}')
for e in diff_ex:
    print('  ', e)

print()
print('=== ۵) توزیع ستون رد ===')
rd = collections.Counter(r[idx['رد']] for r in data)
for k, v in rd.most_common(8):
    print(f'  ×{v}: {k[:60]}')

print()
print('=== ۶) قبول/رد ردیف‌های روشن‌نیست (۵ نمونه) ===')
rn = [r for r in data if r[idx['ریاضی نیست']]]
for r in rn[:5]:
    print(f"  {r[idx['شناسه']]} | قبول={r[idx['قبول']]!r} | رد={r[idx['رد']]!r} | کارت‌تازه={r[idx['کارت تازه']]}")

print()
print('=== ۷) کارت‌های پدرِ پرکاربرد و وابستگی متقابل ===')
# آیا پدری هست که فرزندش پدرِ خودش باشد؟ (حلقه)
fmap = {r[kidx['شناسهٔ مهارت']]: r[kidx['پدر']] for r in kdata}
loops = []
for k, f in fmap.items():
    seen = {k}
    cur = f
    while cur and cur in fmap and cur not in seen:
        seen.add(cur)
        cur = fmap[cur]
    if cur in seen and cur:
        loops.append((k, sorted(seen)))
print('حلقهٔ پدر:', loops if loops else 'هیچ')

print()
print('=== ۸) تولد در برابر اولین کار چک‌لیست ===')
first_use = {}
for r in data:
    m = r[idx['مهارت']]
    if m:
        p = int(ID_RE.match(r[idx['شناسه']]).group(1))
        if m not in first_use or p < first_use[m]:
            first_use[m] = p
mis = []
for cid, (i, r) in cards.items():
    b = int(r[kidx['صفحهٔ تولد']])
    fu = first_use.get(cid)
    if fu is not None and b != fu:
        mis.append((cid, 'تولد=' + str(b), 'اولین کار=' + str(fu)))
print('ناهم‌خوانی تولد با اولین کار:', len(mis))
for e in mis[:10]:
    print('  ', e)
