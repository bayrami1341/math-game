#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""پاس ۱ — شناسنامهٔ G2-1404-RUNTIME.xlsx: هش، شیت‌ها، هدرها، نمونه‌ها"""
import hashlib, os
from openpyxl import load_workbook

XL = '/home/z/my-project/upload/G2-1404-RUNTIME.xlsx'
PR = '/home/z/my-project/upload/G2-ZAI-PROMPT.md'

def sha256(p):
    h = hashlib.sha256()
    with open(p, 'rb') as f:
        for c in iter(lambda: f.read(1 << 20), b''):
            h.update(c)
    return h.hexdigest()

def s(v):
    if v is None:
        return ''
    return ' '.join(str(v).split())

wb = load_workbook(XL, data_only=True)
print('=== FILE ===')
print('xlsx size:', os.path.getsize(XL), 'bytes')
print('xlsx sha256:', sha256(XL))
print('prompt size:', os.path.getsize(PR), 'bytes')
print('prompt sha256:', sha256(PR))
print()
print('=== SHEETS (order matters) ===')
sheets = {}
for i, ws in enumerate(wb.worksheets):
    rows = []
    for r in ws.iter_rows(values_only=True):
        rows.append([s(v) for v in r])
    sheets[ws.title] = rows
    nonempty = [r for r in rows if any(c for c in r)]
    ncol = max((len(r) for r in rows), default=0)
    print(f'[{i}] "{ws.title}"  rows={len(rows)} nonempty={len(nonempty)} cols={ncol}')
print()
for name, rows in sheets.items():
    nonempty = [r for r in rows if any(c for c in r)]
    print(f'===== SHEET: "{name}" — first 6 nonempty rows =====')
    for r in nonempty[:6]:
        print('  ', [c[:70] for c in r])
    print(f'  ... last 2 nonempty rows:')
    for r in nonempty[-2:]:
        print('  ', [c[:70] for c in r])
    print()
