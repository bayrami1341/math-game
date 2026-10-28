#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""پاس ۲ — راستی‌آزمایی کامل G2-1404-RUNTIME.xlsx در برابر G2-ZAI-PROMPT.md"""
import re, json, collections
from openpyxl import load_workbook

XL = '/home/z/my-project/upload/G2-1404-RUNTIME.xlsx'
wb = load_workbook(XL, data_only=True)

def s(v):
    if v is None:
        return ''
    return ' '.join(str(v).split())

def sheet_rows(name):
    ws = wb[name]
    rows = []
    for r in ws.iter_rows(values_only=True):
        rows.append([s(v) for v in r])
    return rows

Z = '\u200c'  # ZWNJ

RAH = sheet_rows('راهنما')
GHA = sheet_rows('قاعده')
CHK = sheet_rows('چک‌لیست')
KRT = sheet_rows('کارت')
NSH = sheet_rows('نشست')

print('############ راهنما — متن کامل ############')
for r in RAH:
    if any(c for c in r):
        print(f'  [{r[0]}] {r[1]}')
print()
print('############ قاعده — متن کامل ############')
for r in GHA:
    if any(c for c in r):
        print(f'  شماره={r[0]} | موضوع={r[1]}')
        print(f'    قاعده: {r[2]}')
print()

# ---------- چک‌لیست ----------
hdr = CHK[0]
idx = {h: i for i, h in enumerate(hdr)}
data = [r for r in CHK[1:] if any(c for c in r)]
print('############ چک‌لیست ############')
print('data rows:', len(data))
ID_RE = re.compile(r'^G2-1404-P(\d+)-T(\d+)$')
ids, bad_ids = [], []
pages = {}
for r in data:
    m = ID_RE.match(r[idx['شناسه']])
    if not m:
        bad_ids.append(r[idx['شناسه']])
        continue
    p, t = int(m.group(1)), int(m.group(2))
    ids.append((p, t, r))
    pages.setdefault(p, []).append((t, r))
print('unique IDs:', len({r[idx['شناسه']] for r in data}))
print('bad IDs:', bad_ids if bad_ids else 'هیچ')
plist = sorted(pages)
print('pages: min=%d max=%d count=%d' % (plist[0], plist[-1], len(plist)))
missing = [p for p in range(1, 144) if p not in pages]
beyond = [p for p in plist if p < 1 or p > 143]
print('missing pages 1..143:', missing if missing else 'هیچ')
print('pages beyond 1..143:', beyond if beyond else 'هیچ')
t_bad = []
for p in plist:
    ts = sorted(t for t, _ in pages[p])
    if ts != list(range(1, len(ts) + 1)):
        t_bad.append((p, ts))
print('T-sequence violations:', t_bad if t_bad else 'هیچ')
# بخش
sec = collections.Counter(r[idx['بخش']] for r in data)
sec_ranges = {}
for r in data:
    sec_ranges.setdefault(r[idx['بخش']], [999, 0])
    p = int(ID_RE.match(r[idx['شناسه']]).group(1))
    sec_ranges[r[idx['بخش']]][0] = min(sec_ranges[r[idx['بخش']]][0], p)
    sec_ranges[r[idx['بخش']]][1] = max(sec_ranges[r[idx['بخش']]][1], p)
print('بخش distribution:', dict(sec))
print('بخش page-ranges:', {k: tuple(v) for k, v in sorted(sec_ranges.items())})
# صفحه/کار ستون‌ها هم‌خوان با شناسه
pg_mis = [r[idx['شناسه']] for r in data if r[idx['صفحهٔ چاپ']] != str(int(ID_RE.match(r[idx['شناسه']]).group(1)))]
tk_mis = [r[idx['شناسه']] for r in data if r[idx['کار']] != str(int(ID_RE.match(r[idx['شناسه']]).group(2)))]
print('صفحهٔ چاپ mismatch:', len(pg_mis), pg_mis[:5])
print('کار mismatch:', len(tk_mis), tk_mis[:5])
# مهارت
sk_used = [r[idx['مهارت']] for r in data if r[idx['مهارت']]]
sk_empty_rows = [r for r in data if not r[idx['مهارت']]]
print('مهارت non-empty:', len(sk_used), '| empty:', len(sk_empty_rows))
print('unique مهارت codes used:', len(set(sk_used)))
# روشن نیست
rn_col = idx['ریاضی نیست']
rn_vals = collections.Counter(r[rn_col] for r in data)
print('«ریاضی نیست» column values:', dict(rn_vals))
rn_rows = [r for r in data if r[rn_col]]
print('rows with روشن‌نیست marker:', len(rn_rows))
print('empty-مهارت rows == روشن‌نیست rows?', {r[idx['شناسه']] for r in sk_empty_rows} == {r[idx['شناسه']] for r in rn_rows})
# وضعیت
st = collections.Counter(r[idx['وضعیت']] for r in data)
print('وضعیت values:', dict(st))
kt = collections.Counter(r[idx['کارت تازه']] for r in data)
print('کارت تازه values:', dict(kt))
sq = collections.Counter(r[idx['سقف']] for r in data)
print('سقف values:', {k: v for k, v in sq.items() if k}, '| empty:', sq.get('', 0))
# پدر
father_col = idx['پدر']
f_vals = collections.Counter(r[father_col] for r in data)
print('پدر (چک‌لیست) distinct:', sorted({k for k in f_vals if k}))
# تشخیص
dia = collections.Counter()
tpl_ok = tpl_bad = 0
tpl_bad_ex = []
for r in data:
    d = r[idx['تشخیص']]
    if r[rn_col]:
        dia['روشن‌نیست: ' + d[:40]] += 1
        continue
    if r[idx['رد']] == 'اگر جواب غیر از قبول بدهد':
        if d == 'نشانهٔ رد دیده نشد. کارت باز نمی' + Z + 'شود.':
            tpl_ok += 1
        else:
            tpl_bad += 1
            tpl_bad_ex.append((r[idx['شناسه']], d))
        dia['الگوی بی‌نشانه'] += 1
    else:
        exp = 'چه دیدی: ' + r[idx['رد']] + '. قدم بعد: ' + r[idx['قبول']]
        if d == exp:
            tpl_ok += 1
        else:
            tpl_bad += 1
            tpl_bad_ex.append((r[idx['شناسه']], d))
        dia['الگوی چه‌دیدی/قدم‌بعد'] += 1
print('تشخیص الگوها:', dict(dia))
print('تشخیص template: ok=%d bad=%d' % (tpl_ok, tpl_bad))
for e in tpl_bad_ex[:5]:
    print('   نمونهٔ بد:', e)
# اگر این باز است این را نزن
dont = collections.Counter(r[idx['اگر این باز است این را نزن']] for r in data)
print('«اگر این باز است این را نزن» distinct:', sorted({k for k in dont if k}))
# منبع
src = collections.Counter(r[idx['منبع']] for r in data)
print('منبع values:', dict(src))

# ---------- کارت ----------
print()
print('############ کارت ############')
kh = KRT[0]
kidx = {h: i for i, h in enumerate(kh)}
kdata = [r for r in KRT[1:] if any(c for c in r)]
print('data rows (cards):', len(kdata))
cards = {r[kidx['شناسهٔ مهارت']]: r for r in kdata}
print('unique card ids:', len(cards))
dups = [k for k, v in collections.Counter(r[kidx['شناسهٔ مهارت']] for r in kdata).items() if v > 1]
print('duplicate card ids:', dups if dups else 'هیچ')
ntype = collections.Counter(r[kidx['نوع کار']] for r in kdata)
print('نوع کار distribution:', dict(ntype))
kst = collections.Counter(r[kidx['وضعیت']] for r in kdata)
print('وضعیت values:', dict(kst))
# پدر refs + تولد
invalid_f = []
birth_bad = []
order_bad = []
birth_gt = []
for r in kdata:
    cid = r[kidx['شناسهٔ مهارت']]
    f = r[kidx['پدر']]
    if f and f not in cards:
        invalid_f.append((cid, f))
    b = r[kidx['صفحهٔ تولد']]
    if not b.isdigit() or not (1 <= int(b) <= 143):
        birth_bad.append((cid, b))
    if f in cards:
        fb = cards[f][kidx['صفحهٔ تولد']]
        if fb.isdigit() and b.isdigit() and int(fb) > int(b):
            birth_gt.append((cid, 'پدر تولد=' + fb, 'فرزند تولد=' + b))
        if kdata.index(cards[f]) > kdata.index(r):
            order_bad.append((cid, f))
print('پدر نامعتبر:', invalid_f if invalid_f else 'هیچ')
print('تولد خارج ۱..۱۴۳:', birth_bad if birth_bad else 'هیچ')
print('پدرِ دیرتر (تولد پدر > تولد فرزند):', birth_gt if birth_gt else 'هیچ')
print('پدرِ تعریف‌شده بعد از فرزند (ترتیب ردیف کارت):', order_bad if order_bad else 'هیچ')
# وجود/عدم وجود
print('SK-SEE exists?', 'SK-SEE' in cards)
countfam = sorted(k for k in cards if k.startswith('SK-COUNT'))
print('SK-COUNT* cards:', countfam)
chops = [k for k in cards if 'CHOP' in k]
for k in chops:
    print(f'  {k}: نام={cards[k][kidx["نام"]]} پدر={cards[k][kidx["پدر"]]} تولد={cards[k][kidx["صفحهٔ تولد"]]}')
print('SK-READ-3:', cards['SK-READ-3'][kidx['نام']], '| پدر:', cards['SK-READ-3'][kidx['پدر']], '| تولد:', cards['SK-READ-3'][kidx['صفحهٔ تولد']])
print('SK-STORY-SUB:', cards['SK-STORY-SUB'][kidx['نام']], '| پدر:', cards['SK-STORY-SUB'][kidx['پدر']])
# dangling/orphan
used = set(sk_used)
dangling = sorted(used - set(cards))
orphans = sorted(set(cards) - used)
print('dangling (استفاده‌شده، تعریف‌نشده):', dangling if dangling else 'هیچ')
print('orphans (تعریف‌شده، استفاده‌نشده):', orphans if orphans else 'هیچ')
# جملهٔ بستن تکراری
clos = collections.Counter(r[kidx['قاعدهٔ بستن']] for r in kdata)
dup_clos = {k: v for k, v in clos.items() if v > 1}
print('قاعدهٔ بستن تکراری:', len(dup_clos))
for k, v in list(dup_clos.items())[:5]:
    print('   ×', v, k[:80])
# چند صفحه به یک کارت
page_span = {}
for p, t, r in ids:
    m = r[idx['مهارت']]
    if m:
        page_span.setdefault(m, set()).add(p)
multi = {k: len(v) for k, v in page_span.items() if len(v) > 1}
print('کارت‌های چندصفحه‌ای:', len(multi), '| بیشترین صفحات:', sorted(multi.items(), key=lambda x: -x[1])[:5])

# ---------- نشست ----------
print()
print('############ نشست ############')
nh = NSH[0]
nidx = {h: i for i, h in enumerate(nh)}
ndata = [r for r in NSH[1:] if any(c for c in r)]
print('data rows:', len(ndata))
nums = [r[nidx['نشست']] for r in ndata]
seq_bad = [i + 1 for i, n in enumerate(nums) if n != str(i + 1)]
print('نشست sequential 1..N:', 'بله' if not seq_bad else f'خیر {seq_bad[:5]}')
pgs = [int(r[nidx['صفحهٔ کلاس']]) for r in ndata]
print('صفحهٔ کلاس: min=%d max=%d' % (min(pgs), max(pgs)))
print('monotonic non-decreasing:', all(a <= b for a, b in zip(pgs, pgs[1:])))
stat = collections.Counter(r[nidx['وضعیت بعد']] for r in ndata)
print('وضعیت بعد distribution:', dict(stat))
res = collections.Counter(r[nidx['نتیجه']] for r in ndata)
print('نتیجه distribution:', dict(res))
rh = collections.Counter(r[nidx['راهنما']] for r in ndata)
print('راهنما distribution:', dict(rh))
# هم‌ترازی ردیف‌به‌ردیف با چک‌لیست
align_pg = sum(1 for i in range(min(len(ndata), len(data))) if ndata[i][nidx['صفحهٔ کلاس']] == data[i][idx['صفحهٔ چاپ']])
align_sk = sum(1 for i in range(min(len(ndata), len(data))) if ndata[i][nidx['مهارت']] == data[i][idx['مهارت']])
print(f'هم‌ترازی ردیفی با چک‌لیست: صفحه {align_pg}/{len(ndata)} | مهارت {align_sk}/{len(ndata)}')
mis_sk = [(ndata[i][nidx['نشست']], data[i][idx['شناسه']]) for i in range(min(len(ndata), len(data))) if ndata[i][nidx['مهارت']] != data[i][idx['مهارت']]]
print('مهارت misalign samples:', mis_sk[:5] if mis_sk else 'هیچ')
# عرضی: وضعیت بعد × مهارت خالی
xt = collections.Counter((('بدون‌مهارت' if not r[nidx['مهارت']] else 'با‌مهارت'), r[nidx['وضعیت بعد']]) for r in ndata)
print('cross مهارت×وضعیت:', dict(xt))

# ---------- ادعاهای پرامپت ----------
print()
print('############ ادعاها در برابر فایل ############')
claims = []
def claim(name, ok, detail):
    claims.append((name, ok, detail))
    print(('  ✅ ' if ok else '  ❌ ') + name + ' — ' + detail)

claim('۲۵۳ کار', len(data) == 253, f'فایل: {len(data)}')
claim('۹۱ کارت', len(kdata) == 91, f'فایل: {len(kdata)}')
claim('۲۵۳ نشست', len(ndata) == 253, f'فایل: {len(ndata)}')
claim('صفحهٔ ۱ تا ۱۴۳، بدون جاماندگی', not missing and plist[0] == 1 and plist[-1] == 143, f'{plist[0]}..{plist[-1]}')
claim('وضعیت همهٔ ردیف‌ها نوشته‌شد', set(st) == {'نوشته' + Z + 'شد'}, str(dict(st)))
claim('نشست: ۱۸۷ یادداشت', stat.get('یادداشت', 0) == 187, f"فایل: {stat.get('یادداشت')}")
claim('نشست: ۲۵ دست‌نخورده', stat.get('دست نخورد', 0) == 25, f"فایل: {stat.get('دست نخورد')}")
claim('نشست: ۴۱ ریاضی‌نیست', stat.get('ریاضی' + Z + 'نیست', 0) == 41, f"فایل: {stat.get('ریاضی' + Z + 'نیست')}")
claim('۴۱ ردیف عدد قطعی نبود', len(rn_rows) == 41, f'فایل: {len(rn_rows)}')
claim('هیچ کارتی باز/بسته نشده', all(('باز' not in r[nidx['نتیجه']] and 'بسته' not in r[nidx['نتیجه']]) or 'کارت نمی' in r[nidx['نتیجه']] for r in ndata), 'همهٔ نتیجه‌ها «کارت نمی‌آید»')
claim('SK-SEE کارت نیست', 'SK-SEE' not in cards, 'عدم حضور')
claim('پرش یکی‌یکی فقط SK-COUNT-5', set(countfam) <= {'SK-COUNT-5'}, str(countfam))
claim('صفحهٔ ۴۶ دو ردیف (تقارن/نیم)', sorted(t for t, _ in pages.get(46, [])) == [1, 2] and pages[46][0][1][idx['مهارت']] == 'SK-SYMM' and pages[46][1][1][idx['مهارت']] == 'SK-HALF', f"{[(t, r[idx['مهارت']]) for t, r in sorted(pages.get(46, []))]}")
claim('حذف‌ها برادرند، هر دو به صدگان', all(cards[k][kidx['پدر']] == 'SK-HUNDREDS' for k in chops), str([(k, cards[k][kidx['پدر']]) for k in chops]))
claim('پدر دیرتر از کارت به دنیا نیامده', not birth_gt, f'{len(birth_gt)} تخلف')
claim('شناسه P46-T1 = صفحهٔ ۴۶ کار ۱', not pg_mis and not tk_mis, 'همهٔ ردیف‌ها هم‌خوان')
claim('تشخیص = دو جمله (الگو)', tpl_bad == 0, f'ok={tpl_ok} bad={tpl_bad}')

print()
print('JSON_SUMMARY:', json.dumps({
    'checklist_rows': len(data), 'cards': len(kdata), 'sessions': len(ndata),
    'pages': [plist[0], plist[-1]], 'status_values': dict(st),
    'session_status': dict(stat), 'roshan_nist_rows': len(rn_rows),
    'missing_pages': missing, 'birth_violations': birth_gt, 'order_violations': order_bad,
    'orphans': orphans, 'dangling': dangling, 'dup_closing': len(dup_clos),
}, ensure_ascii=False))
